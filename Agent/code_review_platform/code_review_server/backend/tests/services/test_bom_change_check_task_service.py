import zipfile
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import Workbook

from app.core.config import settings


def _write_bom_change_report(path: Path, rows: list[tuple[str, int, int, int]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["检查项", "总数", "异常", "人工复查", "跳过"])
    for row in rows:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_group_completeness_report(path: Path, rows: list[tuple[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "61787-1.1"
    ws.append(["物料代码", "关系组", "名称", "规格型号", "客户", "备注"])
    for material_code, remark in rows:
        ws.append([material_code, "T127", "二极管", "BZT52C30-7-F", "关键物料", remark])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_bom_change_prompt_describes_fixed_category_folders_and_matching_rule() -> None:
    from app.services import bom_change_check_task_service

    prompt = bom_change_check_task_service._build_codex_prompt(
        input_path=Path("input"),
        output_path=Path("reports"),
    )

    assert "input/projects" in prompt
    assert "改制前BOM" in prompt
    assert "改制后BOM" in prompt
    assert "更新后BOM" in prompt
    assert "更新BOM" in prompt
    assert "改制表" in prompt
    assert "公司物料表" in prompt
    assert "可以同时放 `关键物料` 和 `无关键物料` 两个替代表" in prompt
    assert "按 BOM 文件名选择" in prompt
    assert "先判断 `无关键物料`" in prompt
    assert "总替代料汇总表有关键物料" in prompt
    assert "总替代料汇总表无关键物料" in prompt
    assert "不要在同一个 workbook 内同时扫描有关键/无关键两个总表 sheet" in prompt
    assert "按压缩包内文件检查" in prompt
    assert "压缩包内没有替代表" in prompt
    assert "只审查替代料是否合规" in prompt
    assert "不要把缺少改制前BOM或改制单判为失败" in prompt
    assert "替代料组齐全差异.xlsx" in prompt
    assert "普通扁平目录输出独立工作簿 `替代料组齐全差异.xlsx`" in prompt
    assert "更新BOM/日期文件夹/BOM文件" in prompt
    assert "每个日期文件夹生成一个 `BOM表改制检查报告-日期.xlsx`" in prompt
    assert "不检查但必须在对应 sheet 写 `备注=跳过`" in prompt
    assert "分层错误、用量与位号数量不一致、位号重复、BOM字段缺失" in prompt
    assert "改为检查 `改制后BOM` 里的每个 BOM" in prompt
    assert "K 开头关系组不检查" in prompt
    assert "O 列包含当前 BOM 客户的物料不要求加入" in prompt
    assert "缺失项只从选中的关键物料/无关键物料替代表生成" in prompt
    assert "`取消替代关系物料` 不用于生成 `替代料组齐全差异.xlsx` 缺失项" in prompt
    assert "只作为兜底放行" in prompt
    assert "归属=保护板1" in prompt
    assert "归属=保护板2" in prompt
    assert "不要用 `接口板1`、`显示板1`、`LED彩灯条1`" in prompt
    assert "61101-1.0" in prompt
    assert "1101-1.0" in prompt
    assert "61629-1.0" in prompt
    assert "1629-1.0" in prompt
    assert "BOM表改制检查报告-515451301.xlsx" in prompt
    assert "不要生成单独的 `产品非电子物料` sheet" in prompt
    assert "不要生成单独的 `BOM自检` sheet" in prompt
    assert "烧录/升级程序" in prompt
    assert "烧录程序" in prompt
    assert "PLM取程序" in prompt
    assert "只有这一行以上的改制物品才归类为 `产品非电子物料`" in prompt
    assert "不要因为某行缺少物料代码或位号就直接归类" in prompt
    assert "不能直接从改制单抽取成报告行" in prompt
    assert "工程助理提供软件" in prompt
    assert "分界以上没有位号的标签/PE袋/二维码等物料，即使有物料代码" in prompt
    assert "用规格型号去改制前BOM和改制后BOM里对比" in prompt
    assert "产品非电子物料差异统一输出到 `BOM差异核对`" in prompt
    assert "不要生成、保留、统计或打包无项目号副本 `BOM表改制检查报告.xlsx`" in prompt
    assert "批量汇总" in prompt
    assert "check_bom_change.py" in prompt
    assert "--batch-input" in prompt
    assert "--output-dir" in prompt
    assert "不要重新生成批处理脚本" in prompt
    assert "取消替代关系物料" in prompt
    assert "旧料号" in prompt
    assert "对应替代料为" in prompt
    assert "替代关系要先查选中 workbook 的对应总表 sheet" in prompt
    assert "按 `正常` 放行，不输出到 `替代料关系核对` 明细" in prompt
    assert "备注写 `替代料汇总表未找到该替代关系`" in prompt
    assert "`替代料关系核对` 的 `变更内容` 留空" in prompt
    assert "两个不同关系组的 `类型=主料` 行使用同一个主料物料代码，也视为同一个替代族" in prompt
    assert "`K187` 和 `T715` 都以 `1.03.01.000005` 为主料" in prompt
    assert "`位号` 和 `物料代码` 同等有效" in prompt
    assert "只要指定删除位号在改制后 BOM 中已经不存在，就按正常处理" in prompt
    assert "改制单物料代码空白、写错，或改制前 BOM 已经没有该位号" in prompt
    assert "类别/归属包含 `PCBA`" in prompt
    assert "不输出到 `改制单核对`，也不输出到 `BOM差异核对`" in prompt
    assert "BOM差异核对` 必须输出所有 BOM 对比差异" in prompt
    assert "`IC-烧录程序`/烧录 IC/MCU 程序差异" in prompt
    assert "改制前项目的 MCU 写 `删除`，改制后项目的 MCU 写 `增加`" in prompt
    assert "这两条 MCU 行必须挨在一起" in prompt
    assert "`用量` 列必须填写改制单用量" in prompt
    assert "`变更内容` 直接写 `增加` 或 `删除`" in prompt
    assert "不管是否被改制单同时覆盖删除和增加" in prompt
    assert "拆成两行输出：旧物料写 `删除`，新物料写 `增加`" in prompt
    assert "备注写 `取消替代物料，可以忽略`" in prompt
    assert "替代料行用量为 `替代料` 时按位号数量填写" in prompt
    assert "`用量` 必须等于位号数量" in prompt
    assert "允许位号为空，不检查其位号数量或位号重复" in prompt
    assert "名称含 `BOM表导出` 的辅助导出 sheet 必须跳过" in prompt
    assert "所有符合 BOM 表头的真实 BOM sheet" in prompt
    assert "同一个 BOM 的同一个 sheet 内，同一位号不能出现在多个非替代料行" in prompt
    assert "输出到 `BOM差异核对` 并判 `异常`" in prompt
    assert "不同 sheet 之间不要互相判重复位号" in prompt
    assert "`变更内容` 写 `BOM自检`" in prompt
    assert "`用量=替代料` 的替代料行允许和上方绑定主料共用位号" in prompt
    assert "位号格写有 `配套` 的配套物料行" in prompt
    assert "同一个位号同时出现在 BOM 的主料行和替代料行" in prompt
    assert "必须输出主料增加行，不能因为替代料关系有效而隐藏" in prompt
    assert "`改制单核对` 必须按改制单物料明细逐行输出正常和异常行" in prompt
    assert "不合并多条改制单内容" in prompt
    assert "正常增加项也必须显示" in prompt
    assert "其它明细 sheet 只输出 `结论=异常`、`结论=人工复查`、`结论=跳过`" in prompt
    assert "物料代码`、`变更内容`、`位号`、`用量`、`规格型号`、`备注" in prompt
    assert "不要保留行号、证据、检查类、结论等额外明细列" in prompt
    assert "`改制单核对` 只有异常行显示具体异常原因" in prompt
    assert "正常行备注留空" in prompt


def test_read_reports_summary_sums_all_project_workbooks(tmp_path: Path) -> None:
    from app.services import bom_change_check_task_service

    _write_bom_change_report(
        tmp_path / "项目A" / "BOM表改制检查报告-项目A.xlsx",
        [
            ("改制单核对", 5, 1, 0, 0),
            ("BOM差异核对", 3, 2, 1, 0),
            ("BOM自检", 1, 4, 0, 0),
            ("替代料关系核对", 4, 0, 1, 0),
            ("客户限制核对", 2, 1, 0, 1),
        ],
    )
    _write_bom_change_report(
        tmp_path / "项目B" / "BOM表改制检查报告.xlsx",
        [
            ("改制单核对", 2, 3, 0, 0),
            ("BOM差异核对", 1, 0, 2, 0),
            ("BOM自检", 2, 1, 1, 0),
            ("替代料关系核对", 1, 5, 0, 1),
            ("客户限制核对", 1, 0, 0, 2),
        ],
    )

    summary = bom_change_check_task_service._read_reports_summary(tmp_path)

    assert summary == {
        "change_order_abnormal_count": 4,
        "bom_diff_abnormal_count": 7,
        "replacement_abnormal_count": 5,
        "customer_restriction_abnormal_count": 1,
        "manual_review_count": 5,
        "skipped_count": 4,
    }


def test_read_reports_summary_ignores_legacy_default_copy_when_numbered_report_exists(tmp_path: Path) -> None:
    from app.services import bom_change_check_task_service

    _write_bom_change_report(
        tmp_path / "项目A" / "BOM表改制检查报告-515451301.xlsx",
        [("BOM差异核对", 1, 2, 0, 0)],
    )
    _write_bom_change_report(
        tmp_path / "项目A" / "BOM表改制检查报告.xlsx",
        [("BOM差异核对", 1, 9, 0, 0)],
    )

    summary = bom_change_check_task_service._read_reports_summary(tmp_path)

    assert summary["bom_diff_abnormal_count"] == 2


def test_read_reports_summary_uses_batch_failure_reason_when_no_project_report(tmp_path: Path) -> None:
    from app.services import bom_change_check_task_service

    wb = Workbook()
    ws = wb.active
    ws.title = "批量汇总"
    ws.append(["项目", "模式", "状态", "报告路径", "改制表", "匹配到的改制前BOM", "匹配到的改制后BOM", "失败原因"])
    ws.append([
        "改制单-616721001",
        "BOM表改制检查",
        "失败",
        "",
        "改制单-616721001.xlsx",
        "BOM-SST25-1391-1.1-备货-A008.xlsx",
        "BOM-SST26-1672-1.0-001.xlsx",
        "ValueError: BOM-SST26-1672-1.0-001.xlsx 文件名未包含关键物料/无关键物料，无法选择替代表",
    ])
    wb.save(tmp_path / "批量BOM表改制检查汇总.xlsx")

    with pytest.raises(RuntimeError) as exc_info:
        bom_change_check_task_service._read_reports_summary(tmp_path)

    message = str(exc_info.value)
    assert "项目检查失败" in message
    assert "改制单-616721001" in message
    assert "文件名未包含关键物料/无关键物料" in message


def test_bom_change_report_zip_uses_uploaded_archive_name(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "BOM_CHANGE_CHECK_STORAGE_DIR", str(tmp_path))
    from app.services import bom_change_check_task_service

    task_id = UUID("22222222-2222-2222-2222-222222222222")
    input_dir = tmp_path / str(task_id) / "input"
    report_dir = tmp_path / str(task_id) / "reports" / "项目A"
    log_dir = tmp_path / str(task_id) / "logs"
    input_dir.mkdir(parents=True)
    report_dir.mkdir(parents=True)
    log_dir.mkdir(parents=True)
    (input_dir / "616291001.zip").write_bytes(b"zip")
    (tmp_path / str(task_id) / "reports" / "替代料组齐全差异.xlsx").write_bytes(b"group")
    (report_dir / "BOM表改制检查报告-项目A.xlsx").write_bytes(b"xlsx")
    (log_dir / "codex-stdout.log").write_text("ok", encoding="utf-8")

    report_zip = bom_change_check_task_service._create_report_zip(task_id)

    assert report_zip.name == "616291001.zip"
    with zipfile.ZipFile(report_zip) as archive:
        names = set(archive.namelist())
    assert "input/616291001.zip" in names
    assert "reports/替代料组齐全差异.xlsx" in names
    assert "reports/项目A/BOM表改制检查报告-项目A.xlsx" in names
    assert "logs/codex-stdout.log" in names


def test_apply_replacement_table_keeps_project_archive_company_tables(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "BOM_CHANGE_CHECK_STORAGE_DIR", str(tmp_path))
    from app.services import bom_change_check_task_service

    task_id = UUID("55555555-5555-5555-5555-555555555555")
    company_dir = tmp_path / str(task_id) / "input" / "projects" / "7-31" / "公司物料表"
    company_dir.mkdir(parents=True)
    key_table = company_dir / "替代料表20260731关键物料.xlsx"
    no_key_table = company_dir / "替代料表20260731无关键物料.xlsx"
    key_table.write_bytes(b"key")
    no_key_table.write_bytes(b"no-key")

    result = bom_change_check_task_service.apply_replacement_table_to_task_input(task_id)

    assert result["source"] == "project_archive"
    assert key_table.exists()
    assert no_key_table.exists()
    assert "关键物料" in str(result["file_name"])
    assert "无关键物料" in str(result["file_name"])


def test_bom_change_report_zip_skips_legacy_default_copy_when_numbered_report_exists(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "BOM_CHANGE_CHECK_STORAGE_DIR", str(tmp_path))
    from app.services import bom_change_check_task_service

    task_id = UUID("44444444-4444-4444-4444-444444444444")
    input_dir = tmp_path / str(task_id) / "input"
    report_dir = tmp_path / str(task_id) / "reports" / "改制单-515451301"
    input_dir.mkdir(parents=True)
    report_dir.mkdir(parents=True)
    (input_dir / "7-24.rar").write_bytes(b"rar")
    (report_dir / "BOM表改制检查报告-515451301.xlsx").write_bytes(b"numbered")
    (report_dir / "BOM表改制检查报告.xlsx").write_bytes(b"legacy")

    report_zip = bom_change_check_task_service._create_report_zip(task_id)

    with zipfile.ZipFile(report_zip) as archive:
        names = set(archive.namelist())
    assert "reports/改制单-515451301/BOM表改制检查报告-515451301.xlsx" in names
    assert "reports/改制单-515451301/BOM表改制检查报告.xlsx" not in names


def test_partial_bom_change_result_skips_corrupt_workbook_and_keeps_zip(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "BOM_CHANGE_CHECK_STORAGE_DIR", str(tmp_path))
    from app.services import bom_change_check_task_service

    task_id = UUID("33333333-3333-3333-3333-333333333333")
    input_dir = tmp_path / str(task_id) / "input"
    report_dir = tmp_path / str(task_id) / "reports"
    input_dir.mkdir(parents=True)
    (input_dir / "7-21.rar").write_bytes(b"rar")
    _write_bom_change_report(
        report_dir / "项目A" / "BOM表改制检查报告-项目A.xlsx",
        [("替代料关系核对", 1, 2, 0, 0)],
    )
    corrupt = report_dir / "项目B" / "BOM表改制检查报告-项目B.xlsx"
    corrupt.parent.mkdir(parents=True)
    corrupt.write_bytes(b"partial-xlsx")

    result = bom_change_check_task_service._build_task_result(task_id, allow_partial=True)

    assert result["replacement_abnormal_count"] == 2
    assert result["result_files"][0]["file_name"] == "7-21.zip"
    assert Path(str(result["result_files"][0]["file_path"])).is_file()


def test_bom_change_result_exposes_summary_workbooks_and_counts_group_missing_rows(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "BOM_CHANGE_CHECK_STORAGE_DIR", str(tmp_path))
    from app.services import bom_change_check_task_service

    task_id = UUID("44444444-4444-4444-4444-444444444444")
    input_dir = tmp_path / str(task_id) / "input"
    report_dir = tmp_path / str(task_id) / "reports"
    input_dir.mkdir(parents=True)
    (input_dir / "8-28-1.rar").write_bytes(b"rar")
    _write_bom_change_report(
        report_dir / "项目A" / "BOM表改制检查报告-项目A.xlsx",
        [("BOM差异核对", 1, 2, 0, 0)],
    )
    _write_bom_change_report(
        report_dir / bom_change_check_task_service.BATCH_SUMMARY_NAME,
        [("批量汇总", 1, 2, 0, 0)],
    )
    _write_bom_change_report(
        report_dir / bom_change_check_task_service.BOM_DIFF_SUMMARY_NAME,
        [("BOM差异汇总", 1, 2, 0, 0)],
    )
    _write_group_completeness_report(
        report_dir / bom_change_check_task_service.GROUP_COMPLETENESS_REPORT_NAME,
        [("1.06.01.000012", "BOM缺少该物料"), ("", "无缺失项")],
    )

    result = bom_change_check_task_service._build_task_result(task_id)

    assert [item["file_name"] for item in result["result_files"]] == [
        "8-28-1.zip",
        bom_change_check_task_service.BATCH_SUMMARY_NAME,
        bom_change_check_task_service.BOM_DIFF_SUMMARY_NAME,
        bom_change_check_task_service.GROUP_COMPLETENESS_REPORT_NAME,
    ]
    assert result["bom_diff_abnormal_count"] == 3
    for result_file in result["result_files"][1:]:
        assert result_file["content_type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert Path(str(result_file["file_path"])).is_file()


def test_run_task_treats_codex_timeout_after_batch_finished_as_success(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "BOM_CHANGE_CHECK_STORAGE_DIR", str(tmp_path))
    from app.services import bom_change_check_task_service

    task_id = UUID("66666666-6666-6666-6666-666666666666")
    task_input_dir = tmp_path / str(task_id) / "input"
    task_input_dir.mkdir(parents=True)
    (task_input_dir / "8-18.rar").write_bytes(b"rar")
    _write_bom_change_report(
        tmp_path / str(task_id) / "reports" / "改制单-518751002" / "BOM表改制检查报告-518751002.xlsx",
        [("替代料关系核对", 1, 6, 0, 0)],
    )
    (tmp_path / str(task_id) / "reports" / "timing.jsonl").write_text(
        '{"event": "batch_finished", "projects": 1}\n',
        encoding="utf-8",
    )

    def timeout_after_report(*args, **kwargs):
        raise bom_change_check_task_service.CodexProcessTimeout(
            stdout="",
            stderr="stream disconnected",
        )

    monkeypatch.setattr(bom_change_check_task_service, "_run_codex_process", timeout_after_report)

    result = bom_change_check_task_service.run_task(task_id, [])

    assert result["replacement_abnormal_count"] == 6
    assert result["result_files"][0]["file_name"] == "8-18.zip"

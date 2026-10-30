from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from app.services.bom_electronic_edit_processor import (
    CHANGE_RECORD_SHEET_NAME,
    REPORT_SHEET_NAME,
    ChangeGroup,
    ChangeRecordPeople,
    ChangeRow,
    MaterialInfo,
    ReportRow,
    _change_record_text_from_report_rows,
    _extract_version_from_bom_name,
    _sheet_segment,
    format_quantity,
    parse_change_notice,
    process_bom_workbook,
    validate_processed_bom_workbook,
)


def _save_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        ["1.02.01.000061", "0603旧", "旧描述", "1", "R758", "1.02.01.000107", "", "", "1", "R758"],
        ["1.08.02.000052", "JMGL3V10A", "旧MOS1", "1", "Q77", "1.08.02.000158", "", "", "1", "Q77"],
        ["1.08.02.000110", "JMTL3N10A", "旧MOS2", "替代料", "Q77", "/", "/", "/", "/", "/"],
        ["4.03.01.000048", "旧纸箱", "旧纸箱描述", "1/130", "/", "4.03.08.000007", "", "", "1/120", "/"],
        ["/", "/", "/", "/", "/", "4.03.09.000007", "", "", "1/20", "/"],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_material_master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "物料主数据"
    ws["D1"] = "物料代码"
    ws["E1"] = "物料描述"
    ws["F1"] = "型号"
    ws["I1"] = "BMS规则描述"
    ws["J1"] = "BMS封装"
    ws["K1"] = "一级分类"
    rows = [
        ("1.02.01.000107", "电阻", "0603新", "新描述", "R0603", "贴片料"),
        ("1.02.01.000664", "电阻", "AC0603FR-071KL", "车规1K", "R0603", "贴片料"),
        ("1.08.02.000158", "MOS", "PA910BM", "新MOS", "SOT23", "贴片料"),
        ("4.03.08.000007", "外箱", "新外箱", "新外箱描述", "", "包装料"),
        ("4.03.09.000007", "刀卡组", "新刀卡", "新刀卡描述", "", "包装料"),
    ]
    for row_index, (code, name, model, description, package, material_type) in enumerate(rows, start=2):
        ws.cell(row=row_index, column=4).value = code
        ws.cell(row=row_index, column=5).value = name
        ws.cell(row=row_index, column=6).value = model
        ws.cell(row=row_index, column=9).value = description
        ws.cell(row=row_index, column=10).value = package
        ws.cell(row=row_index, column=11).value = material_type
    wb.save(path)


def _save_shifted_material_master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "物料主数据"
    ws.append(["品牌名称（英文）", "一级分类", "物料代码", "物料描述", "型号", "", "", "BMS规格描述", "BMS封装", "一级分类"])
    ws.append(["", "辅料", "1.00.02.000011", "License", "WBMSHV-License 三颗小草", "", "", "License WBMSHV-License 三颗小草 Wifi高压", "", "辅料"])
    wb.save(path)


def _save_add_only_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    row = ["/", "/", "/", "/", "/", "1.00.02.000011", "WBMSHV-License 三颗小草", "License WBMSHV-License 三颗小草 Wifi高压", "1", "/"]
    for column, value in enumerate(row, start=1):
        ws.cell(row=3, column=column).value = value
    wb.save(path)


def _save_add_only_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装", "备注"])
    ws.append([None, None, "2.01.05.001196", "保护板", "BCU-0602E-1.0", None, None, "保护板描述", None, None])
    ws.append([1, None, "1.01.01.002521", "PCB", "SST25-0636-1.0", 1, None, "PCB描述", None, None])
    ws.append([2, "贴片料", "1.02.01.000002", "电阻", "0603WAJ0000T5E", 7, "R1,R2,R3,R4,R5,R6,R7", "电阻描述", "R0603", None])
    ws.append([3, "插件料", "1.20.01.000027", "接线端子", "TLS5446-T-H", 2, "J5,J4", "端子描述", "DIP4", None])
    ws.append([4, "辅料", "1.14.01.000003", "纽扣电池", "CR2032", 1, "配套BT1", "纽扣电池描述", "20X20X3.2", None])
    ws.append([5, "辅料", "4.02.01.000001", "三防漆", "PU1030-40S-1-LF", "19.9ml", "", "三防漆描述", "", None])
    record = wb.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "拟制", "审核", "批准"])
    record.append(["2026/8/1", "1.0", "旧记录", "廖晶华", "邓如锦", "许志杰"])
    wb.save(path)


def _save_add_only_change_notice_for_existing_material(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    row = ["/", "/", "/", "/", "/", "1.03.01.000009", "0402B104K500CT", "贴片陶瓷电容", "4", "C522,C520,C521,C523"]
    for column, value in enumerate(row, start=1):
        ws.cell(row=3, column=column).value = value
    wb.save(path)


def _save_bom_with_existing_add_material(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000029", "电阻", "0603WAJ0104T5E", 1, "RH15", "贴片厚膜电阻", "R0603"])
    ws.append([2, "贴片料", "1.03.01.000009", "电容", "0402B104K500CT", 41, "CJ6,CJ11,CJ15,CJ10", "贴片陶瓷电容", "C0402"])
    ws.append([3, "插件料", "1.20.01.000027", "接线端子", "TLS5446-T-H", 2, "J5,J4", "端子描述", "DIP4"])
    wb.save(path)


def _save_replacement_change_notice_with_existing_target(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    row = ["1.03.01.000029", "0603", "旧电容", "41", "CJ6,CJ11", "1.03.01.000009", "0402", "新电容", "41", "CJ6,CJ11"]
    for column, value in enumerate(row, start=1):
        ws.cell(row=3, column=column).value = value
    wb.save(path)


def _save_bom_with_existing_replacement_target(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.03.01.000009", "电容", "0402B104K500CT", 4, "C522,C520,C521,C523", "贴片陶瓷电容", "C0402"])
    ws.append([2, "贴片料", "1.03.01.000029", "电容", "0603WAJ0104T5E", 41, "CJ6,CJ11", "贴片陶瓷电容", "C0603"])
    wb.save(path)


def _save_layered_add_only_bom(path: Path) -> None:
    wb = Workbook()
    for index, title in enumerate(["总表", "P8S200A-61097-1.0", "DIP-61097-1.0", "SMT-61097-1.0"]):
        ws = wb.active if index == 0 else wb.create_sheet(title)
        ws.title = title
        ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
        ws.append([None, None, "2.01.05.001196", "保护板", "BCU-0602E-1.0", None, None, "保护板描述", None])
        ws.append([1, None, "1.01.01.002521", "PCB", "SST25-0636-1.0", 1, None, "PCB描述", None])
        ws.append([2, "贴片料", "1.02.01.000002", "电阻", "0603WAJ0000T5E", 7, "R1,R2,R3,R4,R5,R6,R7", "电阻描述", "R0603"])
        ws.append([3, "插件料", "1.20.01.000027", "接线端子", "TLS5446-T-H", 2, "J5,J4", "端子描述", "DIP4"])
        ws.append([4, "辅料", "1.14.01.000003", "纽扣电池", "CR2032", 1, "配套BT1", "纽扣电池描述", "20X20X3.2"])
    wb.save(path)


def _save_layered_tvs_replacement_bom(path: Path) -> None:
    wb = Workbook()
    for index, title in enumerate(["总表", "P13S120A-08331-3.0", "DIP-08331-1.0", "SMT-08331-1.0"]):
        ws = wb.active if index == 0 else wb.create_sheet(title)
        ws.title = title
        if index == 0:
            ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
        ws.append([116, "贴片料", "1.06.02.000020", "二极管", "SMCJ100CA", 1, "DP7", "TVS管 100V 111V~123V 1500W 双向 SMC", "SMC"])
        ws.append([None, "贴片料", "1.06.02.000155", "二极管", "SMCJ100CA", "替代料", "DP7", "TVS管 100V 111V~123V 1500W 双向 SMC", "SMC"])
        ws.append([None, "贴片料", "1.06.02.000108", "二极管", "SMCJ100CA", "替代料", "DP7", "TVS管 100V 111-123V 1500W 双向 SMC 萨瑞微", "SMC"])
    wb.save(path)


def _save_no_type_pcb_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    row = ["/", "/", "/", "/", "/", "1.01.01.009999", "SST-PCB", "PCB描述", "1", "/"]
    for column, value in enumerate(row, start=1):
        ws.cell(row=3, column=column).value = value
    wb.save(path)


def _save_no_type_pcb_master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "物料主数据"
    ws.append(["品牌名称（英文）", "一级分类", "物料代码", "物料描述", "型号", "", "", "BMS规格描述", "BMS封装", "一级分类"])
    ws.append(["", "", "1.01.01.009999", "PCB", "SST-PCB", "", "", "PCB描述", "", ""])
    wb.save(path)


def _save_material_master_with_bms_package_and_blank_package(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "物料主数据"
    ws.append([
        "品牌名称（英文）",
        "一级分类",
        "物料代码",
        "物料描述",
        "型号",
        "",
        "",
        "BMS规格描述",
        "BMS封装",
        "一级分类",
        "采购说明（BMS）",
        "每箱数量",
        "默认仓库",
        "英文描述",
        "库存计量单位",
        "封装",
    ])
    ws.append([
        "瞬雷",
        "贴片料",
        "1.06.02.000194",
        "二极管",
        "SMDJ130A",
        "",
        "",
        "TVS管 130V 144~159V 3000W 单向 SMC",
        "SMC",
        "贴片料",
        "",
        "",
        "BMS公司仓",
        "",
        "",
        "",
    ])
    wb.save(path)


def _save_add_tvs_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    row = ["/", "/", "/", "/", "/", "1.06.02.000194", "SMDJ130A", "TVS管 130V 144~159V 3000W 单向 SMC", "1", "DP7"]
    for column, value in enumerate(row, start=1):
        ws.cell(row=3, column=column).value = value
    wb.save(path)


def _save_unchanged_quantity_locator_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        [
            "1.08.02.000130",
            "DP012N10TGNI3",
            "NMOS DP012N10TG NI3 100V",
            "用量不变",
            "位号不变",
            "1.08.02.000179",
            "CC360N10HS",
            "",
            "用量不变",
            "位号不变",
        ],
        [
            "1.08.02.000101",
            "DSU014N10N3A",
            "NMOS DSU014N10N 3A 100V",
            "用量不变",
            "位号不变",
            "",
            "",
            "",
            "",
            "",
        ],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_unchanged_quantity_locator_master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "物料主数据"
    ws["D1"] = "物料代码"
    ws["E1"] = "物料描述"
    ws["F1"] = "型号"
    ws["I1"] = "BMS规则描述"
    ws["J1"] = "BMS封装"
    ws["K1"] = "一级分类"
    ws.cell(row=2, column=4).value = "1.08.02.000179"
    ws.cell(row=2, column=5).value = "MOS"
    ws.cell(row=2, column=6).value = "CC360N10HS"
    ws.cell(row=2, column=9).value = "MOS CC360N10HS 100V"
    ws.cell(row=2, column=10).value = "TO-263"
    ws.cell(row=2, column=11).value = "贴片料"
    wb.save(path)


def _save_unchanged_quantity_locator_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.08.02.000130", "MOS", "DP012N10TGNI3", 2, "Q1,Q2", "NMOS DP012N10TG NI3 100V", "TO-263"])
    ws.append([2, "贴片料", "1.08.02.000101", "MOS", "DSU014N10N3A", 4, "Q3,Q4,Q5,Q6", "NMOS DSU014N10N 3A 100V", "TO-263"])
    wb.save(path)


def _save_unchanged_quantity_locator_bom_with_existing_target(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.08.02.000130", "MOS", "DP012N10TGNI3", 2, "Q1,Q2", "NMOS DP012N10TG NI3 100V", "TO-263"])
    ws.append([2, "贴片料", "1.08.02.000179", "MOS", "CC360N10HS", 1, "Q7", "MOS CC360N10HS 100V", "TO-263"])
    ws.append([3, "贴片料", "1.08.02.000101", "MOS", "DSU014N10N3A", 4, "Q3,Q4,Q5,Q6", "NMOS DSU014N10N 3A 100V", "TO-263"])
    wb.save(path)


def _save_same_locator_replacement_group_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        [
            "1.08.02.000176",
            "DP012N10TGNI3",
            "NMOS DP012N10TGNI3 100V 1.05mΩ VTH：2.5V-2.8V TOLL-8L",
            "用量不变",
            "位号不变",
            "1.08.02.000179",
            "CC360N10HS",
            "",
            "用量不变",
            "位号不变",
        ],
        [
            "1.08.02.000130",
            "DP012N10TGNI3",
            "NMOS DP012N10TGNI3 100V 1.05mΩ 二档电压2.8~3.0V TOLL-8L",
            "用量不变",
            "位号不变",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "1.08.02.000152",
            "DP012N10TGNI3",
            "NMOS DP012N10TGNI3 100V 1.05mΩ 三档驱动电压3.0-3.2V TOLL-8L",
            "用量不变",
            "位号不变",
            "",
            "",
            "",
            "",
            "",
        ],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_same_locator_replacement_group_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    locators = "QP48,QP43,QP44,QP45"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([77, "贴片料", "1.08.02.000176", "MOS管", "DP012N10TGNI3", 34, locators, "NMOS DP012N10TGNI3 100V 1.05mΩ VTH：2.5V-2.8V TOLL-8L", "TOLL-8L"])
    ws.append([None, "贴片料", "1.08.02.000130", "MOS管", "DP012N10TGNI3", "替代料", locators, "NMOS DP012N10TGNI3 100V 1.05mΩ TOLL-8L", "TOLL-8L"])
    ws.append([None, "贴片料", "1.08.02.000152", "MOS管", "DP012N10TGNI3", "替代料", locators, "NMOS DP012N10TGNI3 100V 1.05mΩ 三档驱动电压3.0-3.2V TOLL-8L", "TOLL-8L"])
    ws.append([78, "贴片料", "1.08.04.000008", "MOS管", "JST2305A", 1, "QL1", "PMOS JST2305A", "SOT23"])
    wb.save(path)


def _save_same_locator_replacement_group_bom_with_existing_target(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    locators = "QP50,QP51,QP39,QP27,QP25,QP2"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([113, "贴片料", "1.08.02.000130", "MOS管", "DP012N10TGNI3", 6, locators, "NMOS DP012N10TGNI3 100V 1.05mΩ 二档电压2.8~3.0V TOLL-8L", "TOLL-8L"])
    ws.append([None, "贴片料", "1.08.02.000152", "MOS管", "DP012N10TGNI3", "替代料", locators, "NMOS DP012N10TGNI3 100V 1.05mΩ 三档驱动电压3.0-3.2V TOLL-8L", "TOLL-8L"])
    ws.append([None, "贴片料", "1.08.02.000179", "MOS管", "CC360N10HS", "替代料", locators, "NMOS CC360N10HS 100V 1.0mΩ TOLL-8L", "TOLL-8L"])
    ws.append([None, "贴片料", "1.08.02.000176", "MOS管", "DP012N10TGNI3", "替代料", locators, "NMOS DP012N10TGNI3 100V 1.05mΩ VTH：2.5V-2.8V TOLL-8L", "TOLL-8L"])
    record = wb.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "拟制", "审核", "批准"])
    record.append(["2026/8/1", "1.0", "旧记录", "廖晶华", "邓如锦", "许志杰"])
    wb.save(path)


def _save_partial_locator_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        ["1.06.02.000020", "SMAJ30A", "旧主料", "1", "DP7", "1.06.02.000194", "SMDJ130A", "新TVS", "1", "DP7"],
        ["1.06.02.000155", "SMAJ30A", "旧替代1", "替代料", "DP7", "/", "/", "/", "/", "/"],
        ["1.06.02.000108", "SMAJ30A", "旧替代2", "替代料", "DP7", "/", "/", "/", "/", "/"],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_partial_locator_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料编码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.06.02.000020", "二极管", "SMAJ30A", 1, "DP7", "旧主料", "SMA"])
    ws.append([2, "贴片料", "1.02.01.000001", "电阻", "0603", 1, "R1", "不相关", "R0603"])
    smt = wb.create_sheet("SMT-TEST")
    smt.append(["序号", "物料类型", "物料编码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    smt.append([1, "贴片料", "1.06.02.000108", "二极管", "SMAJ30A", 1, "DP7", "旧替代2", "SMA"])
    smt.append([None, "贴片料", "1.06.02.000020", "二极管", "SMAJ30A", "替代料", "DP7", "旧主料替代", "SMA"])
    smt.append([2, "贴片料", "1.02.01.000001", "电阻", "0603", 1, "R1", "不相关", "R0603"])
    wb.save(path)


def _save_already_changed_locator_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料编码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([3, "贴片料", "1.06.02.000194", "二极管", "SMDJ130A", 1, "DP7", "新TVS", "SMC"])
    ws.append([4, "贴片料", "1.02.01.000001", "电阻", "0603", 1, "R1", "不相关", "R0603"])
    record = wb.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "拟制", "审核", "批准"])
    record.append(["2026/8/1", "3.1", "已有记录", "廖晶华", "邓如锦", "许志杰"])
    wb.save(path)


def _save_product_sheet_locator_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "P16S120A-11870-3.1"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([108, "贴片料", "1.06.02.000108", "二极管", "SMCJ100CA", 1, "DP7", "TVS管 100V 111-123V 1500W 双向 SMC 萨瑞微", "SMC"])
    ws.append([None, "贴片料", "1.06.02.000020", "二极管", "SMCJ100CA", "替代料", "DP7", "TVS管 100V 111V 1500W 双向  SMC 硕凯", "SMC"])
    wb.save(path)


def _save_two_sheet_model_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "P16S120A-11870-3.1"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([107, "贴片料", "1.06.02.000019", "二极管", "SMAJ30A", 1, "TVS17", "TVS管 30V 33.3V 400W 单向 SMA", "SMAS"])
    record = wb.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "拟制", "审核", "批准"])
    wb.save(path)


def _save_bom_with_material_category_and_name_headers(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料分类", "物料代码", "物料名称", "规格型号", "用量", "位号", "规格描述", "封装", "备注", "类型"])
    ws.append([110, "贴片料", "1.06.02.000152", "二极管", "SMAJ30A", 1, "TVS17", "旧描述", "SMA", "", "主料"])
    ws.append([111, "贴片料", "1.06.02.000019", "二极管", "SMAJ30A", "替代料", "TVS17", "旧替代", "SMA", "", "替代料"])
    wb.save(path)


def _save_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    ws.append(headers)
    ws.append(["1.02.01.000061", "0603旧", "旧描述", 1, "R758"])
    ws.append(["1.08.02.000052", "JMGL3V10A", "旧MOS1", 1, "Q77"])
    ws.append(["1.08.02.000110", "JMTL3N10A", "旧MOS2", "替代料", "Q77"])
    ws.append(["4.03.01.000048", "旧纸箱", "旧纸箱描述", "1/130", "/"])
    ws2 = wb.create_sheet("DIP")
    ws2.append(headers)
    ws2.append(["1.02.01.000061", "不一致", "旧描述", 1, "R758"])
    wb.save(path)


def _save_bom_with_material_code_header(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append(["1", "贴片料", "1.02.01.000061", "电阻", "0603旧", 1, "R758", "旧描述", "旧封装"])
    wb.save(path)


def _save_packaging_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [[
        "4.03.01.000048",
        "沛城纸箱585*550*290mm",
        "旧纸箱描述",
        "1/130",
        "/",
        "4.03.08.000007",
        "",
        "",
        "1/120",
        "/",
    ], [
        "/",
        "/",
        "/",
        "/",
        "/",
        "4.03.09.000007",
        "",
        "",
        "1/20",
        "0.05",
    ]]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_locator_force_change_notice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        ["1.02.01.000014", "0603WAF1001T5E", "旧1K", "1", "R762", "1.02.01.000664", "", "", "1", "R762"],
        ["1.08.02.000052", "JMGL3V10A", "旧MOS1", "1", "Q77", "1.08.02.000158", "", "", "1", "Q77"],
        ["1.08.02.000110", "JMTL3N10A", "旧MOS2", "替代料", "Q77", "/", "/", "/", "/", "/"],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_change_notice_with_inventory_rows(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        ["1.02.01.000061", "0603旧", "旧描述", "1", "R758", "1.02.01.000107", "", "", "1", "R758"],
        ["1.02.01.000061", "0603旧", "公司仓+外协厂", "在库+在途", "", "426288+230000", "", "", "", ""],
        ["1.08.02.000052", "JMGL3V10A", "公司仓+外协厂", "在库+在途", "", "95+0", "", "", "", ""],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(path)


def _save_locator_force_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料编码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([15, "贴片料", "1.02.01.000014", "电阻", "0603WAF1001T5E", 5, "RJ23,RJ24,RJ27,RJ22,R762", "旧1K", "R0603"])
    ws.append([16, "贴片料", "1.02.01.000015", "电阻", "0402WGJ0102TCE", 2, "R124,R75", "后续1K", "R0402"])
    ws.append([None, "贴片料", "1.08.02.000110", "MOS管", "JMTL3N10A", 1, "Q77", "旧MOS2", "SOT23"])
    ws.append([None, "贴片料", "1.08.02.000052", "中电压NMOS", "JMGL3V10A", "替代料", "Q77", "旧MOS1", "SOT23"])
    ws.append([99, "贴片料", "1.02.01.000213", "电阻", "0402WGF100JTCE", "替代料", "R768", "替代料不占序号", "R0402"])
    record = wb.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "拟制", "审核", "批准"])
    record.append(["2026/7/15", "1.1", "旧记录", "廖晶华", "邓如锦", "许志杰"])
    wb.save(path)


def _save_packaging_bom(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料编码", "名称", "规格型号", "用量", "位号", "规格描述"])
    ws.append([
        9,
        "包装料",
        "4.03.01.000048",
        "旧纸箱",
        "沛城纸箱585*550*290mm",
        "1/130",
        "",
        "旧纸箱描述",
    ])
    ws.append([
        12,
        "包装料",
        "4.04.01.000025",
        "气泡袋",
        "气泡袋",
        1,
        "",
        "气泡袋描述",
    ])
    record = wb.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "拟制", "审核", "批准"])
    record.append(["2026/7/15", "1.3", "更改R1479=1.02.01.000664", "廖晶华", "邓如锦", "许志杰"])
    wb.save(path)


def _save_packaging_bom_with_merged_blank_record_rows(path: Path) -> None:
    _save_packaging_bom(path)
    wb = load_workbook(path)
    record = wb[CHANGE_RECORD_SHEET_NAME]
    record.merge_cells("A3:F3")
    record.merge_cells("A4:F4")
    record.merge_cells("A5:F5")
    wb.save(path)


def test_format_quantity_handles_fractions() -> None:
    assert format_quantity("1/20") == 0.05
    assert format_quantity("1/120") == 0.0083
    assert format_quantity("替代料") == "替代料"


def test_process_bom_workbook_modifies_groups_and_writes_report(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_change_notice(change_notice)
    _save_bom(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 8
    assert result.abnormal_count == 0
    wb = load_workbook(output)
    ws = wb["总表"]
    assert ws.cell(row=2, column=1).value == "1.02.01.000107"
    assert ws.cell(row=2, column=2).value == "0603新"
    assert ws.cell(row=3, column=1).value == "1.08.02.000158"
    assert ws.cell(row=4, column=1).value == "4.03.08.000007"
    assert ws.cell(row=4, column=4).value == 0.0083
    assert ws.cell(row=5, column=1).value == "4.03.09.000007"
    assert ws.cell(row=5, column=4).value == 0.05
    assert ws.cell(row=2, column=1).fill.fgColor.rgb == "00BDD7EE"

    report = wb[REPORT_SHEET_NAME]
    report_values = list(report.iter_rows(values_only=True))
    assert not any(row[3] == "异常" for row in report_values)


def test_process_bom_workbook_accepts_material_code_header(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_change_notice(change_notice)
    _save_bom_with_material_code_header(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.processed_sheet_count == 1
    wb = load_workbook(output)
    ws = wb["总表"]
    assert ws.cell(row=2, column=3).value == "1.02.01.000107"
    assert ws.cell(row=2, column=5).value == "0603新"
    assert ws.cell(row=2, column=8).value == "新描述"
    assert ws.cell(row=2, column=9).value == "R0603"


def test_process_bom_workbook_combines_replace_checklist_rows(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_change_notice(change_notice)
    _save_bom(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    row = next(item for item in result.checklist_rows if item.before_code == "1.02.01.000061")
    assert row.action == "删除"
    assert row.after_action == "添加"
    assert row.after_code == "1.02.01.000107"
    assert row.status == "已完成"


def test_process_bom_workbook_can_skip_report_sheet(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_change_notice(change_notice)
    _save_bom(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        include_report_sheet=False,
    )

    assert result.modified_count > 0
    wb = load_workbook(output)
    assert REPORT_SHEET_NAME not in wb.sheetnames


def test_process_bom_workbook_normalizes_packaging_quantity_and_blank_locator(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_packaging_change_notice(change_notice)
    _save_packaging_bom(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.abnormal_count == 0
    assert result.modified_count == 3
    wb = load_workbook(output)
    ws = wb["总表"]
    assert ws.cell(row=2, column=1).value == 10
    assert ws.cell(row=2, column=2).value == "包装料"
    assert ws.cell(row=2, column=3).value == "4.03.08.000007"
    assert ws.cell(row=2, column=4).value == "外箱"
    assert ws.cell(row=2, column=6).value == 0.0083
    assert ws.cell(row=2, column=6).number_format == "General"
    assert ws.cell(row=2, column=7).value in (None, "")
    assert ws.cell(row=3, column=1).value == 11
    assert ws.cell(row=3, column=2).value == "包装料"
    assert ws.cell(row=3, column=3).value == "4.03.09.000007"
    assert ws.cell(row=3, column=4).value == "刀卡组"
    assert ws.cell(row=3, column=6).value == 0.05
    assert ws.cell(row=3, column=6).number_format == "General"
    assert ws.cell(row=3, column=7).value in (None, "")
    assert ws.cell(row=4, column=1).value == 12
    report = wb[REPORT_SHEET_NAME]
    assert [report.cell(row=row_index, column=13).value for row_index in range(3, 5)] == [
        "0.0083",
        "0.05",
    ]


def test_change_record_describes_same_material_quantity_increase() -> None:
    group = ChangeGroup(
        key="code:4.03.02.000021:1",
        locator_type="物料编码",
        locator_value="4.03.02.000021",
        rows=[
            ChangeRow(
                before_code="4.03.02.000021",
                before_model="标签纸",
                before_description="标签",
                before_quantity="1",
                before_locator="/",
                after_code="4.03.02.000021",
                after_model="标签纸",
                after_description="标签",
                after_quantity="2",
                after_locator="/",
            )
        ],
    )
    report_rows = [
        ReportRow(
            sheet_name="总表",
            locator_type=group.locator_type,
            locator_value=group.locator_value,
            action="删除",
            before_code="4.03.02.000021",
            before_quantity="1",
            before_locator="/",
            result="成功",
        ),
        ReportRow(
            sheet_name="总表",
            locator_type=group.locator_type,
            locator_value=group.locator_value,
            action="添加",
            after_code="4.03.02.000021",
            after_quantity="2",
            after_locator="/",
            result="成功",
        ),
    ]

    assert (
        _change_record_text_from_report_rows([group], report_rows)
        == "增加4.03.02.000021的用量1"
    )


def test_process_bom_workbook_appends_change_record(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST25-0698-1.3-阿瑞斯-004.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_packaging_change_notice(change_notice)
    _save_packaging_bom(bom)
    _save_material_master(master)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 7, 31),
    )

    wb = load_workbook(output, data_only=True)
    record = wb[CHANGE_RECORD_SHEET_NAME]
    assert record.cell(row=3, column=1).value.date().isoformat() == "2026-07-31"
    assert record.cell(row=3, column=2).value == "1.3"
    assert record.cell(row=3, column=3).value == "增加包材"
    assert record.cell(row=3, column=4).value == "廖晶华"
    assert record.cell(row=3, column=5).value == "邓如锦"
    assert record.cell(row=3, column=6).value == "许志杰"


def test_process_bom_workbook_appends_change_record_after_merged_blank_rows(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST21-B470-5.0-002.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_packaging_change_notice(change_notice)
    _save_packaging_bom_with_merged_blank_record_rows(bom)
    _save_material_master(master)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 8, 4),
    )

    wb = load_workbook(output, data_only=True)
    record = wb[CHANGE_RECORD_SHEET_NAME]
    assert record.cell(row=6, column=1).value.date().isoformat() == "2026-08-04"
    assert record.cell(row=6, column=2).value == "5.0"
    assert record.cell(row=6, column=3).value == "增加包材"


def test_extract_version_from_bom_name_prefers_rev_version() -> None:
    assert _extract_version_from_bom_name(Path("BOM-SST22-1133-REV2.0-20240511.xlsx")) == "2.0"
    assert _extract_version_from_bom_name(Path("BOM-SST22-1133-REV1.1-20220517.xlsx")) == "1.1"
    assert _extract_version_from_bom_name(Path("bom-sst22-1133-rev1.2-20220517.xlsx")) == "1.2"
    assert _extract_version_from_bom_name(Path("BOM-SST23-1399-1.0-003.xlsx")) == "1.0"
    assert _extract_version_from_bom_name(Path("BOM-SST25-0698-1.3-阿瑞斯-004.xlsx")) == "1.3"
    assert _extract_version_from_bom_name(Path("bom.xlsx")) == ""


def test_process_bom_workbook_change_record_version_matches_rev_file_name(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST22-1133-REV1.1-20220516.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "BOM-SST22-1133-REV1.1-20220517.xlsx"
    _save_packaging_change_notice(change_notice)
    _save_packaging_bom(bom)
    _save_material_master(master)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 10, 30),
    )

    wb = load_workbook(output, data_only=True)
    record = wb[CHANGE_RECORD_SHEET_NAME]
    assert record.cell(row=3, column=1).value.date().isoformat() == "2026-10-30"
    assert record.cell(row=3, column=2).value == "1.1"
    assert record.cell(row=3, column=3).value == "增加包材"


def test_process_bom_workbook_splits_multi_locator_and_force_updates_locator_group(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_locator_force_change_notice(change_notice)
    _save_locator_force_bom(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.abnormal_count == 0
    wb = load_workbook(output)
    ws = wb["总表"]
    assert ws.cell(row=2, column=3).value == "1.02.01.000014"
    assert ws.cell(row=2, column=6).value == 4
    assert ws.cell(row=2, column=7).value == "RJ23,RJ24,RJ27,RJ22"
    assert ws.cell(row=3, column=3).value == "1.02.01.000664"
    assert ws.cell(row=3, column=6).value == 1
    assert ws.cell(row=3, column=7).value == "R762"
    assert ws.cell(row=3, column=8).value == "车规1K"
    assert ws.cell(row=3, column=9).value == "R0603"
    assert ws.cell(row=2, column=1).value == 15
    assert ws.cell(row=3, column=1).value == 16
    assert ws.cell(row=4, column=1).value == 17
    assert ws.cell(row=5, column=3).value == "1.08.02.000158"
    assert ws.cell(row=5, column=6).value == 1
    assert ws.cell(row=5, column=7).value == "Q77"
    assert ws.cell(row=5, column=1).value == 18
    assert ws.cell(row=5, column=9).value == "SOT23"
    assert ws.max_row == 5
    assert not any(
        ws.cell(row=row_index, column=6).value == "替代料"
        and ws.cell(row=row_index, column=7).value == "Q77"
        for row_index in range(2, ws.max_row + 1)
    )


def test_process_bom_workbook_skips_ecN_when_target_is_already_the_only_material(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_locator_force_change_notice(change_notice)
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.08.02.000158", "MOS", "PA910BM", 1, "Q77", "新MOS", "SOT23"])
    wb.save(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 0
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.max_row == 2
    assert ws.cell(row=2, column=3).value == "1.08.02.000158"
    assert ws.cell(row=2, column=7).value == "Q77"


def test_process_bom_workbook_keeps_target_when_target_is_repeated_on_before_side(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_locator_force_change_notice(change_notice)
    wb = load_workbook(change_notice)
    ws = wb.active
    ws.insert_rows(5)
    values = [
        "1.08.02.000158", "PA910BM", "目标MOS", "1", "Q77",
        "1.08.02.000158", "PA910BM", "目标MOS", "1", "Q77",
    ]
    for column, value in enumerate(values, start=1):
        ws.cell(row=5, column=column).value = value
    wb.save(change_notice)

    _save_locator_force_bom(bom)
    wb = load_workbook(bom)
    ws = wb["总表"]
    ws.insert_rows(4)
    ws.cell(row=4, column=1).value = 17
    ws.cell(row=4, column=2).value = "贴片料"
    ws.cell(row=4, column=3).value = "1.08.02.000158"
    ws.cell(row=4, column=4).value = "MOS管"
    ws.cell(row=4, column=5).value = "PA910BM"
    ws.cell(row=4, column=6).value = 1
    ws.cell(row=4, column=7).value = "Q77"
    ws.cell(row=4, column=8).value = "目标MOS"
    ws.cell(row=4, column=9).value = "SOT23"
    wb.save(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    q77_rows = [
        row_index
        for row_index in range(2, ws.max_row + 1)
        if ws.cell(row=row_index, column=7).value == "Q77"
    ]
    assert len(q77_rows) == 1
    assert ws.cell(row=q77_rows[0], column=3).value == "1.08.02.000158"


def test_process_bom_workbook_writes_locator_material_change_record(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST25-1050-1.1-004.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_locator_force_change_notice(change_notice)
    _save_locator_force_bom(bom)
    _save_material_master(master)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 8, 3),
    )

    wb = load_workbook(output, data_only=True)
    record = wb[CHANGE_RECORD_SHEET_NAME]
    assert record.cell(row=3, column=3).value == (
        "添加物料R762=1.02.01.000664、Q77=1.08.02.000158；"
        "删除R762=1.02.01.000014、Q77=1.08.02.000110、Q77=1.08.02.000052"
    )


def test_process_bom_workbook_updates_partial_before_materials_in_same_locator_group(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_partial_locator_change_notice(change_notice)
    _save_partial_locator_bom(bom)
    _save_material_master_with_bms_package_and_blank_package(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.abnormal_count == 0
    assert result.modified_count == 5
    wb = load_workbook(output, data_only=True)
    total = wb["总表"]
    assert total.cell(row=2, column=3).value == "1.06.02.000194"
    assert total.cell(row=2, column=6).value == 1
    assert total.cell(row=2, column=7).value == "DP7"
    smt = wb["SMT-TEST"]
    assert smt.cell(row=2, column=3).value == "1.06.02.000194"
    assert smt.cell(row=3, column=3).value == "1.02.01.000001"
    report_rows = list(wb[REPORT_SHEET_NAME].iter_rows(values_only=True))
    assert not any(row[3] == "异常" for row in report_rows)
    deleted_codes = [row[4] for row in report_rows if row[3] == "删除"]
    assert deleted_codes == ["1.06.02.000020", "1.06.02.000108", "1.06.02.000020"]


def test_process_bom_workbook_splits_partial_locator_change_and_syncs_alternatives(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    ws.append(["1.02.01.000014", "旧电阻", "", 2, "R1,R2", "1.02.01.000664", "新电阻", "", 2, "R1,R2"])
    wb.save(change_notice)
    _save_material_master(master)
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000014", "电阻", "旧电阻", 3, "R1,R2,R3", "旧描述", "R0603"])
    ws.append([None, "贴片料", "1.02.01.000539", "电阻", "旧替代1", "替代料", "R1,R2,R3", "旧替代描述", "R0603"])
    ws.append([None, "贴片料", "1.02.01.000788", "电阻", "旧替代2", "替代料", "R1,R2,R3", "旧替代描述", "R0603"])
    ws.append([2, "贴片料", "1.02.01.000003", "电阻", "保留电阻", 1, "R4", "保留描述", "R0603"])
    wb.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert [(ws.cell(row=index, column=3).value, ws.cell(row=index, column=6).value, ws.cell(row=index, column=7).value) for index in range(2, 6)] == [
        ("1.02.01.000014", 1, "R3"),
        ("1.02.01.000539", "替代料", "R3"),
        ("1.02.01.000788", "替代料", "R3"),
        ("1.02.01.000664", 2, "R1,R2"),
    ]
    assert [ws.cell(row=index, column=1).value for index in (2, 3, 4, 5, 6)] == [1, None, None, 2, 3]


def test_process_bom_workbook_adds_new_material_after_partial_locator_cancellation(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    ws.append(["1.02.01.000014", "旧电阻", "", 1, "R2", "/", "/", "/", "/", "/"])
    ws.append(["/", "/", "/", "/", "/", "1.02.01.000664", "新电阻", "", 1, "R2"])
    wb.save(change_notice)
    _save_material_master(master)
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000014", "电阻", "旧电阻", 2, "R1,R2", "旧描述", "R0603"])
    ws.append([None, "贴片料", "1.02.01.000539", "电阻", "旧替代", "替代料", "R1,R2", "旧替代描述", "R0603"])
    wb.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert [(ws.cell(row=index, column=3).value, ws.cell(row=index, column=6).value, ws.cell(row=index, column=7).value) for index in range(2, 5)] == [
        ("1.02.01.000014", 1, "R1"),
        ("1.02.01.000539", "替代料", "R1"),
        ("1.02.01.000664", 1, "R2"),
    ]


def test_process_bom_workbook_adds_alternative_without_doubling_unchanged_primary(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "\u53d8\u66f4\u524d\u63cf\u8ff0"
    ws["F1"] = "\u53d8\u66f4\u540e\u63cf\u8ff0"
    headers = ["\u7269\u6599\u7f16\u7801", "\u89c4\u683c\u578b\u53f7", "\u89c4\u683c\u63cf\u8ff0", "\u7528\u91cf", "\u4f4d\u53f7"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    ws.append([
        "1.12.04.000016", "ITV9550L5045", "primary", 3, "F2,F3,F1",
        "1.12.04.000016", "ITV9550L5045", "primary", 3, "F2,F3,F1",
    ])
    ws.append([
        "/", "/", "/", "/", "/",
        "1.12.04.000025", "WPF45A14K-3P", "alternative", "\u66ff\u4ee3\u6599", "F2,F3,F1",
    ])
    wb.save(change_notice)

    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["\u7269\u6599\u7f16\u7801", "\u540d\u79f0", "\u89c4\u683c\u578b\u53f7", "\u89c4\u683c\u63cf\u8ff0", "\u7269\u6599\u7c7b\u578b"])
    ws.append(["1.12.04.000016", "fuse", "ITV9550L5045", "primary desc", "\u8d34\u7247\u6599"])
    ws.append(["1.12.04.000025", "fuse", "WPF45A14K-3P", "alternative desc", "\u8d34\u7247\u6599"])
    wb.save(master)

    wb = Workbook()
    ws = wb.active
    ws.title = "\u603b\u8868"
    ws.append(["\u5e8f\u53f7", "\u7269\u6599\u7c7b\u578b", "\u7269\u6599\u7f16\u7801", "\u540d\u79f0", "\u89c4\u683c\u578b\u53f7", "\u7528\u91cf", "\u4f4d\u53f7", "\u89c4\u683c\u63cf\u8ff0", "\u5c01\u88c5"])
    ws.append([121, "\u8d34\u7247\u6599", "1.12.04.000016", "fuse", "ITV9550L5045", 3, "F2,F3,F1", "old primary desc", "SMD"])
    wb.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["\u603b\u8868"]
    assert [(ws.cell(row=index, column=3).value, ws.cell(row=index, column=6).value, ws.cell(row=index, column=7).value) for index in range(2, 4)] == [
        ("1.12.04.000016", 3, "F2,F3,F1"),
        ("1.12.04.000025", "\u66ff\u4ee3\u6599", "F2,F3,F1"),
    ]


def test_process_bom_workbook_keeps_same_code_addition_out_of_cancelled_row(tmp_path: Path) -> None:
    bom = tmp_path / "BOM-SST26-0777-1.0-001.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.06.02.000194", "二极管", "D", 2, "D119,D120", "旧描述", "SOD"])
    record = book.create_sheet(CHANGE_RECORD_SHEET_NAME)
    record.append(["日期", "版本", "变更记录", "制表", "审核", "批准"])
    record.append([date(2026, 8, 1), "0.9", "历史记录", "制表人", "审核人", "批准人"])
    book.save(bom)

    groups = [
        ChangeGroup(
            "cancel-old-locators",
            "位号",
            "D119,D120",
            [ChangeRow("1.06.02.000194", "D", "", "2", "D119,D120", "/", "/", "/", "/", "/")],
        ),
        ChangeGroup(
            "add-new-locator",
            "位号",
            "DP20",
            [ChangeRow("/", "/", "/", "/", "/", "1.06.02.000194", "D", "", "1", "DP20")],
        ),
    ]
    materials = {
        "1.06.02.000194": [MaterialInfo("1.06.02.000194", "二极管", "D", "新描述", "SOD", "贴片料")]
    }

    process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=groups,
        materials=materials,
        change_record_people=ChangeRecordPeople("制表人", "审核人", "批准人"),
        change_record_date=date(2026, 8, 26),
    )

    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    rows = [
        (ws.cell(row=row, column=3).value, ws.cell(row=row, column=6).value, ws.cell(row=row, column=7).value)
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=3).value
    ]
    assert rows == [("1.06.02.000194", 1, "DP20")]
    assert wb[CHANGE_RECORD_SHEET_NAME].cell(row=3, column=3).value == (
        "添加物料DP20=1.06.02.000194；删除D119,D120=1.06.02.000194"
    )


def test_process_bom_workbook_does_not_duplicate_already_changed_locator_material(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST20-A320-3.1-003.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_partial_locator_change_notice(change_notice)
    _save_already_changed_locator_bom(bom)
    _save_material_master_with_bms_package_and_blank_package(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 8, 5),
    )

    assert result.modified_count == 0
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    matching_rows = [
        row_index
        for row_index in range(2, ws.max_row + 1)
        if ws.cell(row=row_index, column=3).value == "1.06.02.000194"
        and ws.cell(row=row_index, column=7).value == "DP7"
    ]
    assert matching_rows == [2]
    record = wb[CHANGE_RECORD_SHEET_NAME]
    assert record.cell(row=3, column=1).value in (None, "")
    assert record.cell(row=3, column=3).value in (None, "")


def test_process_bom_workbook_updates_matched_locator_rows_on_product_sheet(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST21-1870-3.1-001.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_partial_locator_change_notice(change_notice)
    _save_product_sheet_locator_bom(bom)
    _save_material_master_with_bms_package_and_blank_package(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.abnormal_count == 0
    assert result.modified_count == 3
    wb = load_workbook(output, data_only=True)
    ws = wb["P16S120A-11870-3.1"]
    assert ws.cell(row=2, column=3).value == "1.06.02.000194"
    assert ws.cell(row=2, column=6).value == 1
    assert ws.cell(row=2, column=7).value == "DP7"
    assert ws.cell(row=2, column=8).value == "TVS管 130V 144~159V 3000W 单向 SMC"
    assert ws.cell(row=3, column=3).value in (None, "")


def test_process_bom_workbook_treats_two_sheet_model_sheet_as_total(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST21-1870-3.1-001.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_add_tvs_change_notice(change_notice)
    _save_two_sheet_model_bom(bom)
    _save_material_master_with_bms_package_and_blank_package(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["P16S120A-11870-3.1"]
    assert _sheet_has_material_code(wb, "P16S120A-11870-3.1", "1.06.02.000194")
    assert any(
        ws.cell(row=row_index, column=7).value == "DP7"
        for row_index in range(2, ws.max_row + 1)
        if ws.cell(row=row_index, column=3).value == "1.06.02.000194"
    )


def test_sheet_segment_treats_single_or_first_bom_sheet_as_total() -> None:
    assert _sheet_segment("SMT-50448-1.0", ["SMT-50448-1.0"]) == "total"

    sheet_titles = ["BOM-SST25-0448-1.0", "P4S100A-50448-1.0", "DIP-50448-1.0", "SMT-50448-1.0", "变更记录"]

    assert _sheet_segment("BOM-SST25-0448-1.0", sheet_titles) == "total"
    assert _sheet_segment("SMT-50448-1.0", sheet_titles) == "smt"
    assert _sheet_segment("DIP-50448-1.0", sheet_titles) == "dip"
    assert _sheet_segment("P4S100A-50448-1.0", sheet_titles) == "product"
    assert _sheet_segment("变更记录", sheet_titles) == "other"


def test_parse_change_notice_accepts_single_table_addition_format(tmp_path: Path) -> None:
    change_notice = tmp_path / "single-table-change.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "工程变更单"
    ws.append(["工程变更通知单"])
    ws.append(["变更内容", "归属", "变更内容", "位号", "用量", "规格型号", "物料代码"])
    ws.append([None, "保护板", "增加", None, 0.0167, "沛城外箱665*480*220mm", "4.03.08.000001"])
    ws.append([None, "保护板", "增加", None, 1, "防静电气泡袋", "4.04.01.000009"])
    ws.append(["变更实施"])
    wb.save(change_notice)

    groups = parse_change_notice(change_notice)

    assert [(group.locator_type, group.locator_value) for group in groups] == [
        ("物料编码", "4.03.08.000001"),
        ("物料编码", "4.04.01.000009"),
    ]
    assert [(row.after_code, row.after_quantity, row.after_locator) for group in groups for row in group.rows] == [
        ("4.03.08.000001", "0.0167", "/"),
        ("4.04.01.000009", "1", "/"),
    ]


def test_parse_change_notice_accepts_single_table_with_blank_quantity_header(tmp_path: Path) -> None:
    change_notice = tmp_path / "blank-quantity-header-change.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "工程变更单"
    ws.append(["工程变更通知单"])
    ws.append(["变更内容", "归属", "变更内容", "位号", None, "规格型号", "物料代码"])
    ws.append([None, "保护板", "取消", "QP7,QP9", 2, "旧物料", "1.08.02.000121"])
    ws.append([None, "保护板", "增加", "QP7,QP9", 2, "新物料", "1.08.02.000130"])
    wb.save(change_notice)

    groups = parse_change_notice(change_notice)

    assert [(group.locator_type, group.locator_value) for group in groups] == [
        ("位号", "QP7,QP9"),
    ]
    assert [(row.before_code, row.after_code, row.before_quantity, row.after_quantity) for row in groups[0].rows] == [
        ("1.08.02.000121", "/", "2", "/"),
        ("/", "1.08.02.000130", "/", "2"),
    ]


def test_parse_change_notice_accepts_two_sided_table_without_description_column(tmp_path: Path) -> None:
    change_notice = tmp_path / "packaging-change.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "用量", "位号"]
    for column, header in enumerate(headers, start=2):
        ws.cell(row=2, column=column).value = header
        ws.cell(row=2, column=column + 4).value = header
    ws.append(["硬件类", "/", "/", "/", "/", "4.03.01.000038", "纸箱", "1/60", "/"])
    ws.append([None, "/", "/", "/", "/", "4.04.01.000029", "气泡袋", "1", "/"])
    wb.save(change_notice)

    groups = parse_change_notice(change_notice)

    assert [row.after_code for group in groups for row in group.rows] == [
        "4.03.01.000038",
        "4.04.01.000029",
    ]
    assert [row.after_quantity for group in groups for row in group.rows] == ["1/60", "1"]


def test_parse_change_notice_keeps_each_after_quantity_in_multi_addition_group(tmp_path: Path) -> None:
    change_notice = tmp_path / "packaging-change.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    rows = [
        ["4.03.01.000078", "旧纸箱", "旧描述", "1/16", "/", "4.03.08.000007", "外箱", "外箱描述", "1/16", "/"],
        ["/", "/", "/", "/", "/", "4.03.09.000008", "刀卡", "刀卡描述", "1/8", "/"],
        ["/", "/", "/", "/", "/", "4.03.05.000007", "平卡", "平卡描述", "3/16", "/"],
        ["/", "/", "/", "/", "/", "4.05.01.000033", "珍珠棉", "珍珠棉描述", "1/4", "/"],
    ]
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=column).value = value
    wb.save(change_notice)

    groups = parse_change_notice(change_notice)

    assert [row.after_quantity for row in groups[0].rows] == ["1/16", "1/8", "3/16", "1/4"]


def test_process_bom_workbook_matches_single_table_locator_regardless_of_order(tmp_path: Path) -> None:
    change_notice = tmp_path / "single-table-replacement.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["变更内容", "归属", "变更内容", "位号", "用量", "规格型号", "物料代码"])
    ws.append([None, "保护板", "取消", "D14,D57,D58,D59", 4, "TVS 120V", "1.06.02.000198"])
    ws.append([None, "保护板", "增加", "D14,D57,D58,D59", 4, "TVS 133V", "1.06.02.000226"])
    wb.save(change_notice)
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.06.02.000198", "二极管", "5.0SMDJ110CA", 4, "D59,D58,D57,D14", "旧 TVS", "SMC"])
    wb.save(bom)
    wb = Workbook()
    ws = wb.active
    ws.append(["", "", "", "物料代码", "物料描述", "规格型号", "", "", "BMS规则描述", "BMS封装", "一级分类"])
    ws.append(["", "", "", "1.06.02.000226", "二极管", "5.0SMDJ120CA", "", "", "新 TVS", "SMC", "贴片料"])
    wb.save(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.cell(row=2, column=3).value == "1.06.02.000226"
    assert set(str(ws.cell(row=2, column=7).value).split(",")) == {"D14", "D57", "D58", "D59"}


def test_process_bom_workbook_deletes_only_matched_row_without_clearing_remaining_rows(tmp_path: Path) -> None:
    change_notice = tmp_path / "delete-only-change.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["变更内容", "归属", "变更内容", "位号", "用量", "规格型号", "物料代码"])
    ws.append([None, "保护板", "取消", "R1", 1, "旧电阻", "1.02.01.000001"])
    wb.save(change_notice)
    _save_material_master(master)
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000001", "电阻", "旧电阻", 1, "R1", "旧描述", "R0603"])
    ws.append([2, "贴片料", "1.02.01.000002", "电阻", "保留电阻", 1, "R2", "保留描述", "R0603"])
    ws.append([3, "贴片料", "1.02.01.000003", "电阻", "保留电阻", 1, "R3", "保留描述", "R0603"])
    wb.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.max_row == 3
    assert [ws.cell(row=row, column=3).value for row in range(2, 4)] == ["1.02.01.000002", "1.02.01.000003"]


def test_process_bom_workbook_deletes_only_target_locator_without_inserting_zero_rows(tmp_path: Path) -> None:
    change_notice = tmp_path / "partial-delete-change.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["变更内容", "归属", "变更内容", "位号", "用量", "规格型号", "物料代码"])
    ws.append([None, "保护板", "取消", "R1", 1, "旧电阻", "1.02.01.000001"])
    wb.save(change_notice)
    _save_material_master(master)
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000001", "电阻", "旧电阻", 2, "R1,R2", "旧描述", "R0603"])
    ws.append([2, "贴片料", "1.02.01.000002", "电阻", "保留电阻", 1, "R3", "保留描述", "R0603"])
    wb.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.max_row == 3
    assert ws.cell(row=2, column=3).value == "1.02.01.000001"
    assert ws.cell(row=2, column=6).value == 1
    assert ws.cell(row=2, column=7).value == "R2"
    assert ws.cell(row=3, column=3).value == "1.02.01.000002"


def test_parse_change_notice_treats_single_table_cancel_as_delete_by_code_and_locator(tmp_path: Path) -> None:
    change_notice = tmp_path / "single-table-replacement.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "工程变更单"
    ws.append(["变更内容", "归属", "变更内容", "位号", "用量", "规格型号", "物料代码"])
    ws.append([None, "保护板", "取消", "D14,D57,D58,D59", 4, "TVS 120V", "1.06.02.000198"])
    ws.append([None, "保护板", "增加", "D14,D57,D58,D59", 4, "TVS 133V", "1.06.02.000226"])
    wb.save(change_notice)

    groups = parse_change_notice(change_notice)

    assert [(group.locator_type, group.locator_value) for group in groups] == [
        ("位号", "D14,D57,D58,D59"),
    ]
    assert [(row.before_code, row.after_code) for row in groups[0].rows] == [
        ("1.06.02.000198", "/"),
        ("/", "1.06.02.000226"),
    ]


def test_parse_change_notice_keeps_unchanged_quantity_and_locator_markers(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    _save_unchanged_quantity_locator_change_notice(change_notice)

    groups = parse_change_notice(change_notice)

    assert [group.locator_type for group in groups] == ["物料编码", "物料编码"]
    assert [group.locator_value for group in groups] == ["1.08.02.000130", "1.08.02.000101"]
    assert [group.rows[0].after_code for group in groups] == ["1.08.02.000179", "1.08.02.000179"]
    assert [group.rows[0].after_quantity for group in groups] == ["用量不变", "用量不变"]
    assert [group.rows[0].after_locator for group in groups] == ["位号不变", "位号不变"]


def test_process_bom_workbook_replaces_material_and_preserves_unchanged_quantity_locator(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_unchanged_quantity_locator_change_notice(change_notice)
    _save_unchanged_quantity_locator_bom(bom)
    _save_unchanged_quantity_locator_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 4
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.cell(row=2, column=2).value == "贴片料"
    assert ws.cell(row=2, column=3).value == "1.08.02.000179"
    assert ws.cell(row=2, column=5).value == "CC360N10HS"
    assert ws.cell(row=2, column=6).value == 2
    assert ws.cell(row=2, column=7).value == "Q1,Q2"
    assert ws.cell(row=2, column=8).value == "MOS CC360N10HS 100V"
    assert ws.cell(row=2, column=9).value == "TO-263"
    assert ws.cell(row=3, column=3).value == "1.08.02.000179"
    assert ws.cell(row=3, column=2).value == "贴片料"
    assert ws.cell(row=3, column=6).value == 4
    assert ws.cell(row=3, column=7).value == "Q3,Q4,Q5,Q6"


def test_validate_processed_bom_workbook_reports_old_code_left_in_processed_sheet(tmp_path: Path) -> None:
    notice = tmp_path / "notice.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "bom.xlsx"
    _save_change_notice(notice)
    _save_material_master(master)
    _save_bom(bom)

    groups = parse_change_notice(notice)
    issues = validate_processed_bom_workbook(
        bom_path=bom,
        groups=groups,
        successful_report_rows=[
            ReportRow(
                sheet_name="总表",
                locator_type="位号",
                locator_value="R758",
                action="删除",
                before_code="1.02.01.000061",
                before_locator="R758",
                result="成功",
            )
        ],
    )

    assert any(
        issue.sheet_name == "总表"
        and issue.before_code == "1.02.01.000061"
        and issue.after_code == "1.02.01.000107"
        and issue.reason == "变更前物料仍存在"
        for issue in issues
    )


def test_process_bom_workbook_matches_unchanged_quantity_locator_by_code_only(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_unchanged_quantity_locator_change_notice(change_notice)
    _save_unchanged_quantity_locator_master(master)

    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append(
        [
            1,
            "贴片料",
            "1.08.02.000130",
            "MOS",
            "=VLOOKUP(C2,[1]Table!$C:$M,3,0)",
            2,
            "Q1,Q2",
            "=VLOOKUP(C2,[1]Table!$C:$M,6,0)",
            "TO-263",
        ]
    )
    wb.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.cell(row=2, column=3).value == "1.08.02.000179"
    assert ws.cell(row=2, column=5).value == "CC360N10HS"
    assert ws.cell(row=2, column=6).value == 2
    assert ws.cell(row=2, column=7).value == "Q1,Q2"
    assert ws.cell(row=2, column=8).value == "MOS CC360N10HS 100V"


def test_process_bom_workbook_ignores_stale_description_when_code_quantity_and_locator_match(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "变更单"
    ws["A1"] = "变更前描述"
    ws["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=2, column=index).value = header
        ws.cell(row=2, column=index + 5).value = header
    values = [
        "1.02.01.000665",
        "AC0805FR-0720K5L",
        "贴片厚膜电阻 20.5K ±1% 1/8W （车规级） R0805",
        "1",
        "位号不变",
        "1.02.01.000695",
        "ADE05K1001FT",
        "贴片厚膜电阻 1K ±1% 1/8W 车规抗硫化 R0805",
        "1",
        "位号不变",
    ]
    for column, value in enumerate(values, start=1):
        ws.cell(row=3, column=column).value = value
    wb.save(change_notice)

    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append(
        [
            1,
            "贴片料",
            "1.02.01.000665",
            "电阻",
            "AC0805FR-0720K5L",
            1,
            "R763",
            "贴片厚膜电阻 20.5K ±1% 1/8W （样品） R0805",
            "R0805",
        ]
    )
    wb.save(bom)

    wb = Workbook()
    ws = wb.active
    ws.title = "物料主数据"
    ws["D1"] = "物料代码"
    ws["E1"] = "物料描述"
    ws["F1"] = "型号"
    ws["I1"] = "BMS规则描述"
    ws["J1"] = "BMS封装"
    ws["K1"] = "一级分类"
    ws.cell(row=2, column=4).value = "1.02.01.000695"
    ws.cell(row=2, column=5).value = "电阻"
    ws.cell(row=2, column=6).value = "ADE05K1001FT"
    ws.cell(row=2, column=9).value = "贴片厚膜电阻 1K ±1% 1/8W 车规抗硫化 R0805"
    ws.cell(row=2, column=10).value = "R0805"
    ws.cell(row=2, column=11).value = "贴片料"
    wb.save(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.cell(row=2, column=3).value == "1.02.01.000695"
    assert ws.cell(row=2, column=5).value == "ADE05K1001FT"
    assert ws.cell(row=2, column=6).value == 1
    assert ws.cell(row=2, column=7).value == "R763"
    assert ws.cell(row=2, column=8).value == "贴片厚膜电阻 1K ±1% 1/8W 车规抗硫化 R0805"


def test_process_bom_workbook_replaces_old_material_when_target_already_exists(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_unchanged_quantity_locator_change_notice(change_notice)
    _save_unchanged_quantity_locator_bom_with_existing_target(bom)
    _save_unchanged_quantity_locator_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 4
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.max_row == 4
    assert [ws.cell(row=row_index, column=3).value for row_index in range(2, 5)] == [
        "1.08.02.000179",
        "1.08.02.000179",
        "1.08.02.000179",
    ]
    assert ws.cell(row=2, column=6).value == 2
    assert ws.cell(row=2, column=7).value == "Q1,Q2"
    assert ws.cell(row=3, column=6).value == 1
    assert ws.cell(row=3, column=7).value == "Q7"
    assert ws.cell(row=4, column=6).value == 4
    assert ws.cell(row=4, column=7).value == "Q3,Q4,Q5,Q6"


def test_process_bom_workbook_merges_same_locator_replacement_group_to_one_target(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_same_locator_replacement_group_change_notice(change_notice)
    _save_same_locator_replacement_group_bom(bom)
    _save_unchanged_quantity_locator_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 4
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.max_row == 3
    assert ws.cell(row=2, column=3).value == "1.08.02.000179"
    assert ws.cell(row=2, column=5).value == "CC360N10HS"
    assert ws.cell(row=2, column=6).value == 34
    assert ws.cell(row=2, column=7).value == "QP48,QP43,QP44,QP45"
    assert ws.cell(row=3, column=3).value == "1.08.04.000008"
    report = wb[REPORT_SHEET_NAME]
    actions = [report.cell(row=row_index, column=4).value for row_index in range(2, report.max_row + 1)]
    assert actions == ["删除", "删除", "删除", "添加"]


def test_process_bom_workbook_merges_existing_target_in_same_locator_group(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_same_locator_replacement_group_change_notice(change_notice)
    _save_same_locator_replacement_group_bom_with_existing_target(bom)
    _save_unchanged_quantity_locator_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 4
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    material_codes = [ws.cell(row=row_index, column=3).value for row_index in range(2, ws.max_row + 1)]
    assert material_codes == ["1.08.02.000179"]
    assert ws.cell(row=2, column=6).value == 6
    assert ws.cell(row=2, column=7).value == "QP50,QP51,QP39,QP27,QP25,QP2"


def test_process_bom_workbook_change_record_uses_actual_locator_for_unchanged_locator_group(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST26-0212-1.0-001.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_same_locator_replacement_group_change_notice(change_notice)
    _save_same_locator_replacement_group_bom_with_existing_target(bom)
    _save_unchanged_quantity_locator_master(master)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 8, 6),
    )

    wb = load_workbook(output, data_only=True)
    record = wb[CHANGE_RECORD_SHEET_NAME]
    locator = "QP50,QP51,QP39,QP27,QP25,QP2"
    assert record.cell(row=3, column=3).value == (
        f"添加物料{locator}=1.08.02.000179；"
        f"删除{locator}=1.08.02.000130、{locator}=1.08.02.000152、{locator}=1.08.02.000176"
    )


def test_process_bom_workbook_inserts_add_only_material_without_locator(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-BCU-0602E-1.0-A001.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_add_only_change_notice(change_notice)
    _save_add_only_bom(bom)
    _save_shifted_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_record_people=ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        ),
        change_record_date=date(2026, 8, 3),
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.cell(row=6, column=1).value == 4
    assert ws.cell(row=7, column=1).value == 5
    assert ws.cell(row=8, column=1).value == 6
    assert ws.cell(row=8, column=2).value == "辅料"
    assert ws.cell(row=8, column=3).value == "1.00.02.000011"
    assert ws.cell(row=8, column=4).value == "License"
    assert ws.cell(row=8, column=5).value == "WBMSHV-License 三颗小草"
    assert ws.cell(row=8, column=6).value == 1
    assert ws.cell(row=8, column=7).value in (None, "")
    record = wb[CHANGE_RECORD_SHEET_NAME]
    assert record.cell(row=3, column=3).value == "添加物料1.00.02.000011"


def test_process_bom_workbook_inserts_packaging_at_bottom_of_bom_data(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000001", "电阻", "R", 1, "R1", "", ""])
    ws.append([2, "辅料", "4.01.02.000002", "黄胶", "K-1668", "替代料", None, "", ""])
    ws.append([3, "辅料", "4.02.01.000001", "三防漆", "PU1030", 6, None, "", ""])
    ws.append([4, "辅料", "4.05.01.000002", "珍珠棉", "430*86*15", 0.1429, None, "", ""])
    ws.append([5, "插件料", "1.18.02.000001", "排针", "P", 1, "J1", "", ""])
    auxiliary_border = Border(bottom=Side(style="thin"))
    next_section_border = Border(bottom=Side(style="dashed"))
    for column in range(1, 10):
        ws.cell(row=5, column=column).border = auxiliary_border
        ws.cell(row=6, column=column).border = next_section_border
    book.save(bom)

    groups = [
        ChangeGroup(
            "add-packaging",
            "物料编码",
            "4.03.01.000038",
            [ChangeRow("/", "/", "/", "/", "/", "4.03.01.000038", "纸箱", "", "1/60", "/")],
        )
    ]
    materials = {
        "4.03.01.000038": [MaterialInfo("4.03.01.000038", "纸箱", "纸箱", "", "", "包装料")]
    }

    process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=groups,
        materials=materials,
    )

    ws = load_workbook(output, data_only=True)["总表"]
    assert ws.cell(row=7, column=3).value == "4.03.01.000038"
    assert ws.cell(row=6, column=3).value == "1.18.02.000001"
    assert ws.cell(row=7, column=3).border.bottom.style == "dashed"


def test_process_bom_workbook_inserts_auxiliary_at_bottom_of_bom_data(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000001", "电阻", "R", 1, "R1", "", ""])
    ws.append([2, "辅料", "4.03.02.000001", "旧辅料", "A", 1, None, "", ""])
    ws.append([3, "插件料", "1.18.02.000001", "排针", "P", 1, "J1", "", ""])
    book.save(bom)

    group = ChangeGroup(
        "add-auxiliary",
        "物料编码",
        "1.00.02.000005",
        [ChangeRow("/", "/", "/", "/", "/", "1.00.02.000005", "License", "", "1", "/")],
    )
    materials = {
        "1.00.02.000005": [
            MaterialInfo("1.00.02.000005", "License", "BBMSLV-License", "License描述", "", "辅料")
        ]
    }

    process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=[group],
        materials=materials,
    )

    ws = load_workbook(output, data_only=True)["总表"]
    assert ws.cell(row=5, column=3).value == "1.00.02.000005"
    assert ws.cell(row=4, column=3).value == "1.18.02.000001"


def test_process_bom_workbook_merges_addition_into_existing_material_code(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_add_only_change_notice_for_existing_material(change_notice)
    _save_bom_with_existing_add_material(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    matching_rows = [
        row_index
        for row_index in range(2, ws.max_row + 1)
        if ws.cell(row=row_index, column=3).value == "1.03.01.000009"
    ]
    assert matching_rows == [3]
    assert ws.cell(row=3, column=6).value == 45
    assert ws.cell(row=3, column=7).value == "CJ6,CJ11,CJ15,CJ10,C522,C520,C521,C523"


def test_process_bom_workbook_consolidates_duplicate_added_material_code(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_add_only_change_notice_for_existing_material(change_notice)
    _save_bom_with_existing_add_material(bom)
    _save_material_master(master)
    workbook = load_workbook(bom)
    ws = workbook["总表"]
    ws.append([4, "贴片料", "1.03.01.000009", "电容", "0402B104K500CT", 4, "C522,C520,C521,C523", "贴片陶瓷电容", "C0402"])
    workbook.save(bom)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    matching_rows = [
        row_index
        for row_index in range(2, ws.max_row + 1)
        if ws.cell(row=row_index, column=3).value == "1.03.01.000009"
    ]
    assert len(matching_rows) == 1
    row_index = matching_rows[0]
    assert ws.cell(row=row_index, column=6).value == 45
    assert ws.cell(row=row_index, column=7).value == "CJ6,CJ11,CJ15,CJ10,C522,C520,C521,C523"


def test_process_bom_workbook_merges_replacement_into_existing_target_material(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_replacement_change_notice_with_existing_target(change_notice)
    _save_bom_with_existing_replacement_target(bom)
    _save_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert [ws.cell(row=row, column=3).value for row in range(2, ws.max_row + 1)] == ["1.03.01.000009"]
    assert ws.cell(row=2, column=6).value == 45
    assert ws.cell(row=2, column=7).value == "C522,C520,C521,C523,CJ6,CJ11"


def test_process_bom_workbook_removes_changed_locators_from_old_alternative_when_target_already_has_them(
    tmp_path: Path,
) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([
        44,
        "贴片料",
        "1.03.01.000048",
        "电容",
        "0402B103K500CT",
        7,
        "CM3,CM2,CA3,CA1,CM11,CJ1,C7",
        "主料",
        "C0402",
    ])
    ws.append([
        None,
        "贴片料",
        "1.03.01.000047",
        "电容",
        "CC0402KRX7R9BB103",
        "替代料",
        "CM3,CM2,CA3,CA1,CM11,CJ1,C7",
        "旧替代料",
        "C0402",
    ])
    ws.append([
        45,
        "贴片料",
        "1.03.01.000482",
        "电容",
        "0402X474K500CT",
        3,
        "CM2,CM3,CM11",
        "目标料已存在",
        "C0402",
    ])
    book.save(bom)

    group = ChangeGroup(
        "replace-cm2-cm3-cm11",
        "位号",
        "CM2,CM3,CM11",
        [
            ChangeRow(
                "1.03.01.000047",
                "CC0402KRX7R9BB103",
                "",
                "3",
                "CM2,CM3,CM11",
                "1.03.01.000482",
                "TCC0402X7R474K500AT",
                "",
                "3",
                "CM2,CM3,CM11",
            )
        ],
    )
    materials = {
        "1.03.01.000482": [MaterialInfo("1.03.01.000482", "电容", "TCC0402X7R474K500AT", "目标料", "C0402", "贴片料")]
    }

    result = process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=[group],
        materials=materials,
    )

    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    rows = {
        ws.cell(row=row, column=3).value: (ws.cell(row=row, column=6).value, ws.cell(row=row, column=7).value)
        for row in range(2, ws.max_row + 1)
    }
    assert rows["1.03.01.000048"] == (4, "CA3,CA1,CJ1,C7")
    assert rows["1.03.01.000047"] == ("替代料", "CA3,CA1,CJ1,C7")
    assert rows["1.03.01.000482"] == (3, "CM2,CM3,CM11")


def test_process_bom_workbook_removes_changed_locators_from_old_alternative_when_adding_target(
    tmp_path: Path,
) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([44, "贴片料", "1.03.01.000048", "电容", "0402B103K500CT", 7, "CM3,CM2,CA3,CA1,CM11,CJ1,C7", "主料", "C0402"])
    ws.append([
        None,
        "贴片料",
        "1.03.01.000047",
        "电容",
        "CC0402KRX7R9BB103",
        "替代料",
        "CM3,CM2,CA3,CA1,CM11,CJ1,C7",
        "旧替代料",
        "C0402",
    ])
    book.save(bom)

    group = ChangeGroup(
        "replace-cm2-cm3-cm11",
        "位号",
        "CM2,CM3,CM11",
        [
            ChangeRow(
                "1.03.01.000047",
                "CC0402KRX7R9BB103",
                "",
                "3",
                "CM2,CM3,CM11",
                "1.03.01.000482",
                "TCC0402X7R474K500AT",
                "",
                "3",
                "CM2,CM3,CM11",
            )
        ],
    )
    materials = {
        "1.03.01.000482": [MaterialInfo("1.03.01.000482", "电容", "TCC0402X7R474K500AT", "目标料", "C0402", "贴片料")]
    }

    result = process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=[group],
        materials=materials,
    )

    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    rows = {
        ws.cell(row=row, column=3).value: (ws.cell(row=row, column=6).value, ws.cell(row=row, column=7).value)
        for row in range(2, ws.max_row + 1)
    }
    assert rows["1.03.01.000048"] == (4, "CA3,CA1,CJ1,C7")
    assert rows["1.03.01.000047"] == ("替代料", "CA3,CA1,CJ1,C7")
    assert rows["1.03.01.000482"] == (3, "CM2,CM3,CM11")


def test_process_bom_workbook_adds_alternative_using_existing_primary_locator_and_quantity(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料编码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.08.05.000014", "MOS", "KJ20P10K", 1, "Q1", "主料", "TO252"])
    workbook.save(bom)
    workbook.close()
    _save_material_master(master)

    groups = [ChangeGroup(
        "alternative-add-Q1",
        "物料编码",
        "1.08.05.000014",
        [ChangeRow(
            "1.08.05.000014", "KJ20P10K", "", "用量不变", "位号不变",
            "1.08.05.000021", "DP700P10DTL", "", "替代料", "位号不变",
        )],
    )]
    materials = {
        "1.08.05.000021": [MaterialInfo("1.08.05.000021", "MOS", "DP700P10DTL", "替代料", "TO252", "贴片料")],
    }

    result = process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=master,
        output_path=output,
        change_groups=groups,
        materials=materials,
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    try:
        ws = wb["总表"]
        assert ws.max_row == 3
        assert [ws.cell(row=3, column=index).value for index in (3, 6, 7)] == [
            "1.08.05.000021", "替代料", "Q1"
        ]
        assert ws.cell(row=2, column=6).value == 1
        assert ws.cell(row=2, column=7).value == "Q1"
    finally:
        wb.close()


def test_process_bom_workbook_removes_whole_alternative_group_when_target_already_exists(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_partial_locator_change_notice(change_notice)
    _save_material_master_with_bms_package_and_blank_package(master)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.06.02.000155", "二极管", "SMAJ30A", 1, "DP7", "旧主料", "SMC"])
    ws.append([None, "贴片料", "1.06.02.000108", "二极管", "SMAJ30A", "替代料", "DP7", "旧替代1", "SMC"])
    ws.append([None, "贴片料", "1.06.02.000020", "二极管", "SMAJ30A", "替代料", "DP7", "旧替代2", "SMC"])
    ws.append([2, "贴片料", "1.06.02.000194", "二极管", "SMDJ130A", 3, "DP8,DP9,DP10", "新TVS", "SMC"])
    workbook.save(bom)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert [ws.cell(row=row, column=3).value for row in range(2, ws.max_row + 1)] == ["1.06.02.000194"]
    assert ws.cell(row=2, column=6).value == 4
    assert ws.cell(row=2, column=7).value == "DP8,DP9,DP10,DP7"


def test_process_bom_workbook_merges_same_target_locator_changes_without_summing_quantity(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_material_master_with_bms_package_and_blank_package(master)

    notice_book = Workbook()
    notice_sheet = notice_book.active
    notice_sheet.append(["变更前描述", None, None, None, None, "变更后描述"])
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for column, header in enumerate(headers, start=1):
        notice_sheet.cell(2, column).value = header
        notice_sheet.cell(2, column + 5).value = header
    notice_sheet.append(["1.06.02.000150", "SMAJ30A", "旧DP15", 1, "DP15", "1.06.02.000194", "SMDJ130A", "新TVS", 2, "DP15,DP7"])
    notice_sheet.append(["1.06.02.000020", "SMAJ30A", "旧DP7", 1, "DP7", "1.06.02.000194", "SMDJ130A", "新TVS", 2, "DP15,DP7"])
    notice_book.save(change_notice)

    bom_book = Workbook()
    ws = bom_book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.06.02.000150", "二极管", "SMAJ30A", 1, "DP15", "旧DP15", "SMC"])
    ws.append([2, "贴片料", "1.06.02.000020", "二极管", "SMAJ30A", 1, "DP7", "旧DP7", "SMC"])
    bom_book.save(bom)

    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert [ws.cell(row=row, column=3).value for row in range(2, ws.max_row + 1)] == ["1.06.02.000194"]
    assert ws.cell(row=2, column=6).value == 2
    assert ws.cell(row=2, column=7).value == "DP15,DP7"


def test_process_bom_workbook_moves_only_ecn_locators_from_alternative_group(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([
        42,
        "贴片料",
        "1.03.01.000048",
        "电容",
        "0402B103K500CT",
        6,
        "CM3,CM2,CA3,CA1,CM11,C37",
        "贴片陶瓷电容 10nF",
        "C0402",
    ])
    ws.append([
        None,
        "贴片料",
        "1.03.01.000047",
        "电容",
        "CC0402KRX7R9BB103",
        "替代料",
        "CM3,CM2,CA3,CA1,CM11,C37",
        "贴片陶瓷电容 10nF",
        "C0402",
    ])
    ws.append([43, "贴片料", "1.03.01.000053", "电容", "0603B105K500CT", 1, "CM22", "", "C0603"])
    book.save(bom)

    group = ChangeGroup(
        "replace-cap-locators",
        "位号",
        "CM2,CM3,CM11",
        [
            ChangeRow(
                "1.03.01.000047",
                "CC0402KRX7R9BB103",
                "",
                "3",
                "CM2,CM3,CM11",
                "1.03.01.000482",
                "0402X474K500CT",
                "",
                "3",
                "CM2,CM3,CM11",
            ),
            ChangeRow("1.03.01.000048", "0402B103K500CT", "", "替代料", "CM2,CM3,CM11", "/", "/", "/", "/", "/"),
            ChangeRow("1.03.01.000411", "TCC0402X7R103K500AT", "", "替代料", "CM2,CM3,CM11", "/", "/", "/", "/", "/"),
            ChangeRow("1.03.01.000365", "0402B103K500NT", "", "替代料", "CM2,CM3,CM11", "/", "/", "/", "/", "/"),
        ],
    )
    materials = {
        "1.03.01.000482": [
            MaterialInfo("1.03.01.000482", "电容", "0402X474K500CT", "贴片陶瓷电容 470nF", "C0402", "贴片料")
        ]
    }

    result = process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=[group],
        materials=materials,
    )

    assert result.abnormal_count == 0
    ws = load_workbook(output, data_only=True)["总表"]
    rows = {
        ws.cell(row=row, column=3).value: (ws.cell(row=row, column=6).value, ws.cell(row=row, column=7).value)
        for row in range(2, ws.max_row + 1)
    }
    assert rows["1.03.01.000048"] == (3, "CA3,CA1,C37")
    assert rows["1.03.01.000047"] == ("替代料", "CA3,CA1,C37")
    assert rows["1.03.01.000482"] == (3, "CM2,CM3,CM11")


def test_process_bom_workbook_preserves_multiple_additions_at_the_same_insert_row(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"

    notice_book = Workbook()
    notice_sheet = notice_book.active
    notice_sheet["A1"] = "变更前描述"
    notice_sheet["F1"] = "变更后描述"
    headers = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
    for column, header in enumerate(headers, start=1):
        notice_sheet.cell(2, column).value = header
        notice_sheet.cell(2, column + 5).value = header
    notice_sheet.append(["/", "/", "/", "/", "/", "1.18.02.000060", "P20", "", 2, "J4,J193"])
    notice_sheet.append(["/", "/", "/", "/", "/", "1.18.02.000061", "P21", "", 1, "J194"])
    notice_book.save(change_notice)

    bom_book = Workbook()
    ws = bom_book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "插件料", "1.18.02.000034", "排针", "P6", 1, "J3", "旧排针", "DIP2X6"])
    ws.append([2, "辅料", "4.01.01.000001", "胶", "K-704", 1, None, "胶", None])
    bom_book.save(bom)

    materials = {
        "1.18.02.000060": [MaterialInfo("1.18.02.000060", "排针", "P20", "", "DIP2X10", "插件料")],
        "1.18.02.000061": [MaterialInfo("1.18.02.000061", "排针", "P21", "", "DIP2X11", "插件料")],
    }
    process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=tmp_path / "unused.xlsx",
        output_path=output,
        materials=materials,
    )

    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    material_codes = [ws.cell(row=row, column=3).value for row in range(2, ws.max_row + 1)]
    assert "1.18.02.000060" in material_codes
    assert "1.18.02.000061" in material_codes


def test_process_bom_workbook_does_not_delete_an_earlier_addition_with_a_later_replacement(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "插件料", "1.18.02.000034", "排针", "P6", 2, "J4,J3", "旧排针", "DIP2X6"])
    ws.append([2, "插件料", "1.21.01.000272", "插座", "OLD", 1, "J173", "旧插座", "CNN4"])
    book.save(bom)

    groups = [
        ChangeGroup("cancel-j4", "位号", "J4", [ChangeRow("1.18.02.000034", "P6", "", "1", "J4", "/", "/", "/", "/", "/")]),
        ChangeGroup("add-j4", "位号", "J4,J193", [ChangeRow("/", "/", "/", "/", "/", "1.18.02.000060", "P20", "", "2", "J4,J193")]),
        ChangeGroup("replace-j173", "位号", "J173", [ChangeRow("1.21.01.000272", "OLD", "", "1", "J173", "1.21.01.000542", "NEW", "", "1", "J173")]),
    ]
    materials = {
        "1.18.02.000060": [MaterialInfo("1.18.02.000060", "排针", "P20", "", "DIP2X10", "插件料")],
        "1.21.01.000542": [MaterialInfo("1.21.01.000542", "插座", "NEW", "", "CNN4", "插件料")],
    }
    process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=groups,
        materials=materials,
    )

    ws = load_workbook(output, data_only=True)["总表"]
    rows = {ws.cell(row=row, column=3).value: (ws.cell(row=row, column=6).value, ws.cell(row=row, column=7).value) for row in range(2, ws.max_row + 1)}
    assert rows["1.18.02.000060"] == (2, "J4,J193")
    assert rows["1.21.01.000542"] == (1, "J173")
    assert "1.21.01.000272" not in rows


def test_process_bom_workbook_merges_addition_into_existing_alternative_material_group(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "贴片料", "1.02.01.000836", "电阻", "0805", 1, "R1551", "主料", "R0805"])
    ws.append([None, "贴片料", "1.02.01.000017", "电阻", "0805", "替代料", "R1551", "目标料", "R0805"])
    ws.append([2, "贴片料", "1.02.01.000024", "电阻", "0805", 1, "R94", "待替换", "R0805"])
    book.save(bom)

    group = ChangeGroup(
        "replace-r94",
        "位号",
        "R94",
        [ChangeRow("1.02.01.000024", "0805", "", "1", "R94", "1.02.01.000017", "0805", "", "1", "R94")],
    )
    materials = {"1.02.01.000017": [MaterialInfo("1.02.01.000017", "电阻", "0805", "目标料", "R0805", "贴片料")]}
    process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=[group],
        materials=materials,
    )

    ws = load_workbook(output, data_only=True)["总表"]
    assert [ws.cell(row=row, column=3).value for row in range(2, ws.max_row + 1)] == ["1.02.01.000836", "1.02.01.000017"]
    assert ws.cell(row=2, column=6).value == 2
    assert ws.cell(row=2, column=7).value == "R1551,R94"
    assert ws.cell(row=3, column=7).value == "R1551,R94"


def test_process_bom_workbook_increments_existing_addition_without_locator(tmp_path: Path) -> None:
    bom = tmp_path / "bom.xlsx"
    output = tmp_path / "output.xlsx"
    book = Workbook()
    ws = book.active
    ws.title = "总表"
    ws.append(["序号", "物料类型", "物料代码", "名称", "规格型号", "用量", "位号", "规格描述", "封装"])
    ws.append([1, "辅料", "4.03.02.000029", "标签", "标签28*28", 1, None, "标签", None])
    book.save(bom)

    group = ChangeGroup(
        "add-label",
        "物料编码",
        "4.03.02.000029",
        [ChangeRow("/", "/", "/", "/", "/", "4.03.02.000029", "标签28*28", "", "1", "/")],
    )
    materials = {"4.03.02.000029": [MaterialInfo("4.03.02.000029", "标签", "标签28*28", "标签", "", "辅料")]}
    result = process_bom_workbook(
        change_notice_path=tmp_path / "unused.xlsx",
        bom_path=bom,
        material_master_path=tmp_path / "unused-master.xlsx",
        output_path=output,
        change_groups=[group],
        materials=materials,
    )

    ws = load_workbook(output, data_only=True)["总表"]
    assert result.modified_count == 1
    assert ws.cell(row=2, column=6).value == 2


def _sheet_has_material_code(wb, sheet_name: str, code: str) -> bool:
    return any(cell.value == code for row in wb[sheet_name].iter_rows() for cell in row)


def test_process_bom_workbook_adds_auxiliary_only_to_total_and_product_sheets(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-BCU-0602E-1.0-A001.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_add_only_change_notice(change_notice)
    _save_layered_add_only_bom(bom)
    _save_shifted_material_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    assert _sheet_has_material_code(wb, "总表", "1.00.02.000011")
    assert _sheet_has_material_code(wb, "P8S200A-61097-1.0", "1.00.02.000011")
    assert not _sheet_has_material_code(wb, "DIP-61097-1.0", "1.00.02.000011")
    assert not _sheet_has_material_code(wb, "SMT-61097-1.0", "1.00.02.000011")


def test_process_bom_workbook_adds_no_type_pcb_only_to_total_and_smt_sheets(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST26-1097-1.0-007.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_no_type_pcb_change_notice(change_notice)
    _save_layered_add_only_bom(bom)
    _save_no_type_pcb_master(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 2
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    assert _sheet_has_material_code(wb, "总表", "1.01.01.009999")
    assert _sheet_has_material_code(wb, "SMT-61097-1.0", "1.01.01.009999")
    assert not _sheet_has_material_code(wb, "P8S200A-61097-1.0", "1.01.01.009999")
    assert not _sheet_has_material_code(wb, "DIP-61097-1.0", "1.01.01.009999")


def test_process_bom_workbook_replaces_smt_material_in_total_and_smt_sheets(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST20-8331-3.0-001.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_partial_locator_change_notice(change_notice)
    _save_layered_tvs_replacement_bom(bom)
    _save_material_master_with_bms_package_and_blank_package(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    assert _sheet_has_material_code(wb, "总表", "1.06.02.000194")
    assert _sheet_has_material_code(wb, "SMT-08331-1.0", "1.06.02.000194")
    assert not _sheet_has_material_code(wb, "总表", "1.06.02.000020")
    assert not _sheet_has_material_code(wb, "SMT-08331-1.0", "1.06.02.000020")
    assert _sheet_has_material_code(wb, "P13S120A-08331-3.0", "1.06.02.000020")
    assert _sheet_has_material_code(wb, "DIP-08331-1.0", "1.06.02.000020")


def test_process_bom_workbook_fills_category_name_and_bms_package_aliases(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    bom = tmp_path / "BOM-SST24-1574-1.0-005.xlsx"
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    _save_add_tvs_change_notice(change_notice)
    _save_bom_with_material_category_and_name_headers(bom)
    _save_material_master_with_bms_package_and_blank_package(master)

    result = process_bom_workbook(
        change_notice_path=change_notice,
        bom_path=bom,
        material_master_path=master,
        output_path=output,
    )

    assert result.modified_count == 1
    assert result.abnormal_count == 0
    wb = load_workbook(output, data_only=True)
    ws = wb["总表"]
    assert ws.cell(row=4, column=2).value == "贴片料"
    assert ws.cell(row=4, column=3).value == "1.06.02.000194"
    assert ws.cell(row=4, column=4).value == "二极管"
    assert ws.cell(row=4, column=9).value == "SMC"


def test_parse_change_notice_ignores_inventory_rows_after_change_detail(tmp_path: Path) -> None:
    change_notice = tmp_path / "change.xlsx"
    _save_change_notice_with_inventory_rows(change_notice)

    groups = parse_change_notice(change_notice)

    assert [(group.locator_type, group.locator_value) for group in groups] == [("位号", "R758")]

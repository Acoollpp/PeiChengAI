import shutil
import uuid
import zipfile
from pathlib import Path
from threading import Lock
from time import sleep

import pytest
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, col, select

from app.models import BomElectronicEditTask, BomElectronicEditTaskStatus
from app.services.bom_electronic_edit_processor import ChecklistRow, ProcessResult
from app.services.bom_electronic_edit_task_service import (
    INTERRUPTED_STATUSES,
    _change_record_people_from_task_name,
    _copy_bom_to_temporary_workfile,
    _result_workbook_path,
    _write_checklist_summary,
    ensure_checklist_summary_in_result_zip,
    mark_interrupted_tasks_failed,
    run_task,
)
from tests.utils.user import create_random_user


def test_result_workbook_path_increments_trailing_sequence() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(
        task_id,
        Path("BOM-SST25-0687-1.1-阿瑞斯-004.xlsx"),
    )

    assert path.name == "BOM-SST25-0687-1.1-阿瑞斯-005.xlsx"


def test_copy_bom_to_temporary_workfile_creates_isolated_copy(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    source = tmp_path / "BOM.xlsx"
    source.write_bytes(b"bom")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )

    temporary = _copy_bom_to_temporary_workfile(task_id, source)

    assert temporary != source
    assert temporary.read_bytes() == b"bom"
    assert temporary.parent == tmp_path / "storage" / str(task_id) / "work" / "processing"


def test_result_workbook_path_preserves_sequence_width() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-IOB-40797B-1.0-001.xlsx"))

    assert path.name == "BOM-IOB-40797B-1.0-002.xlsx"


def test_result_workbook_path_increments_letter_sequence() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-SST25-0637-2.0-A001.xlsx"))

    assert path.name == "BOM-SST25-0637-2.0-A002.xlsx"


def test_result_workbook_path_converts_xls_output_to_xlsx() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-SST20-8331-3.1-001.xls"))

    assert path.name == "BOM-SST20-8331-3.1-002.xlsx"


def test_result_workbook_path_adds_initial_sequence_when_no_trailing_sequence() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-SST26-1900-1.0.xlsx"))

    assert path.name == "BOM-SST26-1900-1.0-001.xlsx"


def test_result_workbook_path_increments_sequence_before_trailing_text() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-SST26-1848-1.0-华振-001关键物料可替代.xlsx"))

    assert path.name == "BOM-SST26-1848-1.0-华振-002关键物料可替代.xlsx"


def test_result_workbook_path_replaces_trailing_date_with_sequence() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-SST22-0290-REV2.0-20240403.xlsx"))

    assert path.name == "BOM-SST22-0290-REV2.0-001.xlsx"


def test_result_workbook_path_replaces_rev_date_suffix_with_sequence() -> None:
    task_id = uuid.uuid4()

    path = _result_workbook_path(task_id, Path("BOM-SST22-1133-REV1.1-20220517.xlsx"))

    assert path.name == "BOM-SST22-1133-REV1.1-001.xlsx"


def test_change_record_people_from_task_name_uses_liao_jinghua_rule() -> None:
    people = _change_record_people_from_task_name("7-31修改BOM廖晶华")

    assert people.preparer == "廖晶华"
    assert people.reviewer == "邓如锦"
    assert people.approver == "许志杰"


def test_change_record_people_from_task_name_uses_deng_rujin_rule() -> None:
    people = _change_record_people_from_task_name("7-31修改BOM邓如锦")

    assert people.preparer == "邓如锦"
    assert people.reviewer == "廖晶华"
    assert people.approver == "许志杰"


def test_change_record_people_from_task_name_uses_liu_demei_rule() -> None:
    people = _change_record_people_from_task_name("8-4修改BOM刘德媚")

    assert people.preparer == "刘德媚"
    assert people.reviewer == "/"
    assert people.approver == "/"


def test_change_record_people_from_task_name_requires_known_person() -> None:
    with pytest.raises(ValueError, match="任务名称未包含廖晶华、邓如锦或刘德媚"):
        _change_record_people_from_task_name("7-31修改BOM")


def test_run_task_converts_xls_bom_before_processing(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "BOM-SST20-8331-3.1-001.xls"
    change_notice.write_bytes(b"change")
    master.write_bytes(b"master")
    bom.write_bytes(b"xls")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )
    converted_paths: list[Path] = []
    processed_paths: list[Path] = []
    passed_change_groups: list[object] = []
    passed_materials: list[object] = []

    def fake_convert(source: Path, target: Path) -> None:
        converted_paths.append(target)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(source.read_bytes())

    def fake_process(**kwargs) -> ProcessResult:
        processed_paths.append(kwargs["bom_path"])
        passed_change_groups.append(kwargs["change_groups"])
        passed_materials.append(kwargs["materials"])
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"result")
        return ProcessResult(
            output_path=output_path,
            modified_count=1,
            abnormal_count=0,
            report_count=1,
            processed_sheet_count=1,
        )

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service._convert_xls_to_xlsx_with_excel",
        fake_convert,
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {"1.01.01.000001": ["parsed-material"]},
    )

    result = run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            {"role": "bom", "file_path": str(bom)},
        ],
        task_name="8-4廖晶华",
    )

    assert converted_paths == [tmp_path / "storage" / str(task_id) / "work" / "converted" / "BOM-SST20-8331-3.1-001.xlsx"]
    assert processed_paths == converted_paths
    assert passed_change_groups == [["parsed-change"]]
    assert passed_materials == [{"1.01.01.000001": ["parsed-material"]}]
    assert result["result_files"][0]["file_name"] == "BOM电子档修改结果.zip"
    assert Path(str(result["result_files"][0]["file_path"])).is_file()


def test_run_task_records_single_bom_processing_error_and_continues(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    broken_bom = tmp_path / "BOM-SST22-0432-1.1-002.xls"
    next_bom = tmp_path / "BOM-SST23-0442-1.1-006.xlsx"
    for path in (change_notice, master, broken_bom, next_bom):
        path.write_bytes(b"input")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )
    monkeypatch.setattr("app.services.bom_electronic_edit_task_service.time.sleep", lambda _: None)
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {},
    )

    def fake_convert(source: Path, _target: Path) -> None:
        raise RuntimeError(f"转换 .xls 文件失败：{source.name}；NoCOMClassIdentified")

    def fake_process(**kwargs) -> ProcessResult:
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"result")
        return ProcessResult(output_path, 1, 0, 1, 1)

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service._convert_xls_to_xlsx_with_excel",
        fake_convert,
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )

    result = run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            {"role": "bom", "file_path": str(broken_bom)},
            {"role": "bom", "file_path": str(next_bom)},
        ],
        task_name="9-24邓如锦",
    )

    assert result["modified_count"] == 1
    assert result["abnormal_count"] == 1
    assert result["report_count"] == 2
    assert result["processed_sheet_count"] == 1
    zip_path = Path(str(result["result_files"][0]["file_path"]))
    with zipfile.ZipFile(zip_path) as archive:
        names = archive.namelist()
    assert "异常_BOM-SST22-0432-1.1-002.xlsx" in names
    assert "BOM-SST23-0442-1.1-007.xlsx" in names
    assert "BOM电子档修改清单汇总.xlsx" in names

    summary = load_workbook(tmp_path / "storage" / str(task_id) / "reports" / "BOM电子档修改清单汇总.xlsx", data_only=True)
    checklist = summary["清单汇总"]
    assert checklist.cell(row=2, column=14).value == "异常"
    assert "NoCOMClassIdentified" in checklist.cell(row=2, column=15).value


def test_run_task_retries_locked_bom_and_continues_with_next_file(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    locked_bom = tmp_path / "locked.xlsx"
    next_bom = tmp_path / "next.xlsx"
    for path in (change_notice, master, locked_bom, next_bom):
        path.write_bytes(b"input")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )
    monkeypatch.setattr("app.services.bom_electronic_edit_task_service.time.sleep", lambda _: None)
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {},
    )
    attempts: list[str] = []

    def fake_process(**kwargs) -> ProcessResult:
        name = kwargs["bom_path"].name
        attempts.append(name)
        if name == "locked.xlsx":
            raise PermissionError("locked")
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"result")
        return ProcessResult(output_path, 1, 0, 1, 1)

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )
    original_copy = shutil.copyfile

    def fake_copy(source: str | Path, destination: str | Path, *args: object, **kwargs: object) -> str:
        if Path(source).name == "locked.xlsx":
            raise PermissionError("locked")
        return original_copy(source, destination, *args, **kwargs)

    monkeypatch.setattr("app.services.bom_electronic_edit_task_service.shutil.copyfile", fake_copy)

    result = run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            {"role": "bom", "file_path": str(locked_bom)},
            {"role": "bom", "file_path": str(next_bom)},
        ],
        task_name="8-6邓如锦",
    )

    assert attempts.count("locked.xlsx") == 0
    assert attempts.count("next.xlsx") == 1
    assert result["modified_count"] == 1
    assert result["abnormal_count"] == 1
    assert (tmp_path / "storage" / str(task_id) / "reports" / "异常_locked.xlsx").is_file()


def test_run_task_processes_boms_with_at_most_two_workers(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    boms = [tmp_path / f"bom-{index}.xlsx" for index in range(3)]
    for path in [change_notice, master, *boms]:
        path.write_bytes(b"input")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )
    monkeypatch.setattr("app.services.bom_electronic_edit_task_service.time.sleep", lambda _: None)
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {},
    )
    lock = Lock()
    active = 0
    maximum_active = 0

    def fake_process(**kwargs) -> ProcessResult:
        nonlocal active, maximum_active
        with lock:
            active += 1
            maximum_active = max(maximum_active, active)
        sleep(0.05)
        with lock:
            active -= 1
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"result")
        return ProcessResult(output_path, 1, 0, 1, 1)

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )

    run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            *[{"role": "bom", "file_path": str(bom)} for bom in boms],
        ],
        task_name="8-6邓如锦",
    )

    assert maximum_active == 2


def test_run_task_keeps_original_bom_filename_when_nothing_changed(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "BOM-SST20-A320-3.1-003.xlsx"
    change_notice.write_bytes(b"change")
    master.write_bytes(b"master")
    bom.write_bytes(b"bom")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )

    def fake_process(**kwargs) -> ProcessResult:
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"unchanged")
        return ProcessResult(
            output_path=output_path,
            modified_count=0,
            abnormal_count=0,
            report_count=0,
            processed_sheet_count=1,
        )

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {"1.01.01.000001": ["parsed-material"]},
    )

    result = run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            {"role": "bom", "file_path": str(bom)},
        ],
        task_name="8-5廖晶华",
    )

    zip_path = Path(str(result["result_files"][0]["file_path"]))
    with zipfile.ZipFile(zip_path) as archive:
        assert archive.namelist() == ["BOM-SST20-A320-3.1-003.xlsx", "BOM电子档修改清单汇总.xlsx"]


def test_run_task_writes_standalone_checklist_summary_workbook(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "BOM-SST20-8331-3.0-001.xlsx"
    change_notice.write_bytes(b"change")
    master.write_bytes(b"master")
    bom.write_bytes(b"bom")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )

    def fake_process(**kwargs) -> ProcessResult:
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"result")
        return ProcessResult(
            output_path=output_path,
            modified_count=1,
            abnormal_count=1,
            report_count=2,
            processed_sheet_count=2,
            checklist_rows=[
                ChecklistRow(
                    sheet_name="总表",
                    sheet_type="总表",
                    locator_value="DP7",
                    action="添加",
                    after_code="1.06.02.000194",
                    after_model="SMDJ130A",
                    after_quantity="1",
                    after_locator="DP7",
                    after_action="添加",
                    status="已完成",
                ),
                ChecklistRow(
                    sheet_name="SMT-08331-1.0",
                    sheet_type="SMT段",
                    locator_value="DP7",
                    action="变更",
                    before_code="1.06.02.000020",
                    before_model="SMCJ100CA",
                    before_quantity="1",
                    before_locator="DP7",
                    after_code="1.06.02.000194",
                    after_model="SMDJ130A",
                    after_quantity="1",
                    after_locator="DP7",
                    status="未更改",
                    reason="BOM仍存在变更前物料",
                ),
            ],
        )

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {"1.06.02.000194": ["parsed-material"]},
    )

    result = run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            {"role": "bom", "file_path": str(bom)},
            {"role": "task_options", "include_report_sheet": False},
        ],
        task_name="8-5邓如锦",
    )

    zip_path = Path(str(result["result_files"][0]["file_path"]))
    with zipfile.ZipFile(zip_path) as archive:
        names = archive.namelist()
        assert "BOM电子档修改清单汇总.xlsx" in names
        archive.extract("BOM电子档修改清单汇总.xlsx", tmp_path)

    summary = load_workbook(tmp_path / "BOM电子档修改清单汇总.xlsx", data_only=True)
    checklist = summary["清单汇总"]
    assert [cell.value for cell in checklist[1]] == [
        "BOM文件",
        "Sheet",
        "Sheet类型",
        "位号/定位值",
        "动作",
        "变更前物料编码",
        "变更前型号",
        "变更前用量",
        "变更前位号",
        "变更后物料编码",
        "变更后型号",
        "变更后用量",
        "变更后位号",
        "状态",
        "原因",
    ]
    assert checklist.cell(row=2, column=1).value == "BOM-SST20-8331-3.0-002.xlsx"
    assert checklist.cell(row=2, column=3).value == "总表"
    assert checklist.cell(row=2, column=5).value == "添加"
    assert checklist.cell(row=2, column=14).value == "已完成"
    assert checklist.cell(row=3, column=3).value == "SMT段"
    assert checklist.cell(row=3, column=14).value == "未更改"
    bom_summary = summary["BOM汇总"]
    assert [cell.value for cell in bom_summary[1]] == [
        "BOM文件",
        "状态",
        "总表完成数",
        "分层完成数",
        "未更改数",
        "异常数",
        "判定结论",
        "备注",
    ]
    assert bom_summary.cell(row=2, column=2).value == "存在未更改"
    assert bom_summary.cell(row=2, column=3).value == 1
    assert bom_summary.cell(row=2, column=4).value == 0
    assert bom_summary.cell(row=2, column=5).value == 1


def test_run_task_treats_single_bom_sheet_as_total_in_checklist_summary(tmp_path: Path, monkeypatch) -> None:
    task_id = uuid.uuid4()
    change_notice = tmp_path / "change.xlsx"
    master = tmp_path / "master.xlsx"
    bom = tmp_path / "BOM-SST24-0766-1.3-沛城-001.xlsx"
    change_notice.write_bytes(b"change")
    master.write_bytes(b"master")
    bom.write_bytes(b"bom")
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )

    def fake_process(**kwargs) -> ProcessResult:
        output_path = kwargs["output_path"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "P16S100A-40766-1.3"
        ws.append(["物料代码", "规格型号", "规格描述", "用量", "位号"])
        wb.save(output_path)
        return ProcessResult(
            output_path=output_path,
            modified_count=1,
            abnormal_count=0,
            report_count=1,
            processed_sheet_count=1,
            checklist_rows=[
                ChecklistRow(
                    sheet_name="P16S100A-40766-1.3",
                    sheet_type="成品段",
                    locator_value="1.02.01.000665",
                    action="删除",
                    before_code="1.02.01.000665",
                    before_model="AC0805FR-0720K5L",
                    before_quantity="1",
                    before_locator="R763",
                    after_code="1.02.01.000695",
                    after_model="ADE05K1001FT",
                    after_quantity="1",
                    after_locator="R763",
                    after_action="添加",
                    status="已完成",
                )
            ],
        )

    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.process_bom_workbook",
        fake_process,
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.parse_change_notice",
        lambda path: ["parsed-change"],
    )
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.bom_electronic_edit_processor.load_material_master",
        lambda path: {"1.02.01.000695": ["parsed-material"]},
    )

    result = run_task(
        task_id,
        [
            {"role": "change_notice", "file_path": str(change_notice)},
            {"role": "material_master", "file_path": str(master)},
            {"role": "bom", "file_path": str(bom)},
        ],
        task_name="8-7廖晶华",
    )

    zip_path = Path(str(result["result_files"][0]["file_path"]))
    with zipfile.ZipFile(zip_path) as archive:
        archive.extract("BOM电子档修改清单汇总.xlsx", tmp_path)

    summary = load_workbook(tmp_path / "BOM电子档修改清单汇总.xlsx", data_only=True)
    checklist = summary["清单汇总"]
    assert checklist.cell(row=2, column=2).value == "P16S100A-40766-1.3"
    assert checklist.cell(row=2, column=3).value == "总表"
    assert checklist.cell(row=2, column=5).value == "删除"
    assert checklist.cell(row=2, column=6).value == "1.02.01.000665"
    assert checklist.cell(row=2, column=10).value == "1.02.01.000695"
    assert checklist.cell(row=2, column=14).value == "已完成"


def test_write_checklist_summary_uses_model_sheet_name_for_bom_display_name(tmp_path: Path) -> None:
    output_path = tmp_path / "BOM-SST24-0766-1.3-沛城-002.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "P16S100A-40766-1.3"
    ws.append(["物料代码", "规格型号", "规格描述", "用量", "位号"])
    wb.save(output_path)
    result = ProcessResult(
        output_path=output_path,
        modified_count=1,
        abnormal_count=0,
        report_count=1,
        processed_sheet_count=1,
        checklist_rows=[
            ChecklistRow(
                sheet_name="总表",
                sheet_type="总表",
                locator_value="1.02.01.000665",
                action="删除",
                before_code="1.02.01.000665",
                before_model="AC0805FR-0720K5L",
                before_quantity="1",
                before_locator="R763",
                after_code="1.02.01.000695",
                after_model="ADE05K1001FT",
                after_quantity="1",
                after_locator="R763",
                status="已完成",
            )
        ],
    )

    summary_path = tmp_path / "summary.xlsx"
    _write_checklist_summary(summary_path, [(output_path.name, result)])

    summary = load_workbook(summary_path, data_only=True)
    checklist = summary["清单汇总"]
    bom_summary = summary["BOM汇总"]
    assert checklist.cell(row=2, column=1).value == "P16S100A-40766-1.3"
    assert checklist.cell(row=2, column=2).value == "P16S100A-40766-1.3"
    assert checklist.cell(row=2, column=3).value == "总表"
    assert bom_summary.cell(row=2, column=1).value == "P16S100A-40766-1.3"


def test_ensure_checklist_summary_rewrites_existing_stale_summary(tmp_path: Path) -> None:
    report_path = tmp_path / "BOM-SST24-0766-1.3-沛城-002.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "P16S100A-40766-1.3"
    ws.append(["物料代码", "规格型号", "规格描述", "用量", "位号"])
    report = wb.create_sheet("BOM电子档修改报告")
    report.append(["Sheet", "定位值", "动作", "变更前物料编码", "变更后物料编码", "结果"])
    report.append(["P16S100A-40766-1.3", "R763", "删除", "1.02.01.000665", "", "成功"])
    wb.save(report_path)

    summary_path = tmp_path / "BOM电子档修改清单汇总.xlsx"
    stale = Workbook()
    stale.active.append(["BOM文件"])
    stale.active.append([report_path.name])
    stale.save(summary_path)
    zip_path = tmp_path / "BOM电子档修改结果.zip"
    with zipfile.ZipFile(zip_path, "w") as archive:
        archive.write(report_path, report_path.name)
        archive.write(summary_path, summary_path.name)

    assert ensure_checklist_summary_in_result_zip(zip_path) is True

    with zipfile.ZipFile(zip_path) as archive:
        archive.extract(summary_path.name, tmp_path / "extracted")
    summary = load_workbook(tmp_path / "extracted" / summary_path.name, data_only=True)
    checklist = summary["清单汇总"]
    assert checklist.cell(row=1, column=1).value == "BOM文件"
    assert checklist.cell(row=2, column=1).value == "P16S100A-40766-1.3"


def test_mark_interrupted_tasks_failed_marks_only_incomplete_tasks(db: Session) -> None:
    # Earlier route tests in the same session can leave interrupted background tasks.
    # This service-level assertion needs a deterministic interrupted-task set.
    for stale_task in db.exec(
        select(BomElectronicEditTask).where(
            col(BomElectronicEditTask.status).in_(INTERRUPTED_STATUSES)
        )
    ).all():
        db.delete(stale_task)
    db.commit()

    user = create_random_user(db)
    running_task = BomElectronicEditTask(
        user_id=user.id,
        task_name="running",
        status=BomElectronicEditTaskStatus.running,
    )
    queued_task = BomElectronicEditTask(
        user_id=user.id,
        task_name="queued",
        status=BomElectronicEditTaskStatus.queued,
    )
    success_task = BomElectronicEditTask(
        user_id=user.id,
        task_name="success",
        status=BomElectronicEditTaskStatus.success,
    )
    db.add(running_task)
    db.add(queued_task)
    db.add(success_task)
    db.commit()

    changed = mark_interrupted_tasks_failed(db)

    db.refresh(running_task)
    db.refresh(queued_task)
    db.refresh(success_task)
    assert changed == 1
    assert running_task.status == BomElectronicEditTaskStatus.failed
    assert running_task.finished_at is not None
    assert running_task.error_message == "服务重启导致任务中断，请重新提交任务"
    assert queued_task.status == BomElectronicEditTaskStatus.queued
    assert queued_task.error_message is None
    assert success_task.status == BomElectronicEditTaskStatus.success
    assert success_task.error_message is None


def test_mark_interrupted_tasks_failed_recovers_completed_reports(
    db: Session,
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "app.services.bom_electronic_edit_task_service.settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR",
        str(tmp_path / "storage"),
    )
    user = create_random_user(db)
    task = BomElectronicEditTask(
        user_id=user.id,
        task_name="running with reports",
        status=BomElectronicEditTaskStatus.running,
        input_files=[
            {
                "file_name": "BOM.xlsx",
                "file_path": str(tmp_path / "BOM.xlsx"),
                "role": "bom",
            }
        ],
    )
    db.add(task)
    db.commit()
    db.refresh(task)

    report_dir = tmp_path / "storage" / str(task.id) / "reports"
    report_dir.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(["BOM"])
    report_path = report_dir / "BOM-001.xlsx"
    workbook.save(report_path)

    changed = mark_interrupted_tasks_failed(db)

    db.refresh(task)
    assert changed == 1
    assert task.status == BomElectronicEditTaskStatus.success
    assert task.error_message is None
    assert task.report_count == 1
    assert len(task.result_files) == 1
    zip_path = Path(task.result_files[0]["file_path"])
    assert zip_path.is_file()
    with zipfile.ZipFile(zip_path) as archive:
        assert archive.namelist() == ["BOM-001.xlsx"]

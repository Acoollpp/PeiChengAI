from datetime import datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import unquote
from uuid import UUID, uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session

from app.core.config import settings
from app.models import Department, ReviewFocusDirection, ReviewTask, UserRole
from tests.utils.user import authentication_token_from_key, create_key_user
from tests.utils.utils import random_lower_string


def _create_department(db: Session, name: str | None = None) -> Department:
    department = Department(name=name or f"方向部门-{random_lower_string()}")
    db.add(department)
    db.commit()
    db.refresh(department)
    return department


def _direction_payload(department_id: UUID | None) -> dict:
    suffix = random_lower_string()
    payload = {
        "title": f"权限审查-{suffix}",
        "content": "重点检查权限边界和异常处理",
        "is_active": True,
    }
    if department_id is not None:
        payload["department_id"] = str(department_id)
    return payload


def _create_focus_direction(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
    *,
    department: Department | None = None,
) -> dict:
    target_department = department or _create_department(db)
    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json=_direction_payload(target_department.id),
    )
    assert response.status_code == 200
    return response.json()


def _department_admin_headers(
    client: TestClient,
    db: Session,
    department: Department | None,
) -> dict[str, str]:
    user, api_key = create_key_user(db)
    user.role = UserRole.department_admin
    user.department_id = department.id if department is not None else None
    db.add(user)
    db.commit()
    return authentication_token_from_key(client=client, key=api_key.key, db=db)


def _create_task_with_focus_snapshot(
    *,
    client: TestClient,
    db: Session,
    department: Department,
    snapshot: str | None,
    minutes_ago: int,
) -> dict:
    user, api_key = create_key_user(db)
    user.department_id = department.id
    db.add(user)
    db.commit()
    db.refresh(user)
    headers = authentication_token_from_key(client=client, key=api_key.key, db=db)
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=headers,
        json={
            "svn_url": "svn://repo/project",
            "old_revision": str(10 + minutes_ago),
            "new_revision": str(20 + minutes_ago),
            "focus_direction_snapshot": snapshot,
        },
    )
    assert response.status_code == 200
    task = response.json()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    used_at = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(
        minutes=minutes_ago
    )
    db_task.created_at = used_at
    db_task.updated_at = used_at
    db.add(db_task)
    db.commit()
    return task


def test_admin_can_create_and_list_focus_directions(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db)
    payload = _direction_payload(department.id)

    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json=payload,
    )

    assert create_response.status_code == 200
    created = create_response.json()
    assert created["title"] == payload["title"]
    assert created["content"] == payload["content"]
    assert created["department_id"] == str(department.id)
    assert created["department_name"] == department.name
    assert created["is_active"] is True
    assert created["usage_count"] == 0
    assert created["created_by_user_id"] is not None

    list_response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions?skip=0&limit=100",
        headers=superuser_token_headers,
    )
    assert list_response.status_code == 200
    content = list_response.json()
    assert content["count"] >= 1
    assert any(direction["id"] == created["id"] for direction in content["data"])


def test_list_focus_directions_returns_dynamic_usage_count(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db, "列表统计部门-" + random_lower_string())
    payload = {
        "title": "列表统计方向-" + random_lower_string(),
        "content": "重点检查权限边界和异常处理-" + random_lower_string(),
        "department_id": str(department.id),
        "is_active": True,
    }
    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json=payload,
    )
    assert create_response.status_code == 200
    created = create_response.json()

    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department,
        snapshot=payload["content"],
        minutes_ago=1,
    )
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department,
        snapshot=f"  {payload['content']}  ",
        minutes_ago=2,
    )

    list_response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions?skip=0&limit=100",
        headers=superuser_token_headers,
    )

    assert list_response.status_code == 200
    directions = list_response.json()["data"]
    listed = next(direction for direction in directions if direction["id"] == created["id"])
    assert listed["usage_count"] == 2


def test_admin_can_create_global_focus_direction_without_department(
    client: TestClient,
    superuser_token_headers: dict[str, str],
) -> None:
    payload = _direction_payload(None)

    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json=payload,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["title"] == payload["title"]
    assert content["content"] == payload["content"]
    assert content["department_id"] is None
    assert content["department_name"] is None


def test_admin_can_create_focus_direction_with_multiple_departments(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department_a = _create_department(db, "多部门A-" + random_lower_string())
    department_b = _create_department(db, "多部门B-" + random_lower_string())

    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "多部门方向",
            "content": "检查跨部门公共逻辑",
            "department_ids": [str(department_a.id), str(department_b.id)],
            "is_active": True,
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["title"] == "多部门方向"
    assert content["department_id"] is None
    assert content["department_name"] is None
    assert {item["id"] for item in content["departments"]} == {
        str(department_a.id),
        str(department_b.id),
    }
    assert {item["name"] for item in content["departments"]} == {
        department_a.name,
        department_b.name,
    }


def test_admin_can_create_global_focus_direction_with_empty_department_ids(
    client: TestClient,
    superuser_token_headers: dict[str, str],
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "全局方向",
            "content": "检查所有部门通用风险",
            "department_ids": [],
            "is_active": True,
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["department_id"] is None
    assert content["department_name"] is None
    assert content["departments"] == []


def test_normal_user_cannot_create_focus_direction(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
) -> None:
    department = _create_department(db)

    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=normal_user_token_headers,
        json=_direction_payload(department.id),
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "没有足够权限"


def test_department_admin_can_create_own_department_focus_direction(
    client: TestClient,
    db: Session,
) -> None:
    department = _create_department(db, "部门管理员方向-" + random_lower_string())
    headers = _department_admin_headers(client, db, department)

    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=headers,
        json=_direction_payload(department.id),
    )

    assert response.status_code == 200
    content = response.json()
    assert content["department_id"] == str(department.id)
    assert [item["id"] for item in content["departments"]] == [str(department.id)]


def test_department_admin_cannot_create_global_or_other_department_focus_direction(
    client: TestClient,
    db: Session,
) -> None:
    own_department = _create_department(db, "方向本部门-" + random_lower_string())
    other_department = _create_department(db, "方向其他部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, own_department)

    global_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=headers,
        json=_direction_payload(None),
    )
    other_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=headers,
        json=_direction_payload(other_department.id),
    )

    assert global_response.status_code == 403
    assert global_response.json()["detail"] == "只能维护本部门审查方向。"
    assert other_response.status_code == 403
    assert other_response.json()["detail"] == "只能维护本部门审查方向。"


def test_department_admin_lists_only_own_department_focus_directions(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    own_department = _create_department(db, "列表本部门-" + random_lower_string())
    other_department = _create_department(db, "列表其他部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, own_department)
    own_direction = _create_focus_direction(
        client, db, superuser_token_headers, department=own_department
    )
    other_direction = _create_focus_direction(
        client, db, superuser_token_headers, department=other_department
    )
    global_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "部门管理员不可见全局方向-" + random_lower_string(),
            "content": "检查所有部门通用风险",
            "department_ids": [],
            "is_active": True,
        },
    )
    shared_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "部门管理员可见共享方向-" + random_lower_string(),
            "content": "检查跨部门公共逻辑",
            "department_ids": [str(own_department.id), str(other_department.id)],
            "is_active": True,
        },
    )
    assert global_response.status_code == 200
    assert shared_response.status_code == 200

    response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=headers,
        params={
            "department_scope": "global",
            "department_id": str(other_department.id),
        },
    )

    assert response.status_code == 200
    ids = {item["id"] for item in response.json()["data"]}
    assert own_direction["id"] in ids
    assert shared_response.json()["id"] in ids
    assert other_direction["id"] not in ids
    assert global_response.json()["id"] not in ids


def test_department_admin_can_update_and_delete_own_department_focus_direction(
    client: TestClient,
    db: Session,
) -> None:
    department = _create_department(db, "维护本部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, department)
    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=headers,
        json=_direction_payload(department.id),
    )
    assert create_response.status_code == 200
    direction = create_response.json()

    update_response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=headers,
        json={
            "title": "部门管理员更新方向-" + random_lower_string(),
            "department_ids": [str(department.id)],
        },
    )

    assert update_response.status_code == 200
    assert update_response.json()["department_id"] == str(department.id)

    delete_response = client.delete(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=headers,
    )

    assert delete_response.status_code == 200
    assert delete_response.json()["message"] == "重点审查方向已停用"
    db_direction = db.get(ReviewFocusDirection, UUID(direction["id"]))
    assert db_direction is not None
    assert db_direction.is_active is False


def test_department_admin_cannot_move_own_focus_direction_out_of_department(
    client: TestClient,
    db: Session,
) -> None:
    own_department = _create_department(db, "移动本部门-" + random_lower_string())
    other_department = _create_department(db, "移动其他部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, own_department)
    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=headers,
        json=_direction_payload(own_department.id),
    )
    assert create_response.status_code == 200
    direction = create_response.json()

    global_response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=headers,
        json={"department_ids": []},
    )
    other_response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=headers,
        json={"department_id": str(other_department.id)},
    )

    assert global_response.status_code == 403
    assert global_response.json()["detail"] == "只能维护本部门审查方向。"
    assert other_response.status_code == 403
    assert other_response.json()["detail"] == "只能维护本部门审查方向。"


def test_department_admin_cannot_edit_shared_focus_direction(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    own_department = _create_department(db, "共享本部门-" + random_lower_string())
    other_department = _create_department(db, "共享其他部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, own_department)
    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "共享方向-" + random_lower_string(),
            "content": "检查跨部门公共逻辑",
            "department_ids": [str(own_department.id), str(other_department.id)],
            "is_active": True,
        },
    )
    assert create_response.status_code == 200
    direction = create_response.json()

    response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=headers,
        json={"title": "部门管理员不应更新共享方向"},
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "共享方向只能由超级管理员维护。"


def test_department_admin_cannot_delete_shared_focus_direction(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    own_department = _create_department(db, "删除共享本部门-" + random_lower_string())
    other_department = _create_department(db, "删除共享其他部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, own_department)
    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "删除共享方向-" + random_lower_string(),
            "content": "检查跨部门公共逻辑",
            "department_ids": [str(own_department.id), str(other_department.id)],
            "is_active": True,
        },
    )
    assert create_response.status_code == 200

    response = client.delete(
        f"{settings.API_V1_STR}/review-focus-directions/{create_response.json()['id']}",
        headers=headers,
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "共享方向只能由超级管理员维护。"


def test_department_admin_cannot_access_usage_stats_platform_apis(
    client: TestClient,
    db: Session,
) -> None:
    department = _create_department(db, "统计权限部门-" + random_lower_string())
    headers = _department_admin_headers(client, db, department)
    payload = {
        "title": "统计生成方向",
        "content": "检查权限边界",
        "department_id": str(department.id),
    }

    responses = [
        client.get(
            f"{settings.API_V1_STR}/review-focus-directions/usage-stats",
            headers=headers,
        ),
        client.patch(
            f"{settings.API_V1_STR}/review-focus-directions/usage-stats/state",
            headers=headers,
            json={"content": "检查权限边界", "is_used": True},
        ),
        client.get(
            f"{settings.API_V1_STR}/review-focus-directions/usage-stats/export",
            headers=headers,
        ),
        client.post(
            f"{settings.API_V1_STR}/review-focus-directions/from-usage",
            headers=headers,
            json=payload,
        ),
    ]

    assert [response.status_code for response in responses] == [403, 403, 403, 403]


def test_admin_can_update_focus_direction(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    direction = _create_focus_direction(client, db, superuser_token_headers)
    department = _create_department(db)
    update_payload = {
        "title": f"性能审查-{random_lower_string()}",
        "content": "重点检查缓存穿透和慢查询",
        "department_id": str(department.id),
        "is_active": False,
    }

    response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=superuser_token_headers,
        json=update_payload,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["id"] == direction["id"]
    assert content["title"] == update_payload["title"]
    assert content["content"] == update_payload["content"]
    assert content["department_id"] == str(department.id)
    assert content["department_name"] == department.name
    assert content["is_active"] is False
    assert content["updated_at"] != direction["updated_at"]


def test_admin_can_clear_focus_direction_department(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    direction = _create_focus_direction(client, db, superuser_token_headers)

    response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=superuser_token_headers,
        json={"department_id": None},
    )

    assert response.status_code == 200
    content = response.json()
    assert content["id"] == direction["id"]
    assert content["department_id"] is None
    assert content["department_name"] is None

    db_direction = db.get(ReviewFocusDirection, UUID(direction["id"]))
    assert db_direction is not None
    assert db_direction.department_id is None


def test_admin_update_without_department_keeps_existing_department(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db)
    direction = _create_focus_direction(
        client, db, superuser_token_headers, department=department
    )

    response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=superuser_token_headers,
        json={"title": f"安全审查-{random_lower_string()}"},
    )

    assert response.status_code == 200
    content = response.json()
    assert content["id"] == direction["id"]
    assert content["department_id"] == str(department.id)
    assert content["department_name"] == department.name

    db_direction = db.get(ReviewFocusDirection, UUID(direction["id"]))
    assert db_direction is not None
    assert db_direction.department_id == department.id


def test_delete_focus_direction_soft_deletes(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    direction = _create_focus_direction(client, db, superuser_token_headers)

    response = client.delete(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=superuser_token_headers,
    )

    assert response.status_code == 200
    assert response.json()["message"] == "重点审查方向已停用"
    db_direction = db.get(ReviewFocusDirection, UUID(direction["id"]))
    assert db_direction is not None
    assert db_direction.is_active is False
    assert db_direction.updated_at is not None


def test_missing_focus_direction_returns_404(
    client: TestClient,
    superuser_token_headers: dict[str, str],
) -> None:
    response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{uuid4()}",
        headers=superuser_token_headers,
        json={"title": "不存在方向"},
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "重点审查方向不存在"


def test_create_or_update_with_missing_department_returns_400(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    create_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json=_direction_payload(uuid4()),
    )

    assert create_response.status_code == 400
    assert create_response.json()["detail"] == "部门不存在"

    direction = _create_focus_direction(client, db, superuser_token_headers)
    update_response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/{direction['id']}",
        headers=superuser_token_headers,
        json={"department_id": str(uuid4())},
    )

    assert update_response.status_code == 400
    assert update_response.json()["detail"] == "部门不存在"


def test_admin_can_filter_focus_directions(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department_a = _create_department(db, "筛选部门A-" + random_lower_string())
    department_b = _create_department(db, "筛选部门B-" + random_lower_string())
    active_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "缓存穿透审查",
            "content": "检查缓存穿透和慢查询",
            "department_ids": [str(department_a.id)],
            "is_active": True,
        },
    )
    inactive_response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        json={
            "title": "权限边界审查",
            "content": "检查权限边界",
            "department_ids": [str(department_b.id)],
            "is_active": False,
        },
    )
    assert active_response.status_code == 200
    assert inactive_response.status_code == 200

    response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions",
        headers=superuser_token_headers,
        params={
            "search": "缓存",
            "department_scope": "department",
            "department_id": str(department_a.id),
            "is_active": True,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["count"] == 1
    assert payload["data"][0]["title"] == "缓存穿透审查"


def test_admin_can_read_usage_stats_grouped_by_normalized_content(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    unique_marker = random_lower_string()
    department_a = _create_department(db, "研发一部-" + random_lower_string())
    department_b = _create_department(db, "研发二部-" + random_lower_string())
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department_a,
        snapshot=f"  检查   权限\n边界  {unique_marker}",
        minutes_ago=30,
    )
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department_b,
        snapshot=f"检查 权限 边界 {unique_marker}",
        minutes_ago=10,
    )
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department_a,
        snapshot=f"\t检查 SQL 注入 {unique_marker}\t",
        minutes_ago=20,
    )
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department_a,
        snapshot="   \n  ",
        minutes_ago=5,
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats",
        headers=superuser_token_headers,
        params={"search": unique_marker},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["count"] == 2
    assert payload["data"][0]["content"] == f"检查 权限 边界 {unique_marker}"
    assert payload["data"][0]["usage_count"] == 2
    assert payload["data"][0]["latest_used_at"] is not None
    assert payload["data"][0]["departments"] == sorted(
        [department_a.name, department_b.name]
    )
    assert payload["data"][1]["content"] == f"检查 SQL 注入 {unique_marker}"
    assert payload["data"][1]["usage_count"] == 1
    assert payload["data"][1]["departments"] == [department_a.name]


def test_admin_can_mark_usage_stat_used_state(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db, "统计状态部门-" + random_lower_string())
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department,
        snapshot="检查空指针和越界",
        minutes_ago=1,
    )

    stats_response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats",
        headers=superuser_token_headers,
        params={"search": "空指针"},
    )
    assert stats_response.status_code == 200
    stat = stats_response.json()["data"][0]
    assert stat["content"] == "检查空指针和越界"
    assert stat["is_used"] is False

    mark_response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats/state",
        headers=superuser_token_headers,
        json={"content": " 检查空指针和越界 ", "is_used": True},
    )
    assert mark_response.status_code == 200
    assert mark_response.json()["message"] == "使用状态更新成功"

    used_response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats",
        headers=superuser_token_headers,
        params={"search": "空指针", "is_used": True},
    )
    assert used_response.status_code == 200
    assert used_response.json()["count"] == 1
    assert used_response.json()["data"][0]["is_used"] is True


def test_admin_can_filter_usage_stats_by_department_and_used_state(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department_a = _create_department(db, "统计筛选A-" + random_lower_string())
    department_b = _create_department(db, "统计筛选B-" + random_lower_string())
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department_a,
        snapshot="检查状态机跳转",
        minutes_ago=2,
    )
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department_b,
        snapshot="检查协议兼容",
        minutes_ago=1,
    )
    mark_response = client.patch(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats/state",
        headers=superuser_token_headers,
        json={"content": "检查状态机跳转", "is_used": True},
    )
    assert mark_response.status_code == 200

    response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats",
        headers=superuser_token_headers,
        params={
            "department_scope": "department",
            "department_id": str(department_a.id),
            "is_used": True,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["count"] == 1
    assert payload["data"][0]["content"] == "检查状态机跳转"


def test_admin_can_export_filtered_usage_stats_excel(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db, "导出部门-" + random_lower_string())
    other_department = _create_department(db, "导出未命中部门-" + random_lower_string())
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=department,
        snapshot="导出检查项",
        minutes_ago=1,
    )
    _create_task_with_focus_snapshot(
        client=client,
        db=db,
        department=other_department,
        snapshot="不应导出的检查项",
        minutes_ago=2,
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-focus-directions/usage-stats/export",
        headers=superuser_token_headers,
        params={"search": "导出检查"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "重点审查方向用户输入统计_" in unquote(
        response.headers["content-disposition"]
    )
    assert len(response.content) > 100
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("内容", "使用次数", "涉及部门", "最新使用时间", "是否已使用")
    assert len(rows) == 2
    assert rows[1][0] == "导出检查项"
    assert rows[1][1] == 1
    assert rows[1][2] == department.name
    assert rows[1][4] == "未使用"


def test_admin_can_create_focus_direction_from_usage(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db)

    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions/from-usage",
        headers=superuser_token_headers,
        json={
            "title": "权限边界审查",
            "content": "  检查   权限\n边界  ",
            "department_id": str(department.id),
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["title"] == "权限边界审查"
    assert content["content"] == "检查 权限 边界"
    assert content["department_id"] == str(department.id)
    assert content["department_name"] == department.name
    assert content["is_active"] is True


def test_admin_can_create_global_focus_direction_from_usage(
    client: TestClient,
    superuser_token_headers: dict[str, str],
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-focus-directions/from-usage",
        headers=superuser_token_headers,
        json={
            "title": "全局权限边界审查",
            "content": "  检查   权限\n边界  ",
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["title"] == "全局权限边界审查"
    assert content["content"] == "检查 权限 边界"
    assert content["department_id"] is None
    assert content["department_name"] is None
    assert content["is_active"] is True

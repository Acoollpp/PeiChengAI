import base64
import hashlib
import json
import logging
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, time, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import text
from sqlmodel import Session, select

from app.api.routes import review_tasks as review_tasks_route
from app.core.config import settings
from app.models import (
    MAX_REVIEW_TASK_ISSUE_DESCRIPTION_LENGTH,
    MAX_REVIEW_TASK_ISSUE_DISPLAY_SNAPSHOT_LENGTH,
    MAX_REVIEW_TASK_ISSUE_FEEDBACK_LENGTH,
    MAX_REVIEW_TASK_ISSUE_TEXT_LENGTH,
    Department,
    ReviewTask,
    ReviewTaskIssue,
    ReviewTaskIssueSnapshotCreate,
    ReviewTaskReportFeedback,
    ReviewTaskRequirementFile,
    ReviewTaskStatus,
    User,
    UserRole,
)
from app.services import report_service
from tests.utils.user import (
    authentication_token_from_key,
    create_key_user,
)
from tests.utils.utils import random_lower_string

TASK_STATS_TIMEZONE = ZoneInfo("Asia/Shanghai")
VALID_PNG_BYTES = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
    "0000000c49444154789c63f8cfc0000003010100c9fe92ef0000000049454e44ae"
    "426082"
)


def _load_shared_markdown_image_corpus() -> dict[str, object]:
    fixture_path = (
        Path(__file__).resolve().parents[5]
        / "tests"
        / "fixtures"
        / "report-image-markdown-corpus.json"
    )
    return json.loads(fixture_path.read_text(encoding="utf-8"))


def _load_shared_image_validation_corpus() -> list[dict[str, Any]]:
    fixture_path = (
        Path(__file__).resolve().parents[5]
        / "tests"
        / "fixtures"
        / "report-image-validation-corpus.json"
    )
    return json.loads(fixture_path.read_text(encoding="utf-8"))


IMAGE_VALIDATION_CORPUS = _load_shared_image_validation_corpus()
VALID_IMAGE_CASES = [
    {**image_case, "valid_base64": valid_base64}
    for image_case in IMAGE_VALIDATION_CORPUS
    for valid_base64 in [
        image_case["valid_base64"],
        *image_case.get("additional_valid_base64", []),
    ]
]
ANIMATED_IMAGE_CASES = [
    image_case
    for image_case in IMAGE_VALIDATION_CORPUS
    if "animated_base64" in image_case
]
INVALID_IMAGE_CASES = [
    {**image_case, "invalid_base64": invalid_base64}
    for image_case in IMAGE_VALIDATION_CORPUS
    for invalid_base64 in [
        image_case["invalid_base64"],
        *image_case.get("additional_invalid_base64", []),
        *image_case.get("server_only_invalid_base64", []),
    ]
]


def _valid_issue_snapshot(order: int = 0) -> dict[str, object]:
    return {
        "issue_text": "高风险：测试问题",
        "description_text": "用于验证服务端限制。",
        "display_snapshot": "### 高风险：测试问题",
        "order": order,
        "source": "template_1_heading",
        "location": {
            "start_offset": 0,
            "end_offset": 8,
            "feedback_insertion_offset": 8,
            "anchor_line": 0,
            "heading_level": 3,
        },
    }


def _revision_number(value: str) -> int:
    return int(value.strip().removeprefix("r").removeprefix("R"))


def _create_review_task(
    client: TestClient,
    headers: dict[str, str],
    old_revision: str = "1",
    new_revision: str | None = None,
) -> dict:
    old_number = _revision_number(old_revision)
    normalized_new_revision = new_revision or f"r{max(old_number + 1, 1)}"
    data = {
        "svn_url": "https://svn.example.com/project/trunk",
        "old_revision": old_revision,
        "new_revision": normalized_new_revision,
    }
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=headers,
        json=data,
    )
    assert response.status_code == 200
    return response.json()


def _create_custom_review_task(
    client: TestClient,
    headers: dict[str, str],
    *,
    rule: str,
    old_revision: str = "1",
    new_revision: str = "2",
    mode: str = "dual_revision",
    has_word_material: bool = False,
    has_excel_material: bool = False,
    has_pdf_material: bool = False,
    svn_url: str = "https://svn.example.com/custom-review/trunk",
) -> dict:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=headers,
        json={
            "svn_url": svn_url,
            "old_revision": old_revision,
            "new_revision": new_revision,
            "skill_name": "bms-code-reviewer",
            "review_item_ids": ["DS-99"],
            "custom_review_rule": rule,
            "custom_review_mode": mode,
            "has_custom_review_word_material": has_word_material,
            "has_custom_review_excel_material": has_excel_material,
            "has_custom_review_pdf_material": has_pdf_material,
        },
    )
    assert response.status_code == 200
    return response.json()


def _create_new_protocol_task_with_issue(
    client: TestClient,
    headers: dict[str, str],
) -> tuple[dict, dict]:
    task = _create_review_task(client, headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([_valid_issue_snapshot()], ensure_ascii=False),
        },
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=headers,
    )
    assert issues_response.status_code == 200
    return task, issues_response.json()["data"][0]


def _create_named_svn_user_task(
    *,
    client: TestClient,
    db: Session,
    full_name: str,
    svn_username: str,
    old_revision: str = "1",
) -> dict:
    user, api_key = create_key_user(db)
    user.full_name = full_name
    user.svn_username = f"{svn_username}-{random_lower_string()}"
    db.add(user)
    db.commit()
    db.refresh(user)
    user_headers = authentication_token_from_key(client=client, key=api_key.key, db=db)
    task = _create_review_task(client, user_headers, old_revision=old_revision)
    return {
        "task": task,
        "headers": user_headers,
        "user": user,
        "svn_username": user.svn_username,
    }


def _create_department(db: Session, name: str | None = None) -> Department:
    department = Department(name=name or f"任务部门-{random_lower_string()}")
    db.add(department)
    db.commit()
    db.refresh(department)
    return department


def _create_department_user(
    *,
    client: TestClient,
    db: Session,
    department: Department | None,
    role: UserRole = UserRole.user,
    full_name: str | None = None,
    svn_username: str | None = None,
) -> tuple[dict[str, str], User]:
    user, api_key = create_key_user(db)
    user.role = role
    user.department_id = department.id if department else None
    user.full_name = full_name or f"部门用户-{random_lower_string()}"
    user.svn_username = svn_username or f"dept-svn-{random_lower_string()}"
    db.add(user)
    db.commit()
    db.refresh(user)
    return authentication_token_from_key(client=client, key=api_key.key, db=db), user


def _shanghai_datetime(days_delta: int, hour: int = 12) -> datetime:
    today = datetime.now(TASK_STATS_TIMEZONE).date()
    return (
        datetime.combine(
            today + timedelta(days=days_delta),
            time(hour=hour),
            tzinfo=TASK_STATS_TIMEZONE,
        )
        .astimezone(timezone.utc)
        .replace(tzinfo=None)
    )


def _today_elapsed_datetime(minutes_ago: int) -> datetime:
    current = datetime.now(TASK_STATS_TIMEZONE)
    today_start = datetime.combine(
        current.date(),
        time.min,
        tzinfo=TASK_STATS_TIMEZONE,
    )
    elapsed = max(1, int((current - today_start).total_seconds() // 60) - minutes_ago)
    return (
        (today_start + timedelta(minutes=elapsed))
        .astimezone(timezone.utc)
        .replace(tzinfo=None)
    )


def _set_review_task_stats_fields(
    db: Session,
    task_id: str,
    *,
    created_at: datetime,
    status: ReviewTaskStatus,
    started_at: datetime | None = None,
    finished_at: datetime | None = None,
    report_path: str | None = None,
) -> None:
    db_task = db.get(ReviewTask, UUID(task_id))
    assert db_task is not None
    db_task.created_at = created_at
    db_task.updated_at = created_at
    db_task.status = status
    db_task.started_at = started_at
    db_task.finished_at = finished_at
    db_task.report_path = report_path
    db.add(db_task)
    db.commit()


def _set_review_task_pdf_report(
    db: Session,
    task_id: str,
    *,
    pdf_report_path: str,
    pdf_report_name: str = "report.pdf",
) -> None:
    db_task = db.get(ReviewTask, UUID(task_id))
    assert db_task is not None
    db_task.pdf_report_path = pdf_report_path
    db_task.pdf_report_name = pdf_report_name
    db.add(db_task)
    db.commit()


def _set_review_task_report(
    db: Session,
    task_id: str,
    *,
    report_path: str = "/tmp/report.md",
    pdf_report_path: str = "/tmp/report.pdf",
) -> None:
    db_task = db.get(ReviewTask, UUID(task_id))
    assert db_task is not None
    db_task.report_path = report_path
    db_task.report_name = "report.md"
    db_task.pdf_report_path = pdf_report_path
    db_task.pdf_report_name = "report.pdf"
    db_task.status = ReviewTaskStatus.success
    db.add(db_task)
    db.commit()


def _submit_report_feedback(
    client: TestClient,
    headers: dict[str, str],
    task_id: str,
    *,
    helpful: bool,
    useful_issues_text: str | None = None,
):
    return client.post(
        f"{settings.API_V1_STR}/review-tasks/{task_id}/report-feedback",
        headers=headers,
        json={
            "helpful": helpful,
            "useful_issues_text": useful_issues_text,
        },
    )


def test_create_review_task_normalizes_revisions(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "1",
            "new_revision": "r2",
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["svn_url"] == "https://svn.example.com/project/trunk"
    assert content["old_revision"] == "r1"
    assert content["new_revision"] == "r2"
    assert content["status"] == "created"


def test_create_review_task_allows_zero_old_revision_for_full_review(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "0",
            "new_revision": "r123",
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["old_revision"] == "r0"
    assert content["new_revision"] == "r123"


def test_create_review_task_rejects_zero_new_revision(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "0",
            "new_revision": "0",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "新版本号不能为 0"


def test_create_review_task_rejects_zero_new_revision_with_nonzero_old(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "r1",
            "new_revision": "r0",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "新版本号不能为 0"


def test_create_review_task_rejects_equal_non_zero_revisions(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "12",
            "new_revision": "r12",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "旧版本号必须小于新版本号"


def test_create_review_task_stores_template_and_skill_name(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "svn://repo/project",
            "old_revision": "r12",
            "new_revision": "r14",
            "template_key": "template_2",
            "skill_name": "通用-代码变更审查-模板二",
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["template_key"] == "template_2"
    assert content["skill_name"] == "通用-代码变更审查-模板二"
    assert content["old_revision"] == "r12"
    assert content["new_revision"] == "r14"


def test_create_review_task_without_template_fields_remains_supported(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "svn://repo/project",
            "old_revision": "12",
            "new_revision": "14",
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["template_key"] is None
    assert content["skill_name"] is None
    assert content["review_item_ids"] == []


def test_create_review_task_stores_and_returns_review_item_ids(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "svn://repo/project",
            "old_revision": "0",
            "new_revision": "r14",
            "skill_name": "bms-code-reviewer",
            "review_item_ids": ["DS-01", "DS-03"],
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["old_revision"] == "0"
    assert content["review_item_ids"] == ["DS-01", "DS-03"]
    task = db.get(ReviewTask, UUID(content["id"]))
    assert task is not None
    assert task.review_item_ids == ["DS-01", "DS-03"]

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{content['id']}",
        headers=normal_user_token_headers,
    )
    assert read_response.status_code == 200
    assert read_response.json()["review_item_ids"] == ["DS-01", "DS-03"]


def test_create_review_task_stores_ds99_input_snapshot_without_material_metadata(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "svn://repo/custom-review",
            "old_revision": "r12",
            "new_revision": "r14",
            "skill_name": "bms-code-reviewer",
            "review_item_ids": ["DS-03", "DS-99"],
            "custom_review_rule": "  保护条件必须具有明确恢复路径。  ",
            "custom_review_mode": "dual_revision",
            "has_custom_review_word_material": True,
            "has_custom_review_excel_material": False,
            "has_custom_review_pdf_material": True,
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["review_item_ids"] == ["DS-03", "DS-99"]
    assert content["custom_review_rule"] == "保护条件必须具有明确恢复路径。"
    assert content["custom_review_mode"] == "dual_revision"
    assert content["has_custom_review_word_material"] is True
    assert content["has_custom_review_excel_material"] is False
    assert content["has_custom_review_pdf_material"] is True
    task = db.get(ReviewTask, UUID(content["id"]))
    assert task is not None
    assert task.custom_review_rule == "保护条件必须具有明确恢复路径。"
    assert task.has_custom_review_word_material is True


@pytest.mark.parametrize(
    "payload, detail",
    [
        (
            {"review_item_ids": ["DS-99"]},
            "DS-99 自定义审查必须填写规则正文",
        ),
        (
            {
                "review_item_ids": ["DS-99"],
                "custom_review_rule": "检查保护条件",
                "custom_review_mode": "new_revision_full",
            },
            "DS-99 自定义审查材料状态缺失",
        ),
        (
            {"custom_review_rule": "不应脱离 DS-99 保存"},
            "未选择 DS-99 时不能提交自定义审查输入",
        ),
    ],
)
def test_create_review_task_validates_ds99_snapshot_contract(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    payload: dict[str, object],
    detail: str,
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "svn://repo/custom-review",
            "old_revision": "r12",
            "new_revision": "r14",
            **payload,
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == detail


def test_read_legacy_review_task_uses_database_default_for_review_item_ids(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    me_response = client.get(
        f"{settings.API_V1_STR}/users/me",
        headers=normal_user_token_headers,
    )
    assert me_response.status_code == 200
    task_id = uuid4()
    db.execute(
        text(
            """
            INSERT INTO review_task (
                id, user_id, svn_url, old_revision, new_revision, status
            ) VALUES (
                :id, :user_id, :svn_url, :old_revision, :new_revision, 'success'
            )
            """
        ),
        {
            "id": task_id,
            "user_id": UUID(me_response.json()["id"]),
            "svn_url": "svn://repo/legacy",
            "old_revision": "r0",
            "new_revision": "r14",
        },
    )
    db.commit()

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task_id}",
        headers=normal_user_token_headers,
    )

    assert read_response.status_code == 200
    assert read_response.json()["review_item_ids"] == []


def test_create_review_task_legacy_payload_returns_new_fields_as_empty(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "svn://repo/project",
            "old_revision": "12",
            "new_revision": "14",
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["review_item_ids"] == []
    assert content["custom_review_rule"] is None
    assert content["custom_review_mode"] is None
    assert content["has_custom_review_word_material"] is None
    assert content["has_custom_review_excel_material"] is None
    assert content["has_custom_review_pdf_material"] is None
    assert content["department_id"] is None
    assert content["department_name"] is None
    assert content["focus_direction_id"] is None
    assert content["focus_direction_snapshot"] is None
    assert content["requirement_file_name"] is None
    assert content["requirement_file_size"] is None
    assert content["requirement_file_sha256"] is None


def test_create_review_task_stores_department_and_focus_snapshot(
    client: TestClient,
    db: Session,
) -> None:
    department = Department(name=f"审查部门-{random_lower_string()}")
    db.add(department)
    db.commit()
    db.refresh(department)
    user, api_key = create_key_user(db)
    user.department_id = department.id
    db.add(user)
    db.commit()
    db.refresh(user)
    headers = authentication_token_from_key(client=client, key=api_key.key, db=db)
    focus_snapshot = "重点检查权限、缓存一致性和异常处理"

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=headers,
        json={
            "svn_url": "svn://repo/project",
            "old_revision": "12",
            "new_revision": "14",
            "focus_direction_snapshot": focus_snapshot,
        },
    )

    assert response.status_code == 200
    content = response.json()
    assert content["department_id"] == str(department.id)
    assert content["department_name"] == department.name
    assert content["focus_direction_snapshot"] == focus_snapshot

    db_task = db.get(ReviewTask, UUID(content["id"]))
    assert db_task is not None
    assert db_task.department_id == department.id
    assert db_task.department_name == department.name
    assert db_task.focus_direction_snapshot == focus_snapshot


def test_create_review_task_rejects_same_revision(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "1",
            "new_revision": "r1",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "旧版本号必须小于新版本号"


def test_create_review_task_rejects_reverse_revision_range(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    response = client.post(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
        json={
            "svn_url": "https://svn.example.com/project/trunk",
            "old_revision": "r200",
            "new_revision": "r100",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "旧版本号必须小于新版本号"


def test_upload_and_read_review_task_report(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("report.md", b"# Review\n\nLooks good.\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    uploaded = upload_response.json()
    assert uploaded["status"] == "success"
    assert uploaded["report_path"]

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
    )
    assert read_response.status_code == 200
    assert read_response.headers["content-type"].startswith("text/markdown")
    assert read_response.text == "# Review\n\nLooks good.\n"


def test_新协议读取报告保留_CRLF_换行以维持问题快照定位(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    report_content = (
        "### 高风险：测试问题\r\n\r\n问题描述：用于验证 CRLF 报告的问题快照定位。\r\n"
    ).encode()
    report_markdown = report_content.decode()
    issue = _valid_issue_snapshot()
    issue["location"] = {
        "start_offset": 0,
        "end_offset": len(report_markdown),
        "feedback_insertion_offset": len(report_markdown),
        "anchor_line": 0,
        "heading_level": 3,
    }

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
        },
        files={"upload": ("report.md", report_content, "text/markdown")},
    )

    assert upload_response.status_code == 200
    assert (tmp_path / f"{task['id']}.md").read_bytes() == report_content

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
    )

    assert read_response.status_code == 200
    assert read_response.content == report_content


def test_新协议上传后返回有序问题快照和反馈进度(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issues = [
        {
            "issue_text": "高风险：空指针解引用",
            "description_text": "未校验返回值即解引用。",
            "display_snapshot": "### 高风险：空指针解引用\n\n未校验返回值即解引用。",
            "order": 0,
            "source": "template_1_heading",
            "location": {
                "start_offset": 0,
                "end_offset": 8,
                "feedback_insertion_offset": 8,
                "anchor_line": 0,
                "heading_level": 3,
            },
        },
        {
            "issue_text": "中风险：数组越界",
            "description_text": "循环上界包含数组长度。",
            "display_snapshot": "### 中风险：数组越界\n\n循环上界包含数组长度。",
            "order": 1,
            "source": "template_1_heading",
            "location": {
                "start_offset": 0,
                "end_offset": 8,
                "feedback_insertion_offset": 8,
                "anchor_line": 0,
                "heading_level": 3,
            },
        },
    ]

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps(issues, ensure_ascii=False),
        },
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )

    assert upload_response.status_code == 200
    uploaded = upload_response.json()
    assert uploaded["feedback_protocol_version"] == 1
    assert uploaded["issue_count"] == 2
    assert uploaded["evaluated_issue_count"] == 0
    assert uploaded["has_issue_feedback_draft"] is False
    assert uploaded["highest_issue_severity"] == "high"

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )

    assert issues_response.status_code == 200
    content = issues_response.json()
    assert content["count"] == 2
    assert content["evaluated_count"] == 0
    assert [item["order"] for item in content["data"]] == [0, 1]
    assert [item["issue_text"] for item in content["data"]] == [
        "高风险：空指针解引用",
        "中风险：数组越界",
    ]
    assert all(item["id"] for item in content["data"])
    assert content["data"][0]["description_text"] == "未校验返回值即解引用。"
    assert content["data"][0]["display_snapshot"].startswith("### 高风险")
    assert content["data"][0]["location"] == issues[0]["location"]

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["issue_count"] == 2
    assert detail_response.json()["evaluated_issue_count"] == 0
    assert detail_response.json()["has_issue_feedback_draft"] is False
    assert detail_response.json()["highest_issue_severity"] == "high"

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["feedback_protocol_version"] == 1
    assert listed_task["issue_count"] == 2
    assert listed_task["evaluated_issue_count"] == 0
    assert listed_task["has_issue_feedback_draft"] is False
    assert listed_task["highest_issue_severity"] == "high"


@pytest.mark.parametrize(
    ("issue_text", "expected_severity"),
    [
        ("P01 [高风险] 固定模板问题", "high"),
        ("P01 [中风险] 固定模板问题", "medium"),
        ("P01 [低风险] 固定模板问题", "low"),
        ("P01 [建议] 固定模板问题", "suggestion"),
        ("CR-001 [高] 真实服务端固定模板问题", "high"),
        ("CR-004 [中] 真实服务端固定模板问题", "medium"),
        ("CR-007 [低] 真实服务端固定模板问题", "low"),
        ("CR-004 [中] 枚举值 [高] 未处理", "medium"),
        ("1. P01 [高风险] 外层数字序号历史问题", "high"),
        ("一、CR-001 [中] 外层中文序号历史问题", "medium"),
        ("P01 [高风险]", None),
        ("CR-001 [中]", None),
        ("PC-BMS-001 [严重]", None),
        ("DS-02-001 [低]", None),
        ("1. P01 [高风险]", None),
        ("一、CR-001 [中]", None),
        ("PC-BMS-001 [严重] 历史平台报告问题", "high"),
        ("PC-BMS-002 [重要] 历史平台报告问题", "medium"),
        ("PC-BMS-003 [一般] 历史平台报告问题", "low"),
        ("PC-BMS-004 [建议] 历史平台报告问题", "suggestion"),
        ("PC-BMS-005 [高风险] 历史风险词问题", "high"),
        ("PC-BMS-006 [高] 历史风险词问题", "high"),
        ("PC-BMS-007 [中风险] 历史风险词问题", "medium"),
        ("PC-BMS-008 [中] 历史风险词问题", "medium"),
        ("PC-BMS-009 [低风险] 历史风险词问题", "low"),
        ("PC-BMS-010 [低] 历史风险词问题", "low"),
        ("DS-02-001 [高风险] 历史审查项编号问题", "high"),
        ("DS-02-002 [高] 历史审查项编号问题", "high"),
        ("DS-02-003 [中风险] 历史审查项编号问题", "medium"),
        ("DS-02-004 [中] 历史审查项编号问题", "medium"),
        ("DS-02-005 [低风险] 历史审查项编号问题", "low"),
        ("DS-02-006 [低] 历史审查项编号问题", "low"),
        ("DS-02-007 [建议] 历史审查项编号问题", "suggestion"),
        ("P01 [重要] 错误交叉组合", None),
        ("CR-001 [严重] 错误交叉组合", None),
        ("DS-02-008 [严重] 错误交叉组合", None),
        ("DS-03-001 [高] 其他审查项编号", None),
        ("PC-BMS-01 [高] 合同固定前的非三位编号", "high"),
        ("高风险：旧模板问题", "high"),
        ("中风险：旧模板问题", "medium"),
        ("中风险：正文提到高风险：但标题仍为中风险", "medium"),
        ("低风险：旧模板问题", "low"),
        ("建议：旧模板问题", "suggestion"),
    ],
)
def test_列表和详情聚合新协议问题快照的最高严重度(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    issue_text: str,
    expected_severity: str | None,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue = _valid_issue_snapshot()
    issue["issue_text"] = issue_text

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
        },
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    assert upload_response.json()["highest_issue_severity"] == expected_severity

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["highest_issue_severity"] == expected_severity

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["highest_issue_severity"] == expected_severity


def test_历史编号问题的列表详情与反馈进度保持同步(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue_texts = [
        "DS-02-001 [高] FAULT 分支条件由比较表达式误改为赋值表达式",
        "DS-02-002 [中] 新增 SLEEP 状态分支的 if 缺少 else 分支",
        "DS-02-003 [低] 新增控制语句未采用 Allman 大括号风格",
    ]
    issues = []
    for order, issue_text in enumerate(issue_texts):
        issue = _valid_issue_snapshot(order)
        issue["issue_text"] = issue_text
        issues.append(issue)

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps(issues, ensure_ascii=False),
        },
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    assert upload_response.json()["highest_issue_severity"] == "high"

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    stored_issues = issues_response.json()["data"]
    assert [issue["issue_text"] for issue in stored_issues] == issue_texts

    draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{stored_issues[1]['id']}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": True, "feedback_text": "已核对中风险问题"},
    )
    assert draft_response.status_code == 200

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["issue_count"] == 3
    assert detail_response.json()["evaluated_issue_count"] == 1
    assert detail_response.json()["has_issue_feedback_draft"] is True
    assert detail_response.json()["highest_issue_severity"] == "high"

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["issue_count"] == 3
    assert listed_task["evaluated_issue_count"] == 1
    assert listed_task["has_issue_feedback_draft"] is True
    assert listed_task["highest_issue_severity"] == "high"


def test_无严重度_BMS_快照在上传列表和详情中保留且最高严重度为空(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue_texts = [
        "STATE-17 休眠高风险路径仍可打开充电 MOS",
        "编号-任意-42 建议路径缺少故障恢复一致性校验",
    ]
    issues = []
    for order, issue_text in enumerate(issue_texts):
        issue = _valid_issue_snapshot(order)
        issue["issue_text"] = issue_text
        issue["source"] = "bms_code_reviewer_heading"
        issues.append(issue)

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps(issues, ensure_ascii=False),
        },
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )

    assert upload_response.status_code == 200
    assert upload_response.json()["issue_count"] == 2
    assert upload_response.json()["highest_issue_severity"] is None
    assert (tmp_path / f"{task['id']}.md").exists()

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    assert issues_response.json()["count"] == 2
    assert [
        item["issue_text"] for item in issues_response.json()["data"]
    ] == issue_texts

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["issue_count"] == 2
    assert detail_response.json()["highest_issue_severity"] is None

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["issue_count"] == 2
    assert listed_task["evaluated_issue_count"] == 0
    assert listed_task["has_issue_feedback_draft"] is False
    assert listed_task["highest_issue_severity"] is None


@pytest.mark.parametrize(
    "issue_texts",
    [
        ["STATE-17 任意标题", "编号-任意-42 标题含建议但不声明严重度"],
        ["P01 [高风险] 历史 P 编号问题", "CR-002 [建议] 历史 CR 编号建议"],
        ["DS-02-001 [高] 历史 DS-02 问题", "DS-02-002 [低] 历史低风险问题"],
        ["PC-BMS-001 [严重] canonical 问题", "PC-BMS-002 [一般] canonical 一般问题"],
    ],
)
def test_全量_BMS_source_不依赖标题格式并保持一致(
    issue_texts: list[str],
) -> None:
    issues = []
    for order, issue_text in enumerate(issue_texts):
        issue = _valid_issue_snapshot(order)
        issue["issue_text"] = issue_text
        issue["source"] = "bms_code_reviewer_heading"
        issues.append(ReviewTaskIssueSnapshotCreate.model_validate(issue))

    assert (
        review_tasks_route._has_consistent_review_issue_sources(
            tuple(issue.source for issue in issues)
        )
        is True
    )


def test_上传混合source保留报告资产和快照但在反馈入口统一安全降级(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    bms_issue = _valid_issue_snapshot(0)
    bms_issue["issue_text"] = "STATE-ISSUE53-MIXED 无严重度问题"
    bms_issue["source"] = "bms_code_reviewer_heading"
    legacy_issue = _valid_issue_snapshot(1)
    legacy_issue["issue_text"] = "高风险：ISSUE53-MIXED 旧模板问题"
    report_content = b"# Review\n\n![evidence](images/evidence.png)\n"

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([bms_issue, legacy_issue], ensure_ascii=False),
            "asset_paths": json.dumps(["images/evidence.png"]),
        },
        files=[
            ("upload", ("report.md", report_content, "text/markdown")),
            ("assets", ("evidence.png", VALID_PNG_BYTES, "image/png")),
        ],
    )

    assert upload_response.status_code == 200
    assert upload_response.json()["feedback_protocol_version"] == 1
    assert upload_response.json()["issue_count"] == 0
    assert upload_response.json()["highest_issue_severity"] is None
    assert (tmp_path / f"{task['id']}.md").read_bytes() == report_content
    assert (
        report_service.asset_path_for_task(
            UUID(task["id"]), "images/evidence.png"
        ).read_bytes()
        == VALID_PNG_BYTES
    )

    task_id = UUID(task["id"])
    stored_issues = db.exec(
        select(ReviewTaskIssue)
        .where(ReviewTaskIssue.task_id == task_id)
        .order_by(ReviewTaskIssue.order)
    ).all()
    assert [issue.issue_text for issue in stored_issues] == [
        bms_issue["issue_text"],
        legacy_issue["issue_text"],
    ]

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["issue_count"] == 0
    assert detail_response.json()["highest_issue_severity"] is None

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["issue_count"] == 0
    assert listed_task["highest_issue_severity"] is None

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    assert issues_response.json()["data"] == []
    assert issues_response.json()["count"] == 0

    draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{stored_issues[0].id}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": False, "feedback_text": "不得编辑混合 source"},
    )
    assert draft_response.status_code == 404
    assert draft_response.json()["detail"] == "Review task issue not found"

    pdf_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={"upload": ("混合source.pdf", b"%PDF-1.4\nmixed", "application/pdf")},
    )
    assert pdf_response.status_code == 409
    assert pdf_response.json()["detail"] == "issue_feedback_not_supported"
    assert not (tmp_path / f"{task['id']}.pdf").exists()

    export_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
    )
    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.content))
    rows = list(workbook["问题反馈明细"].iter_rows(values_only=True))
    workbook.close()
    exported_issue_texts = {str(row[5]) for row in rows[1:]}
    assert str(bms_issue["issue_text"]) not in exported_issue_texts
    assert str(legacy_issue["issue_text"]) not in exported_issue_texts


def test_数据库历史混合_source_快照在六个反馈入口统一安全降级(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    task_id = UUID(task["id"])
    task_record = db.get(ReviewTask, task_id)
    assert task_record is not None
    task_record.feedback_protocol_version = 1
    db.add(task_record)
    issue_texts = [
        "STATE-17 无严重度 BMS 问题",
        "高风险：混入旧模板 source 的问题",
    ]
    for order, issue_text in enumerate(issue_texts):
        snapshot = _valid_issue_snapshot(order)
        db.add(
            ReviewTaskIssue(
                task_id=task_id,
                issue_text=issue_text,
                description_text=str(snapshot["description_text"]),
                display_snapshot=str(snapshot["display_snapshot"]),
                order=order,
                source=(
                    "bms_code_reviewer_heading" if order == 0 else "template_1_heading"
                ),
                location=snapshot["location"],
                helpful=True if order == 0 else None,
                feedback_text="历史反馈草稿" if order == 0 else None,
            )
        )
    db.commit()

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["issue_count"] == 0
    assert detail_response.json()["evaluated_issue_count"] == 0
    assert detail_response.json()["has_issue_feedback_draft"] is False
    assert detail_response.json()["highest_issue_severity"] is None

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["issue_count"] == 0
    assert listed_task["evaluated_issue_count"] == 0
    assert listed_task["has_issue_feedback_draft"] is False
    assert listed_task["highest_issue_severity"] is None

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    assert issues_response.json() == {
        "data": [],
        "count": 0,
        "evaluated_count": 0,
        "feedback_submitted_at": None,
    }

    stored_issues = db.exec(
        select(ReviewTaskIssue)
        .where(ReviewTaskIssue.task_id == task_id)
        .order_by(ReviewTaskIssue.order)
    ).all()
    assert len(stored_issues) == 2
    task_record.feedback_submitted_at = datetime.now(timezone.utc)
    db.add(task_record)
    db.commit()

    draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{stored_issues[0].id}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": False, "feedback_text": "不得编辑无效历史快照"},
    )
    assert draft_response.status_code == 404
    assert draft_response.json()["detail"] == "Review task issue not found"

    pdf_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={
            "upload": (
                "无效历史反馈.pdf",
                b"%PDF-1.4\ninvalid-history",
                "application/pdf",
            )
        },
    )
    assert pdf_response.status_code == 409
    assert pdf_response.json()["detail"] == "issue_feedback_not_supported"
    assert not (tmp_path / f"{task['id']}.pdf").exists()

    export_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
    )
    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.content))
    rows = list(workbook["问题反馈明细"].iter_rows(values_only=True))
    workbook.close()
    exported_issue_texts = {str(row[5]) for row in rows[1:]}
    assert exported_issue_texts.isdisjoint(issue_texts)

    db.expire_all()
    preserved_issues = db.exec(
        select(ReviewTaskIssue)
        .where(ReviewTaskIssue.task_id == task_id)
        .order_by(ReviewTaskIssue.order)
    ).all()
    assert len(preserved_issues) == 2
    assert preserved_issues[0].feedback_text == "历史反馈草稿"


def test_有效零问题新协议仍可读取空列表并提交最终_PDF(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={"feedback_protocol_version": "1", "issues": "[]"},
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    assert issues_response.json()["count"] == 0

    pdf_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={
            "upload": ("零问题报告.pdf", b"%PDF-1.4\nzero-issues", "application/pdf")
        },
    )
    assert pdf_response.status_code == 200
    assert pdf_response.json()["feedback_submitted_at"] is not None
    assert (tmp_path / f"{task['id']}.pdf").exists()


@pytest.mark.parametrize(
    ("case", "expected_detail"),
    [
        ("JSON总大小", "问题快照总大小超过限制"),
        ("问题数量", "问题快照数量超过限制"),
        ("单字段长度", "问题快照格式无效"),
        ("定位越界", "问题快照定位范围无效"),
        ("代理对中间", "问题快照定位范围无效"),
        ("行锚点不一致", "问题快照定位范围无效"),
    ],
)
def test_新协议上传执行问题快照资源和定位限制(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    case: str,
    expected_detail: str,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue = {
        "issue_text": "高风险：测试问题",
        "description_text": "用于验证服务端限制。",
        "display_snapshot": "### 高风险：测试问题",
        "order": 0,
        "source": "template_1_heading",
        "location": {
            "start_offset": 0,
            "end_offset": 8,
            "feedback_insertion_offset": 8,
            "anchor_line": 0,
            "heading_level": 3,
        },
    }
    report_content = "# Review\n"
    if case == "JSON总大小":
        monkeypatch.setattr(review_tasks_route, "MAX_REPORT_ISSUES_JSON_BYTES", 1)
    elif case == "问题数量":
        monkeypatch.setattr(review_tasks_route, "MAX_REPORT_ISSUE_COUNT", 0)
    elif case == "单字段长度":
        issue["issue_text"] = "x" * (8 * 1024 + 1)
    elif case == "定位越界":
        issue["location"] = {
            **issue["location"],
            "end_offset": 10,
            "feedback_insertion_offset": 10,
        }
    elif case == "代理对中间":
        report_content = "# 😀\n"
        issue["location"] = {
            **issue["location"],
            "end_offset": 3,
            "feedback_insertion_offset": 3,
        }
    elif case == "行锚点不一致":
        issue["location"] = {
            **issue["location"],
            "anchor_line": 1,
        }

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
        },
        files={"upload": ("report.md", report_content.encode(), "text/markdown")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == expected_detail
    assert not (tmp_path / f"{task['id']}.md").exists()


def test_问题快照解析在数量和_JSON_精确边界成功且超限失败(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    at_count_limit = [
        _valid_issue_snapshot(order)
        for order in range(review_tasks_route.MAX_REPORT_ISSUE_COUNT)
    ]
    assert (
        len(
            review_tasks_route._parse_issue_snapshots(
                json.dumps(at_count_limit, ensure_ascii=False)
            )
        )
        == review_tasks_route.MAX_REPORT_ISSUE_COUNT
    )

    over_count_limit = [
        *at_count_limit,
        _valid_issue_snapshot(review_tasks_route.MAX_REPORT_ISSUE_COUNT),
    ]
    with pytest.raises(HTTPException, match="问题快照数量超过限制"):
        review_tasks_route._parse_issue_snapshots(
            json.dumps(over_count_limit, ensure_ascii=False)
        )

    raw_issue = json.dumps([_valid_issue_snapshot()], ensure_ascii=False)
    raw_issue_bytes = len(raw_issue.encode("utf-8"))
    monkeypatch.setattr(
        review_tasks_route,
        "MAX_REPORT_ISSUES_JSON_BYTES",
        raw_issue_bytes,
    )
    assert len(review_tasks_route._parse_issue_snapshots(raw_issue)) == 1
    monkeypatch.setattr(
        review_tasks_route,
        "MAX_REPORT_ISSUES_JSON_BYTES",
        raw_issue_bytes - 1,
    )
    with pytest.raises(HTTPException, match="问题快照总大小超过限制"):
        review_tasks_route._parse_issue_snapshots(raw_issue)


@pytest.mark.parametrize(
    ("field_name", "max_length"),
    [
        ("issue_text", MAX_REVIEW_TASK_ISSUE_TEXT_LENGTH),
        ("description_text", MAX_REVIEW_TASK_ISSUE_DESCRIPTION_LENGTH),
        ("display_snapshot", MAX_REVIEW_TASK_ISSUE_DISPLAY_SNAPSHOT_LENGTH),
    ],
)
def test_问题快照文本字段在精确长度边界成功且超限失败(
    field_name: str,
    max_length: int,
) -> None:
    payload = _valid_issue_snapshot()
    payload[field_name] = "x" * max_length
    ReviewTaskIssueSnapshotCreate.model_validate(payload)

    payload[field_name] = "x" * (max_length + 1)
    with pytest.raises(ValidationError):
        ReviewTaskIssueSnapshotCreate.model_validate(payload)


def test_新协议上传按_UTF16_code_unit_验证问题定位(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue = _valid_issue_snapshot()
    issue["location"] = {
        "start_offset": 0,
        "end_offset": 5,
        "feedback_insertion_offset": 5,
        "anchor_line": 0,
        "heading_level": 3,
    }

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
        },
        files={"upload": ("report.md", "# 😀\n".encode(), "text/markdown")},
    )

    assert response.status_code == 200


def test_新协议上传保持相对图片目录结构(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    report_content = "# Review\n\n![证据](images/nested/evidence.png)\n".encode()

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps(["images/nested/evidence.png"]),
        },
        files=[
            ("upload", ("report.md", report_content, "text/markdown")),
            ("assets", ("evidence.png", VALID_PNG_BYTES, "image/png")),
        ],
    )

    assert upload_response.status_code == 200

    asset_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/assets/images/nested/evidence.png",
        headers=normal_user_token_headers,
    )

    assert asset_response.status_code == 200
    assert asset_response.content == VALID_PNG_BYTES
    assert asset_response.headers["content-type"].startswith("image/png")


def test_新协议服务端按共享CommonMark语料识别图片(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    corpus = _load_shared_markdown_image_corpus()
    report_content = str(corpus["markdown"])
    asset_paths = list(corpus["expected_paths"])

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps(asset_paths, ensure_ascii=False),
        },
        files=[
            (
                "upload",
                ("report.md", report_content.encode(), "text/markdown"),
            ),
            *[
                ("assets", (Path(asset_path).name, VALID_PNG_BYTES, "image/png"))
                for asset_path in asset_paths
            ],
        ],
    )

    assert response.status_code == 200


def test_删除新协议任务同步清理问题记录和图片资产(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue = {
        "issue_text": "高风险：资源泄漏",
        "description_text": "异常路径未释放句柄。",
        "display_snapshot": "### 高风险：资源泄漏\n\n异常路径未释放句柄。",
        "order": 0,
        "source": "template_1_heading",
        "location": {
            "start_offset": 0,
            "end_offset": 28,
            "feedback_insertion_offset": 28,
            "anchor_line": 0,
            "heading_level": 3,
        },
    }
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
            "asset_paths": json.dumps(["images/evidence.png"]),
        },
        files=[
            (
                "upload",
                (
                    "report.md",
                    b"# Review\n\n![evidence](images/evidence.png)\n",
                    "text/markdown",
                ),
            ),
            ("assets", ("evidence.png", VALID_PNG_BYTES, "image/png")),
        ],
    )
    assert upload_response.status_code == 200
    assert list(tmp_path.rglob("evidence.png"))

    delete_response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert delete_response.status_code == 200
    assert (
        db.exec(
            select(ReviewTaskIssue).where(ReviewTaskIssue.task_id == UUID(task["id"]))
        ).all()
        == []
    )
    assert list(tmp_path.rglob("evidence.png")) == []


def test_删除任务数据库提交失败时恢复任务和文件(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps(["images/evidence.png"]),
        },
        files=[
            (
                "upload",
                (
                    "report.md",
                    b"# Review\n\n![evidence](images/evidence.png)\n",
                    "text/markdown",
                ),
            ),
            ("assets", ("evidence.png", VALID_PNG_BYTES, "image/png")),
        ],
    )
    assert upload_response.status_code == 200
    original_commit = Session.commit

    def fail_commit(_session: Session) -> None:
        raise RuntimeError("模拟删除提交失败")

    monkeypatch.setattr(Session, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="模拟删除提交失败"):
        client.delete(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}",
            headers=normal_user_token_headers,
        )
    monkeypatch.setattr(Session, "commit", original_commit)

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    asset_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/assets/images/evidence.png",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert asset_response.status_code == 200
    assert asset_response.content == VALID_PNG_BYTES


def test_删除任务清理失败会留痕并由后续删除重试(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    first_task = _create_review_task(client, normal_user_token_headers)
    second_task = _create_review_task(client, normal_user_token_headers)
    for task in (first_task, second_task):
        response = client.post(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
            headers=normal_user_token_headers,
            data={"feedback_protocol_version": "1", "issues": "[]"},
            files={"upload": ("report.md", b"# Review\n", "text/markdown")},
        )
        assert response.status_code == 200

    original_purge = report_service.purge_quarantined_task_files

    def fail_purge(_quarantine: report_service.QuarantinedTaskFiles) -> None:
        raise OSError("模拟隔离文件清理失败")

    caplog.set_level(logging.ERROR)
    monkeypatch.setattr(report_service, "purge_quarantined_task_files", fail_purge)
    first_delete = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{first_task['id']}",
        headers=normal_user_token_headers,
    )
    monkeypatch.setattr(
        report_service,
        "purge_quarantined_task_files",
        original_purge,
    )

    assert first_delete.status_code == 200
    assert "待重试清理目录" in caplog.text
    assert report_service.pending_delete_trash_paths()

    second_delete = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{second_task['id']}",
        headers=normal_user_token_headers,
    )

    assert second_delete.status_code == 200
    assert report_service.pending_delete_trash_paths() == []


def test_删除任务进程中断后会恢复数据库仍存在任务的文件(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    interrupted_task = _create_review_task(client, normal_user_token_headers)
    trigger_task = _create_review_task(client, normal_user_token_headers)
    report_content = b"# interrupted deletion\n"
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{interrupted_task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("report.md", report_content, "text/markdown")},
    )
    assert upload_response.status_code == 200
    report_path = Path(upload_response.json()["report_path"])

    report_service.quarantine_task_files(
        UUID(interrupted_task["id"]),
        [report_path],
    )
    assert not report_path.exists()

    trigger_delete = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{trigger_task['id']}",
        headers=normal_user_token_headers,
    )
    recovered_report = client.get(
        f"{settings.API_V1_STR}/review-tasks/{interrupted_task['id']}/report",
        headers=normal_user_token_headers,
    )

    assert trigger_delete.status_code == 200
    assert recovered_report.status_code == 200
    assert recovered_report.content == report_content
    assert report_service.pending_delete_trash_paths() == []


@pytest.mark.parametrize(
    ("asset_path", "asset_content", "expected_detail"),
    [
        ("/images/evidence.png", b"png", "图片资产路径无效"),
        ("C:/images/evidence.png", b"png", "图片资产路径无效"),
        ("images/../evidence.png", b"png", "图片资产路径无效"),
        ("attachments/evidence.png", b"png", "图片资产必须位于 images 目录"),
        ("images/evidence.png", b"", "图片资产不能为空"),
    ],
    ids=["绝对路径", "Windows绝对路径", "目录穿越", "图片目录外", "空资产"],
)
def test_新协议上传拒绝不安全或空图片资产(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    asset_path: str,
    asset_content: bytes,
    expected_detail: str,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps([asset_path]),
        },
        files=[
            (
                "upload",
                (
                    "report.md",
                    f"# Review\n\n![evidence]({asset_path})\n".encode(),
                    "text/markdown",
                ),
            ),
            ("assets", ("evidence.png", asset_content, "image/png")),
        ],
    )

    assert response.status_code == 400
    assert response.json()["detail"] == expected_detail
    assert not (tmp_path / f"{task['id']}.md").exists()


@pytest.mark.parametrize(
    ("report_content", "asset_paths", "uploaded_files", "expected_detail"),
    [
        (
            b"# Review\n\n![evidence](images/evidence.png)\n",
            [],
            [],
            "Markdown 引用的图片资产未完整上传",
        ),
        (
            b"# Review\n",
            ["images/evidence.png"],
            [("assets", ("evidence.png", b"png", "image/png"))],
            "存在未被 Markdown 引用的图片资产",
        ),
        (
            b"# Review\n\n![evidence](images/evidence.png)\n",
            ["images/evidence.png"],
            [("assets", ("evidence.png", b"png", "image/jpeg"))],
            "报告图片格式与内容类型不一致",
        ),
    ],
    ids=["缺失引用资产", "额外资产", "内容类型不匹配"],
)
def test_新协议上传校验Markdown图片引用和格式(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    report_content: bytes,
    asset_paths: list[str],
    uploaded_files: list[tuple[str, tuple[str, bytes, str]]],
    expected_detail: str,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps(asset_paths),
        },
        files=[
            ("upload", ("report.md", report_content, "text/markdown")),
            *uploaded_files,
        ],
    )

    assert response.status_code == 400
    assert response.json()["detail"] == expected_detail
    assert not (tmp_path / f"{task['id']}.md").exists()


@pytest.mark.parametrize(
    "image_case",
    VALID_IMAGE_CASES,
    ids=[
        f"{image_case['name']}-{index}"
        for index, image_case in enumerate(VALID_IMAGE_CASES)
    ],
)
def test_新协议上传接受内容与扩展名一致的真实图片(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    image_case: dict[str, Any],
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    asset_path = f"images/evidence{image_case['extension']}"
    asset_content = base64.b64decode(image_case["valid_base64"])

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps([asset_path]),
        },
        files=[
            (
                "upload",
                (
                    "report.md",
                    f"# Review\n\n![evidence]({asset_path})\n".encode(),
                    "text/markdown",
                ),
            ),
            (
                "assets",
                (Path(asset_path).name, asset_content, image_case["content_type"]),
            ),
        ],
    )

    assert response.status_code == 200


@pytest.mark.parametrize(
    "image_case",
    INVALID_IMAGE_CASES,
    ids=[
        f"{image_case['name']}-{index}"
        for index, image_case in enumerate(INVALID_IMAGE_CASES)
    ],
)
def test_新协议上传拒绝只有正确文件头或含主动内容的图片(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    image_case: dict[str, Any],
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    asset_path = f"images/evidence{image_case['extension']}"
    asset_content = base64.b64decode(image_case["invalid_base64"])

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps([asset_path]),
        },
        files=[
            (
                "upload",
                (
                    "report.md",
                    f"# Review\n\n![evidence]({asset_path})\n".encode(),
                    "text/markdown",
                ),
            ),
            (
                "assets",
                (
                    Path(asset_path).name,
                    asset_content,
                    image_case["content_type"],
                ),
            ),
        ],
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "报告图片内容与文件格式不一致"


@pytest.mark.parametrize(
    "image_case",
    ANIMATED_IMAGE_CASES,
    ids=[image_case["name"] for image_case in ANIMATED_IMAGE_CASES],
)
def test_新协议上传拒绝动画图片(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    image_case: dict[str, str],
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    asset_path = f"images/evidence{image_case['extension']}"

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps([asset_path]),
        },
        files=[
            (
                "upload",
                (
                    "report.md",
                    f"# Review\n\n![evidence]({asset_path})\n".encode(),
                    "text/markdown",
                ),
            ),
            (
                "assets",
                (
                    Path(asset_path).name,
                    base64.b64decode(image_case["animated_base64"]),
                    image_case["content_type"],
                ),
            ),
        ],
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "报告图片内容与文件格式不一致"


@pytest.mark.parametrize(
    ("limit_name", "limit", "report_content", "asset_paths", "asset_files", "detail"),
    [
        (
            "MAX_REPORT_MARKDOWN_BYTES",
            8,
            b"123456789",
            [],
            [],
            "审查报告超过大小限制",
        ),
        (
            "MAX_REPORT_IMAGE_COUNT",
            0,
            b"# Review\n\n![evidence](images/evidence.png)\n",
            ["images/evidence.png"],
            [("assets", ("evidence.png", VALID_PNG_BYTES, "image/png"))],
            "报告图片数量超过限制",
        ),
        (
            "MAX_REPORT_IMAGE_BYTES",
            len(VALID_PNG_BYTES) - 1,
            b"# Review\n\n![evidence](images/evidence.png)\n",
            ["images/evidence.png"],
            [("assets", ("evidence.png", VALID_PNG_BYTES, "image/png"))],
            "单个报告图片超过大小限制",
        ),
        (
            "MAX_REPORT_ASSETS_TOTAL_BYTES",
            len(VALID_PNG_BYTES) - 1,
            b"# Review\n\n![evidence](images/evidence.png)\n",
            ["images/evidence.png"],
            [("assets", ("evidence.png", VALID_PNG_BYTES, "image/png"))],
            "报告图片总大小超过限制",
        ),
    ],
    ids=["Markdown大小", "图片数量", "单图大小", "图片总大小"],
)
def test_新协议上传执行资源限制(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    limit_name: str,
    limit: int,
    report_content: bytes,
    asset_paths: list[str],
    asset_files: list[tuple[str, tuple[str, bytes, str]]],
    detail: str,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    monkeypatch.setattr(report_service, limit_name, limit)
    task = _create_review_task(client, normal_user_token_headers)

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": "[]",
            "asset_paths": json.dumps(asset_paths),
        },
        files=[
            ("upload", ("report.md", report_content, "text/markdown")),
            *asset_files,
        ],
    )

    assert response.status_code == 400
    assert response.json()["detail"] == detail


def test_新协议报告上传后不可重复覆盖或降级(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue = {
        "issue_text": "高风险：稳定问题",
        "description_text": "第一次上传的问题描述。",
        "display_snapshot": "### 高风险：稳定问题",
        "order": 0,
        "source": "template_1_heading",
        "location": {
            "start_offset": 0,
            "end_offset": 8,
            "feedback_insertion_offset": 8,
            "anchor_line": 0,
            "heading_level": 3,
        },
    }
    original_report = b"# Original\n\n![evidence](images/evidence.png)\n"
    first_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
            "asset_paths": json.dumps(["images/evidence.png"]),
        },
        files=[
            ("upload", ("report.md", original_report, "text/markdown")),
            ("assets", ("evidence.png", VALID_PNG_BYTES, "image/png")),
        ],
    )
    assert first_response.status_code == 200
    issue_id = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    ).json()["data"][0]["id"]

    repeated_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={"feedback_protocol_version": "1", "issues": "[]"},
        files={"upload": ("report.md", b"# Replaced\n", "text/markdown")},
    )
    downgrade_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("legacy.md", b"# Legacy\n", "text/markdown")},
    )

    assert repeated_response.status_code == 409
    assert downgrade_response.status_code == 409
    assert (tmp_path / f"{task['id']}.md").read_bytes() == original_report
    assert list(tmp_path.rglob("evidence.png"))[0].read_bytes() == VALID_PNG_BYTES
    issues = db.exec(
        select(ReviewTaskIssue).where(ReviewTaskIssue.task_id == UUID(task["id"]))
    ).all()
    assert len(issues) == 1
    assert str(issues[0].id) == issue_id


def test_新协议上传清理上次崩溃遗留的孤儿报告包后允许重试(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    orphan_report_path = tmp_path / f"{task['id']}.md"
    orphan_asset_path = tmp_path / task["id"] / "assets" / "images" / "old.png"
    orphan_staging_path = tmp_path / f".{task['id']}-report-interrupted" / "report.md"
    orphan_asset_path.parent.mkdir(parents=True)
    orphan_staging_path.parent.mkdir(parents=True)
    orphan_report_path.write_bytes(b"# orphan\n")
    orphan_asset_path.write_bytes(VALID_PNG_BYTES)
    orphan_staging_path.write_bytes(b"# interrupted staging\n")

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={"feedback_protocol_version": "1", "issues": "[]"},
        files={"upload": ("report.md", b"# retried report\n", "text/markdown")},
    )

    assert response.status_code == 200
    assert orphan_report_path.read_bytes() == b"# retried report\n"
    assert not orphan_asset_path.exists()
    assert not orphan_staging_path.exists()


def test_新协议并发上传只允许一个请求成功(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    def upload(index: int) -> int:
        issue = {
            "issue_text": f"并发问题 {index}",
            "description_text": "只应保存一份。",
            "display_snapshot": f"### 并发问题 {index}",
            "order": 0,
            "source": "template_1_heading",
            "location": {
                "start_offset": 0,
                "end_offset": 8,
                "feedback_insertion_offset": 8,
                "anchor_line": 0,
                "heading_level": 3,
            },
        }
        return client.post(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
            headers=normal_user_token_headers,
            data={
                "feedback_protocol_version": "1",
                "issues": json.dumps([issue], ensure_ascii=False),
            },
            files={
                "upload": (
                    "report.md",
                    f"# concurrent {index}\n".encode(),
                    "text/markdown",
                )
            },
        ).status_code

    with ThreadPoolExecutor(max_workers=2) as executor:
        statuses = list(executor.map(upload, (1, 2)))

    assert sorted(statuses) == [200, 409]
    db.expire_all()
    issues = db.exec(
        select(ReviewTaskIssue).where(ReviewTaskIssue.task_id == UUID(task["id"]))
    ).all()
    assert len(issues) == 1


def test_新协议数据库提交失败时不保留报告问题或资产(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    issue = {
        "issue_text": "高风险：事务失败",
        "description_text": "用于验证回滚。",
        "display_snapshot": "### 高风险：事务失败",
        "order": 0,
        "source": "template_1_heading",
        "location": {
            "start_offset": 0,
            "end_offset": 8,
            "feedback_insertion_offset": 8,
            "anchor_line": 0,
            "heading_level": 3,
        },
    }
    original_commit = Session.commit

    def fail_commit(_session: Session) -> None:
        raise RuntimeError("模拟数据库提交失败")

    monkeypatch.setattr(Session, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="模拟数据库提交失败"):
        client.post(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
            headers=normal_user_token_headers,
            data={"feedback_protocol_version": "1", "issues": json.dumps([issue])},
            files={"upload": ("report.md", b"# Review\n", "text/markdown")},
        )
    monkeypatch.setattr(Session, "commit", original_commit)

    assert not (tmp_path / f"{task['id']}.md").exists()
    assert not (tmp_path / task["id"] / "assets").exists()
    db.expire_all()
    assert (
        db.exec(
            select(ReviewTaskIssue).where(ReviewTaskIssue.task_id == UUID(task["id"]))
        ).all()
        == []
    )


def test_零问题报告仍标记为新协议且不创建虚假问题(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={"feedback_protocol_version": "1", "issues": "[]"},
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )

    assert upload_response.status_code == 200
    assert upload_response.json()["feedback_protocol_version"] == 1
    assert upload_response.json()["issue_count"] == 0
    assert upload_response.json()["highest_issue_severity"] is None
    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    assert issues_response.json() == {
        "data": [],
        "count": 0,
        "evaluated_count": 0,
        "feedback_submitted_at": None,
    }


def test_Electron无效BMS载荷按历史报告流程安全降级(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(
        client,
        normal_user_token_headers,
        old_revision="987654321",
        new_revision="987654322",
    )
    invalid_report = (
        "# BMS 代码审查报告\n\n"
        "## 三、审查问题\n\n"
        "### 普通说明\n\n"
        "完整字段缺失，不得声明逐问题反馈协议。\n"
    ).encode()

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("invalid-bms.md", invalid_report, "text/markdown")},
    )

    assert upload_response.status_code == 200
    content = upload_response.json()
    assert content["feedback_protocol_version"] is None
    assert content["issue_count"] == 0
    assert content["evaluated_issue_count"] == 0
    assert content["highest_issue_severity"] is None
    assert (tmp_path / f"{task['id']}.md").read_bytes() == invalid_report

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["feedback_protocol_version"] is None
    assert detail_response.json()["issue_count"] == 0

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["feedback_protocol_version"] is None
    assert listed_task["issue_count"] == 0

    issues_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert issues_response.status_code == 200
    assert issues_response.json() == {
        "data": [],
        "count": 0,
        "evaluated_count": 0,
        "feedback_submitted_at": None,
    }

    draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{uuid4()}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": False, "feedback_text": "不得建立草稿"},
    )
    assert draft_response.status_code == 409
    assert draft_response.json()["detail"] == "issue_feedback_not_supported"

    pdf_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={
            "upload": (
                "普通报告.pdf",
                b"%PDF-1.4\nlegacy-invalid-bms",
                "application/pdf",
            )
        },
    )
    assert pdf_response.status_code == 200
    assert pdf_response.json()["feedback_submitted_at"] is None
    assert (tmp_path / f"{task['id']}.pdf").exists()

    task_id = UUID(task["id"])
    task_record = db.get(ReviewTask, task_id)
    assert task_record is not None
    assert task_record.feedback_protocol_version is None
    assert task_record.feedback_submitted_at is None
    assert (
        db.exec(select(ReviewTaskIssue).where(ReviewTaskIssue.task_id == task_id)).all()
        == []
    )

    export_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
    )
    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.content))
    rows = list(workbook["问题反馈明细"].iter_rows(values_only=True))
    workbook.close()
    assert all(
        (row[2], row[3], row[4])
        != (task["svn_url"], task["old_revision"], task["new_revision"])
        for row in rows[1:]
    )

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
    )
    assert read_response.status_code == 200
    assert read_response.content == invalid_report
    download_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/download",
        headers=normal_user_token_headers,
    )
    assert download_response.status_code == 200
    assert download_response.content == invalid_report


def test_新旧报告混合列表详情和统计保持兼容(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    _, api_key = create_key_user(db)
    user_headers = authentication_token_from_key(
        client=client,
        key=api_key.key,
        db=db,
    )
    legacy_task = _create_review_task(
        client,
        user_headers,
        old_revision="501",
    )
    legacy_upload = client.post(
        f"{settings.API_V1_STR}/review-tasks/{legacy_task['id']}/report",
        headers=user_headers,
        files={"upload": ("legacy.md", b"# Legacy\n", "text/markdown")},
    )
    assert legacy_upload.status_code == 200

    modern_task = _create_review_task(
        client,
        user_headers,
        old_revision="502",
    )
    modern_issue = _valid_issue_snapshot()
    modern_issue["issue_text"] = "P01 [高风险] 固定模板问题"
    modern_upload = client.post(
        f"{settings.API_V1_STR}/review-tasks/{modern_task['id']}/report",
        headers=user_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([modern_issue], ensure_ascii=False),
        },
        files={"upload": ("modern.md", b"# Modern\n", "text/markdown")},
    )
    assert modern_upload.status_code == 200

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=user_headers,
        params={"mine": True},
    )
    assert list_response.status_code == 200
    listed_by_id = {task["id"]: task for task in list_response.json()["data"]}
    assert listed_by_id[legacy_task["id"]]["feedback_protocol_version"] is None
    assert listed_by_id[legacy_task["id"]]["issue_count"] == 0
    assert listed_by_id[legacy_task["id"]]["highest_issue_severity"] is None
    assert listed_by_id[modern_task["id"]]["feedback_protocol_version"] == 1
    assert listed_by_id[modern_task["id"]]["issue_count"] == 1
    assert listed_by_id[modern_task["id"]]["highest_issue_severity"] == "high"

    for task_id, expected_protocol, expected_count, expected_severity in (
        (legacy_task["id"], None, 0, None),
        (modern_task["id"], 1, 1, "high"),
    ):
        detail_response = client.get(
            f"{settings.API_V1_STR}/review-tasks/{task_id}",
            headers=user_headers,
        )
        assert detail_response.status_code == 200
        detail = detail_response.json()
        assert detail["feedback_protocol_version"] == expected_protocol
        assert detail["issue_count"] == expected_count
        assert detail["highest_issue_severity"] == expected_severity

    stats_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=user_headers,
        params={"range": "all"},
    )
    assert stats_response.status_code == 200
    stats = stats_response.json()
    assert stats["total"] == 2
    assert stats["success"] == 2
    assert stats["failed_or_error"] == 0


def test_历史报告不能通过新协议重新解析或覆盖(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    legacy_report = b"# Legacy\n"
    first_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("legacy.md", legacy_report, "text/markdown")},
    )
    assert first_response.status_code == 200

    upgrade_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        data={"feedback_protocol_version": "1", "issues": "[]"},
        files={"upload": ("new.md", b"# New protocol\n", "text/markdown")},
    )

    assert upgrade_response.status_code == 409
    assert (tmp_path / f"{task['id']}.md").read_bytes() == legacy_report


def test_问题快照读取遵循普通用户部门管理员和系统管理员权限(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department = _create_department(db)
    owner_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
    )
    same_department_admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    other_user_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=_create_department(db),
    )
    admin, admin_api_key = create_key_user(db)
    admin.role = UserRole.admin
    db.add(admin)
    db.commit()
    admin_headers = authentication_token_from_key(
        client=client, key=admin_api_key.key, db=db
    )
    task = _create_review_task(client, owner_headers)
    issue = {
        "issue_text": "中风险：权限测试",
        "description_text": "用于验证任务数据权限。",
        "display_snapshot": "### 中风险：权限测试",
        "order": 0,
        "source": "template_1_heading",
        "location": {
            "start_offset": 0,
            "end_offset": 8,
            "feedback_insertion_offset": 8,
            "anchor_line": 0,
            "heading_level": 3,
        },
    }
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=owner_headers,
        data={
            "feedback_protocol_version": "1",
            "issues": json.dumps([issue], ensure_ascii=False),
        },
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200

    for allowed_headers in (
        owner_headers,
        same_department_admin_headers,
        admin_headers,
    ):
        response = client.get(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
            headers=allowed_headers,
        )
        assert response.status_code == 200
        assert response.json()["count"] == 1

    forbidden_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=other_user_headers,
    )
    assert forbidden_response.status_code == 403


def test_逐问题反馈Excel清理XML非法字符但保留合法空白(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    caplog: pytest.LogCaptureFixture,
) -> None:
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    marker = f"https://svn.example.com/illegal-excel-{random_lower_string()}"
    task_record = db.get(ReviewTask, UUID(task["id"]))
    issue_record = db.get(ReviewTaskIssue, UUID(issue["id"]))
    assert task_record is not None
    assert issue_record is not None
    issue_description = "if、\x0cor、do\n保留换行\t保留制表符\r保留回车"
    task_record.svn_url = marker
    issue_record.description_text = issue_description
    db.add(task_record)
    db.add(issue_record)
    db.commit()

    with caplog.at_level(logging.WARNING, logger="app.api.routes.review_tasks"):
        response = client.get(
            f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
            headers=normal_user_token_headers,
        )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook["问题反馈明细"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    row = next(exported for exported in rows if exported[2] == marker)
    assert row[7] == "if、or、do\n保留换行\t保留制表符\r保留回车"
    assert "\x0c" not in row[7]
    assert str(issue_record.id) in caplog.text
    assert "field=description_text" in caplog.text
    assert "count=1" in caplog.text
    assert issue_description not in caplog.text


def test_逐问题反馈Excel导出任务选择的审查项名称(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    cases = [
        (
            ["DS-01", "DS-03"],
            "DS-01 需求变更审查；DS-03 MOS 充放电状态机审查",
        ),
        ([], "未记录"),
        (["DS-88"], "DS-88"),
    ]
    expected_by_svn: dict[str, str] = {}
    for review_item_ids, expected in cases:
        task, _ = _create_new_protocol_task_with_issue(
            client,
            normal_user_token_headers,
        )
        marker = f"https://svn.example.com/review-items-{random_lower_string()}"
        task_record = db.get(ReviewTask, UUID(task["id"]))
        assert task_record is not None
        task_record.svn_url = marker
        task_record.review_item_ids = review_item_ids
        db.add(task_record)
        expected_by_svn[marker] = expected
    db.commit()

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    rows = list(workbook["问题反馈明细"].iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == (
        "提交人",
        "日期",
        "SVN路径",
        "旧版本号",
        "新版本号",
        "选择的审查项",
        "AI报告问题原文",
        "问题描述原文",
        "是否有帮助",
        "审查问题反馈",
    )
    for row in rows[1:]:
        marker = str(row[2])
        if marker in expected_by_svn:
            assert row[5] == expected_by_svn.pop(marker)
    assert expected_by_svn == {}


def test_逐问题反馈Excel按上海时区开始时间筛选(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    before_task, _ = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    boundary_task, _ = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    before_record = db.get(ReviewTask, UUID(before_task["id"]))
    boundary_record = db.get(ReviewTask, UUID(boundary_task["id"]))
    assert before_record is not None
    assert boundary_record is not None
    before_marker = f"https://svn.example.com/before-{random_lower_string()}"
    boundary_marker = f"https://svn.example.com/boundary-{random_lower_string()}"
    before_record.svn_url = before_marker
    before_record.started_at = datetime(2026, 8, 13, 15, 59, 59)
    boundary_record.svn_url = boundary_marker
    boundary_record.started_at = datetime(2026, 8, 13, 16, 0, 0)
    db.add(before_record)
    db.add(boundary_record)
    db.commit()

    filtered_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
        params={"started_from": "2026-08-14"},
    )
    unfiltered_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
    )
    invalid_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=normal_user_token_headers,
        params={"started_from": "not-a-date"},
    )

    assert filtered_response.status_code == 200
    filtered_workbook = load_workbook(
        BytesIO(filtered_response.content),
        read_only=True,
    )
    filtered_rows = list(
        filtered_workbook["问题反馈明细"].iter_rows(min_row=2, values_only=True)
    )
    filtered_workbook.close()
    filtered_paths = {str(row[2]) for row in filtered_rows}
    assert boundary_marker in filtered_paths
    assert before_marker not in filtered_paths

    assert unfiltered_response.status_code == 200
    unfiltered_workbook = load_workbook(
        BytesIO(unfiltered_response.content),
        read_only=True,
    )
    unfiltered_rows = list(
        unfiltered_workbook["问题反馈明细"].iter_rows(min_row=2, values_only=True)
    )
    unfiltered_workbook.close()
    unfiltered_paths = {str(row[2]) for row in unfiltered_rows}
    assert {before_marker, boundary_marker}.issubset(unfiltered_paths)

    assert invalid_response.status_code == 422
    assert invalid_response.json()["detail"] == "Invalid started_from"


def test_逐问题反馈Excel固定十列并排除历史报告级反馈(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department = _create_department(db)
    owner_headers, owner = _create_department_user(
        client=client,
        db=db,
        department=department,
        full_name="张工",
        svn_username="zhang.svn",
    )
    task, issue = _create_new_protocol_task_with_issue(client, owner_headers)
    task_record = db.get(ReviewTask, UUID(task["id"]))
    first_issue = db.get(ReviewTaskIssue, UUID(issue["id"]))
    assert task_record is not None
    assert first_issue is not None
    task_record.created_at = datetime(2026, 7, 29, 16, 30, tzinfo=timezone.utc)
    task_record.review_item_ids = ["DS-03"]
    first_issue.issue_text = "高风险：空指针解引用"
    first_issue.description_text = "未校验返回值即解引用。"
    first_issue.helpful = True
    first_issue.feedback_text = "已增加空值检查并完成回归测试。"
    db.add(
        ReviewTaskIssue(
            task_id=task_record.id,
            issue_text="中风险：边界条件误报",
            description_text="产品约束保证输入不会达到该分支。",
            display_snapshot="### 中风险：边界条件误报",
            order=1,
            source="template_1_heading",
            location=_valid_issue_snapshot()["location"],
            helpful=False,
            feedback_text=None,
        )
    )
    db.add(
        ReviewTaskIssue(
            task_id=task_record.id,
            issue_text="=1+1",
            description_text="异常日志没有记录请求编号。",
            display_snapshot="### 低风险：日志信息不足",
            order=2,
            source="template_1_heading",
            location=_valid_issue_snapshot()["location"],
            helpful=None,
            feedback_text=None,
        )
    )

    legacy_task = _create_review_task(client, owner_headers)
    legacy_upload = client.post(
        f"{settings.API_V1_STR}/review-tasks/{legacy_task['id']}/report",
        headers=owner_headers,
        files={"upload": ("legacy.md", b"# Legacy\n", "text/markdown")},
    )
    assert legacy_upload.status_code == 200
    legacy_task_id = UUID(legacy_task["id"])
    db.add(
        ReviewTaskReportFeedback(
            task_id=legacy_task_id,
            user_id=owner.id,
            helpful=True,
            useful_issues_text="历史报告级反馈不得导出。",
        )
    )
    db.add(
        ReviewTaskIssue(
            task_id=legacy_task_id,
            issue_text="历史任务伪问题",
            description_text="即使存在异常遗留问题记录也不得导出。",
            display_snapshot="### 历史任务伪问题",
            order=0,
            source="template_1_heading",
            location=_valid_issue_snapshot()["location"],
            helpful=True,
            feedback_text="不应出现",
        )
    )
    db.add(task_record)
    db.add(first_issue)
    db.commit()

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
        headers=owner_headers,
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "问题反馈明细_" in unquote(response.headers["content-disposition"])
    assert unquote(response.headers["content-disposition"]).endswith(".xlsx")
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["问题反馈明细"]
    sheet = workbook["问题反馈明细"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "提交人",
        "日期",
        "SVN路径",
        "旧版本号",
        "新版本号",
        "选择的审查项",
        "AI报告问题原文",
        "问题描述原文",
        "是否有帮助",
        "审查问题反馈",
    )
    assert len(rows) == 4
    assert rows[1] == (
        "张工",
        datetime(2026, 7, 30),
        task["svn_url"],
        task["old_revision"],
        task["new_revision"],
        "DS-03 MOS 充放电状态机审查",
        "高风险：空指针解引用",
        "未校验返回值即解引用。",
        "有帮助",
        "已增加空值检查并完成回归测试。",
    )
    assert rows[2][5:] == (
        "DS-03 MOS 充放电状态机审查",
        "中风险：边界条件误报",
        "产品约束保证输入不会达到该分支。",
        "没帮助",
        None,
    )
    assert rows[3][5:] == (
        "DS-03 MOS 充放电状态机审查",
        "=1+1",
        "异常日志没有记录请求编号。",
        "未反馈",
        None,
    )
    assert sheet.auto_filter.ref == "A1:J4"
    assert sheet["G4"].data_type == "s"
    assert all("历史" not in str(cell) for row in rows for cell in row)


def test_逐问题反馈Excel遵循个人部门和全平台任务读取范围(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department_a = _create_department(db)
    department_b = _create_department(db)
    owner_a_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        full_name="甲部门提交人",
    )
    owner_b_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_b,
        full_name="乙部门提交人",
    )
    department_admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    platform_admin, platform_admin_key = create_key_user(db)
    platform_admin.role = UserRole.admin
    db.add(platform_admin)
    db.commit()
    platform_admin_headers = authentication_token_from_key(
        client=client,
        key=platform_admin_key.key,
        db=db,
    )

    task_a, issue_a = _create_new_protocol_task_with_issue(client, owner_a_headers)
    task_b, issue_b = _create_new_protocol_task_with_issue(client, owner_b_headers)
    task_a_record = db.get(ReviewTask, UUID(task_a["id"]))
    task_b_record = db.get(ReviewTask, UUID(task_b["id"]))
    issue_a_record = db.get(ReviewTaskIssue, UUID(issue_a["id"]))
    issue_b_record = db.get(ReviewTaskIssue, UUID(issue_b["id"]))
    assert task_a_record is not None
    assert task_b_record is not None
    assert issue_a_record is not None
    assert issue_b_record is not None
    marker_a = f"https://svn.example.com/export-a-{random_lower_string()}"
    marker_b = f"https://svn.example.com/export-b-{random_lower_string()}"
    task_a_record.svn_url = marker_a
    task_b_record.svn_url = marker_b
    issue_a_record.issue_text = "甲部门问题"
    issue_b_record.issue_text = "乙部门问题"
    db.add(task_a_record)
    db.add(task_b_record)
    db.add(issue_a_record)
    db.add(issue_b_record)
    db.commit()

    def exported_svn_paths(headers: dict[str, str]) -> set[str]:
        response = client.get(
            f"{settings.API_V1_STR}/review-tasks/issue-feedback/export",
            headers=headers,
        )
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = workbook["问题反馈明细"].iter_rows(min_row=2, values_only=True)
        return {str(row[2]) for row in rows}

    assert exported_svn_paths(owner_a_headers) == {marker_a}
    assert exported_svn_paths(department_admin_headers) == {marker_a}
    platform_paths = exported_svn_paths(platform_admin_headers)
    assert marker_a in platform_paths
    assert marker_b in platform_paths


def test_任务提交人可幂等保存并恢复逐问题反馈草稿(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    url = (
        f"{settings.API_V1_STR}/review-tasks/{task['id']}"
        f"/issues/{issue['id']}/feedback-draft"
    )
    payload = {"helpful": True, "feedback_text": "帮助定位了空指针并完成修复。"}

    first_response = client.put(
        url,
        headers=normal_user_token_headers,
        json=payload,
    )
    assert first_response.status_code == 200
    assert first_response.json()["helpful"] is True
    assert first_response.json()["feedback_text"] == payload["feedback_text"]
    first_updated_at = first_response.json()["updated_at"]

    repeated_response = client.put(
        url,
        headers=normal_user_token_headers,
        json=payload,
    )
    assert repeated_response.status_code == 200
    assert repeated_response.json()["id"] == issue["id"]
    assert repeated_response.json()["updated_at"] == first_updated_at

    restored_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/issues",
        headers=normal_user_token_headers,
    )
    assert restored_response.status_code == 200
    restored = restored_response.json()
    assert restored["evaluated_count"] == 1
    assert restored["data"][0]["helpful"] is True
    assert restored["data"][0]["feedback_text"] == payload["feedback_text"]
    assert (
        len(
            db.exec(
                select(ReviewTaskIssue).where(
                    ReviewTaskIssue.task_id == UUID(task["id"])
                )
            ).all()
        )
        == 1
    )


def test_未评价问题也可单独保存说明且空白草稿归一为空(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    url = (
        f"{settings.API_V1_STR}/review-tasks/{task['id']}"
        f"/issues/{issue['id']}/feedback-draft"
    )

    text_response = client.put(
        url,
        headers=normal_user_token_headers,
        json={"helpful": None, "feedback_text": "先记录处理上下文"},
    )
    assert text_response.status_code == 200
    assert text_response.json()["helpful"] is None
    assert text_response.json()["feedback_text"] == "先记录处理上下文"

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["evaluated_issue_count"] == 0
    assert detail_response.json()["has_issue_feedback_draft"] is True

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )
    assert list_response.status_code == 200
    listed_task = next(
        item for item in list_response.json()["data"] if item["id"] == task["id"]
    )
    assert listed_task["evaluated_issue_count"] == 0
    assert listed_task["has_issue_feedback_draft"] is True

    blank_response = client.put(
        url,
        headers=normal_user_token_headers,
        json={"helpful": None, "feedback_text": "  \n "},
    )
    assert blank_response.status_code == 200
    assert blank_response.json()["feedback_text"] is None

    cleared_detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert cleared_detail_response.status_code == 200
    assert cleared_detail_response.json()["has_issue_feedback_draft"] is False


def test_逐问题反馈草稿限制长度且仅任务提交人可写(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    url = (
        f"{settings.API_V1_STR}/review-tasks/{task['id']}"
        f"/issues/{issue['id']}/feedback-draft"
    )
    admin, admin_api_key = create_key_user(db)
    admin.role = UserRole.admin
    db.add(admin)
    db.commit()
    admin_headers = authentication_token_from_key(
        client=client,
        key=admin_api_key.key,
        db=db,
    )

    forbidden_response = client.put(
        url,
        headers=admin_headers,
        json={"helpful": False, "feedback_text": "管理员不能代填"},
    )
    assert forbidden_response.status_code == 403

    too_long_response = client.put(
        url,
        headers=normal_user_token_headers,
        json={
            "helpful": False,
            "feedback_text": "x" * (MAX_REVIEW_TASK_ISSUE_FEEDBACK_LENGTH + 1),
        },
    )
    assert too_long_response.status_code == 422


def test_历史任务和已正式提交任务拒绝保存逐问题反馈草稿(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    legacy_task = _create_review_task(client, normal_user_token_headers)
    legacy_url = (
        f"{settings.API_V1_STR}/review-tasks/{legacy_task['id']}"
        f"/issues/{uuid4()}/feedback-draft"
    )
    legacy_response = client.put(
        legacy_url,
        headers=normal_user_token_headers,
        json={"helpful": True, "feedback_text": None},
    )
    assert legacy_response.status_code == 409
    assert legacy_response.json()["detail"] == "issue_feedback_not_supported"

    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    db_task.feedback_submitted_at = datetime.now(timezone.utc)
    db.add(db_task)
    db.commit()

    submitted_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{issue['id']}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": False, "feedback_text": "不得覆盖"},
    )
    assert submitted_response.status_code == 409
    assert submitted_response.json()["detail"] == "issue_feedback_submitted"


def test_新协议仅在全部评价且最终PDF保存成功后锁定反馈(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    original_markdown_path = tmp_path / f"{task['id']}.md"
    original_markdown = original_markdown_path.read_bytes()
    pdf_url = f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report"

    incomplete_response = client.post(
        pdf_url,
        headers=normal_user_token_headers,
        files={"upload": ("最终报告.pdf", b"%PDF-1.4\nincomplete", "application/pdf")},
    )
    assert incomplete_response.status_code == 409
    assert incomplete_response.json()["detail"] == "issue_feedback_incomplete"
    assert not (tmp_path / f"{task['id']}.pdf").exists()

    draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{issue['id']}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": True, "feedback_text": "已补充空值判断"},
    )
    assert draft_response.status_code == 200

    final_pdf = b"%PDF-1.4\nfinal-feedback-pdf"
    submitted_response = client.post(
        pdf_url,
        headers=normal_user_token_headers,
        files={"upload": ("最终报告.pdf", final_pdf, "application/pdf")},
    )
    assert submitted_response.status_code == 200
    submitted = submitted_response.json()
    assert submitted["feedback_submitted_at"] is not None
    assert submitted["pdf_report_name"] == "最终报告.pdf"
    assert (tmp_path / f"{task['id']}.pdf").read_bytes() == final_pdf
    assert original_markdown_path.read_bytes() == original_markdown

    locked_draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{issue['id']}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": False, "feedback_text": "不得修改"},
    )
    assert locked_draft_response.status_code == 409
    assert locked_draft_response.json()["detail"] == "issue_feedback_submitted"

    duplicate_response = client.post(
        pdf_url,
        headers=normal_user_token_headers,
        files={"upload": ("重复.pdf", b"%PDF-1.4\nduplicate", "application/pdf")},
    )
    assert duplicate_response.status_code == 409
    assert duplicate_response.json()["detail"] == "issue_feedback_submitted"
    assert (tmp_path / f"{task['id']}.pdf").read_bytes() == final_pdf

    db.expire_all()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    assert db_task.feedback_submitted_at is not None


def test_新协议反馈说明为空白时拒绝上传最终PDF并保持草稿可编辑(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    db_issue = db.get(ReviewTaskIssue, UUID(issue["id"]))
    assert db_issue is not None
    db_issue.helpful = True
    db_issue.feedback_text = "  \n "
    db.add(db_issue)
    db.commit()

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={"upload": ("最终报告.pdf", b"%PDF-1.4\nfinal", "application/pdf")},
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "issue_feedback_text_required"
    assert not (tmp_path / f"{task['id']}.pdf").exists()

    draft_response = client.put(
        (
            f"{settings.API_V1_STR}/review-tasks/{task['id']}"
            f"/issues/{issue['id']}/feedback-draft"
        ),
        headers=normal_user_token_headers,
        json={"helpful": True, "feedback_text": "已补充处理结果"},
    )
    assert draft_response.status_code == 200
    assert draft_response.json()["feedback_text"] == "已补充处理结果"

    db.expire_all()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    assert db_task.feedback_submitted_at is None


def test_新协议并发提交最终PDF只允许一次完成锁定(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    assert (
        client.put(
            (
                f"{settings.API_V1_STR}/review-tasks/{task['id']}"
                f"/issues/{issue['id']}/feedback-draft"
            ),
            headers=normal_user_token_headers,
            json={"helpful": True, "feedback_text": "并发提交验证"},
        ).status_code
        == 200
    )

    def submit(index: int) -> int:
        return client.post(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
            headers=normal_user_token_headers,
            files={
                "upload": (
                    f"最终报告-{index}.pdf",
                    f"%PDF-1.4\nfinal-{index}".encode(),
                    "application/pdf",
                )
            },
        ).status_code

    with ThreadPoolExecutor(max_workers=2) as executor:
        statuses = list(executor.map(submit, (1, 2)))

    assert sorted(statuses) == [200, 409]
    db.expire_all()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    assert db_task.feedback_submitted_at is not None
    assert (tmp_path / f"{task['id']}.pdf").read_bytes() in {
        b"%PDF-1.4\nfinal-1",
        b"%PDF-1.4\nfinal-2",
    }


def test_新协议PDF内容或保存失败时保持反馈草稿可编辑(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    draft_url = (
        f"{settings.API_V1_STR}/review-tasks/{task['id']}"
        f"/issues/{issue['id']}/feedback-draft"
    )
    assert (
        client.put(
            draft_url,
            headers=normal_user_token_headers,
            json={"helpful": False, "feedback_text": "草稿内容"},
        ).status_code
        == 200
    )
    pdf_url = f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report"

    invalid_response = client.post(
        pdf_url,
        headers=normal_user_token_headers,
        files={"upload": ("最终报告.pdf", b"not-a-pdf", "application/pdf")},
    )
    assert invalid_response.status_code == 400

    async def fail_pdf_save(_task_id: UUID, _upload: Any) -> Path:
        raise OSError("模拟PDF保存失败")

    monkeypatch.setattr(report_service, "save_pdf_upload", fail_pdf_save)
    with pytest.raises(OSError, match="模拟PDF保存失败"):
        client.post(
            pdf_url,
            headers=normal_user_token_headers,
            files={"upload": ("最终报告.pdf", b"%PDF-1.4\nretry", "application/pdf")},
        )

    db.expire_all()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    assert db_task.feedback_submitted_at is None
    assert db_task.pdf_report_path is None
    retry_draft_response = client.put(
        draft_url,
        headers=normal_user_token_headers,
        json={"helpful": True, "feedback_text": "失败后仍可编辑"},
    )
    assert retry_draft_response.status_code == 200


def test_只有任务提交人可以上传新协议最终PDF并锁定反馈(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    draft_url = (
        f"{settings.API_V1_STR}/review-tasks/{task['id']}"
        f"/issues/{issue['id']}/feedback-draft"
    )
    assert (
        client.put(
            draft_url,
            headers=normal_user_token_headers,
            json={"helpful": True, "feedback_text": "权限校验草稿"},
        ).status_code
        == 200
    )
    admin, admin_api_key = create_key_user(db)
    admin.role = UserRole.admin
    db.add(admin)
    db.commit()
    admin_headers = authentication_token_from_key(
        client=client,
        key=admin_api_key.key,
        db=db,
    )

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=admin_headers,
        files={"upload": ("最终报告.pdf", b"%PDF-1.4\nadmin", "application/pdf")},
    )

    assert response.status_code == 403
    assert not (tmp_path / f"{task['id']}.pdf").exists()
    db.expire_all()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    assert db_task.feedback_submitted_at is None


def test_新协议最终PDF数据库提交失败时清理文件且不锁定(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task, issue = _create_new_protocol_task_with_issue(
        client,
        normal_user_token_headers,
    )
    draft_url = (
        f"{settings.API_V1_STR}/review-tasks/{task['id']}"
        f"/issues/{issue['id']}/feedback-draft"
    )
    assert (
        client.put(
            draft_url,
            headers=normal_user_token_headers,
            json={"helpful": True, "feedback_text": "数据库失败验证"},
        ).status_code
        == 200
    )
    original_commit = Session.commit

    def fail_commit(_session: Session) -> None:
        raise RuntimeError("模拟最终反馈提交失败")

    monkeypatch.setattr(Session, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="模拟最终反馈提交失败"):
        client.post(
            f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
            headers=normal_user_token_headers,
            files={"upload": ("最终报告.pdf", b"%PDF-1.4\nfinal", "application/pdf")},
        )
    monkeypatch.setattr(Session, "commit", original_commit)

    assert not (tmp_path / f"{task['id']}.pdf").exists()
    db.expire_all()
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    assert db_task.feedback_submitted_at is None
    assert db_task.pdf_report_path is None


def test_restart_failed_task_status_clears_previous_error_and_finish_time(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    task = _create_review_task(client, normal_user_token_headers)
    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    db_task.status = ReviewTaskStatus.failed
    db_task.error_message = "Codex 执行失败"
    db_task.started_at = _today_elapsed_datetime(20)
    db_task.finished_at = _today_elapsed_datetime(10)
    db.add(db_task)
    db.commit()

    response = client.patch(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/status",
        headers=normal_user_token_headers,
        json={"status": "checking_environment", "error_message": None},
    )

    assert response.status_code == 200
    content = response.json()
    assert content["status"] == "checking_environment"
    assert content["error_message"] is None
    assert content["finished_at"] is None

    db.refresh(db_task)
    assert db_task.status == ReviewTaskStatus.checking_environment
    assert db_task.error_message is None
    assert db_task.finished_at is None


def test_download_review_task_report_uses_uploaded_report_filename(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers, old_revision="10")
    report_name = "r10_r2审查报告_20260610_142530.md"
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": (report_name, b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    assert upload_response.json()["report_name"] == report_name

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/download",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert report_name in unquote(response.headers["content-disposition"])
    assert response.content == b"# Review\n"


def test_download_review_task_report_falls_back_to_default_report_filename(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers, old_revision="10")
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200

    db_task = db.get(ReviewTask, UUID(task["id"]))
    assert db_task is not None
    db_task.report_name = None
    db.add(db_task)
    db.commit()

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/download",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert "代码审查报告.md" in unquote(response.headers["content-disposition"])


def test_upload_and_download_review_task_pdf_report(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={"upload": ("代码审查报告.pdf", b"%PDF-1.4\npdf\n", "application/pdf")},
    )

    assert upload_response.status_code == 200
    uploaded = upload_response.json()
    assert uploaded["pdf_report_path"]
    assert uploaded["pdf_report_name"] == "代码审查报告.pdf"

    download_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf/download",
        headers=normal_user_token_headers,
    )

    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith("application/pdf")
    assert "代码审查报告.pdf" in unquote(
        download_response.headers["content-disposition"]
    )
    assert download_response.content == b"%PDF-1.4\npdf\n"


def test_download_review_task_pdf_report_returns_404_when_missing(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
) -> None:
    task = _create_review_task(client, normal_user_token_headers)

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf/download",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "PDF report not found"


def test_delete_review_task_removes_existing_pdf_report_file(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=normal_user_token_headers,
        files={"upload": ("report.pdf", b"%PDF-1.4\npdf\n", "application/pdf")},
    )
    assert upload_response.status_code == 200
    pdf_path = Path(upload_response.json()["pdf_report_path"])
    assert pdf_path.is_file()

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert not pdf_path.exists()


def test_download_review_task_pdf_report_requires_task_permission(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    other = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="Other User",
        svn_username="other",
    )
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{other['task']['id']}/pdf-report",
        headers=other["headers"],
        files={"upload": ("report.pdf", b"%PDF-1.4\npdf\n", "application/pdf")},
    )
    assert upload_response.status_code == 200

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{other['task']['id']}/pdf/download",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 403


def test_upload_review_task_requirement_file_stores_metadata_and_record(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    content = b"requirement rows"

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={
            "upload": (
                "产品开发需求表.xls",
                content,
                "application/vnd.ms-excel",
            )
        },
    )

    assert response.status_code == 200
    uploaded = response.json()
    expected_hash = hashlib.sha256(content).hexdigest()
    assert uploaded["requirement_file_name"] == "产品开发需求表.xls"
    assert uploaded["requirement_file_size"] == len(content)
    assert uploaded["requirement_file_sha256"] == expected_hash

    requirement_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    assert requirement_file.file_name == "产品开发需求表.xls"
    assert requirement_file.file_size == len(content)
    assert requirement_file.sha256 == expected_hash
    assert Path(requirement_file.file_path).is_file()


def test_upload_review_task_requirement_file_replaces_old_file_and_record(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    first_content = b"first requirement rows"
    second_content = b"second requirement rows"

    first_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.xls", first_content, "application/vnd.ms-excel")},
    )
    assert first_response.status_code == 200
    first_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    first_file_id = first_file.id
    first_file_path = Path(first_file.file_path)
    assert first_file_path.read_bytes() == first_content

    second_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.xls", second_content, "application/vnd.ms-excel")},
    )

    assert second_response.status_code == 200
    db.expire_all()
    requirement_files = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).all()
    assert len(requirement_files) == 1
    second_file = requirement_files[0]
    assert second_file.id != first_file_id
    assert second_file.sha256 == hashlib.sha256(second_content).hexdigest()
    assert not first_file_path.exists()
    assert Path(second_file.file_path).read_bytes() == second_content
    assert db.get(ReviewTaskRequirementFile, first_file_id) is None

    updated_task = db.get(ReviewTask, UUID(task["id"]))
    assert updated_task is not None
    assert updated_task.requirement_file_id == second_file.id
    assert (
        updated_task.requirement_file_sha256
        == hashlib.sha256(second_content).hexdigest()
    )
    assert updated_task.requirement_file_size == len(second_content)


def test_upload_review_task_requirement_file_replaces_duplicate_same_content(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    content = b"same requirement rows"

    first_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.xls", content, "application/vnd.ms-excel")},
    )
    assert first_response.status_code == 200
    first_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    first_file_id = first_file.id
    shared_path = Path(first_file.file_path)
    assert shared_path.read_bytes() == content

    second_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.xls", content, "application/vnd.ms-excel")},
    )

    assert second_response.status_code == 200
    db.expire_all()
    requirement_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    assert requirement_file.id != first_file_id
    assert Path(requirement_file.file_path) == shared_path
    assert shared_path.read_bytes() == content
    assert db.get(ReviewTaskRequirementFile, first_file_id) is None


def test_upload_review_task_requirement_file_sanitizes_unsafe_and_long_name(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    long_name = "../secret/\x00bad\nname" + ("很长" * 160) + " .xlsx"
    content = b"requirement rows"

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={
            "upload": (
                long_name,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    requirement_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    stored_name = requirement_file.file_name
    assert "/" not in stored_name
    assert "\\" not in stored_name
    assert "\x00" not in stored_name
    assert not any(ord(char) < 32 or ord(char) == 127 for char in stored_name)
    assert len(stored_name.encode("utf-8")) <= 255
    assert stored_name.endswith(".xlsx")
    assert Path(requirement_file.file_path).is_file()


def test_upload_review_task_requirement_file_rejects_unsupported_extension(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.txt", b"rows", "text/plain")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "产品开发需求表仅支持 .xls 或 .xlsx 文件。"


def test_upload_review_task_requirement_file_rejects_empty_excel(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)

    response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={
            "upload": (
                "产品开发需求表.xlsx",
                b"",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "产品开发需求表不能为空。"


def test_delete_review_task_removes_requirement_file_and_record(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.xls", b"rows", "application/vnd.ms-excel")},
    )
    assert upload_response.status_code == 200
    requirement_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    requirement_file_id = requirement_file.id
    requirement_path = Path(requirement_file.file_path)
    assert requirement_path.is_file()

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert not requirement_path.exists()
    db.expire_all()
    assert db.get(ReviewTaskRequirementFile, requirement_file_id) is None
    assert db.get(ReviewTask, UUID(task["id"])) is None


def test_delete_review_task_removes_requirement_file_without_calling_unlink(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=normal_user_token_headers,
        files={"upload": ("需求表.xlsx", b"rows", "application/vnd.ms-excel")},
    )
    assert upload_response.status_code == 200
    requirement_file = db.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == UUID(task["id"])
        )
    ).one()
    requirement_file_id = requirement_file.id
    requirement_path = Path(requirement_file.file_path)

    original_unlink = Path.unlink

    def fail_requirement_unlink(self: Path, missing_ok: bool = False) -> None:
        if self == requirement_path:
            raise OSError("cannot delete")
        original_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", fail_requirement_unlink)

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    db.expire_all()
    assert db.get(ReviewTask, UUID(task["id"])) is None
    assert db.get(ReviewTaskRequirementFile, requirement_file_id) is None
    assert not requirement_path.exists()


def test_delete_review_task_removes_task_row(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    task = _create_review_task(client, normal_user_token_headers)

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert response.json() == {"message": "Review task deleted"}

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert read_response.status_code == 404


def test_delete_review_task_removes_existing_report_file(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    report_path = Path(upload_response.json()["report_path"])
    assert report_path.is_file()

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert response.json() == {"message": "Review task deleted"}
    assert not report_path.exists()

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert read_response.status_code == 404


def test_delete_review_task_succeeds_when_report_file_is_missing(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    report_path = Path(upload_response.json()["report_path"])
    report_path.unlink()

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    assert response.json() == {"message": "Review task deleted"}


def test_delete_review_task_removes_report_without_calling_unlink(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    task = _create_review_task(client, normal_user_token_headers)
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=normal_user_token_headers,
        files={"upload": ("report.md", b"# Review\n", "text/markdown")},
    )
    assert upload_response.status_code == 200
    report_path = Path(upload_response.json()["report_path"])
    assert report_path.is_file()

    original_unlink = Path.unlink

    def fail_report_unlink(self: Path, missing_ok: bool = False) -> None:
        if self == report_path:
            raise OSError("cannot delete")
        original_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", fail_report_unlink)

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert read_response.status_code == 404
    assert not report_path.exists()


def test_normal_user_cannot_delete_other_user_task(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    _, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client, key=other_api_key.key, db=db
    )
    task = _create_review_task(client, other_headers)

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Not enough permissions"


def test_normal_user_cannot_read_other_user_task(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    _, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client, key=other_api_key.key, db=db
    )
    task = _create_review_task(client, other_headers)

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Not enough permissions"


def test_admin_or_superuser_can_list_all_tasks(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    user_task = _create_review_task(client, normal_user_token_headers, old_revision="3")
    _, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client, key=other_api_key.key, db=db
    )
    other_task = _create_review_task(client, other_headers, old_revision="5")
    admin, admin_api_key = create_key_user(db)
    admin.role = UserRole.admin
    db.add(admin)
    db.commit()
    admin_headers = authentication_token_from_key(
        client=client, key=admin_api_key.key, db=db
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
    )

    assert response.status_code == 200
    content = response.json()
    listed_ids = {task["id"] for task in content["data"]}
    assert user_task["id"] in listed_ids
    assert other_task["id"] in listed_ids
    assert content["count"] >= 2

    admin_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=admin_headers,
    )

    assert admin_response.status_code == 200
    admin_listed_ids = {task["id"] for task in admin_response.json()["data"]}
    assert user_task["id"] in admin_listed_ids
    assert other_task["id"] in admin_listed_ids


def test_admin_task_list_includes_user_identity_fields(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    user, api_key = create_key_user(db)
    user.full_name = "张三"
    svn_username = f"zhangsan-{random_lower_string()}"
    user.svn_username = svn_username
    db.add(user)
    db.commit()
    db.refresh(user)
    user_headers = authentication_token_from_key(client=client, key=api_key.key, db=db)
    task = _create_review_task(client, user_headers, old_revision="9")

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
    )

    assert response.status_code == 200
    content = response.json()
    listed_task = next(item for item in content["data"] if item["id"] == task["id"])
    assert listed_task["user_name"] == "张三"
    assert listed_task["svn_username"] == svn_username


def test_admin_can_filter_review_tasks_by_user_name(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    zhang = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="张三",
        svn_username="zhang",
        old_revision="11",
    )
    li = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="李四",
        svn_username="lisi",
        old_revision="12",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
        params={"user_name": "张"},
    )

    assert response.status_code == 200
    content = response.json()
    listed_ids = {task["id"] for task in content["data"]}
    assert zhang["task"]["id"] in listed_ids
    assert li["task"]["id"] not in listed_ids


def test_admin_can_list_own_review_tasks_before_pagination(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    own_task = _create_review_task(
        client,
        superuser_token_headers,
        old_revision="31",
    )
    other = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="其他用户",
        svn_username="other-user",
        old_revision="32",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
        params={"mine": True, "limit": 1},
    )

    assert response.status_code == 200
    content = response.json()
    assert content["count"] == 1
    assert [task["id"] for task in content["data"]] == [own_task["id"]]
    assert other["task"]["id"] not in {task["id"] for task in content["data"]}


def test_admin_can_filter_review_tasks_by_svn_username(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    wang = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="王五",
        svn_username="wang",
        old_revision="13",
    )
    zhao = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="赵六",
        svn_username="zhao",
        old_revision="14",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
        params={"svn_username": "wang"},
    )

    assert response.status_code == 200
    content = response.json()
    listed_ids = {task["id"] for task in content["data"]}
    assert wang["task"]["id"] in listed_ids
    assert zhao["task"]["id"] not in listed_ids


def test_review_task_filters_are_combined_with_and(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    matching = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="张三",
        svn_username="zhang",
        old_revision="15",
    )
    name_only = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="张四",
        svn_username="li",
        old_revision="16",
    )
    svn_only = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="王五",
        svn_username="zhang",
        old_revision="17",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
        params={"user_name": "张", "svn_username": matching["svn_username"]},
    )

    assert response.status_code == 200
    content = response.json()
    listed_ids = {task["id"] for task in content["data"]}
    assert matching["task"]["id"] in listed_ids
    assert name_only["task"]["id"] not in listed_ids
    assert svn_only["task"]["id"] not in listed_ids


def test_review_task_filters_include_status_feedback_and_started_from(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    matching = _create_review_task(client, normal_user_token_headers, old_revision="21")
    helpful = _create_review_task(client, normal_user_token_headers, old_revision="22")
    wrong_status = _create_review_task(
        client, normal_user_token_headers, old_revision="23"
    )
    older_unfeedback = _create_review_task(
        client, normal_user_token_headers, old_revision="24"
    )
    started_from = _today_elapsed_datetime(90)
    recent_started = _today_elapsed_datetime(30)
    older_started = started_from - timedelta(minutes=10)

    _set_review_task_stats_fields(
        db,
        matching["id"],
        created_at=recent_started,
        status=ReviewTaskStatus.success,
        started_at=recent_started,
        finished_at=recent_started + timedelta(minutes=5),
        report_path="/tmp/matching.md",
    )
    _set_review_task_report(db, helpful["id"])
    _submit_report_feedback(
        client,
        normal_user_token_headers,
        helpful["id"],
        helpful=True,
        useful_issues_text="1. 命中有效问题",
    )
    _set_review_task_stats_fields(
        db,
        helpful["id"],
        created_at=recent_started,
        status=ReviewTaskStatus.success,
        started_at=recent_started,
        finished_at=recent_started + timedelta(minutes=5),
        report_path="/tmp/helpful.md",
    )
    _set_review_task_stats_fields(
        db,
        wrong_status["id"],
        created_at=recent_started,
        status=ReviewTaskStatus.failed,
        started_at=recent_started,
        finished_at=recent_started + timedelta(minutes=5),
    )
    _set_review_task_stats_fields(
        db,
        older_unfeedback["id"],
        created_at=older_started,
        status=ReviewTaskStatus.success,
        started_at=older_started,
        finished_at=older_started + timedelta(minutes=5),
        report_path="/tmp/older.md",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=superuser_token_headers,
        params={
            "status": "success",
            "report_feedback": "none",
            "started_from": started_from.isoformat(),
        },
    )

    assert response.status_code == 200
    listed_ids = {task["id"] for task in response.json()["data"]}
    assert matching["id"] in listed_ids
    assert helpful["id"] not in listed_ids
    assert wrong_status["id"] not in listed_ids
    assert older_unfeedback["id"] not in listed_ids


def test_normal_user_filters_only_search_own_review_tasks(
    client: TestClient,
    db: Session,
) -> None:
    own = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="张三",
        svn_username="zhang",
        old_revision="18",
    )
    other = _create_named_svn_user_task(
        client=client,
        db=db,
        full_name="张四",
        svn_username="zhang",
        old_revision="19",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=own["headers"],
        params={"user_name": "张", "svn_username": "zhang"},
    )

    assert response.status_code == 200
    content = response.json()
    listed_ids = {task["id"] for task in content["data"]}
    assert own["task"]["id"] in listed_ids
    assert other["task"]["id"] not in listed_ids


def test_normal_user_task_stats_only_count_own_tasks(
    client: TestClient,
    db: Session,
) -> None:
    _, own_api_key = create_key_user(db)
    own_headers = authentication_token_from_key(
        client=client, key=own_api_key.key, db=db
    )
    own_task = _create_review_task(client, own_headers, old_revision="31")
    client.patch(
        f"{settings.API_V1_STR}/review-tasks/{own_task['id']}/status",
        headers=own_headers,
        json={"status": "success"},
    )
    _, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client, key=other_api_key.key, db=db
    )
    other_task = _create_review_task(client, other_headers, old_revision="32")
    client.patch(
        f"{settings.API_V1_STR}/review-tasks/{other_task['id']}/status",
        headers=other_headers,
        json={"status": "failed"},
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=own_headers,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["total"] == 1
    assert content["success"] == 1
    assert content["running"] == 0
    assert content["failed_or_error"] == 0
    assert content["by_status"]["success"] == 1
    assert content["by_status"]["failed"] == 0


def test_admin_task_stats_count_all_tasks(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    baseline_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=superuser_token_headers,
    )
    assert baseline_response.status_code == 200
    baseline = baseline_response.json()

    _, success_api_key = create_key_user(db)
    success_headers = authentication_token_from_key(
        client=client, key=success_api_key.key, db=db
    )
    success_task = _create_review_task(client, success_headers, old_revision="33")
    client.patch(
        f"{settings.API_V1_STR}/review-tasks/{success_task['id']}/status",
        headers=success_headers,
        json={"status": "success"},
    )
    _, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client, key=other_api_key.key, db=db
    )
    running_task = _create_review_task(client, other_headers, old_revision="34")
    client.patch(
        f"{settings.API_V1_STR}/review-tasks/{running_task['id']}/status",
        headers=other_headers,
        json={"status": "running_codex"},
    )
    failed_task = _create_review_task(client, other_headers, old_revision="35")
    client.patch(
        f"{settings.API_V1_STR}/review-tasks/{failed_task['id']}/status",
        headers=other_headers,
        json={"status": "environment_error"},
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=superuser_token_headers,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["total"] - baseline["total"] == 3
    assert content["success"] - baseline["success"] == 1
    assert content["running"] - baseline["running"] == 1
    assert content["failed_or_error"] - baseline["failed_or_error"] == 1
    assert content["by_status"]["success"] - baseline["by_status"]["success"] == 1
    assert (
        content["by_status"]["running_codex"] - baseline["by_status"]["running_codex"]
        == 1
    )
    assert (
        content["by_status"]["environment_error"]
        - baseline["by_status"]["environment_error"]
        == 1
    )


def test_task_stats_support_today_yesterday_and_all_ranges(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fixed_now = datetime(2026, 7, 14, 12, tzinfo=timezone.utc)
    monkeypatch.setattr(review_tasks_route, "_utcnow", lambda: fixed_now)
    today_date = fixed_now.astimezone(TASK_STATS_TIMEZONE).date()
    week_start = today_date - timedelta(days=today_date.weekday())
    last_week_start = week_start - timedelta(days=7)

    def task_time(days_delta: int, hour: int) -> datetime:
        return (
            datetime.combine(
                today_date + timedelta(days=days_delta),
                time(hour=hour),
                tzinfo=TASK_STATS_TIMEZONE,
            )
            .astimezone(timezone.utc)
            .replace(tzinfo=None)
        )

    user, api_key = create_key_user(db)
    user.full_name = "任务统计用户"
    db.add(user)
    db.commit()
    headers = authentication_token_from_key(client=client, key=api_key.key, db=db)

    today_task = _create_review_task(client, headers, old_revision="301")
    _set_review_task_stats_fields(
        db,
        today_task["id"],
        created_at=task_time(0, 19),
        status=ReviewTaskStatus.success,
        started_at=task_time(0, 19),
        finished_at=task_time(0, 19) + timedelta(minutes=10),
        report_path=str(tmp_path / "today.md"),
    )
    yesterday_task = _create_review_task(client, headers, old_revision="302")
    _set_review_task_stats_fields(
        db,
        yesterday_task["id"],
        created_at=task_time(-1, 10),
        status=ReviewTaskStatus.failed,
        report_path=str(tmp_path / "yesterday.md"),
    )
    older_task = _create_review_task(client, headers, old_revision="303")
    _set_review_task_stats_fields(
        db,
        older_task["id"],
        created_at=task_time(-2, 10),
        status=ReviewTaskStatus.running_codex,
        report_path=str(tmp_path / "older.md"),
    )
    last_week_task = _create_review_task(client, headers, old_revision="304")
    last_week_date = today_date - timedelta(days=7)
    _set_review_task_stats_fields(
        db,
        last_week_task["id"],
        created_at=task_time(-7, 10),
        status=ReviewTaskStatus.success,
        report_path=str(tmp_path / "last-week.md"),
    )

    today_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
        params={"range": "today"},
    )
    yesterday_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
        params={"range": "yesterday"},
    )
    this_week_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
        params={"range": "this_week"},
    )
    last_week_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
        params={"range": "last_week"},
    )
    all_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
        params={"range": "all"},
    )

    assert today_response.status_code == 200
    today_content = today_response.json()
    assert today_content["range"] == "today"
    assert today_content["total"] == 1
    assert today_content["success"] == 1
    assert today_content["running"] == 0
    assert today_content["failed_or_error"] == 0
    assert today_content["reports_by_user"] == []
    assert today_content["trend"][-1]["count"] == 1

    assert yesterday_response.status_code == 200
    yesterday_content = yesterday_response.json()
    assert yesterday_content["range"] == "yesterday"
    assert yesterday_content["total"] == 1
    assert yesterday_content["success"] == 0
    assert yesterday_content["failed_or_error"] == 1
    assert yesterday_content["reports_by_user"] == []
    assert yesterday_content["trend"][-1]["count"] == 1

    assert this_week_response.status_code == 200
    this_week_content = this_week_response.json()
    assert this_week_content["range"] == "this_week"
    task_statuses_by_date = [
        (today_date, ReviewTaskStatus.success),
        (today_date - timedelta(days=1), ReviewTaskStatus.failed),
        (today_date - timedelta(days=2), ReviewTaskStatus.running_codex),
        (last_week_date, ReviewTaskStatus.success),
    ]
    this_week_statuses = [
        status for task_date, status in task_statuses_by_date if task_date >= week_start
    ]
    assert this_week_content["total"] == len(this_week_statuses)
    assert this_week_content["success"] == this_week_statuses.count(
        ReviewTaskStatus.success
    )
    assert this_week_content["running"] == this_week_statuses.count(
        ReviewTaskStatus.running_codex
    )
    assert this_week_content["failed_or_error"] == this_week_statuses.count(
        ReviewTaskStatus.failed
    )

    assert last_week_response.status_code == 200
    last_week_content = last_week_response.json()
    assert last_week_content["range"] == "last_week"
    last_week_statuses = [
        status
        for task_date, status in task_statuses_by_date
        if last_week_start <= task_date < week_start
    ]
    assert last_week_content["total"] == len(last_week_statuses)
    assert last_week_content["success"] == last_week_statuses.count(
        ReviewTaskStatus.success
    )
    assert last_week_content["running"] == last_week_statuses.count(
        ReviewTaskStatus.running_codex
    )

    assert all_response.status_code == 200
    all_content = all_response.json()
    assert all_content["range"] == "all"
    assert all_content["total"] == 4
    assert all_content["success"] == 2
    assert all_content["running"] == 1
    assert all_content["failed_or_error"] == 1
    assert all_content["reports_by_user"] == []
    assert all_content["trend"] == [
        {"label": task_date.isoformat(), "count": 1}
        for task_date, _ in sorted(task_statuses_by_date)
    ]


def test_task_stats_average_duration_uses_valid_finished_tasks(
    client: TestClient,
    db: Session,
) -> None:
    own_user, own_api_key = create_key_user(db)
    own_headers = authentication_token_from_key(
        client=client, key=own_api_key.key, db=db
    )
    other_user, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client, key=other_api_key.key, db=db
    )

    valid_short = _create_review_task(client, own_headers, old_revision="401")
    _set_review_task_stats_fields(
        db,
        valid_short["id"],
        created_at=_today_elapsed_datetime(80),
        status=ReviewTaskStatus.success,
        started_at=_today_elapsed_datetime(80),
        finished_at=_today_elapsed_datetime(80) + timedelta(minutes=10),
    )
    valid_long = _create_review_task(client, own_headers, old_revision="402")
    _set_review_task_stats_fields(
        db,
        valid_long["id"],
        created_at=_today_elapsed_datetime(70),
        status=ReviewTaskStatus.failed,
        started_at=_today_elapsed_datetime(70),
        finished_at=_today_elapsed_datetime(70) + timedelta(minutes=20),
    )
    missing_started = _create_review_task(client, own_headers, old_revision="403")
    _set_review_task_stats_fields(
        db,
        missing_started["id"],
        created_at=_today_elapsed_datetime(60),
        status=ReviewTaskStatus.success,
        finished_at=_today_elapsed_datetime(60) + timedelta(minutes=5),
    )
    negative_duration = _create_review_task(client, own_headers, old_revision="404")
    _set_review_task_stats_fields(
        db,
        negative_duration["id"],
        created_at=_today_elapsed_datetime(50),
        status=ReviewTaskStatus.success,
        started_at=_today_elapsed_datetime(50),
        finished_at=_today_elapsed_datetime(50) - timedelta(minutes=5),
    )
    other_task = _create_review_task(client, other_headers, old_revision="405")
    _set_review_task_stats_fields(
        db,
        other_task["id"],
        created_at=_today_elapsed_datetime(40),
        status=ReviewTaskStatus.success,
        started_at=_today_elapsed_datetime(40),
        finished_at=_today_elapsed_datetime(40) + timedelta(hours=2),
    )
    yesterday_valid = _create_review_task(client, own_headers, old_revision="406")
    _set_review_task_stats_fields(
        db,
        yesterday_valid["id"],
        created_at=_shanghai_datetime(-1, 9),
        status=ReviewTaskStatus.success,
        started_at=_shanghai_datetime(-1, 9),
        finished_at=_shanghai_datetime(-1, 9) + timedelta(minutes=30),
    )

    today_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=own_headers,
        params={"range": "today"},
    )
    yesterday_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=own_headers,
        params={"range": "yesterday"},
    )
    all_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=own_headers,
        params={"range": "all"},
    )

    assert today_response.status_code == 200
    today_content = today_response.json()
    assert today_content["range"] == "today"
    assert today_content["total"] == 4
    assert today_content["average_duration_seconds"] == 900

    assert yesterday_response.status_code == 200
    yesterday_content = yesterday_response.json()
    assert yesterday_content["range"] == "yesterday"
    assert yesterday_content["total"] == 1
    assert yesterday_content["average_duration_seconds"] == 1800

    assert all_response.status_code == 200
    all_content = all_response.json()
    assert all_content["range"] == "all"
    assert all_content["total"] >= 5
    assert all_content["average_duration_seconds"] == 1200
    assert other_user.id != own_user.id


def test_admin_task_stats_include_all_nonzero_report_counts_by_user(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    report_user_prefix = f"报告用户-{random_lower_string()}"

    for index in range(21):
        user, api_key = create_key_user(db)
        user.full_name = f"{report_user_prefix} {index:02d}"
        db.add(user)
        db.commit()
        headers = authentication_token_from_key(client=client, key=api_key.key, db=db)

        for report_index in range(index + 1):
            task = _create_review_task(
                client,
                headers,
                old_revision=f"{index + 1000}{report_index}",
            )
            upload_response = client.post(
                f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
                headers=headers,
                files={
                    "upload": (
                        "report.md",
                        b"# Review\n",
                        "text/markdown",
                    )
                },
            )
            assert upload_response.status_code == 200
            if index == 20 and report_index == 0:
                _set_review_task_stats_fields(
                    db,
                    task["id"],
                    created_at=_shanghai_datetime(-1, 10),
                    status=ReviewTaskStatus.success,
                    report_path=upload_response.json()["report_path"],
                )

    no_report_user, no_report_api_key = create_key_user(db)
    no_report_user.full_name = "没有报告用户"
    db.add(no_report_user)
    db.commit()
    no_report_headers = authentication_token_from_key(
        client=client, key=no_report_api_key.key, db=db
    )
    _create_review_task(client, no_report_headers, old_revision="9999")

    failed_report_user, failed_report_api_key = create_key_user(db)
    failed_report_user.full_name = f"{report_user_prefix} 失败报告用户"
    db.add(failed_report_user)
    db.commit()
    failed_report_headers = authentication_token_from_key(
        client=client, key=failed_report_api_key.key, db=db
    )
    failed_task = _create_review_task(
        client,
        failed_report_headers,
        old_revision="10000",
    )
    _set_review_task_stats_fields(
        db,
        failed_task["id"],
        created_at=_shanghai_datetime(0, 11),
        status=ReviewTaskStatus.failed,
        report_path=str(tmp_path / "failed-report.md"),
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=superuser_token_headers,
    )

    assert response.status_code == 200
    reports_by_user = [
        user_stat
        for user_stat in response.json()["reports_by_user"]
        if user_stat["user_name"].startswith(report_user_prefix)
    ]
    assert len(reports_by_user) == 21
    report_user_20 = next(
        user_stat
        for user_stat in reports_by_user
        if user_stat["user_name"] == f"{report_user_prefix} 20"
    )
    assert report_user_20 == {
        "user_name": f"{report_user_prefix} 20",
        "svn_username": None,
        "report_count": 20,
    }
    assert reports_by_user[-1]["user_name"] == f"{report_user_prefix} 00"
    assert reports_by_user[-1]["report_count"] == 1
    assert "没有报告用户" not in {
        user_stat["user_name"] for user_stat in reports_by_user
    }
    assert f"{report_user_prefix} 失败报告用户" not in {
        user_stat["user_name"] for user_stat in reports_by_user
    }

    yesterday_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=superuser_token_headers,
        params={"range": "yesterday"},
    )

    assert yesterday_response.status_code == 200
    yesterday_report_users = [
        user_stat
        for user_stat in yesterday_response.json()["reports_by_user"]
        if user_stat["user_name"] == f"{report_user_prefix} 20"
    ]
    assert yesterday_report_users == [
        {
            "user_name": f"{report_user_prefix} 20",
            "svn_username": None,
            "report_count": 1,
        }
    ]


def test_normal_user_task_stats_do_not_include_report_counts_by_user(
    client: TestClient,
    db: Session,
) -> None:
    _, api_key = create_key_user(db)
    headers = authentication_token_from_key(client=client, key=api_key.key, db=db)

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
    )

    assert response.status_code == 200
    assert response.json()["reports_by_user"] == []


def test_task_stats_include_all_status_keys_with_zero_counts(
    client: TestClient, db: Session
) -> None:
    _, api_key = create_key_user(db)
    headers = authentication_token_from_key(client=client, key=api_key.key, db=db)

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=headers,
    )

    assert response.status_code == 200
    content = response.json()
    assert set(content["by_status"]) == {
        "created",
        "checking_environment",
        "running_codex",
        "uploading_report",
        "success",
        "failed",
        "environment_error",
        "cancelled",
    }
    for count in content["by_status"].values():
        assert isinstance(count, int)
        assert count == 0
    assert content["total"] == 0


def test_role_admin_can_read_other_user_task_detail(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    db: Session,
) -> None:
    user_task = _create_review_task(client, normal_user_token_headers, old_revision="7")
    admin, admin_api_key = create_key_user(db)
    admin.role = UserRole.admin
    db.add(admin)
    db.commit()
    admin_headers = authentication_token_from_key(
        client=client, key=admin_api_key.key, db=db
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{user_task['id']}",
        headers=admin_headers,
    )

    assert response.status_code == 200
    assert response.json()["id"] == user_task["id"]


def test_department_admin_lists_only_own_department_tasks(
    client: TestClient,
    db: Session,
) -> None:
    department_a = _create_department(db, f"审查部门A-{random_lower_string()}")
    department_b = _create_department(db, f"审查部门B-{random_lower_string()}")
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    same_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        full_name="同部门成员",
        svn_username=f"same-dept-{random_lower_string()}",
    )
    other_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_b,
        full_name="其他部门成员",
        svn_username=f"other-dept-{random_lower_string()}",
    )
    same_task = _create_review_task(client, same_headers, old_revision="7101")
    other_task = _create_review_task(client, other_headers, old_revision="7102")

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=admin_headers,
    )

    assert response.status_code == 200
    listed_ids = {task["id"] for task in response.json()["data"]}
    assert same_task["id"] in listed_ids
    assert other_task["id"] not in listed_ids


def test_department_admin_filters_do_not_expand_beyond_department(
    client: TestClient,
    db: Session,
) -> None:
    department_a = _create_department(db, f"搜索部门A-{random_lower_string()}")
    department_b = _create_department(db, f"搜索部门B-{random_lower_string()}")
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    same_headers, same_user = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        full_name="搜索目标",
        svn_username=f"target-{random_lower_string()}",
    )
    other_headers, other_user = _create_department_user(
        client=client,
        db=db,
        department=department_b,
        full_name="搜索目标",
        svn_username=f"target-{random_lower_string()}",
    )
    same_task = _create_review_task(client, same_headers, old_revision="7111")
    other_task = _create_review_task(client, other_headers, old_revision="7112")

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=admin_headers,
        params={"user_name": "搜索目标", "svn_username": "target"},
    )

    assert response.status_code == 200
    listed_ids = {task["id"] for task in response.json()["data"]}
    assert same_task["id"] in listed_ids
    assert other_task["id"] not in listed_ids
    assert same_user.department_id == department_a.id
    assert other_user.department_id == department_b.id


def test_department_admin_can_read_and_download_same_department_member_report(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department = _create_department(db)
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    member_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
    )
    task = _create_review_task(client, member_headers, old_revision="7121")
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=member_headers,
        files={"upload": ("report.md", b"# Same Department\n", "text/markdown")},
    )
    assert upload_response.status_code == 200

    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=admin_headers,
    )
    download_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/download",
        headers=admin_headers,
    )

    assert read_response.status_code == 200
    assert read_response.text == "# Same Department\n"
    assert download_response.status_code == 200
    assert download_response.content == b"# Same Department\n"


def test_department_admin_cannot_read_other_department_task_or_report(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department_a = _create_department(db, f"读取部门A-{random_lower_string()}")
    department_b = _create_department(db, f"读取部门B-{random_lower_string()}")
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    other_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_b,
    )
    task = _create_review_task(client, other_headers, old_revision="7131")
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=other_headers,
        files={"upload": ("report.md", b"# Other Department\n", "text/markdown")},
    )
    assert upload_response.status_code == 200

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=admin_headers,
    )
    report_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=admin_headers,
    )
    download_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/download",
        headers=admin_headers,
    )

    assert detail_response.status_code == 403
    assert report_response.status_code == 403
    assert download_response.status_code == 403


def test_department_admin_can_download_same_department_member_pdf_report(
    client: TestClient,
    db: Session,
    tmp_path: Path,
) -> None:
    department = _create_department(db)
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    member_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
    )
    task = _create_review_task(client, member_headers, old_revision="7132")
    pdf_path = tmp_path / "same-department.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 same department\n")
    _set_review_task_pdf_report(
        db,
        task["id"],
        pdf_report_path=str(pdf_path),
        pdf_report_name="same-department.pdf",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf/download",
        headers=admin_headers,
    )

    assert response.status_code == 200
    assert response.content == b"%PDF-1.4 same department\n"
    assert "same-department.pdf" in unquote(response.headers["content-disposition"])


def test_department_admin_cannot_download_other_department_pdf_report(
    client: TestClient,
    db: Session,
    tmp_path: Path,
) -> None:
    department_a = _create_department(db, f"PDF部门A-{random_lower_string()}")
    department_b = _create_department(db, f"PDF部门B-{random_lower_string()}")
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    other_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_b,
    )
    task = _create_review_task(client, other_headers, old_revision="7133")
    pdf_path = tmp_path / "other-department.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 other department\n")
    _set_review_task_pdf_report(
        db,
        task["id"],
        pdf_report_path=str(pdf_path),
        pdf_report_name="other-department.pdf",
    )

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf/download",
        headers=admin_headers,
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Not enough permissions"


def test_department_admin_cannot_delete_same_department_member_task(
    client: TestClient,
    db: Session,
) -> None:
    department = _create_department(db)
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    member_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
    )
    task = _create_review_task(client, member_headers, old_revision="7141")

    response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=admin_headers,
    )
    read_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=member_headers,
    )

    assert response.status_code == 403
    assert read_response.status_code == 200


def test_department_admin_cannot_write_same_department_member_task(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department = _create_department(db)
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    member_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
    )
    task = _create_review_task(client, member_headers, old_revision="7142")

    status_response = client.patch(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/status",
        headers=admin_headers,
        json={"status": "running_codex"},
    )
    markdown_upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=admin_headers,
        files={"upload": ("report.md", b"# Forbidden\n", "text/markdown")},
    )
    pdf_upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/pdf-report",
        headers=admin_headers,
        files={"upload": ("report.pdf", b"%PDF-1.4\n", "application/pdf")},
    )
    requirement_upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/requirement-file",
        headers=admin_headers,
        files={"upload": ("requirement.txt", b"forbidden", "text/plain")},
    )

    assert status_response.status_code == 403
    assert status_response.json()["detail"] == "Not enough permissions"
    assert markdown_upload_response.status_code == 403
    assert markdown_upload_response.json()["detail"] == "Not enough permissions"
    assert pdf_upload_response.status_code == 403
    assert pdf_upload_response.json()["detail"] == "Not enough permissions"
    assert requirement_upload_response.status_code == 403
    assert requirement_upload_response.json()["detail"] == "Not enough permissions"


def test_department_admin_can_update_own_task_status(
    client: TestClient,
    db: Session,
) -> None:
    department = _create_department(db)
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    task = _create_review_task(client, admin_headers, old_revision="7151")

    response = client.patch(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/status",
        headers=admin_headers,
        json={"status": "running_codex"},
    )

    assert response.status_code == 200
    assert response.json()["status"] == "running_codex"


def test_department_admin_task_stats_scope_and_reports_by_user(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department_a = _create_department(db, f"统计部门A-{random_lower_string()}")
    department_b = _create_department(db, f"统计部门B-{random_lower_string()}")
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    same_headers, same_user = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        full_name="统计同部门成员",
        svn_username=f"stats-same-{random_lower_string()}",
    )
    other_headers, other_user = _create_department_user(
        client=client,
        db=db,
        department=department_b,
        full_name="统计其他部门成员",
        svn_username=f"stats-other-{random_lower_string()}",
    )
    same_task = _create_review_task(client, same_headers, old_revision="7161")
    same_upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{same_task['id']}/report",
        headers=same_headers,
        files={"upload": ("same.md", b"# Same\n", "text/markdown")},
    )
    assert same_upload_response.status_code == 200
    other_task = _create_review_task(client, other_headers, old_revision="7162")
    other_upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{other_task['id']}/report",
        headers=other_headers,
        files={"upload": ("other.md", b"# Other\n", "text/markdown")},
    )
    assert other_upload_response.status_code == 200

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=admin_headers,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["total"] == 1
    assert content["success"] == 1
    assert content["reports_by_user"] == [
        {
            "user_name": same_user.full_name,
            "svn_username": same_user.svn_username,
            "report_count": 1,
        }
    ]
    assert other_user.full_name not in {
        user_stat["user_name"] for user_stat in content["reports_by_user"]
    }


def test_unbound_department_admin_read_scope_returns_403(
    client: TestClient,
    db: Session,
) -> None:
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=None,
        role=UserRole.department_admin,
    )

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=admin_headers,
    )
    stats_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=admin_headers,
    )

    assert list_response.status_code == 403
    assert list_response.json()["detail"] == "当前账号未绑定部门，请联系超级管理员。"
    assert stats_response.status_code == 403
    assert stats_response.json()["detail"] == "当前账号未绑定部门，请联系超级管理员。"


def test_inactive_department_admin_read_scope_returns_403(
    client: TestClient,
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(tmp_path))
    department = _create_department(db)
    department.is_active = False
    db.add(department)
    db.commit()
    db.refresh(department)
    admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
    )
    member_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department,
    )
    task = _create_review_task(client, member_headers, old_revision="7171")
    upload_response = client.post(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/report",
        headers=member_headers,
        files={"upload": ("report.md", b"# Inactive Department\n", "text/markdown")},
    )
    assert upload_response.status_code == 200

    list_response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=admin_headers,
    )
    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=admin_headers,
    )
    download_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}/download",
        headers=admin_headers,
    )
    stats_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/stats",
        headers=admin_headers,
    )

    expected_detail = "当前账号未绑定部门，请联系超级管理员。"
    assert list_response.status_code == 403
    assert list_response.json()["detail"] == expected_detail
    assert detail_response.status_code == 403
    assert detail_response.json()["detail"] == expected_detail
    assert download_response.status_code == 403
    assert download_response.json()["detail"] == expected_detail
    assert stats_response.status_code == 403
    assert stats_response.json()["detail"] == expected_detail


def test_submit_report_feedback_helpful_requires_useful_text(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    task = _create_review_task(client, normal_user_token_headers)
    _set_review_task_report(db, task["id"])

    response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        task["id"],
        helpful=True,
        useful_issues_text="   ",
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "有帮助时请至少填写一条有用信息"
    assert (
        db.exec(
            select(ReviewTaskReportFeedback).where(
                ReviewTaskReportFeedback.task_id == UUID(task["id"])
            )
        ).all()
        == []
    )


def test_submit_report_feedback_helpful_success_and_task_public_fields(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    task = _create_review_task(client, normal_user_token_headers)
    _set_review_task_report(db, task["id"])

    response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        task["id"],
        helpful=True,
        useful_issues_text="1. 空指针风险\n2. 越界访问风险",
    )

    assert response.status_code == 200
    content = response.json()
    assert content["task_id"] == task["id"]
    assert content["helpful"] is True
    assert content["useful_issues_text"] == "1. 空指针风险\n2. 越界访问风险"

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/{task['id']}",
        headers=normal_user_token_headers,
    )
    assert detail_response.status_code == 200
    detail = detail_response.json()
    assert detail["report_feedback_helpful"] is True
    assert (
        detail["report_feedback_useful_issues_text"] == "1. 空指针风险\n2. 越界访问风险"
    )


def test_submit_report_feedback_not_helpful_allows_empty_text(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    task = _create_review_task(client, normal_user_token_headers)
    _set_review_task_report(db, task["id"])

    response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        task["id"],
        helpful=False,
        useful_issues_text=None,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["helpful"] is False
    assert content["useful_issues_text"] in (None, "")


def test_submit_report_feedback_rejects_duplicate(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    task = _create_review_task(client, normal_user_token_headers)
    _set_review_task_report(db, task["id"])
    first_response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        task["id"],
        helpful=False,
        useful_issues_text=None,
    )
    assert first_response.status_code == 200

    duplicate_response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        task["id"],
        helpful=False,
        useful_issues_text=None,
    )

    assert duplicate_response.status_code == 409
    assert duplicate_response.json()["detail"] == "report_feedback_exists"


def test_submit_report_feedback_requires_report_and_pdf(
    client: TestClient, normal_user_token_headers: dict[str, str]
) -> None:
    task = _create_review_task(client, normal_user_token_headers)

    response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        task["id"],
        helpful=False,
        useful_issues_text=None,
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "no_report"


def test_report_feedback_department_admin_can_submit_same_department_task(
    client: TestClient, db: Session
) -> None:
    department = _create_department(db)
    member_headers, member = _create_department_user(
        client=client,
        db=db,
        department=department,
        full_name="成员",
    )
    admin_headers, _admin = _create_department_user(
        client=client,
        db=db,
        department=department,
        role=UserRole.department_admin,
        full_name="部门管理员",
    )
    task = _create_review_task(client, member_headers)
    _set_review_task_report(db, task["id"])

    response = _submit_report_feedback(
        client,
        admin_headers,
        task["id"],
        helpful=False,
        useful_issues_text=None,
    )

    assert response.status_code == 200
    assert response.json()["user_id"] != str(member.id)


def test_report_feedback_department_admin_cannot_submit_other_or_null_department_task(
    client: TestClient, db: Session
) -> None:
    own_department = _create_department(db, "本部门")
    other_department = _create_department(db, "其他部门")
    admin_headers, _admin = _create_department_user(
        client=client,
        db=db,
        department=own_department,
        role=UserRole.department_admin,
    )
    other_headers, _other_user = _create_department_user(
        client=client,
        db=db,
        department=other_department,
    )
    null_department_headers, _null_user = _create_department_user(
        client=client,
        db=db,
        department=None,
    )
    other_task = _create_review_task(client, other_headers)
    null_task = _create_review_task(client, null_department_headers)
    _set_review_task_report(db, other_task["id"])
    _set_review_task_report(db, null_task["id"])

    other_response = _submit_report_feedback(
        client,
        admin_headers,
        other_task["id"],
        helpful=False,
        useful_issues_text=None,
    )
    null_response = _submit_report_feedback(
        client,
        admin_headers,
        null_task["id"],
        helpful=False,
        useful_issues_text=None,
    )

    assert other_response.status_code == 403
    assert null_response.status_code == 403


def test_review_task_list_includes_report_feedback_fields(
    client: TestClient, normal_user_token_headers: dict[str, str], db: Session
) -> None:
    first_task = _create_review_task(client, normal_user_token_headers)
    second_task = _create_review_task(client, normal_user_token_headers)
    _set_review_task_report(db, first_task["id"])
    _set_review_task_report(db, second_task["id"])
    feedback_response = _submit_report_feedback(
        client,
        normal_user_token_headers,
        first_task["id"],
        helpful=True,
        useful_issues_text="1. 有价值内容",
    )
    assert feedback_response.status_code == 200

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks",
        headers=normal_user_token_headers,
    )

    assert response.status_code == 200
    by_id = {item["id"]: item for item in response.json()["data"]}
    assert by_id[first_task["id"]]["report_feedback_helpful"] is True
    assert (
        by_id[first_task["id"]]["report_feedback_useful_issues_text"] == "1. 有价值内容"
    )
    assert by_id[second_task["id"]]["report_feedback_helpful"] is None
    assert by_id[second_task["id"]]["report_feedback_useful_issues_text"] is None


def test_weekly_insights_uses_previous_week_and_current_user_scope(
    client: TestClient,
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fixed_now = datetime(2026, 7, 27, 9, tzinfo=timezone.utc)
    monkeypatch.setattr(review_tasks_route, "_utcnow", lambda: fixed_now)

    def task_time(day: int, hour: int = 10) -> datetime:
        return (
            datetime(
                2026,
                7,
                day,
                hour,
                tzinfo=TASK_STATS_TIMEZONE,
            )
            .astimezone(timezone.utc)
            .replace(tzinfo=None)
        )

    user, api_key = create_key_user(db)
    headers = authentication_token_from_key(client=client, key=api_key.key, db=db)
    helpful_task = _create_review_task(client, headers, old_revision="701")
    not_helpful_task = _create_review_task(client, headers, old_revision="702")
    no_feedback_task = _create_review_task(client, headers, old_revision="703")
    for task, created_at in (
        (helpful_task, task_time(20)),
        (not_helpful_task, task_time(21)),
        (no_feedback_task, task_time(22)),
    ):
        _set_review_task_stats_fields(
            db,
            task["id"],
            created_at=created_at,
            status=ReviewTaskStatus.success,
        )

    db.add(
        ReviewTaskReportFeedback(
            task_id=UUID(helpful_task["id"]),
            user_id=user.id,
            helpful=True,
            useful_issues_text=(
                "1. 高风险：CAN 超时后未执行重发\n2. 中风险：Flash 写入地址未校验"
            ),
        )
    )
    db.add(
        ReviewTaskReportFeedback(
            task_id=UUID(not_helpful_task["id"]),
            user_id=user.id,
            helpful=False,
        )
    )
    db.commit()

    other_user, other_api_key = create_key_user(db)
    other_headers = authentication_token_from_key(
        client=client,
        key=other_api_key.key,
        db=db,
    )
    other_task = _create_review_task(client, other_headers, old_revision="704")
    _set_review_task_stats_fields(
        db,
        other_task["id"],
        created_at=task_time(20),
        status=ReviewTaskStatus.success,
    )
    assert other_user.id != user.id

    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/weekly-insights",
        headers=headers,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["period_start"] == "2026-07-20"
    assert content["period_end"] == "2026-07-26"
    assert content["scope_label"] == "个人"
    assert content["success_tasks"] == 3
    assert content["helpful"] == 1
    assert content["not_helpful"] == 1
    assert content["no_feedback"] == 1
    assert content["feedback_coverage_rate"] == pytest.approx(2 / 3)
    assert content["explicit_useful_rate"] == 0.5
    assert content["valid_useful_tasks"] == 1
    assert content["valid_issue_count"] == 2
    assert content["risk_counts"]["high"] == 1
    assert content["risk_counts"]["medium"] == 1
    assert len(content["daily"]) == 7
    assert sum(item["success"] for item in content["daily"]) == 3
    by_subcategory = {item["subcategory"]: item for item in content["categories"]}
    assert by_subcategory["超时重发审核"]["useful_count"] == 1
    assert by_subcategory["Flash操作审核"]["useful_count"] == 1


def test_custom_review_inputs_superuser_filters_and_export_share_scope(
    client: TestClient,
    db: Session,
    superuser_token_headers: dict[str, str],
) -> None:
    department = _create_department(db, f"自定义统计部门-{random_lower_string()}")
    member_headers, member = _create_department_user(
        client=client,
        db=db,
        department=department,
        full_name="自定义统计用户",
        svn_username="custom-stat-user",
    )
    matched_task = _create_custom_review_task(
        client,
        member_headers,
        rule="保护条件必须具有明确恢复路径",
        old_revision="201",
        new_revision="202",
        has_word_material=True,
        has_excel_material=False,
        has_pdf_material=True,
        svn_url="https://svn.example.com/custom-stat/matched",
    )
    matched_record = db.get(ReviewTask, UUID(matched_task["id"]))
    assert matched_record is not None
    matched_record.status = ReviewTaskStatus.failed
    db.add(matched_record)
    db.commit()

    _create_custom_review_task(
        client,
        member_headers,
        rule="其他自定义规则",
        old_revision="301",
        new_revision="302",
        mode="new_revision_full",
        svn_url="https://svn.example.com/custom-stat/other",
    )
    _create_review_task(client, member_headers, old_revision="401")

    today = datetime.now(TASK_STATS_TIMEZONE).date().isoformat()
    query = {
        "keyword": "恢复路径",
        "created_from": today,
        "created_to": today,
        "user_name": member.full_name,
        "department_name": department.name,
        "svn_url": "matched",
        "revision": "201",
        "mode": "dual_revision",
        "has_word_material": "true",
        "has_excel_material": "false",
        "has_pdf_material": "true",
        "status": "failed",
    }
    response = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs",
        headers=superuser_token_headers,
        params=query,
    )

    assert response.status_code == 200
    content = response.json()
    assert content["count"] == 1
    record = content["data"][0]
    assert record["id"] == matched_task["id"]
    assert record["custom_review_rule"] == "保护条件必须具有明确恢复路径"
    assert record["user_name"] == member.full_name
    assert record["department_name"] == department.name
    assert record["status"] == "failed"

    detail_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs/{matched_task['id']}",
        headers=superuser_token_headers,
    )
    assert detail_response.status_code == 200
    assert detail_response.json()["custom_review_rule"] == record["custom_review_rule"]

    export_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs/export",
        headers=superuser_token_headers,
        params=query,
    )
    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.content))
    sheet = workbook["自定义审查输入统计"]
    assert sheet.max_row == 2
    assert sheet.cell(2, 13).value == "保护条件必须具有明确恢复路径"
    assert sheet.cell(2, 9).value == "已提供"
    assert sheet.cell(2, 10).value == "未提供"
    assert sheet.cell(2, 11).value == "已提供"
    workbook.close()


def test_custom_review_inputs_enforce_management_scope_and_ignore_empty_history(
    client: TestClient,
    db: Session,
) -> None:
    department_a = _create_department(db, f"自定义范围A-{random_lower_string()}")
    department_b = _create_department(db, f"自定义范围B-{random_lower_string()}")
    department_admin_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        role=UserRole.department_admin,
    )
    same_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
        full_name="本部门自定义用户",
    )
    other_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_b,
        full_name="其他部门自定义用户",
    )
    normal_headers, _ = _create_department_user(
        client=client,
        db=db,
        department=department_a,
    )
    same_task = _create_custom_review_task(
        client, same_headers, rule="本部门规则", old_revision="501", new_revision="502"
    )
    cancelled_task = _create_custom_review_task(
        client,
        same_headers,
        rule="已取消任务规则",
        old_revision="503",
        new_revision="504",
    )
    cancelled_response = client.patch(
        f"{settings.API_V1_STR}/review-tasks/{cancelled_task['id']}/status",
        headers=same_headers,
        json={"status": "cancelled"},
    )
    assert cancelled_response.status_code == 200
    assert cancelled_response.json()["status"] == "cancelled"
    other_task = _create_custom_review_task(
        client,
        other_headers,
        rule="其他部门规则",
        old_revision="601",
        new_revision="602",
    )
    empty_history = _create_review_task(client, same_headers, old_revision="701")
    empty_record = db.get(ReviewTask, UUID(empty_history["id"]))
    assert empty_record is not None
    empty_record.custom_review_rule = "   "
    db.add(empty_record)
    db.commit()

    scoped_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs",
        headers=department_admin_headers,
        params={"department_name": department_b.name},
    )
    assert scoped_response.status_code == 200
    assert scoped_response.json()["count"] == 0

    own_scope_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs",
        headers=department_admin_headers,
    )
    assert own_scope_response.status_code == 200
    own_scope_ids = {item["id"] for item in own_scope_response.json()["data"]}
    assert own_scope_ids == {same_task["id"], cancelled_task["id"]}
    assert any(
        item["id"] == cancelled_task["id"] and item["status"] == "cancelled"
        for item in own_scope_response.json()["data"]
    )

    forbidden_detail = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs/{other_task['id']}",
        headers=department_admin_headers,
    )
    assert forbidden_detail.status_code == 404
    scoped_export = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs/export",
        headers=department_admin_headers,
    )
    workbook = load_workbook(BytesIO(scoped_export.content))
    sheet = workbook["自定义审查输入统计"]
    assert sheet.max_row == 3
    assert {sheet.cell(row, 13).value for row in (2, 3)} == {
        "本部门规则",
        "已取消任务规则",
    }
    assert "cancelled" in {sheet.cell(row, 12).value for row in (2, 3)}
    workbook.close()

    deleted_response = client.delete(
        f"{settings.API_V1_STR}/review-tasks/{same_task['id']}",
        headers=same_headers,
    )
    assert deleted_response.status_code == 200
    after_delete_response = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs",
        headers=department_admin_headers,
    )
    assert [item["id"] for item in after_delete_response.json()["data"]] == [
        cancelled_task["id"]
    ]
    after_delete_export = client.get(
        f"{settings.API_V1_STR}/review-tasks/custom-review-inputs/export",
        headers=department_admin_headers,
    )
    workbook = load_workbook(BytesIO(after_delete_export.content))
    sheet = workbook["自定义审查输入统计"]
    assert sheet.max_row == 2
    assert sheet.cell(2, 13).value == "已取消任务规则"
    workbook.close()

    for path in (
        "/custom-review-inputs",
        "/custom-review-inputs/export",
        f"/custom-review-inputs/{same_task['id']}",
    ):
        response = client.get(
            f"{settings.API_V1_STR}/review-tasks{path}", headers=normal_headers
        )
        assert response.status_code == 403

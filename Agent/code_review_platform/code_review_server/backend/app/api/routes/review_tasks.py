import json
import logging
import re
import uuid
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, PlainTextResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import ValidationError
from sqlalchemy import case, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.sql.elements import ColumnElement
from sqlmodel import col, delete, func, select

from app.api.deps import (
    CurrentUser,
    SessionDep,
    ensure_department_admin_has_department,
    has_admin_privileges,
    has_management_privileges,
)
from app.models import (
    CustomReviewInputPublic,
    CustomReviewInputsPublic,
    CustomReviewMode,
    Department,
    ReviewIssueSeverity,
    ReviewTask,
    ReviewTaskCreate,
    ReviewTaskIssue,
    ReviewTaskIssueFeedbackDraftUpdate,
    ReviewTaskIssuePublic,
    ReviewTaskIssueSnapshotCreate,
    ReviewTaskIssuesPublic,
    ReviewTaskPublic,
    ReviewTaskReportFeedback,
    ReviewTaskReportFeedbackCreate,
    ReviewTaskReportFeedbackPublic,
    ReviewTaskReportUserStatsPublic,
    ReviewTaskRequirementFile,
    ReviewTasksPublic,
    ReviewTaskStatsPublic,
    ReviewTaskStatsRange,
    ReviewTaskStatus,
    ReviewTaskStatusUpdate,
    ReviewTaskTrendPointPublic,
    ReviewTaskWeeklyInsightsPublic,
    User,
    UserRole,
)
from app.services import report_service
from app.services.weekly_insight_service import (
    build_weekly_insights,
    previous_complete_week_bounds,
)

router = APIRouter(prefix="/review-tasks", tags=["review-tasks"])
logger = logging.getLogger(__name__)
REVIEW_REPORT_FILE_NAME = "代码审查报告.md"
PDF_REPORT_FILE_NAME = "代码审查报告.pdf"
MAX_REPORT_ISSUE_COUNT = 200
MAX_REPORT_ISSUES_JSON_BYTES = 1024 * 1024
TASK_STATS_TIMEZONE = ZoneInfo("Asia/Shanghai")
ISSUE_FEEDBACK_EXPORT_HEADERS = (
    "提交人",
    "日期",
    "SVN路径",
    "旧版本号",
    "新版本号",
    "选择的审查项",
    "AI报告问题原文",
    "问题描述原文",
    "是否有帮助",
    "审查问题反馈",
)
ISSUE_FEEDBACK_EXPORT_COLUMN_WIDTHS = (18, 12, 42, 12, 12, 38, 42, 56, 14, 56)
CUSTOM_REVIEW_INPUT_EXPORT_HEADERS = (
    "提交时间",
    "用户",
    "SVN账号",
    "部门",
    "SVN地址",
    "旧版本号",
    "新版本号",
    "审查模式",
    "已提供Word材料",
    "已提供Excel材料",
    "已提供PDF材料",
    "任务状态",
    "自定义审查规则",
)
CUSTOM_REVIEW_INPUT_EXPORT_COLUMN_WIDTHS = (
    22,
    18,
    18,
    18,
    44,
    12,
    12,
    18,
    16,
    16,
    16,
    18,
    80,
)
EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_REVISION_RE = re.compile(r"^r?(\d+)$")
EXCEL_ILLEGAL_CHARACTERS_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")
REVIEW_ITEM_EXPORT_NAMES = {
    "DS-01": "需求变更审查",
    "DS-02": "编码规范审查",
    "DS-03": "MOS 充放电状态机审查",
    "DS-04": "AFE 驱动审核",
    "DS-05": "变更点风险检查",
    "DS-06": "规格书审查",
    "DS-07": "休眠唤醒状态机审查",
    "DS-99": "自定义审查",
}

_RUNNING_STATUSES = {
    ReviewTaskStatus.checking_environment,
    ReviewTaskStatus.running_codex,
    ReviewTaskStatus.uploading_report,
}

_FAILED_OR_ERROR_STATUSES = {
    ReviewTaskStatus.failed,
    ReviewTaskStatus.environment_error,
    ReviewTaskStatus.cancelled,
}

ReportFeedbackFilter = Literal["helpful", "not_helpful", "none"]


@dataclass(frozen=True)
class _ReviewIssueSeverityRule:
    severity: ReviewIssueSeverity
    rank: int
    canonical_label: str
    historical_labels: tuple[str, ...]
    legacy_prefixes: tuple[str, ...]


class _ReviewIssueSourceConsistencyStatus(str, Enum):
    valid = "valid"
    invalid = "invalid"


@dataclass(frozen=True)
class _ReviewIssueSourceConsistencyState:
    status: _ReviewIssueSourceConsistencyStatus

    @property
    def is_valid(self) -> bool:
        return self.status == _ReviewIssueSourceConsistencyStatus.valid


@dataclass(frozen=True)
class _EffectiveReviewTaskIssues:
    source_consistency_state: _ReviewIssueSourceConsistencyState
    issues: tuple[ReviewTaskIssue, ...]


_REVIEW_ISSUE_SEVERITY_RULES = (
    _ReviewIssueSeverityRule(
        severity=ReviewIssueSeverity.high,
        rank=4,
        canonical_label="严重",
        historical_labels=("高风险", "高"),
        legacy_prefixes=("高风险：", "高风险:"),
    ),
    _ReviewIssueSeverityRule(
        severity=ReviewIssueSeverity.medium,
        rank=3,
        canonical_label="重要",
        historical_labels=("中风险", "中"),
        legacy_prefixes=("中风险：", "中风险:"),
    ),
    _ReviewIssueSeverityRule(
        severity=ReviewIssueSeverity.low,
        rank=2,
        canonical_label="一般",
        historical_labels=("低风险", "低"),
        legacy_prefixes=("低风险：", "低风险:"),
    ),
    _ReviewIssueSeverityRule(
        severity=ReviewIssueSeverity.suggestion,
        rank=1,
        canonical_label="建议",
        historical_labels=("建议",),
        legacy_prefixes=("建议：", "建议:"),
    ),
)


def _unique_labels(labels: Sequence[str]) -> tuple[str, ...]:
    return tuple(dict.fromkeys(labels))


def _numbered_issue_severity_pattern(
    rule: _ReviewIssueSeverityRule,
    *,
    postgres: bool,
) -> str:
    historical_label_pattern = "|".join(
        re.escape(label) for label in rule.historical_labels
    )
    pc_bms_label_pattern = "|".join(
        re.escape(label)
        for label in _unique_labels((rule.canonical_label, *rule.historical_labels))
    )
    if postgres:
        return (
            r"^(([0-9]+[.)、．]|[一二三四五六七八九十百]+[、.．)])"
            r"[[:space:]]*)?((P[0-9]+|CR-[0-9]+)[[:space:]]+\[("
            f"{historical_label_pattern}"
            r")\]|PC-BMS-[0-9]+[[:space:]]+\[("
            f"{pc_bms_label_pattern}"
            r")\]|DS-02-[0-9]+[[:space:]]+\[("
            f"{historical_label_pattern}"
            r")\])[[:space:]]+[^[:space:]]"
        )
    return (
        r"^(?:(?:\d+[.)、．]|[一二三四五六七八九十百]+[、.．)])\s*)?"
        rf"(?:(?:P\d+|CR-\d+)\s+\[(?:{historical_label_pattern})\]"
        rf"|PC-BMS-\d+\s+\[(?:{pc_bms_label_pattern})\]"
        rf"|DS-02-\d+\s+\[(?:{historical_label_pattern})\])\s+\S"
    )


def _review_issue_severity_sql_condition(
    issue_text: Any,
    rule: _ReviewIssueSeverityRule,
) -> ColumnElement[bool]:
    return or_(
        issue_text.op("~")(_numbered_issue_severity_pattern(rule, postgres=True)),
        *(issue_text.startswith(prefix) for prefix in rule.legacy_prefixes),
    )


def _review_issue_text_matches_severity(
    issue_text: str,
    rule: _ReviewIssueSeverityRule,
) -> bool:
    return bool(
        re.match(
            _numbered_issue_severity_pattern(rule, postgres=False),
            issue_text,
        )
    ) or issue_text.startswith(rule.legacy_prefixes)


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _ensure_department_admin_active_department(
    *, session: SessionDep, current_user: CurrentUser
) -> None:
    if current_user.role != UserRole.department_admin:
        return
    ensure_department_admin_has_department(current_user)
    department = session.get(Department, current_user.department_id)
    if department is None or not department.is_active:
        raise HTTPException(
            status_code=403,
            detail="当前账号未绑定部门，请联系超级管理员。",
        )


def _task_read_scope_filters(
    *, session: SessionDep, current_user: CurrentUser
) -> list[Any]:
    if has_admin_privileges(current_user):
        return []
    if current_user.role == UserRole.department_admin:
        _ensure_department_admin_active_department(
            session=session, current_user=current_user
        )
        return [ReviewTask.department_id == current_user.department_id]
    return [ReviewTask.user_id == current_user.id]


def _custom_review_input_scope_filters(
    *, session: SessionDep, current_user: CurrentUser
) -> list[Any]:
    if not has_management_privileges(current_user):
        raise HTTPException(status_code=403, detail="没有足够权限")
    if current_user.role == UserRole.department_admin:
        _ensure_department_admin_active_department(
            session=session, current_user=current_user
        )
        return [ReviewTask.department_id == current_user.department_id]
    return []


def _parse_created_date(value: str | None, *, is_end: bool) -> datetime | None:
    if not value or not value.strip():
        return None
    try:
        local_date = date.fromisoformat(value.strip())
    except ValueError:
        raise HTTPException(
            status_code=422, detail="日期格式必须为 YYYY-MM-DD"
        ) from None
    local_datetime = datetime.combine(
        local_date + timedelta(days=1) if is_end else local_date,
        time.min,
        tzinfo=TASK_STATS_TIMEZONE,
    )
    return local_datetime.astimezone(timezone.utc).replace(tzinfo=None)


def _custom_review_input_filters(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    keyword: str | None = None,
    created_from: str | None = None,
    created_to: str | None = None,
    user_name: str | None = None,
    department_name: str | None = None,
    svn_url: str | None = None,
    revision: str | None = None,
    mode: CustomReviewMode | None = None,
    has_word_material: bool | None = None,
    has_excel_material: bool | None = None,
    has_pdf_material: bool | None = None,
    status: ReviewTaskStatus | None = None,
) -> list[Any]:
    filters = [
        *_custom_review_input_scope_filters(session=session, current_user=current_user),
        col(ReviewTask.custom_review_rule).is_not(None),
        func.length(func.trim(col(ReviewTask.custom_review_rule))) > 0,
    ]
    normalized_keyword = keyword.strip() if keyword else ""
    if normalized_keyword:
        filters.append(
            col(ReviewTask.custom_review_rule).ilike(f"%{normalized_keyword}%")
        )
    normalized_user_name = user_name.strip() if user_name else ""
    if normalized_user_name:
        filters.append(
            or_(
                col(User.full_name).ilike(f"%{normalized_user_name}%"),
                col(User.svn_username).ilike(f"%{normalized_user_name}%"),
            )
        )
    normalized_department_name = department_name.strip() if department_name else ""
    if normalized_department_name:
        filters.append(
            col(ReviewTask.department_name).ilike(f"%{normalized_department_name}%")
        )
    normalized_svn_url = svn_url.strip() if svn_url else ""
    if normalized_svn_url:
        filters.append(col(ReviewTask.svn_url).ilike(f"%{normalized_svn_url}%"))
    normalized_revision = revision.strip() if revision else ""
    if normalized_revision:
        filters.append(
            or_(
                col(ReviewTask.old_revision).ilike(f"%{normalized_revision}%"),
                col(ReviewTask.new_revision).ilike(f"%{normalized_revision}%"),
            )
        )
    created_from_datetime = _parse_created_date(created_from, is_end=False)
    if created_from_datetime is not None:
        filters.append(col(ReviewTask.created_at) >= created_from_datetime)
    created_to_datetime = _parse_created_date(created_to, is_end=True)
    if created_to_datetime is not None:
        filters.append(col(ReviewTask.created_at) < created_to_datetime)
    if mode is not None:
        filters.append(col(ReviewTask.custom_review_mode) == mode)
    if has_word_material is not None:
        filters.append(
            col(ReviewTask.has_custom_review_word_material) == has_word_material
        )
    if has_excel_material is not None:
        filters.append(
            col(ReviewTask.has_custom_review_excel_material) == has_excel_material
        )
    if has_pdf_material is not None:
        filters.append(
            col(ReviewTask.has_custom_review_pdf_material) == has_pdf_material
        )
    if status is not None:
        filters.append(col(ReviewTask.status) == status)
    return filters


def _custom_review_input_public(
    task: ReviewTask, user: User | None
) -> CustomReviewInputPublic:
    return CustomReviewInputPublic.model_validate(
        task,
        update={
            "user_name": user.full_name if user else None,
            "svn_username": user.svn_username if user else None,
        },
    )


def _custom_review_material_label(value: bool | None) -> str:
    return "已提供" if value is True else "未提供"


def _custom_review_export_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    aware_value = (
        value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)
    )
    return aware_value.astimezone(TASK_STATS_TIMEZONE).replace(tzinfo=None)


def _style_custom_review_input_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:M{max(sheet.max_row, 1)}"
    sheet.row_dimensions[1].height = 28
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(CUSTOM_REVIEW_INPUT_EXPORT_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    row_border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = row_border


def _issue_feedback_export_date(value: datetime | None) -> date | None:
    if value is None:
        return None
    aware_value = (
        value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)
    )
    return aware_value.astimezone(TASK_STATS_TIMEZONE).date()


def _issue_feedback_export_label(helpful: bool | None) -> str:
    if helpful is True:
        return "有帮助"
    if helpful is False:
        return "没帮助"
    return "未反馈"


def _issue_feedback_review_items_label(review_item_ids: list[str]) -> str:
    if not review_item_ids:
        return "未记录"
    return "；".join(
        f"{item_id} {REVIEW_ITEM_EXPORT_NAMES[item_id]}"
        if item_id in REVIEW_ITEM_EXPORT_NAMES
        else item_id
        for item_id in review_item_ids
    )


def _issue_feedback_excel_text(
    value: str,
    *,
    task_id: uuid.UUID,
    issue_id: uuid.UUID,
    field_name: str,
) -> str:
    cleaned, count = EXCEL_ILLEGAL_CHARACTERS_RE.subn("", value)
    if count:
        logger.warning(
            "清理问题反馈Excel非法字符 task_id=%s issue_id=%s field=%s count=%d",
            task_id,
            issue_id,
            field_name,
            count,
        )
    return cleaned


def _style_issue_feedback_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:J{max(sheet.max_row, 1)}"
    sheet.print_area = f"A1:J{max(sheet.max_row, 1)}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4
    sheet.row_dimensions[1].height = 28
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(ISSUE_FEEDBACK_EXPORT_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    row_border = Border(bottom=Side(style="thin", color="D9E2F3"))
    alternate_fill = PatternFill(fill_type="solid", fgColor="F7FAFC")
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = row_border
            if cell.row % 2 == 0:
                cell.fill = alternate_fill
        for column_index in (2, 4, 5, 9):
            row[column_index - 1].alignment = Alignment(
                horizontal="center",
                vertical="top",
                wrap_text=True,
            )
    for cell in sheet["B"][1:]:
        cell.number_format = "yyyy-mm-dd"


def _range_bounds(
    stats_range: ReviewTaskStatsRange,
    now: datetime | None = None,
) -> tuple[datetime | None, datetime | None]:
    if stats_range == ReviewTaskStatsRange.all:
        return None, None

    current = (now or _utcnow()).astimezone(TASK_STATS_TIMEZONE)
    today_start = datetime.combine(
        current.date(),
        time.min,
        tzinfo=TASK_STATS_TIMEZONE,
    )
    week_start = today_start - timedelta(days=current.weekday())
    if stats_range == ReviewTaskStatsRange.this_week:
        start = week_start
        end = current
    elif stats_range == ReviewTaskStatsRange.last_week:
        start = week_start - timedelta(days=7)
        end = week_start
    elif stats_range == ReviewTaskStatsRange.today:
        start = today_start
        end = current
    else:
        start = today_start - timedelta(days=1)
        end = today_start

    return (
        start.astimezone(timezone.utc).replace(tzinfo=None),
        end.astimezone(timezone.utc).replace(tzinfo=None),
    )


def _range_filters(stats_range: ReviewTaskStatsRange) -> list[Any]:
    start, end = _range_bounds(stats_range)
    filters: list[Any] = []
    if start is not None:
        filters.append(ReviewTask.created_at >= start)
    if end is not None:
        filters.append(ReviewTask.created_at < end)
    return filters


def _parse_started_from(value: str | None) -> datetime | None:
    if not value or not value.strip():
        return None

    normalized = value.strip()
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", normalized):
            local_start = datetime.combine(
                date.fromisoformat(normalized),
                time.min,
                tzinfo=TASK_STATS_TIMEZONE,
            )
            return local_start.astimezone(timezone.utc).replace(tzinfo=None)

        parsed = datetime.fromisoformat(normalized.replace("Z", "+00:00"))
    except ValueError:
        raise HTTPException(status_code=422, detail="Invalid started_from") from None

    if parsed.tzinfo is None:
        return parsed
    return parsed.astimezone(timezone.utc).replace(tzinfo=None)


def _unlink_file_after_commit(file_path: Path) -> None:
    if not file_path.is_file():
        return
    try:
        file_path.unlink()
    except OSError as exc:
        logger.warning(
            "Failed to remove review task file after database commit: %s",
            file_path,
            exc_info=exc,
        )


def _trend_points(
    *,
    session: SessionDep,
    filters: list[Any],
) -> list[ReviewTaskTrendPointPublic]:
    rows = session.exec(
        select(ReviewTask.created_at)
        .where(*filters)
        .order_by(col(ReviewTask.created_at))
    ).all()

    counts_by_date: dict[str, int] = {}
    for created_at in rows:
        if created_at is None:
            continue
        label = created_at.astimezone(TASK_STATS_TIMEZONE).date().isoformat()
        counts_by_date[label] = counts_by_date.get(label, 0) + 1

    return [
        ReviewTaskTrendPointPublic(label=label, count=count)
        for label, count in list(counts_by_date.items())[-30:]
    ]


def _normalize_revision(revision: str) -> str:
    match = _REVISION_RE.fullmatch(revision.strip())
    if not match:
        raise HTTPException(
            status_code=400,
            detail="Revision must be a number or start with r followed by a number",
        )
    return f"r{match.group(1)}"


def _can_write_task(current_user: CurrentUser, task: ReviewTask) -> bool:
    return has_admin_privileges(current_user) or task.user_id == current_user.id


def _get_task_for_read(
    *, session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> ReviewTask:
    task = session.get(ReviewTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="Review task not found")
    if has_admin_privileges(current_user):
        return task
    if current_user.role == UserRole.department_admin:
        _ensure_department_admin_active_department(
            session=session, current_user=current_user
        )
        if task.department_id == current_user.department_id:
            return task
        raise HTTPException(status_code=403, detail="Not enough permissions")
    if task.user_id == current_user.id:
        return task
    raise HTTPException(status_code=403, detail="Not enough permissions")


def _get_task_for_write(
    *, session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> ReviewTask:
    task = session.get(ReviewTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="Review task not found")
    if not _can_write_task(current_user, task):
        raise HTTPException(status_code=403, detail="Not enough permissions")
    return task


def _get_task_for_write_locked(
    *, session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> ReviewTask:
    task = session.exec(
        select(ReviewTask)
        .where(col(ReviewTask.id) == task_id)
        .with_for_update()
        .execution_options(populate_existing=True)
    ).one_or_none()
    if not task:
        raise HTTPException(status_code=404, detail="Review task not found")
    if not _can_write_task(current_user, task):
        raise HTTPException(status_code=403, detail="Not enough permissions")
    return task


def _get_task_for_issue_feedback_write_locked(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    invalid_source_status_code: int = 409,
    invalid_source_detail: str = "issue_feedback_not_supported",
) -> ReviewTask:
    task = session.exec(
        select(ReviewTask)
        .where(col(ReviewTask.id) == task_id)
        .with_for_update()
        .execution_options(populate_existing=True)
    ).one_or_none()
    if not task:
        raise HTTPException(status_code=404, detail="Review task not found")
    if task.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="只有任务提交人可以填写问题反馈")
    if task.feedback_protocol_version != 1:
        raise HTTPException(status_code=409, detail="issue_feedback_not_supported")
    source_consistency_state = _review_issue_source_consistency_states_by_task_ids(
        session=session,
        task_ids=[task.id],
    )[task.id]
    if not source_consistency_state.is_valid:
        raise HTTPException(
            status_code=invalid_source_status_code,
            detail=invalid_source_detail,
        )
    if task.feedback_submitted_at is not None:
        raise HTTPException(status_code=409, detail="issue_feedback_submitted")
    return task


def _read_report_path(task: ReviewTask) -> Path:
    report_path = Path(task.report_path) if task.report_path else None
    if not report_path or not report_path.is_file():
        raise HTTPException(status_code=404, detail="Report not found")
    return report_path


def _read_pdf_report_path(task: ReviewTask) -> Path:
    report_path = Path(task.pdf_report_path) if task.pdf_report_path else None
    if not report_path or not report_path.is_file():
        raise HTTPException(status_code=404, detail="PDF report not found")
    return report_path


def _ensure_task_has_reports(task: ReviewTask) -> None:
    if not task.report_path:
        raise HTTPException(status_code=409, detail="no_report")
    if not task.pdf_report_path:
        raise HTTPException(status_code=409, detail="no_pdf_report")


def _normalize_useful_issues_text(
    *, helpful: bool, useful_issues_text: str | None
) -> str | None:
    normalized = useful_issues_text.strip() if useful_issues_text else ""
    if helpful and not normalized:
        raise HTTPException(status_code=400, detail="有帮助时请至少填写一条有用信息")
    return normalized or None


def _task_public(
    task: ReviewTask,
    user: User | None = None,
    report_feedback: ReviewTaskReportFeedback | None = None,
    issue_count: int = 0,
    evaluated_issue_count: int = 0,
    has_issue_feedback_draft: bool = False,
    highest_issue_severity: ReviewIssueSeverity | None = None,
) -> ReviewTaskPublic:
    return ReviewTaskPublic.model_validate(
        task,
        update={
            "user_name": user.full_name if user else None,
            "svn_username": user.svn_username if user else None,
            "report_feedback_helpful": (
                report_feedback.helpful if report_feedback else None
            ),
            "report_feedback_useful_issues_text": (
                report_feedback.useful_issues_text if report_feedback else None
            ),
            "issue_count": issue_count,
            "evaluated_issue_count": evaluated_issue_count,
            "has_issue_feedback_draft": has_issue_feedback_draft,
            "highest_issue_severity": highest_issue_severity,
        },
    )


def _issue_progress_by_task_ids(
    *, session: SessionDep, task_ids: list[uuid.UUID]
) -> dict[uuid.UUID, tuple[int, int, bool, ReviewIssueSeverity | None]]:
    if not task_ids:
        return {}
    source_consistency_states = _review_issue_source_consistency_states_by_task_ids(
        session=session,
        task_ids=task_ids,
    )
    valid_task_ids = [
        task_id
        for task_id, state in source_consistency_states.items()
        if state.is_valid
    ]
    if not valid_task_ids:
        return {}

    issue_text = col(ReviewTaskIssue.issue_text)
    severity_rank = case(
        *(
            (
                _review_issue_severity_sql_condition(issue_text, rule),
                rule.rank,
            )
            for rule in _REVIEW_ISSUE_SEVERITY_RULES
        ),
        else_=0,
    )
    rows = session.exec(
        # SQLModel stubs only cover up to four selected columns.
        select(  # type: ignore[call-overload]
            ReviewTaskIssue.task_id,
            func.count(col(ReviewTaskIssue.id)),
            func.count(col(ReviewTaskIssue.helpful)),
            func.count(col(ReviewTaskIssue.feedback_text)),
            func.max(severity_rank),
        )
        .where(col(ReviewTaskIssue.task_id).in_(valid_task_ids))
        .group_by(col(ReviewTaskIssue.task_id))
    ).all()
    return {
        task_id: (
            int(issue_count),
            int(evaluated_issue_count),
            int(evaluated_issue_count) > 0 or int(feedback_text_count) > 0,
            _review_issue_severity_from_rank(int(highest_severity_rank)),
        )
        for task_id, issue_count, evaluated_issue_count, feedback_text_count, highest_severity_rank in rows
    }


def _review_issue_severity_from_rank(rank: int) -> ReviewIssueSeverity | None:
    return next(
        (rule.severity for rule in _REVIEW_ISSUE_SEVERITY_RULES if rule.rank == rank),
        None,
    )


def _has_consistent_review_issue_sources(
    sources: Sequence[str],
) -> bool:
    # Electron 负责解析 Markdown；服务端仅防止 BMS 与旧模板快照混存。
    bms_issue_count = sum(source == "bms_code_reviewer_heading" for source in sources)
    return bms_issue_count == 0 or bms_issue_count == len(sources)


def _review_issue_source_consistency_states_by_task_ids(
    *,
    session: SessionDep,
    task_ids: Sequence[uuid.UUID],
) -> dict[uuid.UUID, _ReviewIssueSourceConsistencyState]:
    unique_task_ids = list(dict.fromkeys(task_ids))
    if not unique_task_ids:
        return {}
    sources_by_task_id: dict[uuid.UUID, list[str]] = {
        task_id: [] for task_id in unique_task_ids
    }
    rows = session.exec(
        select(
            ReviewTaskIssue.task_id,
            ReviewTaskIssue.source,
            func.count(col(ReviewTaskIssue.id)),
        )
        .where(col(ReviewTaskIssue.task_id).in_(unique_task_ids))
        .group_by(col(ReviewTaskIssue.task_id), col(ReviewTaskIssue.source))
    ).all()
    for task_id, source, _issue_count in rows:
        sources_by_task_id[task_id].append(source)
    return {
        task_id: _ReviewIssueSourceConsistencyState(
            status=(
                _ReviewIssueSourceConsistencyStatus.valid
                if _has_consistent_review_issue_sources(sources)
                else _ReviewIssueSourceConsistencyStatus.invalid
            )
        )
        for task_id, sources in sources_by_task_id.items()
    }


def _read_effective_review_task_issues(
    *,
    session: SessionDep,
    task_id: uuid.UUID,
) -> _EffectiveReviewTaskIssues:
    source_consistency_state = _review_issue_source_consistency_states_by_task_ids(
        session=session,
        task_ids=[task_id],
    )[task_id]
    if not source_consistency_state.is_valid:
        return _EffectiveReviewTaskIssues(
            source_consistency_state=source_consistency_state,
            issues=(),
        )
    issues = tuple(
        session.exec(
            select(ReviewTaskIssue)
            .where(col(ReviewTaskIssue.task_id) == task_id)
            .order_by(col(ReviewTaskIssue.order))
        ).all()
    )
    return _EffectiveReviewTaskIssues(
        source_consistency_state=source_consistency_state,
        issues=issues,
    )


def _highest_issue_severity(
    issues: list[ReviewTaskIssueSnapshotCreate],
) -> ReviewIssueSeverity | None:
    if not _has_consistent_review_issue_sources(
        tuple(issue.source for issue in issues)
    ):
        return None
    for rule in _REVIEW_ISSUE_SEVERITY_RULES:
        if any(
            _review_issue_text_matches_severity(issue.issue_text, rule)
            for issue in issues
        ):
            return rule.severity
    return None


def _parse_issue_snapshots(raw_issues: str) -> list[ReviewTaskIssueSnapshotCreate]:
    if len(raw_issues.encode("utf-8")) > MAX_REPORT_ISSUES_JSON_BYTES:
        raise HTTPException(status_code=400, detail="问题快照总大小超过限制")
    try:
        parsed = json.loads(raw_issues)
        if not isinstance(parsed, list):
            raise ValueError
        if len(parsed) > MAX_REPORT_ISSUE_COUNT:
            raise HTTPException(status_code=400, detail="问题快照数量超过限制")
        issues = [ReviewTaskIssueSnapshotCreate.model_validate(item) for item in parsed]
    except HTTPException:
        raise
    except (json.JSONDecodeError, ValidationError, TypeError, ValueError):
        raise HTTPException(status_code=400, detail="问题快照格式无效") from None

    if [issue.order for issue in issues] != list(range(len(issues))):
        raise HTTPException(status_code=400, detail="问题快照顺序无效")
    return issues


def _validate_issue_snapshot_locations(
    issues: list[ReviewTaskIssueSnapshotCreate],
    report_markdown: str,
) -> None:
    target_offsets = {
        offset
        for issue in issues
        for offset in (
            issue.location.start_offset,
            issue.location.feedback_insertion_offset,
            issue.location.end_offset,
        )
    }
    start_offsets = {issue.location.start_offset for issue in issues}
    valid_offsets: set[int] = set()
    anchor_lines_by_start_offset: dict[int, int] = {}
    utf16_offset = 0
    line_index = 0
    is_line_start = True

    # Electron 的字符串偏移量按 UTF-16 code unit 计算。只记录请求涉及的边界，
    # 避免为最大 5 MiB 报告构建完整偏移表。
    for character in report_markdown:
        if utf16_offset in target_offsets:
            valid_offsets.add(utf16_offset)
        if utf16_offset in start_offsets and is_line_start:
            anchor_lines_by_start_offset[utf16_offset] = line_index
        utf16_offset += 2 if ord(character) > 0xFFFF else 1
        if character == "\n":
            line_index += 1
            is_line_start = True
        else:
            is_line_start = False

    if utf16_offset in target_offsets:
        valid_offsets.add(utf16_offset)
    if utf16_offset in start_offsets and is_line_start:
        anchor_lines_by_start_offset[utf16_offset] = line_index

    for issue in issues:
        location = issue.location
        if (
            not (
                location.start_offset
                <= location.feedback_insertion_offset
                <= location.end_offset
            )
            or location.start_offset not in valid_offsets
            or location.feedback_insertion_offset not in valid_offsets
            or location.end_offset not in valid_offsets
            or anchor_lines_by_start_offset.get(location.start_offset)
            != location.anchor_line
        ):
            raise HTTPException(status_code=400, detail="问题快照定位范围无效")


def _parse_asset_paths(raw_asset_paths: str | None) -> list[str]:
    if raw_asset_paths is None:
        return []
    try:
        parsed = json.loads(raw_asset_paths)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="图片资产路径格式无效") from None
    if not isinstance(parsed, list) or not all(
        isinstance(path, str) for path in parsed
    ):
        raise HTTPException(status_code=400, detail="图片资产路径格式无效")
    return parsed


@router.post("", response_model=ReviewTaskPublic)
def create_review_task(
    *, session: SessionDep, current_user: CurrentUser, task_in: ReviewTaskCreate
) -> Any:
    normalized_review_item_ids = [
        item_id.strip() for item_id in task_in.review_item_ids
    ]
    if any(not item_id or len(item_id) > 64 for item_id in normalized_review_item_ids):
        raise HTTPException(status_code=400, detail="审查项 ID 无效")
    if len(set(normalized_review_item_ids)) != len(normalized_review_item_ids):
        raise HTTPException(status_code=400, detail="审查项 ID 不能重复")

    has_custom_review = "DS-99" in normalized_review_item_ids
    custom_review_fields = (
        task_in.custom_review_rule,
        task_in.custom_review_mode,
        task_in.has_custom_review_word_material,
        task_in.has_custom_review_excel_material,
        task_in.has_custom_review_pdf_material,
    )
    normalized_custom_review_rule: str | None = None
    if has_custom_review:
        normalized_custom_review_rule = (task_in.custom_review_rule or "").strip()
        if not normalized_custom_review_rule:
            raise HTTPException(
                status_code=400, detail="DS-99 自定义审查必须填写规则正文"
            )
        if len(normalized_custom_review_rule) > 1200:
            raise HTTPException(
                status_code=400, detail="DS-99 自定义审查规则正文不能超过 1200 个字符"
            )
        if task_in.custom_review_mode not in {
            CustomReviewMode.dual_revision,
            CustomReviewMode.new_revision_full,
        }:
            raise HTTPException(status_code=400, detail="DS-99 自定义审查模式无效")
        if any(value is None for value in custom_review_fields[2:]):
            raise HTTPException(status_code=400, detail="DS-99 自定义审查材料状态缺失")
    elif any(value is not None for value in custom_review_fields):
        raise HTTPException(
            status_code=400, detail="未选择 DS-99 时不能提交自定义审查输入"
        )

    normalized_old_revision = _normalize_revision(task_in.old_revision)
    new_revision = _normalize_revision(task_in.new_revision)
    old_revision_number = int(normalized_old_revision[1:])
    new_revision_number = int(new_revision[1:])
    if new_revision_number == 0:
        raise HTTPException(status_code=400, detail="新版本号不能为 0")
    if old_revision_number != 0 and old_revision_number >= new_revision_number:
        raise HTTPException(status_code=400, detail="旧版本号必须小于新版本号")

    department = (
        session.get(Department, current_user.department_id)
        if current_user.department_id is not None
        else None
    )
    task = ReviewTask.model_validate(
        task_in,
        update={
            "user_id": current_user.id,
            "old_revision": (
                task_in.old_revision.strip()
                if normalized_review_item_ids
                else normalized_old_revision
            ),
            "new_revision": new_revision,
            "review_item_ids": normalized_review_item_ids,
            "custom_review_rule": normalized_custom_review_rule,
            "custom_review_mode": task_in.custom_review_mode
            if has_custom_review
            else None,
            "has_custom_review_word_material": (
                task_in.has_custom_review_word_material if has_custom_review else None
            ),
            "has_custom_review_excel_material": (
                task_in.has_custom_review_excel_material if has_custom_review else None
            ),
            "has_custom_review_pdf_material": (
                task_in.has_custom_review_pdf_material if has_custom_review else None
            ),
            "department_id": department.id if department else None,
            "department_name": department.name if department else None,
            "focus_direction_snapshot": task_in.focus_direction_snapshot,
        },
    )
    session.add(task)
    session.commit()
    session.refresh(task)
    return task


@router.get("", response_model=ReviewTasksPublic)
def read_review_tasks(
    session: SessionDep,
    current_user: CurrentUser,
    skip: int = 0,
    limit: int = 100,
    mine: bool = False,
    user_name: str | None = None,
    svn_username: str | None = None,
    status: ReviewTaskStatus | None = None,
    report_feedback: ReportFeedbackFilter | None = None,
    started_from: str | None = None,
) -> Any:
    filters = _task_read_scope_filters(session=session, current_user=current_user)
    if mine:
        filters.append(ReviewTask.user_id == current_user.id)
    if user_name and user_name.strip():
        filters.append(col(User.full_name).ilike(f"%{user_name.strip()}%"))
    if svn_username and svn_username.strip():
        filters.append(col(User.svn_username).ilike(f"%{svn_username.strip()}%"))
    if status is not None:
        filters.append(ReviewTask.status == status)
    started_from_datetime = _parse_started_from(started_from)
    if started_from_datetime is not None:
        filters.append(ReviewTask.started_at >= started_from_datetime)
    if report_feedback == "helpful":
        filters.append(ReviewTaskReportFeedback.helpful.is_(True))
    elif report_feedback == "not_helpful":
        filters.append(ReviewTaskReportFeedback.helpful.is_(False))
    elif report_feedback == "none":
        filters.append(ReviewTaskReportFeedback.id.is_(None))

    count_statement = (
        select(func.count())
        .select_from(ReviewTask)
        .join(User, ReviewTask.user_id == User.id)
        .outerjoin(
            ReviewTaskReportFeedback,
            ReviewTaskReportFeedback.task_id == ReviewTask.id,
        )
        .where(*filters)
    )
    statement = (
        select(ReviewTask, User, ReviewTaskReportFeedback)
        .join(User, ReviewTask.user_id == User.id)
        .outerjoin(
            ReviewTaskReportFeedback,
            ReviewTaskReportFeedback.task_id == ReviewTask.id,
        )
        .where(*filters)
    )
    count = session.exec(count_statement).one()
    rows = session.exec(
        statement.order_by(col(ReviewTask.created_at).desc()).offset(skip).limit(limit)
    ).all()
    issue_progress = _issue_progress_by_task_ids(
        session=session,
        task_ids=[task.id for task, _, _ in rows],
    )
    return ReviewTasksPublic(
        data=[
            _task_public(
                task,
                user,
                report_feedback,
                *issue_progress.get(task.id, (0, 0, False, None)),
            )
            for task, user, report_feedback in rows
        ],
        count=count,
    )


@router.get("/custom-review-inputs", response_model=CustomReviewInputsPublic)
def read_custom_review_inputs(
    session: SessionDep,
    current_user: CurrentUser,
    skip: int = 0,
    limit: int = 100,
    keyword: str | None = None,
    created_from: str | None = None,
    created_to: str | None = None,
    user_name: str | None = None,
    department_name: str | None = None,
    svn_url: str | None = None,
    revision: str | None = None,
    mode: CustomReviewMode | None = None,
    has_word_material: bool | None = None,
    has_excel_material: bool | None = None,
    has_pdf_material: bool | None = None,
    status: ReviewTaskStatus | None = None,
) -> CustomReviewInputsPublic:
    filters = _custom_review_input_filters(
        session=session,
        current_user=current_user,
        keyword=keyword,
        created_from=created_from,
        created_to=created_to,
        user_name=user_name,
        department_name=department_name,
        svn_url=svn_url,
        revision=revision,
        mode=mode,
        has_word_material=has_word_material,
        has_excel_material=has_excel_material,
        has_pdf_material=has_pdf_material,
        status=status,
    )
    count = session.exec(
        select(func.count())
        .select_from(ReviewTask)
        .outerjoin(User, col(ReviewTask.user_id) == col(User.id))
        .where(*filters)
    ).one()
    rows = session.exec(
        select(ReviewTask, User)
        .outerjoin(User, col(ReviewTask.user_id) == col(User.id))
        .where(*filters)
        .order_by(col(ReviewTask.created_at).desc(), col(ReviewTask.id).desc())
        .offset(skip)
        .limit(limit)
    ).all()
    return CustomReviewInputsPublic(
        data=[_custom_review_input_public(task, user) for task, user in rows],
        count=count,
    )


@router.get(
    "/custom-review-inputs/export",
    response_class=StreamingResponse,
    responses={
        200: {
            "content": {EXCEL_MEDIA_TYPE: {}},
            "description": "自定义审查输入统计 Excel 文件",
        }
    },
)
def export_custom_review_inputs(
    session: SessionDep,
    current_user: CurrentUser,
    keyword: str | None = None,
    created_from: str | None = None,
    created_to: str | None = None,
    user_name: str | None = None,
    department_name: str | None = None,
    svn_url: str | None = None,
    revision: str | None = None,
    mode: CustomReviewMode | None = None,
    has_word_material: bool | None = None,
    has_excel_material: bool | None = None,
    has_pdf_material: bool | None = None,
    status: ReviewTaskStatus | None = None,
) -> StreamingResponse:
    filters = _custom_review_input_filters(
        session=session,
        current_user=current_user,
        keyword=keyword,
        created_from=created_from,
        created_to=created_to,
        user_name=user_name,
        department_name=department_name,
        svn_url=svn_url,
        revision=revision,
        mode=mode,
        has_word_material=has_word_material,
        has_excel_material=has_excel_material,
        has_pdf_material=has_pdf_material,
        status=status,
    )
    rows = session.exec(
        select(ReviewTask, User)
        .outerjoin(User, col(ReviewTask.user_id) == col(User.id))
        .where(*filters)
        .order_by(col(ReviewTask.created_at).desc(), col(ReviewTask.id).desc())
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "自定义审查输入统计"
    sheet.append(list(CUSTOM_REVIEW_INPUT_EXPORT_HEADERS))
    for task, user in rows:
        sheet.append(
            [
                _custom_review_export_datetime(task.created_at),
                user.full_name if user else None,
                user.svn_username if user else None,
                task.department_name,
                task.svn_url,
                task.old_revision,
                task.new_revision,
                task.custom_review_mode.value
                if task.custom_review_mode is not None
                else None,
                _custom_review_material_label(task.has_custom_review_word_material),
                _custom_review_material_label(task.has_custom_review_excel_material),
                _custom_review_material_label(task.has_custom_review_pdf_material),
                task.status.value,
                task.custom_review_rule,
            ]
        )
        # 自定义输入可能以公式字符开头，导出时必须固定为普通字符串。
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    _style_custom_review_input_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    file_name = f"自定义审查输入统计_{_utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type=EXCEL_MEDIA_TYPE,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@router.get("/custom-review-inputs/{task_id}", response_model=CustomReviewInputPublic)
def read_custom_review_input(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> CustomReviewInputPublic:
    filters = _custom_review_input_filters(
        session=session,
        current_user=current_user,
    )
    row = session.exec(
        select(ReviewTask, User)
        .outerjoin(User, col(ReviewTask.user_id) == col(User.id))
        .where(col(ReviewTask.id) == task_id, *filters)
    ).one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="自定义审查输入不存在")
    task, user = row
    return _custom_review_input_public(task, user)


@router.get(
    "/issue-feedback/export",
    response_class=StreamingResponse,
    responses={
        200: {
            "content": {EXCEL_MEDIA_TYPE: {}},
            "description": "逐问题反馈 Excel 文件",
        }
    },
)
def export_review_task_issue_feedback(
    session: SessionDep,
    current_user: CurrentUser,
    started_from: str | None = None,
) -> StreamingResponse:
    task_filters = [
        *_task_read_scope_filters(session=session, current_user=current_user),
        ReviewTask.feedback_protocol_version == 1,
    ]
    started_from_datetime = _parse_started_from(started_from)
    if started_from_datetime is not None:
        task_filters.append(col(ReviewTask.started_at) >= started_from_datetime)
    task_ids = list(session.exec(select(ReviewTask.id).where(*task_filters)).all())
    source_consistency_states = _review_issue_source_consistency_states_by_task_ids(
        session=session,
        task_ids=task_ids,
    )
    valid_task_ids = [
        task_id
        for task_id, state in source_consistency_states.items()
        if state.is_valid
    ]
    rows = (
        session.exec(
            select(ReviewTask, User, ReviewTaskIssue)
            .join(User, ReviewTask.user_id == User.id)
            .join(ReviewTaskIssue, ReviewTaskIssue.task_id == ReviewTask.id)
            .where(
                *task_filters,
                col(ReviewTask.id).in_(valid_task_ids),
            )
            .order_by(
                col(ReviewTask.created_at).desc(),
                col(ReviewTask.id),
                col(ReviewTaskIssue.order),
            )
        ).all()
        if valid_task_ids
        else []
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "问题反馈明细"
    sheet.append(list(ISSUE_FEEDBACK_EXPORT_HEADERS))
    for task, user, issue in rows:
        row_values: list[tuple[Any, str]] = [
            (user.full_name or user.svn_username or "-", "submitter"),
            (_issue_feedback_export_date(task.created_at), "created_at"),
            (task.svn_url, "svn_url"),
            (task.old_revision, "old_revision"),
            (task.new_revision, "new_revision"),
            (
                _issue_feedback_review_items_label(task.review_item_ids),
                "review_item_ids",
            ),
            (issue.issue_text, "issue_text"),
            (issue.description_text, "description_text"),
            (_issue_feedback_export_label(issue.helpful), "helpful"),
            (issue.feedback_text or None, "feedback_text"),
        ]
        sheet.append(
            [
                _issue_feedback_excel_text(
                    value,
                    task_id=task.id,
                    issue_id=issue.id,
                    field_name=field_name,
                )
                if isinstance(value, str)
                else value
                for value, field_name in row_values
            ]
        )
        # 防止问题原文或用户反馈被 Excel 解释为公式。
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    _style_issue_feedback_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    file_name = f"问题反馈明细_{_utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type=EXCEL_MEDIA_TYPE,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(file_name)}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@router.get("/stats", response_model=ReviewTaskStatsPublic)
def read_review_task_stats(
    session: SessionDep,
    current_user: CurrentUser,
    range: ReviewTaskStatsRange = ReviewTaskStatsRange.today,
) -> Any:
    filters = [
        *_task_read_scope_filters(session=session, current_user=current_user),
        *_range_filters(range),
    ]

    statement = (
        select(ReviewTask.status, func.count())
        .where(*filters)
        .group_by(ReviewTask.status)
    )
    rows = session.exec(statement).all()
    by_status = dict.fromkeys(ReviewTaskStatus, 0)
    for status, count in rows:
        by_status[status] = count

    duration_rows = session.exec(
        select(ReviewTask.started_at, ReviewTask.finished_at).where(
            *filters,
            ReviewTask.started_at.is_not(None),
            ReviewTask.finished_at.is_not(None),
        )
    ).all()
    durations = [
        int((finished_at - started_at).total_seconds())
        for started_at, finished_at in duration_rows
        if (
            started_at is not None
            and finished_at is not None
            and finished_at >= started_at
        )
    ]
    average_duration_seconds = (
        int(sum(durations) / len(durations)) if durations else None
    )

    reports_by_user: list[ReviewTaskReportUserStatsPublic] = []
    if (
        has_admin_privileges(current_user)
        or current_user.role == UserRole.department_admin
    ):
        report_count = func.count(ReviewTask.id)
        report_rows = session.exec(
            select(User.full_name, User.svn_username, report_count)
            .join(User, ReviewTask.user_id == User.id)
            .where(
                *filters,
                ReviewTask.status == ReviewTaskStatus.success,
                ReviewTask.report_path.is_not(None),
            )
            .group_by(User.id, User.full_name, User.svn_username)
            .order_by(report_count.desc(), User.full_name)
        ).all()
        reports_by_user = [
            ReviewTaskReportUserStatsPublic(
                user_name=user_name,
                svn_username=svn_username,
                report_count=count,
            )
            for user_name, svn_username, count in report_rows
        ]

    return ReviewTaskStatsPublic(
        range=range,
        total=sum(by_status.values()),
        success=by_status[ReviewTaskStatus.success],
        running=sum(by_status[status] for status in _RUNNING_STATUSES),
        failed_or_error=sum(by_status[status] for status in _FAILED_OR_ERROR_STATUSES),
        average_duration_seconds=average_duration_seconds,
        by_status=by_status,
        reports_by_user=reports_by_user,
        trend=_trend_points(session=session, filters=filters),
    )


@router.get("/weekly-insights", response_model=ReviewTaskWeeklyInsightsPublic)
def read_review_task_weekly_insights(
    session: SessionDep,
    current_user: CurrentUser,
) -> ReviewTaskWeeklyInsightsPublic:
    generated_at = _utcnow()
    period_start, period_end_exclusive, start_utc, end_utc = (
        previous_complete_week_bounds(generated_at)
    )
    filters = [
        *_task_read_scope_filters(session=session, current_user=current_user),
        ReviewTask.status == ReviewTaskStatus.success,
        ReviewTask.created_at >= start_utc,
        ReviewTask.created_at < end_utc,
    ]
    rows = session.exec(
        select(ReviewTask, ReviewTaskReportFeedback)
        .outerjoin(
            ReviewTaskReportFeedback,
            ReviewTaskReportFeedback.task_id == ReviewTask.id,
        )
        .where(*filters)
        .order_by(col(ReviewTask.created_at))
    ).all()

    if has_admin_privileges(current_user):
        scope_label = "全部部门"
    elif current_user.role == UserRole.department_admin:
        department = session.get(Department, current_user.department_id)
        scope_label = department.name if department else "本部门"
    else:
        scope_label = "个人"

    return build_weekly_insights(
        rows=list(rows),
        period_start=period_start,
        period_end_exclusive=period_end_exclusive,
        generated_at=generated_at,
        scope_label=scope_label,
    )


@router.post(
    "/{task_id}/report-feedback",
    response_model=ReviewTaskReportFeedbackPublic,
)
def submit_review_task_report_feedback(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    feedback_in: ReviewTaskReportFeedbackCreate,
) -> Any:
    task = _get_task_for_read(
        session=session, current_user=current_user, task_id=task_id
    )
    _ensure_task_has_reports(task)
    useful_issues_text = _normalize_useful_issues_text(
        helpful=feedback_in.helpful,
        useful_issues_text=feedback_in.useful_issues_text,
    )
    existing_feedback = session.exec(
        select(ReviewTaskReportFeedback).where(
            ReviewTaskReportFeedback.task_id == task.id
        )
    ).first()
    if existing_feedback:
        raise HTTPException(status_code=409, detail="report_feedback_exists")

    feedback = ReviewTaskReportFeedback(
        task_id=task.id,
        user_id=current_user.id,
        helpful=feedback_in.helpful,
        useful_issues_text=useful_issues_text,
        created_at=_utcnow(),
        updated_at=_utcnow(),
    )
    session.add(feedback)
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(status_code=409, detail="report_feedback_exists") from None
    session.refresh(feedback)
    return feedback


@router.get("/{task_id}", response_model=ReviewTaskPublic)
def read_review_task(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> Any:
    task = _get_task_for_read(
        session=session, current_user=current_user, task_id=task_id
    )
    report_feedback = session.exec(
        select(ReviewTaskReportFeedback).where(
            ReviewTaskReportFeedback.task_id == task.id
        )
    ).first()
    (
        issue_count,
        evaluated_issue_count,
        has_issue_feedback_draft,
        highest_issue_severity,
    ) = _issue_progress_by_task_ids(session=session, task_ids=[task.id]).get(
        task.id, (0, 0, False, None)
    )
    return _task_public(
        task,
        session.get(User, task.user_id),
        report_feedback,
        issue_count,
        evaluated_issue_count,
        has_issue_feedback_draft,
        highest_issue_severity,
    )


@router.delete("/{task_id}")
def delete_review_task(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> dict[str, str]:
    task = _get_task_for_write_locked(
        session=session, current_user=current_user, task_id=task_id
    )
    for pending_path in report_service.retry_pending_task_file_deletions(
        lambda pending_task_id: session.get(ReviewTask, pending_task_id) is not None
    ):
        logger.error("待重试清理目录仍未删除: %s", pending_path)

    requirement_files = session.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == task.id
        )
    ).all()
    report_path = Path(task.report_path) if task.report_path else None
    pdf_report_path = Path(task.pdf_report_path) if task.pdf_report_path else None
    files_to_remove = [
        file_path
        for file_path in (
            report_path,
            pdf_report_path,
            *(
                Path(requirement_file.file_path)
                for requirement_file in requirement_files
            ),
        )
        if file_path is not None
    ]

    quarantine = report_service.quarantine_task_files(task.id, files_to_remove)
    try:
        session.exec(
            delete(ReviewTaskReportFeedback).where(
                col(ReviewTaskReportFeedback.task_id) == task.id
            )
        )
        session.exec(
            delete(ReviewTaskIssue).where(col(ReviewTaskIssue.task_id) == task.id)
        )
        for requirement_file in requirement_files:
            session.delete(requirement_file)
        session.flush()
        session.delete(task)
        session.commit()
    except Exception:
        session.rollback()
        try:
            report_service.restore_quarantined_task_files(quarantine)
        except Exception as restore_error:
            logger.exception("删除任务回滚后恢复已隔离文件失败")
            raise RuntimeError("删除任务失败，且文件恢复失败") from restore_error
        raise

    try:
        report_service.purge_quarantined_task_files(quarantine)
    except OSError:
        logger.exception("任务已删除，待重试清理目录: %s", quarantine.root)
    return {"message": "Review task deleted"}


@router.patch("/{task_id}/status", response_model=ReviewTaskPublic)
def update_review_task_status(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    status_in: ReviewTaskStatusUpdate,
) -> Any:
    task = _get_task_for_write(
        session=session, current_user=current_user, task_id=task_id
    )
    now = _utcnow()
    update_data = status_in.model_dump(exclude_unset=True)
    task.sqlmodel_update(update_data)
    if task.started_at is None and task.status != ReviewTaskStatus.created:
        task.started_at = now
    if task.status in {
        ReviewTaskStatus.success,
        ReviewTaskStatus.failed,
        ReviewTaskStatus.environment_error,
        ReviewTaskStatus.cancelled,
    }:
        task.finished_at = now
    elif task.status in _RUNNING_STATUSES:
        task.finished_at = None
    task.updated_at = now

    session.add(task)
    session.commit()
    session.refresh(task)
    return task


@router.post("/{task_id}/report", response_model=ReviewTaskPublic)
async def upload_review_task_report(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    upload: UploadFile = File(...),
    feedback_protocol_version: int | None = Form(default=None),
    issues: str | None = Form(default=None),
    asset_paths: str | None = Form(default=None),
    assets: list[UploadFile] | None = File(default=None),
) -> Any:
    task = _get_task_for_write_locked(
        session=session, current_user=current_user, task_id=task_id
    )
    if (feedback_protocol_version is None) != (issues is None):
        raise HTTPException(status_code=400, detail="新协议版本和问题快照必须同时提交")
    issue_snapshots = _parse_issue_snapshots(issues) if issues is not None else []
    if feedback_protocol_version is not None and feedback_protocol_version != 1:
        raise HTTPException(status_code=400, detail="不支持的问题反馈协议版本")
    has_consistent_issue_sources = _has_consistent_review_issue_sources(
        tuple(issue.source for issue in issue_snapshots)
    )
    if task.feedback_protocol_version is not None or (
        feedback_protocol_version is not None and task.report_path is not None
    ):
        raise HTTPException(status_code=409, detail="原始报告已上传，不能覆盖")
    parsed_asset_paths = _parse_asset_paths(asset_paths)
    uploaded_assets = assets or []
    if feedback_protocol_version is None and (parsed_asset_paths or uploaded_assets):
        raise HTTPException(status_code=400, detail="历史报告不能提交图片资产")

    prepared_report = None
    if feedback_protocol_version is not None:
        prepared_report = await report_service.prepare_report_upload(
            upload,
            parsed_asset_paths,
            uploaded_assets,
        )
        _validate_issue_snapshot_locations(
            issue_snapshots,
            prepared_report.report_content.decode("utf-8"),
        )
        report_path = report_service.path_for_task(task_id)
    else:
        report_path = await report_service.save_upload(task_id, upload)
    now = _utcnow()
    task.report_path = str(report_path)
    task.report_name = upload.filename or REVIEW_REPORT_FILE_NAME
    task.feedback_protocol_version = feedback_protocol_version
    task.status = ReviewTaskStatus.success
    if task.started_at is None:
        task.started_at = now
    task.finished_at = now
    task.updated_at = now

    if feedback_protocol_version is None:
        session.exec(
            delete(ReviewTaskIssue).where(col(ReviewTaskIssue.task_id) == task.id)
        )
    for issue_snapshot in issue_snapshots:
        session.add(
            ReviewTaskIssue(
                task_id=task.id,
                issue_text=issue_snapshot.issue_text,
                description_text=issue_snapshot.description_text,
                display_snapshot=issue_snapshot.display_snapshot,
                order=issue_snapshot.order,
                source=issue_snapshot.source,
                location=issue_snapshot.location.model_dump(),
            )
        )
    session.add(task)
    report_installed = False
    try:
        session.flush()
        if prepared_report is not None:
            report_service.remove_report_bundle(task_id)
            report_service.install_prepared_report(task_id, prepared_report)
            report_installed = True
        session.commit()
    except Exception:
        session.rollback()
        if report_installed:
            report_service.remove_report_bundle(task_id)
        raise
    session.refresh(task)
    return _task_public(
        task,
        session.get(User, task.user_id),
        issue_count=len(issue_snapshots) if has_consistent_issue_sources else 0,
        highest_issue_severity=_highest_issue_severity(issue_snapshots),
    )


@router.get("/{task_id}/issues", response_model=ReviewTaskIssuesPublic)
def read_review_task_issues(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> ReviewTaskIssuesPublic:
    task = _get_task_for_read(
        session=session, current_user=current_user, task_id=task_id
    )
    effective_issues = _read_effective_review_task_issues(
        session=session,
        task_id=task.id,
    )
    issues = list(effective_issues.issues)
    return ReviewTaskIssuesPublic(
        data=issues,
        count=len(issues),
        evaluated_count=sum(issue.helpful is not None for issue in issues),
        feedback_submitted_at=task.feedback_submitted_at,
    )


@router.put(
    "/{task_id}/issues/{issue_id}/feedback-draft",
    response_model=ReviewTaskIssuePublic,
)
def update_review_task_issue_feedback_draft(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    issue_id: uuid.UUID,
    payload: ReviewTaskIssueFeedbackDraftUpdate,
) -> ReviewTaskIssue:
    task = _get_task_for_issue_feedback_write_locked(
        session=session,
        current_user=current_user,
        task_id=task_id,
        invalid_source_status_code=404,
        invalid_source_detail="Review task issue not found",
    )
    issue = session.exec(
        select(ReviewTaskIssue)
        .where(
            col(ReviewTaskIssue.id) == issue_id,
            col(ReviewTaskIssue.task_id) == task.id,
        )
        .with_for_update()
        .execution_options(populate_existing=True)
    ).one_or_none()
    if issue is None:
        raise HTTPException(status_code=404, detail="Review task issue not found")

    feedback_text = payload.feedback_text
    if feedback_text is not None and not feedback_text.strip():
        feedback_text = None
    if issue.helpful == payload.helpful and issue.feedback_text == feedback_text:
        return issue

    issue.helpful = payload.helpful
    issue.feedback_text = feedback_text
    issue.updated_at = _utcnow()
    session.add(issue)
    session.commit()
    session.refresh(issue)
    return issue


@router.get("/{task_id}/assets/{asset_path:path}")
def read_review_task_asset(
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    asset_path: str,
) -> FileResponse:
    _get_task_for_read(session=session, current_user=current_user, task_id=task_id)
    file_path = report_service.asset_path_for_task(task_id, asset_path)
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail="图片资产不存在")
    return FileResponse(file_path)


@router.post("/{task_id}/pdf-report", response_model=ReviewTaskPublic)
async def upload_review_task_pdf_report(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    upload: UploadFile = File(...),
) -> Any:
    existing_task = session.get(ReviewTask, task_id)
    if existing_task is None:
        raise HTTPException(status_code=404, detail="Review task not found")
    is_issue_feedback_submission = existing_task.feedback_protocol_version == 1
    task = (
        _get_task_for_issue_feedback_write_locked(
            session=session,
            current_user=current_user,
            task_id=task_id,
        )
        if is_issue_feedback_submission
        else _get_task_for_write(
            session=session,
            current_user=current_user,
            task_id=task_id,
        )
    )
    if is_issue_feedback_submission:
        effective_issues = _read_effective_review_task_issues(
            session=session,
            task_id=task.id,
        )
        if not effective_issues.source_consistency_state.is_valid:
            raise HTTPException(status_code=409, detail="issue_feedback_not_supported")
        issues = effective_issues.issues
        if any(issue.helpful is None for issue in issues):
            raise HTTPException(status_code=409, detail="issue_feedback_incomplete")
        if any(
            issue.feedback_text is None or not issue.feedback_text.strip()
            for issue in issues
        ):
            raise HTTPException(
                status_code=409,
                detail="issue_feedback_text_required",
            )

    report_path = await report_service.save_pdf_upload(task_id, upload)
    now = _utcnow()
    task.pdf_report_path = str(report_path)
    task.pdf_report_name = upload.filename or PDF_REPORT_FILE_NAME
    if is_issue_feedback_submission:
        task.feedback_submitted_at = now
    task.updated_at = now

    session.add(task)
    try:
        session.commit()
    except Exception:
        session.rollback()
        if is_issue_feedback_submission:
            report_path.unlink(missing_ok=True)
        raise
    session.refresh(task)
    return task


@router.post("/{task_id}/requirement-file", response_model=ReviewTaskPublic)
async def upload_review_task_requirement_file(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    task_id: uuid.UUID,
    upload: UploadFile = File(...),
) -> Any:
    task = _get_task_for_write(
        session=session, current_user=current_user, task_id=task_id
    )
    old_requirement_files = session.exec(
        select(ReviewTaskRequirementFile).where(
            ReviewTaskRequirementFile.task_id == task.id
        )
    ).all()
    old_requirement_paths = [
        Path(old_requirement_file.file_path)
        for old_requirement_file in old_requirement_files
    ]
    (
        file_path,
        file_size,
        file_hash,
        safe_file_name,
    ) = await report_service.save_requirement_upload(task_id, upload)
    requirement_file = ReviewTaskRequirementFile(
        task_id=task.id,
        file_name=safe_file_name,
        file_path=str(file_path),
        content_type=upload.content_type,
        file_size=file_size,
        sha256=file_hash,
    )
    session.add(requirement_file)
    session.flush()

    for old_requirement_file in old_requirement_files:
        session.delete(old_requirement_file)

    task.requirement_file_id = requirement_file.id
    task.requirement_file_name = safe_file_name
    task.requirement_file_size = file_size
    task.requirement_file_sha256 = file_hash
    task.updated_at = _utcnow()

    session.add(task)
    session.commit()
    session.refresh(task)
    for old_path in old_requirement_paths:
        if old_path != file_path:
            _unlink_file_after_commit(old_path)
    return task


@router.get("/{task_id}/report", response_class=PlainTextResponse)
def read_review_task_report(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> PlainTextResponse:
    task = _get_task_for_read(
        session=session, current_user=current_user, task_id=task_id
    )
    report_path = _read_report_path(task)
    return PlainTextResponse(report_path.read_bytes(), media_type="text/markdown")


@router.get("/{task_id}/download")
def download_review_task_report(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> FileResponse:
    task = _get_task_for_read(
        session=session, current_user=current_user, task_id=task_id
    )
    report_path = _read_report_path(task)
    return FileResponse(
        report_path,
        media_type="text/markdown",
        filename=task.report_name or REVIEW_REPORT_FILE_NAME,
    )


@router.get("/{task_id}/pdf/download")
def download_review_task_pdf_report(
    session: SessionDep, current_user: CurrentUser, task_id: uuid.UUID
) -> FileResponse:
    task = _get_task_for_read(
        session=session, current_user=current_user, task_id=task_id
    )
    report_path = _read_pdf_report_path(task)
    return FileResponse(
        report_path,
        media_type="application/pdf",
        filename=task.pdf_report_name or PDF_REPORT_FILE_NAME,
    )

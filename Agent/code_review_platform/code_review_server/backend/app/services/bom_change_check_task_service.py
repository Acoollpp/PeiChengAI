import json
import re
import shutil
import subprocess
import sys
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from app.core.config import settings
from app.services.codex_command import resolve_codex_command as _resolve_configured_codex_command
from app.services.storage_cleanup import robust_rmtree

REPORT_NAME = "BOM表改制检查报告.xlsx"
REPORT_NAME_PREFIX = "BOM表改制检查报告"
REPORT_ZIP_NAME = "BOM表改制检查报告.zip"
TIMING_LOG_NAME = "timing.jsonl"
BATCH_SUMMARY_NAME = "批量BOM表改制检查汇总.xlsx"
BOM_DIFF_SUMMARY_NAME = "BOM差异汇总.xlsx"
GROUP_COMPLETENESS_REPORT_NAME = "替代料组齐全差异.xlsx"


def _codex_home_dir() -> Path:
    configured = settings.CODEX_HOME_DIR.strip()
    if configured:
        return Path(configured).expanduser()
    return Path.home() / ".codex"


BOM_CHANGE_CHECK_SKILL_PATH = _codex_home_dir() / "skills" / "bom-change-check" / "SKILL.md"
BOM_CHANGE_CHECK_SCRIPT_PATH = BOM_CHANGE_CHECK_SKILL_PATH.parent / "scripts" / "check_bom_change.py"
REPLACEMENT_TABLE_DIR_NAME = "_shared_replacement_table"
REPLACEMENT_TABLE_METADATA_NAME = "metadata.json"
REPLACEMENT_TABLE_ALLOWED_SUFFIXES = {".xlsx", ".xlsm"}
REPLACEMENT_TABLE_ROLES = {
    "key": {
        "label": "关键物料替代表",
        "fallback": "关键物料替代表.xlsx",
        "keyword": "关键物料",
    },
    "no_key": {
        "label": "无关键物料替代表",
        "fallback": "无关键物料替代表.xlsx",
        "keyword": "无关键物料",
    },
}
COMPANY_MATERIAL_FOLDER_NAMES = {"公司物料表", "公司总料表", "总料表", "替代料汇总表"}


class BomChangeCheckTaskTimeout(RuntimeError):
    def __init__(self, message: str, *, result: dict[str, object] | None = None) -> None:
        super().__init__(message)
        self.result = result


class CodexProcessTimeout(RuntimeError):
    def __init__(self, *, stdout: str = "", stderr: str = "") -> None:
        super().__init__("Codex process timed out")
        self.stdout = stdout
        self.stderr = stderr


def task_dir(task_id: uuid.UUID) -> Path:
    return Path(settings.BOM_CHANGE_CHECK_STORAGE_DIR) / str(task_id)


def input_dir(task_id: uuid.UUID) -> Path:
    return task_dir(task_id) / "input"


def logs_dir(task_id: uuid.UUID) -> Path:
    return task_dir(task_id) / "logs"


def reports_dir(task_id: uuid.UUID) -> Path:
    return task_dir(task_id) / "reports"


def replacement_table_dir() -> Path:
    return Path(settings.BOM_CHANGE_CHECK_STORAGE_DIR) / REPLACEMENT_TABLE_DIR_NAME


def replacement_table_metadata_path() -> Path:
    return replacement_table_dir() / REPLACEMENT_TABLE_METADATA_NAME


def remove_task_files(task_id: uuid.UUID) -> None:
    storage_dir = Path(settings.BOM_CHANGE_CHECK_STORAGE_DIR).resolve()
    target_dir = (storage_dir / str(task_id)).resolve()
    if target_dir.name != str(task_id) or target_dir.parent != storage_dir:
        raise RuntimeError("Invalid BOM change check task storage path")
    if target_dir.exists():
        if not target_dir.is_dir():
            raise RuntimeError("BOM change check task storage path is not a directory")
        robust_rmtree(target_dir)


def _safe_filename(filename: str | None, fallback: str) -> str:
    name = Path(filename or fallback).name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip()
    return name or fallback


def _empty_replacement_role_public(role: str) -> dict[str, object]:
    return {
        "exists": False,
        "role": role,
        "label": REPLACEMENT_TABLE_ROLES[role]["label"],
    }


def _replacement_role_public(role: str, metadata: dict[str, object] | None) -> dict[str, object]:
    if not metadata:
        return _empty_replacement_role_public(role)
    return {
        "exists": True,
        "role": role,
        "label": REPLACEMENT_TABLE_ROLES[role]["label"],
        "file_name": metadata.get("file_name", ""),
        "file_size": metadata.get("file_size", 0),
        "content_type": metadata.get("content_type"),
        "uploaded_at": metadata.get("uploaded_at"),
        "uploaded_by": metadata.get("uploaded_by"),
        "note": metadata.get("note", ""),
    }


def _replacement_table_public(metadata: dict[str, object] | None = None) -> dict[str, object]:
    tables = latest_global_replacement_tables() if metadata is None else _normalize_replacement_metadata(metadata)
    role_public = {
        role: _replacement_role_public(role, tables.get(role))
        for role in REPLACEMENT_TABLE_ROLES
    }
    return {
        "exists": any(item.get("exists") for item in role_public.values()),
        "all_exists": all(item.get("exists") for item in role_public.values()),
        "tables": role_public,
    }


def _normalize_replacement_metadata(metadata: dict[str, object] | None) -> dict[str, dict[str, object]]:
    if not metadata:
        return {}
    raw_tables = metadata.get("tables")
    if isinstance(raw_tables, dict):
        tables: dict[str, dict[str, object]] = {}
        for role in REPLACEMENT_TABLE_ROLES:
            item = raw_tables.get(role)
            if isinstance(item, dict) and Path(str(item.get("file_path") or "")).is_file():
                tables[role] = item
        return tables

    # Backward compatibility for the old single-table metadata file.
    file_path = Path(str(metadata.get("file_path") or ""))
    if not file_path.is_file():
        return {}
    file_name = str(metadata.get("file_name") or file_path.name)
    if "无关键物料" in file_name:
        return {"no_key": metadata}
    if "关键物料" in file_name:
        return {"key": metadata}
    return {"legacy": metadata}


def latest_global_replacement_tables() -> dict[str, dict[str, object]]:
    metadata_path = replacement_table_metadata_path()
    if not metadata_path.is_file():
        return {}
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}
    return _normalize_replacement_metadata(metadata)


def latest_global_replacement_table() -> dict[str, object] | None:
    tables = latest_global_replacement_tables()
    for role in ("key", "no_key", "legacy"):
        if role in tables:
            return tables[role]
    return None


def save_global_replacement_table_bytes(
    *,
    filename: str,
    content: bytes,
    content_type: str | None,
    uploaded_by: str,
    role: str = "key",
) -> dict[str, object]:
    if role not in REPLACEMENT_TABLE_ROLES:
        raise HTTPException(status_code=400, detail="替代表类型不正确")
    if not content:
        raise HTTPException(status_code=400, detail=f"{REPLACEMENT_TABLE_ROLES[role]['label']}不能为空")
    safe_name = _safe_filename(filename, REPLACEMENT_TABLE_ROLES[role]["fallback"])
    if Path(safe_name).suffix.lower() not in REPLACEMENT_TABLE_ALLOWED_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"{REPLACEMENT_TABLE_ROLES[role]['label']}必须是 .xlsx 或 .xlsm 文件")
    directory = replacement_table_dir()
    directory.mkdir(parents=True, exist_ok=True)
    old_tables = latest_global_replacement_tables()
    old_role_path = old_tables.get(role, {}).get("file_path")
    if old_role_path:
        old_file = Path(str(old_role_path))
        if old_file.is_file() and old_file.parent == directory:
            old_file.unlink()
    file_path = directory / safe_name
    if file_path.exists():
        stem = file_path.stem
        suffix = file_path.suffix
        file_path = directory / f"{stem}-{role}{suffix}"
    file_path.write_bytes(content)
    role_metadata: dict[str, object] = {
        "file_name": safe_name,
        "file_path": str(file_path),
        "file_size": len(content),
        "content_type": content_type,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "uploaded_by": uploaded_by,
        "role": role,
        "label": REPLACEMENT_TABLE_ROLES[role]["label"],
        "note": f"本次及后续任务默认按照网页上传的最新{REPLACEMENT_TABLE_ROLES[role]['label']}进行检查。",
    }
    tables = {key: value for key, value in old_tables.items() if key in REPLACEMENT_TABLE_ROLES}
    tables[role] = role_metadata
    metadata: dict[str, object] = {"tables": tables}
    replacement_table_metadata_path().write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return role_metadata


async def save_global_replacement_table_upload(
    upload: UploadFile,
    *,
    uploaded_by: str,
    role: str = "key",
) -> dict[str, object]:
    content = await upload.read()
    return save_global_replacement_table_bytes(
        filename=upload.filename or REPLACEMENT_TABLE_ROLES[role]["fallback"],
        content=content,
        content_type=upload.content_type,
        uploaded_by=uploaded_by,
        role=role,
    )


def apply_replacement_table_to_task_input(
    task_id: uuid.UUID,
    *,
    metadata: dict[str, object] | None = None,
) -> dict[str, str | int | None]:
    projects_dir = input_dir(task_id) / "projects"
    projects_dir.mkdir(parents=True, exist_ok=True)
    batch_root = _find_batch_root_for_replacement_table(projects_dir)
    archive_tables = _existing_company_material_tables(batch_root)
    if archive_tables and metadata is None:
        return {
            "file_name": "；".join(path.name for path in archive_tables),
            "file_path": str(archive_tables[0].parent),
            "content_type": "",
            "file_size": sum(path.stat().st_size for path in archive_tables),
            "role": "replacement_table_used",
            "source": "project_archive",
            "note": "按照压缩包公司物料表目录内的关键物料/无关键物料替代表进行检查。",
        }

    if metadata:
        selected_tables = latest_global_replacement_tables()
        selected_tables.update(_normalize_replacement_metadata(metadata))
    else:
        selected_tables = latest_global_replacement_tables()
    if not selected_tables:
        raise HTTPException(status_code=400, detail="请先上传关键物料/无关键物料替代表，或在压缩包公司物料表目录内提供关键物料/无关键物料替代表")
    selected_items = [
        selected_tables[role]
        for role in ("key", "no_key")
        if role in selected_tables
    ]
    if not selected_items and "legacy" in selected_tables:
        selected_items = [selected_tables["legacy"]]
    for selected in selected_items:
        source = Path(str(selected.get("file_path") or ""))
        if not source.is_file():
            raise HTTPException(status_code=400, detail="最新替代表文件不存在，请重新上传")
    for folder_name in COMPANY_MATERIAL_FOLDER_NAMES:
        candidate = batch_root / folder_name
        if candidate.exists():
            if candidate.is_dir():
                shutil.rmtree(candidate)
            else:
                candidate.unlink()
    company_dir = batch_root / "公司物料表"
    company_dir.mkdir(parents=True, exist_ok=True)
    copied_targets: list[Path] = []
    for selected in selected_items:
        source = Path(str(selected.get("file_path") or ""))
        target = company_dir / str(selected.get("file_name") or source.name)
        if target.exists():
            target = company_dir / f"{target.stem}-{selected.get('role', 'table')}{target.suffix}"
        shutil.copy2(source, target)
        copied_targets.append(target)
    notice = company_dir / "网页替代表优先.txt"
    notice.write_text(
        "本次按照网页上传/平台保存的最新关键物料替代表、无关键物料替代表进行检查；压缩包内公司物料表已忽略。\n"
        + "\n".join(
            f"使用文件：{selected.get('label', '替代表')}：{selected.get('file_name', '')}；上传时间：{selected.get('uploaded_at', '')}；上传人：{selected.get('uploaded_by', '')}"
            for selected in selected_items
        )
        + "\n",
        encoding="utf-8",
    )
    return {
        "file_name": "；".join(path.name for path in copied_targets),
        "file_path": str(company_dir),
        "content_type": "",
        "file_size": sum(path.stat().st_size for path in copied_targets),
        "role": "replacement_table_used",
        "source": "web_latest",
        "note": "按照网页上传/平台保存的最新关键物料替代表、无关键物料替代表进行检查，压缩包内公司物料表已忽略。",
    }


def _existing_company_material_tables(batch_root: Path) -> list[Path]:
    for folder_name in COMPANY_MATERIAL_FOLDER_NAMES:
        candidate = batch_root / folder_name
        if not candidate.is_dir():
            continue
        return sorted(
            path
            for path in candidate.glob("*.xls*")
            if path.is_file() and not path.name.startswith("~$")
        )
    return []


def _find_batch_root_for_replacement_table(input_root: Path) -> Path:
    candidates: list[tuple[int, Path]] = []
    roots = [input_root]
    if input_root.exists():
        roots.extend(p for p in input_root.iterdir() if p.is_dir())
    for root in roots:
        try:
            names = [p.name for p in root.iterdir() if p.is_dir()]
        except FileNotFoundError:
            continue
        score = 0
        for tokens in (
            ["改制前BOM"],
            ["改制后BOM", "更新后BOM", "更新BOM"],
            ["改制单", "改制表"],
        ):
            if any(any(token in name for token in tokens) for name in names):
                score += 2
        if any(any(token in name for token in ["公司物料表", "公司总料表", "总料表", "替代料汇总表"]) for name in names):
            score += 1
        if score:
            candidates.append((score, root))
    if not candidates:
        return input_root
    return sorted(candidates, key=lambda item: item[0], reverse=True)[0][1]


def report_zip_path(task_id: uuid.UUID) -> Path:
    return reports_dir(task_id) / _report_zip_name(task_id)


async def save_project_upload(
    task_id: uuid.UUID,
    *,
    upload: UploadFile,
) -> list[dict[str, str | int | None]]:
    content = await upload.read()
    if not content:
        raise HTTPException(status_code=400, detail="项目压缩包不能为空")

    destination = input_dir(task_id)
    destination.mkdir(parents=True, exist_ok=True)
    file_name = _safe_filename(upload.filename, "bom-change-check-upload.zip")
    path = destination / file_name
    path.write_bytes(content)
    suffix = path.suffix.lower()
    if suffix == ".zip":
        _extract_zip(path, destination / "projects")
    elif suffix == ".rar":
        _extract_with_rar(path, destination / "projects")
    elif suffix == ".7z":
        _extract_with_7zip(path, destination / "projects")
    else:
        raise HTTPException(status_code=400, detail="项目压缩包必须是 .zip、.rar 或 .7z 文件")
    return [
        {
            "file_name": file_name,
            "file_path": str(path),
            "content_type": upload.content_type,
            "file_size": len(content),
            "role": "project_upload",
        }
    ]


def _extract_with_7zip(path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    command = _resolve_7z_command()
    _extract_with_archive_tool(path, destination, [command, "x", str(path), f"-o{destination}", "-y"])


def _extract_with_rar(path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    command = _resolve_rar_command()
    _extract_with_archive_tool(path, destination, [command, "x", "-y", str(path), str(destination) + "\\"])


def _extract_with_archive_tool(path: Path, destination: Path, command: list[str]) -> None:
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=300,
            check=False,
        )
    except FileNotFoundError as exc:
        raise HTTPException(
            status_code=400,
            detail=f"服务器未安装可用解压工具，无法解压 {path.suffix.lower()} 文件",
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(status_code=400, detail="压缩包解压超时，请检查文件大小") from exc
    if completed.returncode != 0:
        detail = _format_error("\n".join([completed.stderr or "", completed.stdout or ""]))
        raise HTTPException(
            status_code=400,
            detail=f"压缩包解压失败，请检查文件是否损坏或带密码：{detail or '未知错误'}",
        )
    _validate_extracted_paths(destination)


def _resolve_7z_command() -> str:
    candidates = [
        "7z",
        "7z.exe",
        r"C:\Program Files\7-Zip\7z.exe",
        r"C:\Program Files (x86)\7-Zip\7z.exe",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    raise HTTPException(status_code=400, detail="服务器未安装 7-Zip，无法解压 7z 文件")


def _resolve_rar_command() -> str:
    candidates = [
        r"D:\SoftWare\WinRAR\UnRAR.exe",
        r"D:\SoftWare\WinRAR\Rar.exe",
        r"D:\SoftWare\WinRAR\WinRAR.exe",
        "WinRAR",
        "WinRAR.exe",
        "RAR",
        "RAR.exe",
        "UnRAR",
        "UnRAR.exe",
        r"C:\Program Files\WinRAR\WinRAR.exe",
        r"C:\Program Files\WinRAR\Rar.exe",
        r"C:\Program Files\WinRAR\UnRAR.exe",
        r"C:\Program Files (x86)\WinRAR\WinRAR.exe",
        r"C:\Program Files (x86)\WinRAR\Rar.exe",
        r"C:\Program Files (x86)\WinRAR\UnRAR.exe",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    raise HTTPException(status_code=400, detail="服务器未安装 WinRAR/RAR/UnRAR，无法解压 rar 文件")


def _validate_extracted_paths(destination: Path) -> None:
    root = destination.resolve()
    for extracted in destination.rglob("*"):
        target = extracted.resolve()
        if root not in target.parents and target != root:
            raise HTTPException(status_code=400, detail="压缩包包含不安全路径")


def _extract_zip(path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as archive:
        for item in archive.infolist():
            target = (destination / item.filename).resolve()
            if destination.resolve() not in target.parents and target != destination.resolve():
                raise HTTPException(status_code=400, detail="Zip contains unsafe path")
            if item.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(archive.read(item))


def run_task(task_id: uuid.UUID, input_files: list[dict[str, object]]) -> dict[str, object]:
    task_path = task_dir(task_id)
    input_path = input_dir(task_id)
    output_path = reports_dir(task_id)
    log_path = logs_dir(task_id)
    output_path.mkdir(parents=True, exist_ok=True)
    log_path.mkdir(parents=True, exist_ok=True)
    command = [
        sys.executable,
        str(BOM_CHANGE_CHECK_SCRIPT_PATH),
        "--batch-input",
        str(input_path / "projects"),
        "--output-dir",
        str(output_path),
    ]
    _write_invocation_log(
        log_path,
        command=command,
        input_path=input_path,
        output_path=output_path,
    )
    try:
        completed = _run_codex_process(
            command,
            prompt="",
            cwd=task_path,
            timeout_seconds=settings.BOM_CHANGE_CHECK_CODEX_TIMEOUT_SECONDS,
        )
    except FileNotFoundError as exc:
        raise RuntimeError("服务端未找到 BOM改制检查脚本运行命令，无法执行 BOM改制检查。") from exc
    except CodexProcessTimeout as exc:
        (log_path / "codex-stdout.log").write_text(exc.stdout or "", encoding="utf-8")
        (log_path / "codex-stderr.log").write_text(exc.stderr or "", encoding="utf-8")
        partial_result: dict[str, object] | None = None
        try:
            partial_result = _build_task_result(task_id, allow_partial=True)
        except RuntimeError:
            partial_result = None
        if partial_result and _timing_log_has_batch_finished(output_path):
            return partial_result
        raise BomChangeCheckTaskTimeout(
            "BOM改制检查超时，已保留可用的中间报告，请下载报告包复查。" if partial_result else "BOM改制检查超时，请检查文件大小或稍后重试。",
            result=partial_result,
        ) from exc
    (log_path / "codex-stdout.log").write_text(completed.stdout or "", encoding="utf-8")
    (log_path / "codex-stderr.log").write_text(completed.stderr or "", encoding="utf-8")
    if completed.returncode != 0:
        detail = _format_error("\n".join([completed.stderr or "", completed.stdout or ""]))
        raise RuntimeError(f"Codex 执行 BOM改制检查失败：{detail or '未知错误'}")
    return _build_task_result(task_id)


def _build_task_result(task_id: uuid.UUID, *, allow_partial: bool = False) -> dict[str, object]:
    output_path = reports_dir(task_id)
    summary = _read_reports_summary(output_path, allow_partial=allow_partial)
    report_zip = _create_report_zip(task_id)
    result_files = [
        {
            "file_name": report_zip.name,
            "file_path": str(report_zip),
            "content_type": "application/zip",
            "file_size": report_zip.stat().st_size,
        }
    ]
    batch_summary_path = output_path / BATCH_SUMMARY_NAME
    if batch_summary_path.is_file():
        result_files.append(
            {
                "file_name": batch_summary_path.name,
                "file_path": str(batch_summary_path),
                "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "file_size": batch_summary_path.stat().st_size,
            }
        )
    for extra_report_name in (BOM_DIFF_SUMMARY_NAME, GROUP_COMPLETENESS_REPORT_NAME):
        extra_report_path = output_path / extra_report_name
        if extra_report_path.is_file():
            result_files.append(
                {
                    "file_name": extra_report_path.name,
                    "file_path": str(extra_report_path),
                    "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "file_size": extra_report_path.stat().st_size,
                }
            )
    return {
        "result_files": result_files,
        **summary,
    }


def _timing_log_has_batch_finished(output_path: Path) -> bool:
    timing_path = output_path / TIMING_LOG_NAME
    if not timing_path.is_file():
        return False
    for line in timing_path.read_text(encoding="utf-8", errors="replace").splitlines():
        try:
            payload = json.loads(line)
        except json.JSONDecodeError:
            continue
        if payload.get("event") == "batch_finished":
            return True
    return False


def _build_codex_prompt(
    *,
    input_path: Path,
    output_path: Path,
) -> str:
    return f"""[$bom-change-check]({BOM_CHANGE_CHECK_SKILL_PATH})

请使用 `$bom-change-check` skill 执行批量 BOM改制检查，并生成项目报告和批量汇总。

任务目录：
{task_dir_from_output(output_path)}

输入目录：
{input_path}

解压后的资料根目录：
{input_path / "projects"}

输出要求：
1. 必须完整阅读并遵守 `$bom-change-check` skill 的全部规则。
2. 后端平台只负责调度，实际检查由 Codex 根据 skill 完成；必须调用 skill 内置固定脚本 `check_bom_change.py` 的批量入口，不要重新生成批处理脚本，不要手工逐个 Excel 分析。
3. 固定执行命令格式如下，路径按本任务实际目录替换：
   `python {BOM_CHANGE_CHECK_SCRIPT_PATH} --batch-input "{input_path / "projects"}" --output-dir "{output_path}"`
4. 固定脚本会加载公司物料表/替代表索引；如果同时存在 `关键物料` 和 `无关键物料` 两个替代表，脚本会按每个 BOM 文件名选择对应索引。脚本串行处理项目，每个项目完成后立即写项目报告、批量汇总和 `timing.jsonl`。Codex 只做执行和最终复核，不要把业务规则复制成新的临时代码。
5. `input/projects` 下按固定分类文件夹提交资料，而不是每个项目单独一个文件夹。必须识别这些分类文件夹：
   - `改制前BOM`：存放多个改制前 BOM。
   - `改制后BOM`：存放多个改制后 BOM。
   - `更新后BOM` 或 `更新BOM`：存放只审查替代料是否合规时的更新后 BOM，`更新BOM` 按 `更新后BOM` 处理。
   - `改制表` 或 `改制单`：存放多个改制表。
   - `公司物料表`、`公司总料表`、`总料表` 或 `替代料汇总表`：存放公司物料表/替代料表；可以同时放 `关键物料` 和 `无关键物料` 两个替代表。
6. 如果压缩包内 `公司物料表` 目录已经提供关键物料/无关键物料替代表，且网页本次没有上传替代表，后端会保留并按压缩包内文件检查；如果网页本次上传或压缩包内没有替代表，后端会把网页上传/平台保存的最新 `关键物料替代表` 和 `无关键物料替代表` 复制到 `input/projects/公司物料表/`。`公司物料表/网页替代表优先.txt` 是说明文件，不要当作 Excel 输入。
7. 如果资料里只有 `改制后BOM`/`更新后BOM`/`更新BOM` 和 `公司物料表`/`公司总料表`/`总料表`/`替代料汇总表`，任务模式是“只审查替代料是否合规”。每个更新后 BOM 生成一个项目报告，只输出 `BOM差异核对`、`替代料关系核对`、`客户限制核对`、`summary`，不要把缺少改制前BOM或改制单判为失败。
7.1 每个任务结束时都必须输出替代料组齐全检查结果。优先检查 `更新后BOM`/`更新BOM` 里的每个 BOM；任务没有这两个文件夹时，改为检查 `改制后BOM` 里的每个 BOM。普通扁平目录输出独立工作簿 `替代料组齐全差异.xlsx`，按项目号分 sheet，只输出缺失物料，列顺序为 `物料代码`、`关系组`、`名称`、`规格型号`、`客户`、`备注`。K 开头关系组不检查；O 列包含当前 BOM 客户的物料不要求加入；缺失项只从选中的关键物料/无关键物料替代表生成。
7.1.1 `替代料组齐全差异.xlsx` 必须给被检查目录里的每个 BOM 都保留一个 sheet。该 BOM 有缺失项时列缺失项；没有缺失项时也保留 sheet，并写一行 `备注=无缺失项`，其它明细列留空。BOM 文件名不含 `匹配关键物料表` 或 `匹配无关键物料表` 时，不要判失败，保留该 BOM 的 sheet 并写 `备注=跳过`。
7.1.2 新的更新 BOM 压缩包可以是 `更新BOM/日期文件夹/BOM文件` 结构，例如 `更新BOM/2026-9-7/*.xlsx`。这种结构按日期汇总输出：每个日期文件夹生成一个 `BOM表改制检查报告-日期.xlsx`，每个 BOM 一个 sheet；含 `匹配关键物料表` 的 BOM 用关键物料替代表，含 `匹配无关键物料表` 的 BOM 用无关键物料替代表，两个关键词都没有的 BOM 不检查但必须在对应 sheet 写 `备注=跳过`。有关键词的 BOM 除了检查替代料组是否齐全，还必须扫描所有真实 BOM sheet 的 BOM 自检异常；分层错误、用量与位号数量不一致、位号重复、BOM字段缺失都写到对应 BOM sheet，并分别用简短备注 `分层错误`、`用量与位号数量不一致`、`位号重复`、`BOM字段缺失` 标识。`BOM差异汇总.xlsx` 也必须读取这些日期报告行，不能在日期报告已有差异时仍显示 `无BOM差异项`；汇总表只保留 `分层错误`、`用量与位号数量不一致`、`位号重复` 和 `跳过`，替代料组缺失项如 `BOM缺少该物料` 只保留在原日期报告中，但仍计入网页 `BOM差异异常` 数量。汇总表第一列必须是 `BOM表`，日期报告行填写来源 BOM sheet/文件名，普通项目报告行填写匹配到的改制后/更新后 BOM 文件名，不能只列差异而看不出是哪份 BOM 有问题。这种日期目录模式不用再额外生成 `替代料组齐全差异.xlsx`。
7.1.3 如果压缩包解压后省略了中间的 `更新BOM` 文件夹，直接形成 `任务根目录/2026-9-21/*.xlsx`、`任务根目录/2026-9-22/*.xlsx`，且 `任务根目录/公司物料表/` 为兄弟目录，也必须按日期更新 BOM 模式处理；不能把上一级 `input/projects` 因为存在 `公司物料表` 就误判为批量根目录。
7.2 `更新后BOM`/`更新BOM` 的替代料组齐全检查中，如果 `总替代料汇总表` 里同一个归一化关系编号被重复用于不连续的多块物料，不能跨块合并。每次再次出现基础关系编号时都视为一个新的替代关系块。例如前面 LCD 的 `T1005/T1005-1` 和后面电容的 `T1005/T1005-1` 不能互相要求齐全。
7.3 `取消替代关系物料` 不用于生成 `替代料组齐全差异.xlsx` 缺失项，只作为兜底放行：如果 `更新后BOM`/`更新BOM` 中某个物料没有出现在 `总替代料汇总表`，但出现在 `取消替代关系物料`，不算异常，也不要要求把该取消关系行里的其它料号全部加入 BOM。
8. 完整改制检查模式下，每个改制表生成一个项目报告。改制表中用 `归属=保护板1` 这一行的 `规格型号` 识别改制前 BOM 版本，用 `归属=保护板2` 这一行的 `规格型号` 识别改制后 BOM 版本；不要用 `接口板1`、`显示板1`、`LED彩灯条1` 或标题行识别前后版本。
9. 完整改制检查模式下，匹配重点是“项目号+版本号”：
   - BOM 文件名示例 `BOM-SST26-1101-1.0-备货-006`，匹配键是 `1101-1.0`。
   - 改制表里示例 `P8S150A-61101-1.0`，去掉项目号前导业务前缀 `6` 后匹配键是 `1101-1.0`。
   - BOM 文件名示例 `BOM-SST26-1629-1.0-瑞能-001`，匹配键是 `1629-1.0`。
   - 改制表里示例 `P8S150A-61629-1.0`，去掉项目号前导业务前缀 `6` 后匹配键是 `1629-1.0`。
   - 只有 `更新后BOM` 额外生成的 `替代料组齐全差异.xlsx` 的 sheet 名使用 5 位项目号，例如 `BOM-SST26-1629-1.0-瑞能-001.xlsx` 对应 sheet `61629-1.0`；其它 BOM 匹配和项目报告命名继续使用 4 位匹配规则。
10. 完整改制检查模式下，对每个改制表，用改制前规格型号匹配 `改制前BOM` 文件夹中的 BOM，用改制后规格型号匹配 `改制后BOM` 文件夹中的 BOM，再加上同一个改制表和公司物料表执行 `$bom-change-check`。
10.1 如果变更单表示“增加替代料”：变更前料号已经存在于 BOM，变更前用量和位号标记为不变，变更后物料为替代料时，不得替换或删除原主料行；应在该主料及其已有替代料行之后新增一行。新增替代料沿用原主料的位号，BOM 用量填写 `替代料`，主料原有用量和位号保持不变。新增料号的名称、规格型号、规格描述、封装和物料类型从物料主数据提取。
11. 如果匹配不到唯一的改制前 BOM、改制后 BOM 或公司物料表，必须在批量汇总中标明该改制表失败原因，并继续处理其他改制表；替代料合规模式不要求改制前 BOM 和改制单。
12. 每个项目都要生成一个项目号后缀报告，例如 `BOM表改制检查报告-515451301.xlsx`。完整模式报告必须包含 `改制单核对`、`BOM差异核对`、`替代料关系核对`、`客户限制核对`、`summary` sheet；不要生成单独的 `产品非电子物料` sheet，也不要生成单独的 `BOM自检` sheet。替代料合规模式报告必须包含 `BOM差异核对`、`替代料关系核对`、`客户限制核对`、`summary` sheet。
12. 替代料关系核对时，不再默认用单一旧 `替代料汇总表` 对所有 BOM 检查。公司物料表目录可以放两个替代表：文件名含 `关键物料` 的表和文件名含 `无关键物料` 的表。按 BOM 文件名选择：BOM 文件名含 `无关键物料` 时用无关键物料表；含 `关键物料` 时用关键物料表；判断时先判断 `无关键物料`，避免被 `关键物料` 子串误判。若同时存在两个替代表但 BOM 文件名不含这两个关键词，必须判输入异常，不能随便选。选中关键物料表时，只读 `总替代料汇总表有关键物料`，缺失时读 `替代料表匹配工具用关键物料`；选中无关键物料表时，只读 `总替代料汇总表无关键物料`，缺失时读 `替代料表匹配工具用无关键物料`。不要在同一个 workbook 内同时扫描有关键/无关键两个总表 sheet。替代料关系组先按 `T759`、`T759-1`、`T759-2` 这类前缀归一化；如果两个不同关系组的 `类型=主料` 行使用同一个主料物料代码，也视为同一个替代族，例如 `K187` 和 `T715` 都以 `1.03.01.000005` 为主料时，两个组下的替代料互相替换按正常处理。替代关系要先查选中 workbook 的对应总表 sheet；如果没匹配，再查同一 workbook 内的 `取消替代关系物料` sheet。如果 BOM 中上方主料和替代料两个料号出现在该 sheet 同一行的 `旧料号`、`新料号`、`对应替代料为` 等字段中，这对替代关系按 `正常` 放行，不输出到 `替代料关系核对` 明细；如果两个 sheet 都没有匹配关系，备注写 `替代料汇总表未找到该替代关系`。
13. 改制单核对删除项时，`位号` 和 `物料代码` 同等有效；只要指定删除位号在改制后 BOM 中已经不存在，就按正常处理，即使改制单物料代码空白、写错，或改制前 BOM 已经没有该位号。如果指定删除位号仍在改制后 BOM 任意物料下存在，则判异常。
14. 如果改制单某行的类别/归属包含 `PCBA`，这行改制项及它覆盖到的 BOM 差异都不需要在报告里体现：不输出到 `改制单核对`，也不输出到 `BOM差异核对`。
15. 先在改制单 `改制内容` 区域识别软件分界行：只要文本表达烧录/升级软件含义即可，例如 `烧录/升级程序`、`烧录程序`、`升级程序`、`PLM取程序`、`工程助理提供软件`。只有这一行以上的改制物品才归类为 `产品非电子物料`，不要因为某行缺少物料代码或位号就直接归类。这个归类只作为 BOM 对比匹配依据，不能直接从改制单抽取成报告行，也不能输出到独立 sheet。分界以上有位号的行按普通电子物料一样去改制前 BOM 和改制后 BOM 里对比；分界以上没有位号的标签/PE袋/二维码等物料，即使有物料代码，也用规格型号去改制前BOM和改制后BOM里对比。没有在对应 BOM 侧找到时，输出到 `BOM差异核对` 并判 `异常`。
16. 报告明细 sheet 中，`BOM差异核对` 必须输出所有 BOM 对比差异，包含 `结论=正常` 的差异行；产品非电子物料差异统一输出到 `BOM差异核对`，正常和异常都显示。`改制单核对` 必须按改制单物料明细逐行输出正常和异常行，不合并多条改制单内容。`改制单核对` 是改制单执行清单，不是只列异常；正常增加项也必须显示，例如 `1.02.01.000002 / 增加 / RA54 / 1 / 0603WAJ0000T5E`。其它明细 sheet 只输出 `结论=异常`、`结论=人工复查`、`结论=跳过` 的行，不输出 `正常` 行。
17. `BOM差异核对` 遇到 `IC-烧录程序`/烧录 IC/MCU 程序差异时，不要只输出 `跳过`。必须拆成两条差异：改制前项目的 MCU 写 `删除`，改制后项目的 MCU 写 `增加`，位号、用量、规格型号分别取对应 BOM 侧，备注写 `烧录IC按改制前后MCU拆分显示`。这两条 MCU 行必须挨在一起，且 `增加` 行排在 `删除` 行前面。
18. `BOM差异核对` 中如果位号差异被改制单覆盖，备注不能笼统写“位号变化已被改制单覆盖”；改制单是 `增加` 时写 `已按改制单添加`，改制单是 `删除` 时写 `已按改制单删除`。
19. `BOM差异核对` 中被改制单覆盖的行，`变更内容` 直接写 `增加` 或 `删除`，不要写笼统的 `变更`；同一位号前后料号不同且需要在 `BOM差异核对` 显示时，不管是否被改制单同时覆盖删除和增加，都拆成两行输出：旧物料写 `删除`，新物料写 `增加`，两行都要写清楚料号和规格型号，并且普通换料按 `删除` 在前、`增加` 在后挨在一起。
20. `BOM差异核对` 中，如果前后料号在公司替代料表中存在有效替代关系，这类只是替代料变化的正常差异不要输出到明细；如果是新增的 BOM 替代料行，只要该替代料和它绑定的上一行主料在公司替代料表中存在有效替代关系，也不要输出到明细，即使这个位号同时被改制单覆盖；但改制单明确删除的位号不能因为删除的是替代料行而隐藏，必须按正常删除输出，备注写 `已按改制单删除`，并用该替代料绑定的上一行主料信息显示；如果差异命中 `取消替代关系物料`，必须输出，备注写 `取消替代物料，可以忽略`，同一位号旧料换新料时仍按 `删除`/`增加` 拆两行，不是这种换料拆行的替代料显示行才写 `变更内容=替代料`。
21. 如果同一个位号同时出现在 BOM 的主料行和替代料行，`BOM差异核对` 的位号对比必须优先使用主料行，不能让下方替代料行覆盖主料行。例如 RA54 同时在主料 `1.02.01.000002 / 0603WAJ0000T5E` 和替代料 `1.02.01.000539 / RS-03000JT` 上，改制单要求增加 RA54 主料时，必须输出主料增加行，不能因为替代料关系有效而隐藏。
22. `BOM差异核对` 中被改制单覆盖的 `增加`/`删除` 行，`用量` 列必须填写改制单用量：增加写增加了多少个，删除写删除了多少个；其它显示行也要尽量从 BOM 行填写用量，替代料行用量为 `替代料` 时按位号数量填写。
23. 要对 BOM workbook 做 `BOM自检`：完整改制检查模式检查改制前 BOM 和改制后 BOM；只审查替代料是否合规模式检查更新后 BOM。自检必须扫描同一 workbook 内所有符合 BOM 表头的真实 BOM sheet，不只扫描用于前后差异对比的 `总表` 或第一个 sheet；名称含 `BOM表导出` 的辅助导出 sheet 必须跳过，因为它会把替代料展平为数字用量，不能作为独立 BOM 自检。只有实际有位号的电子物料行才检查 `用量` 是否等于位号数量，`用量` 必须等于位号数量，一个位号对应一个数量；没有位号的电子物料允许位号为空，不检查其位号数量或位号重复，也不检查用量或缺少位号。包装/非电子辅料，包括物料类型为 `包装` 或 `辅料` 以及纸箱、珍珠棉、标签、包装纸、黄胶、白胶等物料，不检查用量、位号、重复位号和分层；但如果 `物料类型` 明确为 `贴片料` 或 `插件料`，不能仅因名称/规格含 `标签` 就当成非电子辅料，仍要检查分层和重复位号；其 `位号` 列中如果写了 `单独生产备货时需配纸箱` 这类说明文字，也不能按位号解析。实际电子物料发现用量与位号数量不一致、同一行位号格内重复同一个位号，或同 sheet 位号重复时（同一个 BOM 的同一个 sheet 内，同一位号不能出现在多个非替代料行），输出到 `BOM差异核对` 并判 `异常`，`变更内容` 写 `BOM自检`，备注写清楚 BOM 侧、sheet 和行号。贴片料中名称或规格型号含“标签”的物料，以及名称或规格型号含 `PCB`、`钢网`、`保护板` 的物料，允许位号为空，不报缺少位号；但这个例外不跳过分层检查，且如果位号格有内容并重复同一个位号，仍要报 `位号重复`。分层检查只检查分层 sheet，不检查 `总表` 等混合 sheet：sheet 名含 `贴片`、`SMT`、`SMD` 视为贴片层，sheet 名含 `插件`、`DIP` 视为插件层，sheet 名含 `组装` 视为组装层，类似 `P16S300A-61905-1.0` 的产品型号 sheet 也视为组装层；实际电子物料中，`物料类型` 含 `贴片料` 必须放贴片层，含 `插件料` 必须放插件层，含 `包装` 或 `辅料` 的非电子物料不参与分层自检，放错层时输出到 `BOM差异核对` 并判 `异常`，`变更内容` 写 `BOM自检`，备注写清楚 `分层错误`、应放层、当前 sheet 和行号。固定 DIP 辅料例外：只有 `1.14.01.000010`、`1.14.01.000004`、`1.14.01.000001`、`4.01.01.000003`、`4.01.01.000001`、`4.01.02.000005`、`4.01.02.000001`、`4.01.02.000002` 这 8 个物料号允许作为 `辅料` 放在 `DIP`/`插件` 层，不报分层错误；这不是所有辅料的通用例外，其他包装/非电子辅料同样不参与自检。`用量=替代料` 的替代料行允许和上方绑定主料共用位号，不按跨行重复位号异常处理；位号格写有 `配套` 的配套物料行，例如 `配套D5,D6,D7,D8,D9,D10`，允许和 LED/主料共用这些位号，不按跨行重复位号异常处理；不同 sheet 之间不要互相判重复位号，因为 `总表` 和 SMT sheet 可能本来就包含同一批物料。
23.1 自检例外必须与固定脚本一致：同一行位号格内重复同一个位号，例如 `JA5,JA2,JA2`，必须报 `位号重复`，即使该行是替代料。图二固定插件料 `1.15.01.000005`、`1.15.01.000003`、`1.15.01.000009`、`1.15.01.000006`、`1.15.01.000008`、`1.16.03.000003` 仅在 `DIP`/`插件` sheet 中跳过分层和用量/位号数量检查，但仍检查重复位号；图三、图四物料 `4.03.01.000001`、`4.05.01.000001`、`4.05.01.000002`、`4.01.03.000001`、`1.18.01.000005`、`1.24.03.000040`、`1.25.02.000006`、`1.25.02.000029`、`4.03.02.000021` 跳过分层、用量/位号数量和重复位号三类自检；`4.03.02.000039` 不在忽略清单内，若作为 `贴片料` 放到非贴片层，必须报 `分层错误`。
24. 报告明细 sheet 只保留这 6 列，顺序必须是：`物料代码`、`变更内容`、`位号`、`用量`、`规格型号`、`备注`。不要保留行号、证据、检查类、结论等额外明细列。
25. `变更内容` 使用改制单操作值 `删除`/`增加`；同一位号从旧料号变成新料号且需要显示时，拆成 `删除` 和 `增加` 两行，不要只写一行 `变更`；只有无法表示成旧料删除加新料增加的 BOM 普通差异才写 `变更`；BOM替代料相关差异写 `替代料`；BOM自检异常写 `BOM自检`；跳过行写 `跳过`；`替代料关系核对` 的 `变更内容` 留空；客户限制写 `客户限制异常`/`客户限制人工复查`。`备注` 只写简短明确的问题类型或跳过原因，例如 `BOM缺少该物料`、`替代料汇总表未找到该替代关系`、`位号重复`；不要把行号、已有组内物料、查找过程等证据写进可见备注。`改制单核对` 只有异常行显示具体异常原因，正常行备注留空。
26. 每个项目报告放到输出目录下的同名项目子文件夹中，文件名必须保留项目号后缀，例如 `reports/改制单-515451301/BOM表改制检查报告-515451301.xlsx`。不要生成、保留、统计或打包无项目号副本 `BOM表改制检查报告.xlsx`；如果同一项目目录同时存在带编号报告和无编号报告，只保留带编号报告，删除或忽略无编号报告。
27. 在输出目录生成批量汇总文件，至少包含每个项目、匹配到的改制前 BOM、匹配到的改制后 BOM、状态、报告路径、改制单异常、BOM差异异常、替代料异常、客户限制异常、人工复查、跳过数量；BOM 自检异常计入 BOM差异异常；替代料合规模式下改制前 BOM、改制单、改制单异常可以留空或为 0。
28. 生成报告后必须用 openpyxl 打开确认可读，并确认 summary 统计与明细一致。
29. 不要把交付报告写入 input 目录；所有交付文件必须写入输出目录：
{output_path}
30. 如果某个项目输入结构异常，批量汇总中标明失败原因，并继续处理其他项目。
"""


def task_dir_from_output(output_path: Path) -> Path:
    return output_path.parent


def _run_codex_process(
    command: list[str],
    *,
    prompt: str,
    cwd: Path,
    timeout_seconds: int,
) -> subprocess.CompletedProcess[str]:
    process = subprocess.Popen(
        command,
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(cwd),
    )
    try:
        stdout, stderr = process.communicate(prompt, timeout=timeout_seconds)
    except subprocess.TimeoutExpired:
        _kill_process_tree(process.pid)
        try:
            stdout, stderr = process.communicate(timeout=5)
        except subprocess.TimeoutExpired:
            stdout, stderr = "", ""
        raise CodexProcessTimeout(stdout=stdout or "", stderr=stderr or "")
    return subprocess.CompletedProcess(
        args=command,
        returncode=process.returncode or 0,
        stdout=stdout,
        stderr=stderr,
    )


def _kill_process_tree(pid: int) -> None:
    subprocess.run(
        ["taskkill", "/PID", str(pid), "/T", "/F"],
        check=False,
        capture_output=True,
        text=True,
    )


def _resolve_codex_command(command: str) -> str:
    return _resolve_configured_codex_command(command, user_profile=_codex_home_dir().parent, which=shutil.which)

def _read_report_summary(report_path: Path) -> dict[str, int]:
    result = {
        "change_order_abnormal_count": 0,
        "bom_diff_abnormal_count": 0,
        "replacement_abnormal_count": 0,
        "customer_restriction_abnormal_count": 0,
        "manual_review_count": 0,
        "skipped_count": 0,
    }
    wb = load_workbook(report_path, read_only=True, data_only=True)
    try:
        if "summary" in wb.sheetnames:
            ws = wb["summary"]
            for row in ws.iter_rows(values_only=True):
                if not row or row[0] not in {"改制单核对", "BOM差异核对", "BOM自检", "替代料关系核对", "客户限制核对"}:
                    continue
                name = str(row[0])
                abnormal = _safe_int(row[2] if len(row) > 2 else 0)
                manual = _safe_int(row[3] if len(row) > 3 else 0)
                skipped = _safe_int(row[4] if len(row) > 4 else 0)
                if name == "改制单核对":
                    result["change_order_abnormal_count"] = abnormal
                elif name in {"BOM差异核对", "BOM自检"}:
                    result["bom_diff_abnormal_count"] += abnormal
                elif name == "替代料关系核对":
                    result["replacement_abnormal_count"] = abnormal
                elif name == "客户限制核对":
                    result["customer_restriction_abnormal_count"] = abnormal
                result["manual_review_count"] += manual
                result["skipped_count"] += skipped
    finally:
        wb.close()
    return result


def _read_reports_summary(output_path: Path, *, allow_partial: bool = False) -> dict[str, int]:
    reports = _project_report_paths(output_path)
    if not reports:
        failure_reason = _read_batch_summary_failure_reason(output_path)
        if failure_reason:
            raise RuntimeError(f"Codex 执行完成，但项目检查失败：{failure_reason}")
        raise RuntimeError("Codex 执行完成，但未生成任何 BOM改制检查报告")
    result = _empty_summary()
    readable_count = 0
    for report in reports:
        try:
            summary = _read_report_summary(report)
        except Exception:
            if allow_partial:
                continue
            raise
        readable_count += 1
        for key, value in summary.items():
            result[key] += value
    if readable_count == 0:
        raise RuntimeError("未找到可读取的 BOM改制检查报告")
    result["bom_diff_abnormal_count"] += _read_group_completeness_abnormal_count(output_path)
    return result


def _read_group_completeness_abnormal_count(output_path: Path) -> int:
    report_path = output_path / GROUP_COMPLETENESS_REPORT_NAME
    if not report_path.is_file():
        return 0
    wb = load_workbook(report_path, read_only=True, data_only=True)
    try:
        total = 0
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            index = {str(value): pos for pos, value in enumerate(headers) if value is not None}
            material_idx = index.get("物料代码")
            remark_idx = index.get("备注")
            if remark_idx is None:
                continue
            for row in rows:
                remark = str(row[remark_idx] or "").strip() if len(row) > remark_idx else ""
                material = (
                    str(row[material_idx] or "").strip()
                    if material_idx is not None and len(row) > material_idx
                    else ""
                )
                if remark and remark != "无缺失项" and (material or remark.startswith("输入异常")):
                    total += 1
        return total
    finally:
        wb.close()


def _read_batch_summary_failure_reason(output_path: Path) -> str:
    summary_path = output_path / BATCH_SUMMARY_NAME
    if not summary_path.is_file():
        return ""
    try:
        wb = load_workbook(summary_path, read_only=True, data_only=True)
    except Exception:
        return ""
    try:
        ws = wb["批量汇总"] if "批量汇总" in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return ""
        index = {str(value): pos for pos, value in enumerate(headers) if value is not None}
        project_idx = index.get("项目")
        status_idx = index.get("状态")
        reason_idx = index.get("失败原因")
        if reason_idx is None:
            return ""
        failures: list[str] = []
        for row in rows:
            reason = str(row[reason_idx] or "").strip() if len(row) > reason_idx else ""
            status = str(row[status_idx] or "").strip() if status_idx is not None and len(row) > status_idx else ""
            if not reason and status != "失败":
                continue
            project = str(row[project_idx] or "").strip() if project_idx is not None and len(row) > project_idx else ""
            failures.append(f"{project}：{reason}" if project and reason else reason or project or status)
        return "；".join(failures)
    finally:
        wb.close()


def _project_report_paths(output_path: Path) -> list[Path]:
    all_reports = sorted(
        path
        for path in output_path.rglob(f"{REPORT_NAME_PREFIX}*.xlsx")
        if path.is_file() and path.name != REPORT_ZIP_NAME and not path.name.startswith("~$")
    )
    coded_dirs = {
        path.parent
        for path in all_reports
        if path.name.startswith(f"{REPORT_NAME_PREFIX}-") and path.name.endswith(".xlsx")
    }
    return [path for path in all_reports if not (path.name == REPORT_NAME and path.parent in coded_dirs)]


def _empty_summary() -> dict[str, int]:
    return {
        "change_order_abnormal_count": 0,
        "bom_diff_abnormal_count": 0,
        "replacement_abnormal_count": 0,
        "customer_restriction_abnormal_count": 0,
        "manual_review_count": 0,
        "skipped_count": 0,
    }


def _safe_int(value: object) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _create_report_zip(task_id: uuid.UUID) -> Path:
    report_zip = report_zip_path(task_id)
    report_zip.parent.mkdir(parents=True, exist_ok=True)
    if report_zip.exists():
        report_zip.unlink()
    root = reports_dir(task_id)
    input_root = input_dir(task_id)
    log_root = logs_dir(task_id)
    with zipfile.ZipFile(report_zip, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for source_root in (input_root, root, log_root):
            if not source_root.exists():
                continue
            for path in source_root.rglob("*"):
                if not path.is_file() or path == report_zip:
                    continue
                if source_root == root and _is_legacy_report_copy(path):
                    continue
                if source_root == input_root:
                    prefix = "input"
                elif source_root == root:
                    prefix = "reports"
                else:
                    prefix = "logs"
                archive.write(path, Path(prefix) / path.relative_to(source_root))
    return report_zip


def _is_legacy_report_copy(path: Path) -> bool:
    if path.name != REPORT_NAME:
        return False
    return any(
        sibling.is_file()
        and sibling.name.startswith(f"{REPORT_NAME_PREFIX}-")
        and sibling.name.endswith(".xlsx")
        and not sibling.name.startswith("~$")
        for sibling in path.parent.glob(f"{REPORT_NAME_PREFIX}-*.xlsx")
    )


def _report_zip_name(task_id: uuid.UUID) -> str:
    source_name = _uploaded_archive_stem(task_id)
    if source_name:
        return f"{source_name}.zip"
    return REPORT_ZIP_NAME


def _uploaded_archive_stem(task_id: uuid.UUID) -> str | None:
    source_dir = input_dir(task_id)
    if not source_dir.is_dir():
        return None
    for suffix in (".zip", ".rar", ".7z"):
        archives = sorted(path for path in source_dir.glob(f"*{suffix}") if path.is_file())
        if archives:
            return archives[0].stem
    return None


def _write_invocation_log(
    log_path: Path,
    *,
    command: list[str],
    input_path: Path,
    output_path: Path,
) -> None:
    payload = {
        "command": command,
        "input_path": str(input_path),
        "input_exists": input_path.exists(),
        "output_path": str(output_path),
        "output_exists": output_path.exists(),
        "skill_path": str(BOM_CHANGE_CHECK_SKILL_PATH),
        "skill_exists": BOM_CHANGE_CHECK_SKILL_PATH.exists(),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    (log_path / "codex-invocation.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _format_error(detail: str) -> str:
    lines = [line.strip() for line in detail.splitlines() if line.strip()]
    return "\n".join(lines[-8:])

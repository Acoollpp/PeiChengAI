from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

BOM_HEADERS = ["物料编码", "规格型号", "规格描述", "用量", "位号"]
HEADER_ALIASES = {
    "序号": ["序号"],
    "物料类型": ["物料类型", "物料分类"],
    "物料编码": ["物料编码", "物料代码"],
    "名称": ["名称", "物料名称", "物料描述"],
    "规格型号": ["规格型号", "型号", "规格"],
    "规格描述": ["规格描述", "BMS规则描述"],
    "封装": ["BMS封装", "封装"],
    "用量": ["用量"],
    "位号": ["位号"],
}
MATERIAL_MASTER_HEADERS = {
    "物料编码": ["物料编码", "物料代码"],
    "名称": ["名称", "物料描述"],
    "规格型号": ["规格型号", "型号"],
    "规格描述": ["规格描述", "BMS规格描述", "BMS规则描述"],
    "封装": ["封装", "BMS封装"],
    "物料类型": ["物料类型", "一级分类"],
}
REPORT_SHEET_NAME = "BOM电子档修改报告"
CHANGE_RECORD_SHEET_NAME = "变更记录"
CHANGE_FILL = PatternFill(fill_type="solid", fgColor="BDD7EE")
UNCHANGED_QUANTITY_MARKER = "用量不变"
UNCHANGED_LOCATOR_MARKER = "位号不变"


@dataclass(frozen=True)
class MaterialInfo:
    code: str
    name: str
    model: str
    description: str
    package: str
    material_type: str


@dataclass(frozen=True)
class ChangeRow:
    before_code: str
    before_model: str
    before_description: str
    before_quantity: str
    before_locator: str
    after_code: str
    after_model: str
    after_description: str
    after_quantity: str
    after_locator: str


@dataclass(frozen=True)
class ChangeGroup:
    key: str
    locator_type: str
    locator_value: str
    rows: list[ChangeRow]


@dataclass
class ReportRow:
    sheet_name: str
    locator_type: str
    locator_value: str
    action: str
    before_code: str = ""
    before_model: str = ""
    before_description: str = ""
    before_quantity: str = ""
    before_locator: str = ""
    after_code: str = ""
    after_model: str = ""
    after_description: str = ""
    after_quantity: str = ""
    after_locator: str = ""
    result: str = ""
    reason: str = ""


@dataclass
class ChecklistRow:
    sheet_name: str
    sheet_type: str
    locator_value: str
    action: str
    before_code: str = ""
    before_model: str = ""
    before_quantity: str = ""
    before_locator: str = ""
    after_code: str = ""
    after_model: str = ""
    after_quantity: str = ""
    after_locator: str = ""
    after_action: str = ""
    status: str = ""
    reason: str = ""


@dataclass(frozen=True)
class ValidationIssue:
    sheet_name: str
    before_code: str
    after_code: str
    locator_value: str
    reason: str


@dataclass(frozen=True)
class ProcessResult:
    output_path: Path
    modified_count: int
    abnormal_count: int
    report_count: int
    processed_sheet_count: int
    checklist_rows: list[ChecklistRow] = field(default_factory=list)


@dataclass(frozen=True)
class ChangeRecordPeople:
    preparer: str
    reviewer: str
    approver: str


@dataclass(frozen=True)
class PendingOperation:
    first_row: int
    rows: list[int]
    group: ChangeGroup
    before_rows: list[ChangeRow]
    after_rows: list[ChangeRow]
    new_materials: list[MaterialInfo]
    is_partial_locator_change: bool
    preserve_first_row: bool = False


def _shift_operations_after_row_deletions(
    operations: list[PendingOperation],
    deleted_rows: set[int],
) -> list[PendingOperation]:
    if not deleted_rows:
        return operations

    def shift(row_index: int) -> int:
        return row_index - sum(1 for deleted_row in deleted_rows if deleted_row < row_index)

    shifted: list[PendingOperation] = []
    for operation in operations:
        remaining_rows = [shift(row_index) for row_index in operation.rows if row_index not in deleted_rows]
        shifted.append(
            PendingOperation(
                first_row=shift(operation.first_row),
                rows=remaining_rows,
                group=operation.group,
                before_rows=operation.before_rows,
                after_rows=operation.after_rows,
                new_materials=operation.new_materials,
                is_partial_locator_change=operation.is_partial_locator_change,
            )
        )
    return shifted


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_blank_or_slash(value: str) -> bool:
    return value.strip() in {"", "/"}


def _looks_material_code(value: str) -> bool:
    return bool(re.fullmatch(r"\d+(?:\.\d+){2,}", value.strip()))


def _is_meaningful_locator(value: str) -> bool:
    text = value.strip()
    if _is_blank_or_slash(text) or text == UNCHANGED_LOCATOR_MARKER:
        return False
    try:
        Decimal(text)
    except Exception:
        return True
    return False


def _is_valid_quantity_marker(value: str) -> bool:
    text = value.strip()
    if text in {"替代料", UNCHANGED_QUANTITY_MARKER}:
        return True
    try:
        Decimal(str(format_quantity(text)))
    except Exception:
        return False
    return True


def _is_valid_after(row: ChangeRow) -> bool:
    return _looks_material_code(row.after_code) and _is_valid_quantity_marker(row.after_quantity)


def _is_valid_before(row: ChangeRow) -> bool:
    return _looks_material_code(row.before_code) and _is_valid_quantity_marker(row.before_quantity)


def format_quantity(value: Any) -> Any:
    text = _text(value)
    if text in {"", "/", "替代料"}:
        return text
    if "/" in text:
        try:
            fraction = Fraction(text.replace(" ", ""))
        except (ValueError, ZeroDivisionError):
            return text
        denominator = fraction.denominator
        while denominator % 2 == 0:
            denominator //= 2
        while denominator % 5 == 0:
            denominator //= 5
        if denominator == 1:
            decimal_text = format(
                Decimal(fraction.numerator) / Decimal(fraction.denominator),
                "f",
            ).rstrip("0").rstrip(".")
        else:
            decimal_text = str(
                (Decimal(fraction.numerator) / Decimal(fraction.denominator)).quantize(
                    Decimal("0.0001"),
                    rounding=ROUND_HALF_UP,
                )
            )
        return float(decimal_text) if "." in decimal_text else int(decimal_text)
    try:
        decimal = Decimal(text)
    except Exception:
        return text
    return float(decimal) if "." in text else int(decimal)


def _quantity_compare_text(value: Any) -> str:
    formatted = format_quantity(value)
    if isinstance(formatted, float):
        decimal = Decimal(str(formatted))
    else:
        text = _text(formatted)
        try:
            decimal = Decimal(text)
        except Exception:
            return text
    if decimal.as_tuple().exponent < -4:
        decimal = decimal.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    return format(decimal.normalize(), "f")


def _locator_compare_text(value: Any) -> str:
    text = _text(value)
    return text if _is_meaningful_locator(text) else "/"


def _locator_bom_value(value: Any) -> str:
    text = _text(value)
    return text if _is_meaningful_locator(text) else ""


def _split_locators(value: Any) -> list[str]:
    text = _text(value)
    if not _is_meaningful_locator(text):
        return []
    return [part.strip() for part in re.split(r"[,，、;；\s]+", text) if part.strip()]


def _join_locators(locators: list[str]) -> str:
    return ",".join(locators)


def _cell_text(ws: Worksheet, row: int, column: int) -> str:
    return _text(ws.cell(row=row, column=column).value)


def _normalize_header(value: str) -> str:
    return "".join(value.split())


def _find_header_in_values(header: str, values: list[str]) -> int | None:
    normalized_values = [_normalize_header(value) for value in values]
    aliases = [_normalize_header(alias) for alias in HEADER_ALIASES.get(header, [header])]
    for alias in aliases:
        if alias in normalized_values:
            return normalized_values.index(alias) + 1
    return None


def _find_header_row(ws: Worksheet, headers: list[str]) -> tuple[int, dict[str, int]] | None:
    for row in range(1, min(ws.max_row, 80) + 1):
        values = [_cell_text(ws, row, column) for column in range(1, ws.max_column + 1)]
        mapping: dict[str, int] = {}
        for header in headers:
            column = _find_header_in_values(header, values)
            if column is not None:
                mapping[header] = column
        if len(mapping) == len(headers):
            return row, mapping
    return None


def _add_optional_columns(ws: Worksheet, header_row: int, columns: dict[str, int]) -> dict[str, int]:
    values = [_cell_text(ws, header_row, column) for column in range(1, ws.max_column + 1)]
    expanded = dict(columns)
    for header in ("序号", "物料类型", "名称", "封装"):
        column = _find_header_in_values(header, values)
        if column is not None:
            expanded[header] = column
    return expanded


def _find_material_master_header(ws: Worksheet) -> tuple[int, dict[str, int]] | None:
    required = ["物料编码", "名称", "规格型号", "规格描述"]
    for row in range(1, min(ws.max_row, 30) + 1):
        values = [_cell_text(ws, row, column) for column in range(1, ws.max_column + 1)]
        mapping: dict[str, int] = {}
        for header, aliases in MATERIAL_MASTER_HEADERS.items():
            column = _find_header_in_values(header, values) if header in HEADER_ALIASES else None
            if column is None:
                normalized_values = [_normalize_header(value) for value in values]
                normalized_aliases = [_normalize_header(alias) for alias in aliases]
                for index, value in enumerate(normalized_values, start=1):
                    if value in normalized_aliases:
                        column = index
                        break
            if column is not None:
                mapping[header] = column
        if all(header in mapping for header in required):
            return row, mapping
    return None


def load_material_master(path: Path) -> dict[str, list[MaterialInfo]]:
    wb = load_workbook(path, data_only=True)
    matches: dict[str, list[MaterialInfo]] = {}
    for ws in wb.worksheets:
        header = _find_material_master_header(ws)
        if header is None:
            continue
        header_row, columns = header
        for row in range(header_row + 1, ws.max_row + 1):
            code = _cell_text(ws, row, columns["物料编码"])
            if not code:
                continue
            matches.setdefault(code, []).append(
                MaterialInfo(
                    code=code,
                    name=_cell_text(ws, row, columns["名称"]),
                    model=_cell_text(ws, row, columns["规格型号"]),
                    description=_cell_text(ws, row, columns["规格描述"]),
                    package=_cell_text(ws, row, columns["封装"]) if "封装" in columns else "",
                    material_type=_cell_text(ws, row, columns["物料类型"]) if "物料类型" in columns else "",
                )
            )
    return matches


def _find_change_columns(ws: Worksheet) -> tuple[int, dict[str, int], dict[str, int]]:
    before_title_col = after_title_col = None
    title_row = None
    for row in range(1, min(ws.max_row, 120) + 1):
        for column in range(1, ws.max_column + 1):
            value = _cell_text(ws, row, column)
            if value == "变更前描述":
                before_title_col = column
                title_row = row
            elif value == "变更后描述":
                after_title_col = column
                title_row = row if title_row is None else title_row
        if before_title_col and after_title_col:
            break
    if not before_title_col or not after_title_col or title_row is None:
        raise ValueError("变更单未找到“变更前描述/变更后描述”明细区")

    header_row = None
    for row in range(title_row + 1, min(ws.max_row, title_row + 8) + 1):
        row_values = [_cell_text(ws, row, column) for column in range(1, ws.max_column + 1)]
        if row_values.count("物料编码") >= 2 and row_values.count("位号") >= 2:
            header_row = row
            break
    if header_row is None:
        raise ValueError("变更单未找到左右两侧明细表头")

    before_map: dict[str, int] = {}
    after_map: dict[str, int] = {}
    for column in range(before_title_col, after_title_col):
        value = _cell_text(ws, header_row, column)
        for header in BOM_HEADERS:
            if _find_header_in_values(header, [value]) is not None:
                before_map[header] = column
    for column in range(after_title_col, ws.max_column + 1):
        value = _cell_text(ws, header_row, column)
        for header in BOM_HEADERS:
            if _find_header_in_values(header, [value]) is not None:
                after_map[header] = column
    # Some approved ECN templates omit the specification-description column.
    # The material master is the source for that value during BOM updates.
    required_headers = ("物料编码", "规格型号", "用量", "位号")
    missing = [header for header in required_headers if header not in before_map or header not in after_map]
    if missing:
        raise ValueError(f"变更单缺少必要明细列：{'、'.join(missing)}")
    return header_row, before_map, after_map


def _single_table_action_score(ws: Worksheet, header_row: int, column: int) -> int:
    actions = {"增加", "删除", "取消"}
    end_row = min(ws.max_row, header_row + 20)
    return sum(1 for row in range(header_row + 1, end_row + 1) if _cell_text(ws, row, column) in actions)


def _find_single_table_change_columns(ws: Worksheet) -> tuple[int, dict[str, int]]:
    header_row = None
    columns: dict[str, int] = {}
    for row_index in range(1, min(ws.max_row, 120) + 1):
        values = [_cell_text(ws, row_index, column) for column in range(1, ws.max_column + 1)]
        material_code_column = None
        for column, value in enumerate(values, start=1):
            if value in {"物料代码", "物料编码"}:
                material_code_column = column
                break
        if "变更内容" not in values or "规格型号" not in values or material_code_column is None:
            continue
        header_row = row_index
        action_columns: list[int] = []
        for column, value in enumerate(values, start=1):
            if value in {"位号", "规格型号"}:
                columns[value] = column
            if value == "变更内容":
                action_columns.append(column)
            if value in {"用量", "数量"}:
                columns["用量"] = column
            if value in {"物料代码", "物料编码"}:
                columns["物料代码"] = column
        if "用量" not in columns and "位号" in columns and "规格型号" in columns:
            between_columns = [
                column
                for column in range(columns["位号"] + 1, columns["规格型号"])
                if _cell_text(ws, row_index, column) == ""
            ]
            if len(between_columns) == 1:
                columns["用量"] = between_columns[0]
        columns["变更内容"] = max(
            action_columns,
            key=lambda column: (_single_table_action_score(ws, row_index, column), column),
        )
        break
    if header_row is None:
        raise ValueError("变更单未找到单表明细区")
    missing = [header for header in ("变更内容", "用量", "规格型号", "物料代码") if header not in columns]
    if missing:
        raise ValueError(f"变更单单表明细区缺少必要列：{'、'.join(missing)}")
    return header_row, columns


def _parse_single_table_change_rows(ws: Worksheet) -> list[ChangeRow]:
    header_row, columns = _find_single_table_change_columns(ws)
    rows: list[ChangeRow] = []
    for row_index in range(header_row + 1, ws.max_row + 1):
        action = _cell_text(ws, row_index, columns["变更内容"])
        if action not in {"增加", "删除", "取消"}:
            continue
        code = _cell_text(ws, row_index, columns["物料代码"])
        if not _looks_material_code(code):
            continue
        model = _cell_text(ws, row_index, columns["规格型号"])
        quantity = _cell_text(ws, row_index, columns["用量"])
        locator = _cell_text(ws, row_index, columns.get("位号", 0)) or "/"
        if action == "增加":
            rows.append(ChangeRow("/", "/", "/", "/", "/", code, model, "", quantity, locator))
        else:
            rows.append(ChangeRow(code, model, "", quantity, locator, "/", "/", "/", "/", "/"))
    if not rows:
        raise ValueError("变更单单表明细区没有可识别的增加或删除物料")
    return rows


def parse_change_notice(path: Path) -> list[ChangeGroup]:
    wb = load_workbook(path, data_only=True)
    parse_error: Exception | None = None
    for ws in wb.worksheets:
        try:
            header_row, before_map, after_map = _find_change_columns(ws)
        except ValueError as exc:
            parse_error = exc
            try:
                return _single_table_change_groups(_parse_single_table_change_rows(ws))
            except ValueError as single_table_error:
                parse_error = single_table_error
                continue
        rows: list[ChangeRow] = []
        last_after_code = ""
        last_after_model = ""
        last_after_description = ""
        last_after_quantity = ""
        last_after_locator = ""
        last_before_locator = ""
        pending_before_indexes: list[int] = []
        merged_before_locator_rows = {
            row_index
            for merged_range in ws.merged_cells.ranges
            if merged_range.min_col <= before_map["位号"] <= merged_range.max_col
            for row_index in range(merged_range.min_row, merged_range.max_row + 1)
        }
        for row_index in range(header_row + 1, ws.max_row + 1):
            before_code = _cell_text(ws, row_index, before_map["物料编码"])
            before_model = _cell_text(ws, row_index, before_map["规格型号"])
            before_description = (
                _cell_text(ws, row_index, before_map["规格描述"])
                if "规格描述" in before_map
                else ""
            )
            before_quantity = _cell_text(ws, row_index, before_map["用量"])
            before_locator = _cell_text(ws, row_index, before_map["位号"])
            after_code = _cell_text(ws, row_index, after_map["物料编码"])
            after_model = _cell_text(ws, row_index, after_map["规格型号"])
            after_description = (
                _cell_text(ws, row_index, after_map["规格描述"])
                if "规格描述" in after_map
                else ""
            )
            after_quantity = _cell_text(ws, row_index, after_map["用量"])
            after_locator = _cell_text(ws, row_index, after_map["位号"])
            if row_index in merged_before_locator_rows:
                if _is_blank_or_slash(after_code):
                    after_code = ""
                if _is_blank_or_slash(after_model):
                    after_model = ""
                if _is_blank_or_slash(after_description):
                    after_description = ""
                if _is_blank_or_slash(after_quantity):
                    after_quantity = ""
                if _is_blank_or_slash(after_locator):
                    after_locator = ""
            has_raw_before = any([before_code, before_model, before_description, before_quantity, before_locator])
            has_raw_after = any([after_code, after_model, after_description, after_quantity, after_locator])
            if not has_raw_before and not has_raw_after:
                continue
            if before_locator:
                last_before_locator = before_locator
            elif before_code and last_before_locator:
                before_locator = last_before_locator
            raw_after_code = after_code
            if after_code:
                last_after_code = after_code
            elif last_after_code:
                after_code = last_after_code
            if after_model:
                last_after_model = after_model
            elif last_after_model:
                after_model = last_after_model
            if after_description:
                last_after_description = after_description
            elif last_after_description:
                after_description = last_after_description
            if after_quantity:
                last_after_quantity = after_quantity
            elif last_after_quantity:
                after_quantity = last_after_quantity
            if after_locator:
                last_after_locator = after_locator
            elif last_after_locator:
                after_locator = last_after_locator
            if raw_after_code and pending_before_indexes:
                for pending_index in pending_before_indexes:
                    pending = rows[pending_index]
                    rows[pending_index] = ChangeRow(
                        before_code=pending.before_code,
                        before_model=pending.before_model,
                        before_description=pending.before_description,
                        before_quantity=pending.before_quantity,
                        before_locator=pending.before_locator,
                        after_code=after_code,
                        after_model=after_model,
                        after_description=after_description,
                        after_quantity=after_quantity,
                        after_locator=after_locator,
                    )
                pending_before_indexes.clear()
                if not before_code:
                    continue
            row = ChangeRow(
                before_code=before_code,
                before_model=before_model,
                before_description=before_description,
                before_quantity=before_quantity,
                before_locator=before_locator,
                after_code=after_code,
                after_model=after_model,
                after_description=after_description,
                after_quantity=after_quantity,
                after_locator=after_locator,
            )
            if any(getattr(row, field) for field in row.__dataclass_fields__):
                rows.append(row)
                if before_code and not raw_after_code:
                    pending_before_indexes.append(len(rows) - 1)
        return _group_change_rows(rows)
    raise ValueError(str(parse_error or "变更单没有可识别的 sheet"))


def _single_table_change_groups(rows: list[ChangeRow]) -> list[ChangeGroup]:
    locator_rows = [
        row
        for row in rows
        if _is_meaningful_locator(row.before_locator) or _is_meaningful_locator(row.after_locator)
    ]
    groups = _group_change_rows(locator_rows)
    groups.extend(
        ChangeGroup(
            key=f"code:{row.after_code or row.before_code}:{len(groups) + 1}",
            locator_type="物料编码",
            locator_value=row.after_code or row.before_code,
            rows=[row],
        )
        for row in rows
        if row not in locator_rows
    )
    return groups


def _group_change_rows(rows: list[ChangeRow]) -> list[ChangeGroup]:
    code_groups: list[tuple[str, list[ChangeRow]]] = []
    by_locator: dict[str, list[ChangeRow]] = {}
    current_code_group: list[ChangeRow] | None = None

    for row in rows:
        has_before = _is_valid_before(row)
        has_after = _is_valid_after(row)
        if not has_before and not has_after:
            continue
        locator = ""
        if _is_meaningful_locator(row.before_locator):
            locator = row.before_locator
        elif _is_meaningful_locator(row.after_locator):
            locator = row.after_locator
        if locator:
            by_locator.setdefault(locator, []).append(row)
            current_code_group = None
        elif has_before:
            current_code_group = [row]
            code_groups.append((row.before_code, current_code_group))
        elif current_code_group is not None:
            current_code_group.append(row)
        elif has_after:
            current_code_group = [row]
            code_groups.append((row.after_code, current_code_group))

    groups: list[ChangeGroup] = []
    for key_value, group_rows in code_groups:
        groups.append(
            ChangeGroup(
                key=f"code:{key_value}:{len(groups) + 1}",
                locator_type="物料编码",
                locator_value=key_value,
                rows=group_rows,
            )
        )
    for locator, locator_rows in by_locator.items():
        groups.append(
            ChangeGroup(
                key=f"ref:{locator}",
                locator_type="位号",
                locator_value=locator,
                rows=locator_rows,
            )
        )
    return groups


def _row_tuple(ws: Worksheet, row_index: int, columns: dict[str, int]) -> tuple[str, str, str, str, str]:
    return (
        _cell_text(ws, row_index, columns["物料编码"]),
        _cell_text(ws, row_index, columns["规格型号"]),
        _cell_text(ws, row_index, columns["规格描述"]),
        _quantity_compare_text(ws.cell(row=row_index, column=columns["用量"]).value),
        _locator_compare_text(ws.cell(row=row_index, column=columns["位号"]).value),
    )


def _matches_before_row(actual: tuple[str, str, str, str, str], before: ChangeRow) -> bool:
    return (
        actual[0] == before.before_code
        and (before.before_quantity == UNCHANGED_QUANTITY_MARKER or actual[3] == _quantity_compare_text(before.before_quantity))
        and (before.before_locator == UNCHANGED_LOCATOR_MARKER or actual[4] == _locator_compare_text(before.before_locator))
    )


def _rows_match_before_rows(actual_rows: list[tuple[str, str, str, str, str]], before_rows: list[ChangeRow]) -> bool:
    remaining = list(actual_rows)
    for before in before_rows:
        match_index = next(
            (index for index, actual in enumerate(remaining) if _matches_before_row(actual, before)),
            None,
        )
        if match_index is None:
            return False
        remaining.pop(match_index)
    return not remaining


def _capture_row_style(ws: Worksheet, row_index: int) -> tuple[float | None, list[dict[str, Any]]]:
    styles: list[dict[str, Any]] = []
    for column in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_index, column=column)
        styles.append(
            {
                "style": copy(cell._style) if cell.has_style else None,
                "number_format": cell.number_format,
                "alignment": copy(cell.alignment),
                "border": copy(cell.border),
            }
        )
    return ws.row_dimensions[row_index].height, styles


def _apply_row_style(
    ws: Worksheet,
    target_row: int,
    captured: tuple[float | None, list[dict[str, Any]]],
) -> None:
    height, styles = captured
    ws.row_dimensions[target_row].height = height
    for column, style in enumerate(styles, start=1):
        target = ws.cell(row=target_row, column=column)
        if style["style"] is not None:
            target._style = copy(style["style"])
        target.number_format = style["number_format"]
        target.alignment = copy(style["alignment"])
        target.border = copy(style["border"])


def _numeric_cell_value(value: Any) -> int | None:
    try:
        decimal = Decimal(_text(value))
    except Exception:
        return None
    if decimal == decimal.to_integral_value():
        return int(decimal)
    return None


def _is_numeric_quantity_value(value: Any) -> bool:
    text = _text(value)
    if _is_blank_or_slash(text) or text == "替代料":
        return False
    try:
        Decimal(_quantity_compare_text(value))
    except Exception:
        return bool(re.search(r"\d", text))
    return True


def _renumber_sequence_for_numeric_quantity_rows(
    ws: Worksheet,
    *,
    header_row: int,
    columns: dict[str, int],
) -> None:
    sequence_column = columns.get("序号")
    if sequence_column is None:
        return

    sequence = 1
    for row_index in range(header_row + 1, ws.max_row + 1):
        if not _is_numeric_quantity_value(ws.cell(row=row_index, column=columns["用量"]).value):
            continue
        current_sequence = _numeric_cell_value(ws.cell(row=row_index, column=sequence_column).value)
        if current_sequence is not None and current_sequence > 0:
            sequence = current_sequence
            break
    for row_index in range(header_row + 1, ws.max_row + 1):
        if _text(ws.cell(row=row_index, column=columns["用量"]).value) == "替代料":
            ws.cell(row=row_index, column=sequence_column).value = None
        if not _is_numeric_quantity_value(ws.cell(row=row_index, column=columns["用量"]).value):
            continue
        ws.cell(row=row_index, column=sequence_column).value = sequence
        sequence += 1


def _set_inserted_sequence_values(
    ws: Worksheet,
    *,
    first_row: int,
    inserted_count: int,
    sequence_column: int | None,
) -> None:
    if sequence_column is None:
        return
    next_sequence = None
    for row_index in range(first_row + inserted_count, ws.max_row + 1):
        next_sequence = _numeric_cell_value(ws.cell(row=row_index, column=sequence_column).value)
        if next_sequence is not None:
            break
    if next_sequence is not None:
        start_sequence = next_sequence - inserted_count
    else:
        previous_sequence = 0
        for row_index in range(first_row - 1, 0, -1):
            previous_sequence = _numeric_cell_value(ws.cell(row=row_index, column=sequence_column).value) or 0
            if previous_sequence:
                break
        start_sequence = previous_sequence + 1
    for offset in range(inserted_count):
        ws.cell(row=first_row + offset, column=sequence_column).value = start_sequence + offset


def _subtract_quantity(value: Any, quantity_to_remove: Any) -> Any:
    try:
        current = Decimal(_quantity_compare_text(value))
        removed = Decimal(_quantity_compare_text(quantity_to_remove))
    except Exception:
        return value
    result = current - removed
    if result == result.to_integral_value():
        return int(result)
    return float(result.normalize())


def _set_bom_values(
    ws: Worksheet,
    row_index: int,
    columns: dict[str, int],
    material: MaterialInfo,
    change_row: ChangeRow,
    *,
    preserved_quantity: Any = None,
    preserved_locator: Any = None,
) -> None:
    quantity = preserved_quantity if change_row.after_quantity == UNCHANGED_QUANTITY_MARKER else format_quantity(change_row.after_quantity)
    if quantity == 0 and preserved_quantity not in (None, "", "/", 0, "0", "替代料"):
        quantity = preserved_quantity
    locator = preserved_locator if change_row.after_locator == UNCHANGED_LOCATOR_MARKER else _locator_bom_value(change_row.after_locator)
    values = {
        "物料编码": material.code,
        "物料类型": material.material_type,
        "名称": material.name,
        "规格型号": material.model,
        "规格描述": material.description,
        "封装": material.package,
        "用量": quantity,
        "位号": locator,
    }
    for header, value in values.items():
        if header in columns:
            ws.cell(row=row_index, column=columns[header]).value = value
    ws.cell(row=row_index, column=columns["用量"]).number_format = "General"
    for column in range(1, ws.max_column + 1):
        ws.cell(row=row_index, column=column).fill = copy(CHANGE_FILL)


def _can_split_locator_row(
    ws: Worksheet,
    *,
    row_index: int,
    columns: dict[str, int],
    group: ChangeGroup,
    before_rows: list[ChangeRow],
) -> bool:
    if group.locator_type != "位号" or len(before_rows) != 1:
        return False
    locators = set(_split_locators(ws.cell(row=row_index, column=columns["位号"]).value))
    target_locators = set(_split_locators(group.locator_value))
    return bool(target_locators) and target_locators.issubset(locators) and locators != target_locators


def _can_split_locator_material_group(
    ws: Worksheet,
    *,
    rows: list[int],
    header_row: int,
    columns: dict[str, int],
    group: ChangeGroup,
    before_rows: list[ChangeRow],
) -> bool:
    if group.locator_type != "位号" or not rows or not before_rows:
        return False
    target_locators = set(_split_locators(group.locator_value))
    if not target_locators:
        return False
    group_rows = _material_group_rows(
        ws,
        row_index=rows[0],
        header_row=header_row,
        columns=columns,
    )
    if not set(rows).issubset(set(group_rows)):
        return False
    primary_locators = set(_split_locators(ws.cell(row=group_rows[0], column=columns["位号"]).value))
    return target_locators.issubset(primary_locators) and primary_locators != target_locators


def _associated_alternative_rows(ws: Worksheet, *, row_index: int, columns: dict[str, int]) -> list[int]:
    rows: list[int] = []
    for candidate_row in range(row_index + 1, ws.max_row + 1):
        if _text(ws.cell(row=candidate_row, column=columns["用量"]).value) != "替代料":
            break
        rows.append(candidate_row)
    return rows


def _material_group_rows(ws: Worksheet, *, row_index: int, header_row: int, columns: dict[str, int]) -> list[int]:
    primary_row = row_index
    while (
        primary_row > header_row + 1
        and _text(ws.cell(row=primary_row, column=columns["用量"]).value) == "替代料"
    ):
        primary_row -= 1
    return [primary_row, *_associated_alternative_rows(ws, row_index=primary_row, columns=columns)]


def _synchronize_alternative_locators(ws: Worksheet, *, row_index: int, columns: dict[str, int]) -> None:
    primary_locators = ws.cell(row=row_index, column=columns["位号"]).value
    for alternative_row in _associated_alternative_rows(ws, row_index=row_index, columns=columns):
        ws.cell(row=alternative_row, column=columns["位号"]).value = primary_locators
        for column in range(1, ws.max_column + 1):
            ws.cell(row=alternative_row, column=column).fill = copy(CHANGE_FILL)


def _synchronize_all_alternative_locators(ws: Worksheet, *, header_row: int, columns: dict[str, int]) -> bool:
    changed = False
    for row_index in range(header_row + 1, ws.max_row + 1):
        if _text(ws.cell(row=row_index, column=columns["用量"]).value) == "替代料":
            continue
        primary_locators = ws.cell(row=row_index, column=columns["位号"]).value
        for alternative_row in _associated_alternative_rows(ws, row_index=row_index, columns=columns):
            if ws.cell(row=alternative_row, column=columns["位号"]).value == primary_locators:
                continue
            ws.cell(row=alternative_row, column=columns["位号"]).value = primary_locators
            for column in range(1, ws.max_column + 1):
                ws.cell(row=alternative_row, column=column).fill = copy(CHANGE_FILL)
            changed = True
    return changed


def _is_alternative_addition_change(
    before_rows: list[ChangeRow],
    after_rows: list[ChangeRow],
) -> bool:
    if len(before_rows) != 1 or len(after_rows) != 1:
        return False
    before = before_rows[0]
    after = after_rows[0]
    return (
        _looks_material_code(before.before_code)
        and _looks_material_code(after.after_code)
        and before.before_quantity == UNCHANGED_QUANTITY_MARKER
        and before.before_locator == UNCHANGED_LOCATOR_MARKER
        and after.after_quantity == "替代料"
        and after.after_locator == UNCHANGED_LOCATOR_MARKER
    )


def _alternative_after_unchanged_primary(
    before_rows: list[ChangeRow],
    after_rows: list[ChangeRow],
) -> ChangeRow | None:
    """Return the alternative row when an ECN leaves its primary material unchanged."""
    if len(before_rows) != 1 or len(after_rows) != 2:
        return None
    primary_before = before_rows[0]
    primary_after_candidates = [
        row for row in after_rows if row.after_code == primary_before.before_code
    ]
    if len(primary_after_candidates) != 1:
        return None
    primary_after = primary_after_candidates[0]
    if (
        _quantity_compare_text(primary_before.before_quantity)
        != _quantity_compare_text(primary_after.after_quantity)
        or _locator_compare_text(primary_before.before_locator)
        != _locator_compare_text(primary_after.after_locator)
    ):
        return None
    alternative = next(
        (row for row in after_rows if row is not primary_after),
        None,
    )
    if (
        alternative is None
        or alternative.after_code == primary_before.before_code
        or alternative.after_quantity != "\u66ff\u4ee3\u6599"
    ):
        return None
    return alternative


def _apply_partial_locator_change(
    ws: Worksheet,
    *,
    row_index: int,
    header_row: int,
    columns: dict[str, int],
    group: ChangeGroup,
    before_rows: list[ChangeRow],
    after_rows: list[ChangeRow],
    new_materials: list[MaterialInfo],
) -> None:
    captured_style = _capture_row_style(ws, row_index)
    target_locators = set(_split_locators(group.locator_value))
    group_rows = _material_group_rows(
        ws,
        row_index=row_index,
        header_row=header_row,
        columns=columns,
    )
    deleted_rows = _remove_locators_from_existing_material_group(
        ws,
        row_index=row_index,
        header_row=header_row,
        columns=columns,
        target_locators=target_locators,
        quantity_to_remove=before_rows[0].before_quantity,
    )
    remaining_group_rows = [
        row - sum(1 for deleted_row in deleted_rows if deleted_row < row)
        for row in group_rows
        if row not in deleted_rows
    ]
    insert_row = max(remaining_group_rows) + 1 if remaining_group_rows else min(deleted_rows or {row_index})
    for deleted_row in sorted(deleted_rows, reverse=True):
        ws.delete_rows(deleted_row, 1)

    if not new_materials:
        return
    ws.insert_rows(insert_row, len(new_materials))
    for offset, (change_row, material) in enumerate(zip(after_rows, new_materials, strict=True)):
        target_row = insert_row + offset
        _apply_row_style(ws, target_row, captured_style)
        if "序号" in columns:
            ws.cell(row=target_row, column=columns["序号"]).value = None
        _set_bom_values(ws, target_row, columns, material, change_row)


def _remove_locators_from_existing_material_group(
    ws: Worksheet,
    *,
    row_index: int,
    header_row: int,
    columns: dict[str, int],
    target_locators: set[str],
    quantity_to_remove: Any,
) -> set[int]:
    group_rows = _material_group_rows(
        ws,
        row_index=row_index,
        header_row=header_row,
        columns=columns,
    )
    primary_row = group_rows[0]
    deleted_rows: set[int] = set()
    primary_locators = _split_locators(ws.cell(row=primary_row, column=columns["位号"]).value)
    remaining_primary_locators = [locator for locator in primary_locators if locator not in target_locators]
    if remaining_primary_locators:
        removed_quantity = len(primary_locators) - len(remaining_primary_locators)
        ws.cell(row=primary_row, column=columns["位号"]).value = _join_locators(remaining_primary_locators)
        ws.cell(row=primary_row, column=columns["用量"]).value = _subtract_quantity(
            ws.cell(row=primary_row, column=columns["用量"]).value,
            removed_quantity or quantity_to_remove,
        )
        ws.cell(row=primary_row, column=columns["用量"]).number_format = "General"
        for column in range(1, ws.max_column + 1):
            ws.cell(row=primary_row, column=column).fill = copy(CHANGE_FILL)
    else:
        deleted_rows.add(primary_row)

    for alternative_row in group_rows[1:]:
        alternative_locators = _split_locators(ws.cell(row=alternative_row, column=columns["位号"]).value)
        remaining_alternative_locators = [locator for locator in alternative_locators if locator not in target_locators]
        if remaining_alternative_locators:
            ws.cell(row=alternative_row, column=columns["位号"]).value = _join_locators(remaining_alternative_locators)
            for column in range(1, ws.max_column + 1):
                ws.cell(row=alternative_row, column=column).fill = copy(CHANGE_FILL)
        else:
            deleted_rows.add(alternative_row)
    return deleted_rows


def _matched_rows(ws: Worksheet, header_row: int, columns: dict[str, int], group: ChangeGroup) -> list[int]:
    rows: list[int] = []
    if group.locator_type == "位号":
        before_codes = {row.before_code for row in group.rows if _looks_material_code(row.before_code)}
        target_locators = set(_split_locators(group.locator_value))
        for row_index in range(header_row + 1, ws.max_row + 1):
            row_locators = set(_split_locators(ws.cell(row=row_index, column=columns["位号"]).value))
            if before_codes:
                if _cell_text(ws, row_index, columns["物料编码"]) not in before_codes:
                    continue
                if target_locators and not target_locators.intersection(row_locators):
                    continue
            elif not target_locators.issubset(row_locators):
                continue
            rows.append(row_index)
    else:
        before_codes = {row.before_code for row in group.rows if _is_valid_before(row)}
        for row_index in range(header_row + 1, ws.max_row + 1):
            if _cell_text(ws, row_index, columns["物料编码"]) in before_codes:
                rows.append(row_index)
    return rows


def _matched_before_rows(
    ws: Worksheet,
    rows: list[int],
    columns: dict[str, int],
    before_rows: list[ChangeRow],
) -> list[ChangeRow]:
    remaining_by_code: dict[str, list[ChangeRow]] = {}
    for row in before_rows:
        remaining_by_code.setdefault(row.before_code, []).append(row)
    matched: list[ChangeRow] = []
    for row_index in rows:
        code = _cell_text(ws, row_index, columns["物料编码"])
        candidates = remaining_by_code.get(code, [])
        if candidates:
            matched.append(candidates.pop(0))
    return matched


def _matched_after_rows(
    ws: Worksheet,
    header_row: int,
    columns: dict[str, int],
    after_rows: list[ChangeRow],
) -> list[int]:
    matched: list[int] = []
    consumed_rows: set[int] = set()
    for change_row in after_rows:
        target_locators = set(_split_locators(change_row.after_locator))
        for row_index in range(header_row + 1, ws.max_row + 1):
            if row_index in consumed_rows:
                continue
            if _cell_text(ws, row_index, columns["物料编码"]) != change_row.after_code:
                continue
            current_locators = set(_split_locators(ws.cell(row=row_index, column=columns["位号"]).value))
            if target_locators and not target_locators.issubset(current_locators):
                continue
            matched.append(row_index)
            consumed_rows.add(row_index)
            break
    return matched


def _find_existing_addition_target_row(
    ws: Worksheet,
    *,
    header_row: int,
    columns: dict[str, int],
    after_rows: list[ChangeRow],
    excluded_rows: set[int] | None = None,
) -> int | None:
    if len(after_rows) != 1:
        return None
    target_code = after_rows[0].after_code
    for row_index in range(header_row + 1, ws.max_row + 1):
        if excluded_rows is not None and row_index in excluded_rows:
            continue
        if _cell_text(ws, row_index, columns["物料编码"]) != target_code:
            continue
        if _is_numeric_quantity_value(ws.cell(row=row_index, column=columns["用量"]).value):
            return row_index
        primary_row = row_index
        while (
            primary_row > header_row + 1
            and _text(ws.cell(row=primary_row, column=columns["用量"]).value) == "替代料"
        ):
            primary_row -= 1
        if _is_numeric_quantity_value(ws.cell(row=primary_row, column=columns["用量"]).value):
            return primary_row
    return None


def _existing_numeric_material_code_rows(
    ws: Worksheet,
    *,
    header_row: int,
    columns: dict[str, int],
    material_code: str,
) -> list[int]:
    return [
        row_index
        for row_index in range(header_row + 1, ws.max_row + 1)
        if _cell_text(ws, row_index, columns["物料编码"]) == material_code
        and _is_numeric_quantity_value(ws.cell(row=row_index, column=columns["用量"]).value)
    ]


def _add_numeric_quantity(current: Any, addition: Any) -> int | float | None:
    try:
        current_decimal = Decimal(_quantity_compare_text(current))
        addition_decimal = Decimal(_quantity_compare_text(addition))
    except Exception:
        return None
    result = current_decimal + addition_decimal
    if result == result.to_integral_value():
        return int(result)
    return float(result.normalize())


def _consolidate_duplicate_material_rows(
    ws: Worksheet,
    *,
    row_indexes: list[int],
    columns: dict[str, int],
) -> tuple[int, int | float, str] | None:
    quantities: list[tuple[int, Decimal]] = []
    for row_index in row_indexes:
        try:
            quantities.append(
                (
                    row_index,
                    Decimal(_quantity_compare_text(ws.cell(row=row_index, column=columns["用量"]).value)),
                )
            )
        except Exception:
            return None
    if len(quantities) < 2:
        return None
    keeper = max(quantities, key=lambda item: item[1])[0]
    total = sum((quantity for _, quantity in quantities), Decimal("0"))
    locators: list[str] = []
    for row_index in [keeper, *[index for index in row_indexes if index != keeper]]:
        for locator in _split_locators(ws.cell(row=row_index, column=columns["位号"]).value):
            if locator not in locators:
                locators.append(locator)
    for row_index in sorted((index for index in row_indexes if index != keeper), reverse=True):
        ws.delete_rows(row_index, 1)
        if row_index < keeper:
            keeper -= 1
    merged_quantity: int | float
    if total == total.to_integral_value():
        merged_quantity = int(total)
    else:
        merged_quantity = float(total.normalize())
    ws.cell(row=keeper, column=columns["用量"]).value = merged_quantity
    ws.cell(row=keeper, column=columns["用量"]).number_format = "General"
    ws.cell(row=keeper, column=columns["位号"]).value = _join_locators(locators)
    for column in range(1, ws.max_column + 1):
        ws.cell(row=keeper, column=column).fill = copy(CHANGE_FILL)
    _synchronize_alternative_locators(ws, row_index=keeper, columns=columns)
    return keeper, merged_quantity, _join_locators(locators)


def validate_processed_bom_workbook(
    *,
    bom_path: Path,
    groups: list[ChangeGroup],
    successful_report_rows: list[ReportRow],
) -> list[ValidationIssue]:
    """Find old material codes that remain at locations successfully processed in this run."""
    successful_replacements = [
        row for row in successful_report_rows if row.action == "删除" and row.result == "成功" and row.before_code
    ]
    if not successful_replacements:
        return []

    wb = load_workbook(bom_path, read_only=True, data_only=True)
    issues: list[ValidationIssue] = []
    try:
        for sheet_name in {row.sheet_name for row in successful_replacements}:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            header = _find_header_row(ws, BOM_HEADERS)
            if header is None:
                continue
            header_row, required_columns = header
            columns = _add_optional_columns(ws, header_row, required_columns)
            for replacement in (row for row in successful_replacements if row.sheet_name == sheet_name):
                for row_index in range(header_row + 1, ws.max_row + 1):
                    before_code = _cell_text(ws, row_index, columns["物料编码"])
                    locator_value = _cell_text(ws, row_index, columns["位号"])
                    if before_code != replacement.before_code:
                        continue
                    if replacement.before_locator not in {"", "/", UNCHANGED_LOCATOR_MARKER}:
                        expected_locators = set(_split_locators(replacement.before_locator))
                        if not expected_locators.intersection(_split_locators(locator_value)):
                            continue
                    issues.append(
                        ValidationIssue(
                            sheet_name=sheet_name,
                            before_code=before_code,
                            after_code=next(
                                (
                                    group_row.after_code
                                    for group in groups
                                    for group_row in group.rows
                                    if group_row.before_code == before_code and _is_valid_after(group_row)
                                ),
                                "",
                            ),
                            locator_value=locator_value,
                            reason="变更前物料仍存在",
                        )
                    )
    finally:
        wb.close()
    return issues


def validate_processed_bom_worksheets(
    *,
    wb: Any,
    groups: list[ChangeGroup],
    successful_report_rows: list[ReportRow],
) -> list[ValidationIssue]:
    """Validate the in-memory BOM before it is saved."""
    successful_replacements = [
        row for row in successful_report_rows if row.action == "删除" and row.result == "成功" and row.before_code
    ]
    if not successful_replacements:
        return []
    issues: list[ValidationIssue] = []
    for sheet_name in {row.sheet_name for row in successful_replacements}:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = _find_header_row(ws, BOM_HEADERS)
        if header is None:
            continue
        header_row, required_columns = header
        columns = _add_optional_columns(ws, header_row, required_columns)
        for replacement in (row for row in successful_replacements if row.sheet_name == sheet_name):
            for row_index in range(header_row + 1, ws.max_row + 1):
                before_code = _cell_text(ws, row_index, columns["物料编码"])
                locator_value = _cell_text(ws, row_index, columns["位号"])
                if before_code != replacement.before_code:
                    continue
                if replacement.before_locator not in {"", "/", UNCHANGED_LOCATOR_MARKER}:
                    if not set(_split_locators(replacement.before_locator)).intersection(_split_locators(locator_value)):
                        continue
                issues.append(
                    ValidationIssue(
                        sheet_name=sheet_name,
                        before_code=before_code,
                        after_code=next(
                            (
                                group_row.after_code
                                for group in groups
                                for group_row in group.rows
                                if group_row.before_code == before_code and _is_valid_after(group_row)
                            ),
                            "",
                        ),
                        locator_value=locator_value,
                        reason="变更前物料仍存在",
                    )
                )
    return issues


def _material_code_sort_key(value: str) -> tuple[int, ...] | None:
    if not _looks_material_code(value):
        return None
    try:
        return tuple(int(part) for part in value.split("."))
    except ValueError:
        return None


def _find_add_only_insert_row(
    ws: Worksheet,
    *,
    header_row: int,
    columns: dict[str, int],
    after_rows: list[ChangeRow],
    new_materials: list[MaterialInfo],
) -> int:
    target_keys = [_material_code_sort_key(row.after_code) for row in after_rows]
    target_keys = [key for key in target_keys if key is not None]
    target_key = min(target_keys) if target_keys else None
    target_types = [material.material_type for material in new_materials if material.material_type]
    target_type = target_types[0] if target_types else ""
    if target_type and "物料类型" in columns:
        packaging_code_prefixes = ("4.03.", "4.04.", "4.05.")
        is_auxiliary_or_packaging = target_type in {"辅料", "包装料"} or any(
            row.after_code.startswith(packaging_code_prefixes) for row in after_rows
        )
        if is_auxiliary_or_packaging:
            last_data_row = header_row
            for row_index in range(header_row + 1, ws.max_row + 1):
                if _cell_text(ws, row_index, columns["物料编码"]):
                    last_data_row = row_index
            return last_data_row + 1
        last_same_type_row = None
        for row_index in range(header_row + 1, ws.max_row + 1):
            material_type = _cell_text(ws, row_index, columns["物料类型"])
            if material_type != target_type:
                continue
            last_same_type_row = row_index
            if not _is_numeric_quantity_value(ws.cell(row=row_index, column=columns["用量"]).value):
                continue
            current_key = _material_code_sort_key(_cell_text(ws, row_index, columns["物料编码"]))
            if target_key is not None and current_key is not None and current_key > target_key:
                return row_index
        if last_same_type_row is not None:
            return last_same_type_row + 1

    last_data_row = header_row + 1
    for row_index in range(header_row + 1, ws.max_row + 1):
        code = _cell_text(ws, row_index, columns["物料编码"])
        if code:
            last_data_row = row_index
        if not _is_numeric_quantity_value(ws.cell(row=row_index, column=columns["用量"]).value):
            continue
        current_key = _material_code_sort_key(code)
        if target_key is not None and current_key is not None and current_key > target_key:
            return row_index
    return last_data_row + 1


def _sheet_segment(title: str, sheet_titles: list[str] | None = None) -> str:
    normalized = title.strip().upper()
    if title == REPORT_SHEET_NAME or title == CHANGE_RECORD_SHEET_NAME:
        return "other"
    if sheet_titles and title == sheet_titles[0]:
        return "total"
    if title == "总表":
        return "total"
    if normalized.startswith("SMT"):
        return "smt"
    if normalized.startswith("DIP"):
        return "dip"
    return "product"


def _allowed_segments_for_material(*, material_type: str, name: str) -> set[str] | None:
    if material_type in {"辅料", "包装料"}:
        return {"total", "product"}
    if material_type == "贴片料":
        return {"total", "smt"}
    if material_type == "插件料":
        return {"total", "dip"}
    if not material_type:
        if name == "保护板":
            return {"total", "product"}
        if name in {"PCB", "钢网"}:
            return {"total", "smt"}
    return None


def _allowed_segments_for_change(
    ws: Worksheet,
    *,
    rows: list[int],
    columns: dict[str, int],
    new_materials: list[MaterialInfo],
) -> set[str] | None:
    allowed_sets: list[set[str]] = []
    for material in new_materials:
        allowed = _allowed_segments_for_material(material_type=material.material_type, name=material.name)
        if allowed is not None:
            allowed_sets.append(allowed)
    if not allowed_sets and rows:
        for row_index in rows:
            allowed = _allowed_segments_for_material(
                material_type=_cell_text(ws, row_index, columns["物料类型"]) if "物料类型" in columns else "",
                name=_cell_text(ws, row_index, columns["名称"]) if "名称" in columns else "",
            )
            if allowed is not None:
                allowed_sets.append(allowed)
    if not allowed_sets:
        return None
    result = set().union(*allowed_sets)
    result.add("total")
    return result


def _has_allowed_segment_match(wb: Any, group: ChangeGroup, allowed_segments: set[str]) -> bool:
    for ws in wb.worksheets:
        segment = _sheet_segment(ws.title, wb.sheetnames)
        if segment not in allowed_segments:
            continue
        header = _find_header_row(ws, BOM_HEADERS)
        if header is None:
            continue
        header_row, columns = header
        if _matched_rows(ws, header_row, columns, group):
            return True
    return False


def _is_contiguous(rows: list[int]) -> bool:
    return not rows or rows == list(range(min(rows), max(rows) + 1))


def _collapse_repeated_after_rows(rows: list[ChangeRow]) -> list[ChangeRow]:
    if len(rows) < 2:
        return rows
    signatures = {
        (row.after_code, row.after_model, row.after_description, row.after_quantity, row.after_locator)
        for row in rows
    }
    return rows[:1] if len(signatures) == 1 else rows


def _operation_after_signature(operation: PendingOperation) -> tuple[str, str, str, str, str] | None:
    if operation.is_partial_locator_change or len(operation.after_rows) != 1 or len(operation.new_materials) != 1:
        return None
    row = operation.after_rows[0]
    return (
        row.after_code,
        row.after_model,
        row.after_description,
        row.after_quantity,
        row.after_locator,
    )


def _operation_locator_signature(
    ws: Worksheet,
    columns: dict[str, int],
    operation: PendingOperation,
) -> tuple[str, ...] | None:
    signatures: list[tuple[str, ...]] = []
    for row_index in operation.rows:
        locators = _split_locators(ws.cell(row=row_index, column=columns["位号"]).value)
        if locators:
            signatures.append(tuple(locators))
    if not signatures:
        return None
    first = signatures[0]
    if any(signature != first for signature in signatures[1:]):
        return None
    return first


def _can_merge_same_target_operations(
    ws: Worksheet,
    columns: dict[str, int],
    left: PendingOperation,
    right: PendingOperation,
) -> bool:
    left_signature = _operation_after_signature(left)
    if left_signature is None or left_signature != _operation_after_signature(right):
        return False
    if not left.rows or not right.rows:
        return False
    left_locators = _operation_locator_signature(ws, columns, left)
    if left_locators is not None and left_locators == _operation_locator_signature(ws, columns, right):
        gap_rows = range(max(left.rows) + 1, min(right.rows))
        return all(_target_row_matches_operation(ws, columns, row_index, left) for row_index in gap_rows)
    if left.group.locator_type != "位号" or right.group.locator_type != "位号":
        return False
    target_locators = set(_split_locators(left.after_rows[0].after_locator))
    return bool(target_locators) and (
        set(_split_locators(left.group.locator_value)).issubset(target_locators)
        and set(_split_locators(right.group.locator_value)).issubset(target_locators)
    )


def _target_row_matches_operation(
    ws: Worksheet,
    columns: dict[str, int],
    row_index: int,
    operation: PendingOperation,
) -> bool:
    if len(operation.after_rows) != 1:
        return False
    if _cell_text(ws, row_index, columns["物料编码"]) != operation.after_rows[0].after_code:
        return False
    operation_locators = _operation_locator_signature(ws, columns, operation)
    if operation_locators is None:
        return False
    return tuple(_split_locators(ws.cell(row=row_index, column=columns["位号"]).value)) == operation_locators


def _target_gap_rows(
    ws: Worksheet,
    columns: dict[str, int],
    left: PendingOperation,
    right: PendingOperation,
) -> list[int]:
    return [
        row_index
        for row_index in range(max(left.rows) + 1, min(right.rows))
        if _target_row_matches_operation(ws, columns, row_index, left)
    ]


def _expand_operation_with_existing_target_rows(
    ws: Worksheet,
    header_row: int,
    columns: dict[str, int],
    operation: PendingOperation,
) -> PendingOperation:
    if _operation_after_signature(operation) is None or not operation.rows:
        return operation
    rows = set(operation.rows)
    start = min(rows)
    while start > header_row + 1 and _target_row_matches_operation(ws, columns, start - 1, operation):
        start -= 1
        rows.add(start)
    end = max(rows)
    while end < ws.max_row and _target_row_matches_operation(ws, columns, end + 1, operation):
        end += 1
        rows.add(end)
    return PendingOperation(
        first_row=min(rows),
        rows=sorted(rows),
        group=operation.group,
        before_rows=operation.before_rows,
        after_rows=operation.after_rows,
        new_materials=operation.new_materials,
        is_partial_locator_change=operation.is_partial_locator_change,
    )


def _merge_same_target_operations(
    ws: Worksheet,
    header_row: int,
    columns: dict[str, int],
    operations: list[PendingOperation],
) -> list[PendingOperation]:
    merged: list[PendingOperation] = []
    for operation in sorted(operations, key=lambda item: item.first_row):
        if merged and _can_merge_same_target_operations(ws, columns, merged[-1], operation):
            previous = merged[-1]
            merged[-1] = PendingOperation(
                first_row=previous.first_row,
                rows=[*previous.rows, *_target_gap_rows(ws, columns, previous, operation), *operation.rows],
                group=previous.group,
                before_rows=[*previous.before_rows, *operation.before_rows],
                after_rows=previous.after_rows,
                new_materials=previous.new_materials,
                is_partial_locator_change=False,
            )
        else:
            merged.append(operation)
    return [_expand_operation_with_existing_target_rows(ws, header_row, columns, operation) for operation in merged]


def _append_abnormal(
    report_rows: list[ReportRow],
    *,
    sheet_name: str,
    group: ChangeGroup,
    reason: str,
) -> None:
    report_rows.append(
        ReportRow(
            sheet_name=sheet_name,
            locator_type=group.locator_type,
            locator_value=group.locator_value,
            action="异常",
            result="未修改",
            reason=reason,
        )
    )


def _append_report_for_change(
    report_rows: list[ReportRow],
    *,
    sheet_name: str,
    group: ChangeGroup,
    before_rows: list[ChangeRow],
    after_rows: list[ChangeRow],
    before_locators: list[Any] | None = None,
    before_quantities: list[Any] | None = None,
    after_locator: Any = None,
    after_quantity: Any = None,
) -> None:
    for index, row in enumerate(before_rows):
        before_locator = before_locators[index] if before_locators is not None and index < len(before_locators) else row.before_locator
        before_quantity = before_quantities[index] if before_quantities is not None and index < len(before_quantities) else row.before_quantity
        report_rows.append(
            ReportRow(
                sheet_name=sheet_name,
                locator_type=group.locator_type,
                locator_value=group.locator_value,
                action="删除",
                before_code=row.before_code,
                before_model=row.before_model,
                before_description=row.before_description,
                before_quantity=_text(format_quantity(before_quantity)),
                before_locator=_locator_compare_text(before_locator),
                result="成功",
            )
        )
    for row in after_rows:
        report_after_quantity = after_quantity if after_quantity is not None else row.after_quantity
        report_after_locator = after_locator if after_locator is not None else row.after_locator
        report_rows.append(
            ReportRow(
                sheet_name=sheet_name,
                locator_type=group.locator_type,
                locator_value=group.locator_value,
                action="添加",
                after_code=row.after_code,
                after_model=row.after_model,
                after_description=row.after_description,
                after_quantity=_text(format_quantity(report_after_quantity)),
                after_locator=_locator_compare_text(report_after_locator),
                result="成功",
            )
        )


def _sheet_type_label(sheet_name: str, sheet_titles: list[str]) -> str:
    segment = _sheet_segment(sheet_name, sheet_titles)
    labels = {
        "total": "总表",
        "smt": "SMT段",
        "dip": "DIP段",
        "product": "成品段",
    }
    return labels.get(segment, "其他")


def _checklist_status_from_report_row(row: ReportRow) -> str:
    if row.action == "异常":
        return "异常"
    if row.result == "成功":
        return "已完成"
    if row.result == "未修改":
        return "未更改"
    return row.result or "未更改"


def _can_merge_report_rows(before: ReportRow, after: ReportRow) -> bool:
    return (
        before.sheet_name == after.sheet_name
        and before.locator_value == after.locator_value
        and before.result == "成功"
        and after.result == "成功"
        and before.action == "删除"
        and after.action == "添加"
    )


def _checklist_rows_from_report_rows(report_rows: list[ReportRow], sheet_titles: list[str]) -> list[ChecklistRow]:
    rows: list[ChecklistRow] = []
    index = 0
    while index < len(report_rows):
        report = report_rows[index]
        if index + 1 < len(report_rows) and _can_merge_report_rows(report, report_rows[index + 1]):
            after = report_rows[index + 1]
            rows.append(
                ChecklistRow(
                    sheet_name=report.sheet_name,
                    sheet_type=_sheet_type_label(report.sheet_name, sheet_titles),
                    locator_value=report.locator_value,
                    action=report.action,
                    before_code=report.before_code,
                    before_model=report.before_model,
                    before_quantity=report.before_quantity,
                    before_locator=report.before_locator,
                    after_code=after.after_code,
                    after_model=after.after_model,
                    after_quantity=after.after_quantity,
                    after_locator=after.after_locator,
                    after_action=after.action,
                    status="已完成",
                )
            )
            index += 2
            continue
        rows.append(
            ChecklistRow(
                sheet_name=report.sheet_name,
                sheet_type=_sheet_type_label(report.sheet_name, sheet_titles),
                locator_value=report.locator_value,
                action=report.action,
                before_code=report.before_code,
                before_model=report.before_model,
                before_quantity=report.before_quantity,
                before_locator=report.before_locator,
                after_code=report.after_code,
                after_model=report.after_model,
                after_quantity=report.after_quantity,
                after_locator=report.after_locator,
                after_action="添加" if report.action == "添加" else "",
                status=_checklist_status_from_report_row(report),
                reason=report.reason,
            )
        )
        index += 1
    return rows


def _write_report_sheet(wb: Any, report_rows: list[ReportRow]) -> None:
    if REPORT_SHEET_NAME in wb.sheetnames:
        del wb[REPORT_SHEET_NAME]
    ws = wb.create_sheet(REPORT_SHEET_NAME)
    headers = [
        "Sheet",
        "定位类型",
        "定位值",
        "动作",
        "变更前物料编码",
        "变更前规格型号",
        "变更前规格描述",
        "变更前用量",
        "变更前位号",
        "变更后物料编码",
        "变更后规格型号",
        "变更后规格描述",
        "变更后用量",
        "变更后位号",
        "结果",
        "原因",
    ]
    ws.append(headers)
    for report in report_rows:
        ws.append([
            report.sheet_name,
            report.locator_type,
            report.locator_value,
            report.action,
            report.before_code,
            report.before_model,
            report.before_description,
            report.before_quantity,
            report.before_locator,
            report.after_code,
            report.after_model,
            report.after_description,
            report.after_quantity,
            report.after_locator,
            report.result,
            report.reason,
        ])
    for index, width in enumerate([18, 12, 16, 10, 18, 18, 28, 12, 14, 18, 18, 28, 12, 14, 10, 40], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _extract_version_from_bom_name(bom_path: Path) -> str:
    rev_match = re.search(r"REV[._\s-]*(\d+(?:\.\d+)*)", bom_path.stem, re.IGNORECASE)
    if rev_match:
        return rev_match.group(1)
    matches = re.findall(r"-(\d+(?:\.\d+)*)", bom_path.stem)
    dotted_matches = [match for match in matches if "." in match]
    if dotted_matches:
        return dotted_matches[-1]
    return matches[-1] if matches else ""


def _same_material_quantity_change_texts(
    pairs: list[tuple[str, str, str, str]],
) -> list[str]:
    """Describe in-place quantity changes for the human-readable change record.

    An ECN may keep the same material code and only raise/lower its quantity.
    The record should describe the delta instead of using the generic
    "increase packaging material" wording.
    """
    texts: list[str] = []
    seen: set[str] = set()
    for before_code, before_quantity, after_code, after_quantity in pairs:
        if before_code != after_code or not before_code:
            continue
        try:
            delta = Decimal(_quantity_compare_text(after_quantity)) - Decimal(
                _quantity_compare_text(before_quantity)
            )
        except Exception:
            continue
        if delta == 0:
            continue
        action = "增加" if delta > 0 else "减少"
        text = f"{action}{after_code}的用量{format_quantity(abs(delta))}"
        if text not in seen:
            seen.add(text)
            texts.append(text)
    return texts


def _change_record_text(groups: list[ChangeGroup]) -> str:
    after_codes = [
        row.after_code
        for group in groups
        for row in group.rows
        if _is_valid_after(row)
    ]
    quantity_texts = _same_material_quantity_change_texts(
        [
            (row.before_code, row.before_quantity, row.after_code, row.after_quantity)
            for group in groups
            for row in group.rows
            if _is_valid_before(row)
            and _is_valid_after(row)
            and _locator_compare_text(row.before_locator)
            == _locator_compare_text(row.after_locator)
        ]
    )
    if quantity_texts:
        return "；".join(quantity_texts)
    if after_codes and all(code.startswith(("4.03.", "4.05.")) for code in after_codes):
        return "增加包材"
    categorized: dict[str, list[str]] = {"更新": [], "添加": [], "删除": []}
    seen: set[tuple[str, str, str]] = set()
    for group in groups:
        for row in group.rows:
            has_before = _is_valid_before(row)
            has_after = _is_valid_after(row)
            if not has_before and not has_after:
                continue
            locator = _locator_compare_text(row.after_locator)
            if locator == "/":
                locator = _locator_compare_text(row.before_locator)
            action = "更新" if has_before and has_after else "添加" if has_after else "删除"
            code = row.after_code if has_after else row.before_code
            item_text = f"{locator}={code}" if locator != "/" else code
            item = (action, item_text, code)
            if item in seen:
                continue
            seen.add(item)
            categorized[action].append(item_text)
    action_labels = {"更新": "更新", "添加": "添加物料", "删除": "删除"}
    parts = [f"{action_labels[action]}{'、'.join(items)}" for action, items in categorized.items() if items]
    if parts:
        return "；".join(parts)
    return f"更新{'、'.join(after_codes)}" if after_codes else ""


def _change_record_text_from_report_rows(groups: list[ChangeGroup], report_rows: list[ReportRow]) -> str:
    after_codes = [
        row.after_code
        for group in groups
        for row in group.rows
        if _is_valid_after(row)
    ]
    # The report intentionally preserves the old quantity when replacing a
    # one-row material in place. Therefore the authoritative quantity delta
    # for same-code ECNs must come from the parsed ECN, not from report rows.
    quantity_texts = _same_material_quantity_change_texts(
        [
            (row.before_code, row.before_quantity, row.after_code, row.after_quantity)
            for group in groups
            for row in group.rows
            if _is_valid_before(row)
            and _is_valid_after(row)
            and _locator_compare_text(row.before_locator)
            == _locator_compare_text(row.after_locator)
        ]
    )
    if quantity_texts:
        return "；".join(quantity_texts)
    successful_rows = [row for row in report_rows if row.result == "成功"]
    if after_codes and all(code.startswith(("4.03.", "4.05.")) for code in after_codes):
        return "增加包材"
    categorized: dict[str, list[str]] = {"添加": [], "删除": []}
    seen: set[tuple[str, str]] = set()
    for row in successful_rows:
        if row.action == "添加" and row.after_code:
            locator = _locator_compare_text(row.after_locator)
            action = "添加"
            item_text = f"{locator}={row.after_code}" if locator != "/" else row.after_code
        elif row.action == "删除" and row.before_code:
            action = "删除"
            locator = _locator_compare_text(row.before_locator)
            item_text = f"{locator}={row.before_code}" if locator != "/" else row.before_code
        else:
            continue
        item = (action, item_text)
        if item in seen:
            continue
        seen.add(item)
        categorized[action].append(item_text)
    order: dict[str, int] = {}
    for index, group in enumerate(groups):
        for row in group.rows:
            locator = _locator_compare_text(row.after_locator)
            if locator == "/":
                locator = _locator_compare_text(row.before_locator)
            if locator != "/":
                order.setdefault(locator, index)
    for items in categorized.values():
        items.sort(key=lambda item: min((order.get(locator, len(order)) for locator in _split_locators(item.split("=", 1)[0])), default=len(order)))
    action_labels = {"添加": "添加物料", "删除": "删除"}
    parts = [f"{action_labels[action]}{'、'.join(items)}" for action, items in categorized.items() if items]
    if parts:
        return "；".join(parts)
    return ""


def _last_change_record_row(ws: Worksheet) -> int:
    last_row = 1
    for row in range(2, ws.max_row + 1):
        values = [_cell_text(ws, row, column) for column in range(1, 7)]
        if any(values):
            last_row = row
    return last_row


def _next_writable_change_record_row(ws: Worksheet, last_row: int) -> int:
    target_row = last_row + 1
    while any(isinstance(ws.cell(row=target_row, column=column), MergedCell) for column in range(1, 7)):
        target_row += 1
    return target_row


def _copy_record_row_style(ws: Worksheet, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for column in range(1, 7):
        source = ws.cell(row=source_row, column=column)
        target = ws.cell(row=target_row, column=column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.border = copy(source.border)
        target.fill = copy(source.fill)


def _append_change_record_sheet(
    wb: Any,
    *,
    bom_path: Path,
    groups: list[ChangeGroup],
    report_rows: list[ReportRow],
    record_date: date,
    people: ChangeRecordPeople,
) -> None:
    if CHANGE_RECORD_SHEET_NAME not in wb.sheetnames:
        return
    ws = wb[CHANGE_RECORD_SHEET_NAME]
    last_row = _last_change_record_row(ws)
    target_row = _next_writable_change_record_row(ws, last_row)
    _copy_record_row_style(ws, last_row, target_row)
    ws.cell(row=target_row, column=1).value = record_date
    ws.cell(row=target_row, column=2).value = _extract_version_from_bom_name(bom_path)
    ws.cell(row=target_row, column=3).value = _change_record_text_from_report_rows(groups, report_rows)
    ws.cell(row=target_row, column=4).value = people.preparer
    ws.cell(row=target_row, column=5).value = people.reviewer
    ws.cell(row=target_row, column=6).value = people.approver


def process_bom_workbook(
    *,
    change_notice_path: Path,
    bom_path: Path,
    material_master_path: Path,
    output_path: Path,
    change_groups: list[ChangeGroup] | None = None,
    materials: dict[str, list[MaterialInfo]] | None = None,
    change_record_people: ChangeRecordPeople | None = None,
    change_record_date: date | None = None,
    include_report_sheet: bool = True,
) -> ProcessResult:
    groups = change_groups if change_groups is not None else parse_change_notice(change_notice_path)
    materials = materials if materials is not None else load_material_master(material_master_path)
    wb = load_workbook(bom_path)
    report_rows: list[ReportRow] = []
    processed_sheet_count = 0
    processed_sheet_titles: list[str] = []
    handled_group_keys: set[str] = set()
    inherited_columns: dict[str, int] | None = None

    for ws in wb.worksheets:
        if ws.title == REPORT_SHEET_NAME:
            continue
        segment = _sheet_segment(ws.title, wb.sheetnames)
        if segment == "other":
            continue
        header = _find_header_row(ws, BOM_HEADERS)
        if header is None:
            if inherited_columns is None:
                continue
            header_row = 0
            columns = dict(inherited_columns)
        else:
            header_row, required_columns = header
            columns = _add_optional_columns(ws, header_row, required_columns)
            if segment == "total":
                inherited_columns = dict(columns)
        processed_sheet_count += 1
        processed_sheet_titles.append(ws.title)
        operations: list[PendingOperation] = []
        sheet_modified = False
        deletion_rows: set[int] = set()
        for group in groups:
            if any(_is_valid_after(row) for row in group.rows):
                continue
            for row_index in _matched_rows(ws, header_row, columns, group):
                deletion_rows.update(
                    _material_group_rows(
                        ws,
                        row_index=row_index,
                        header_row=header_row,
                        columns=columns,
                    )
                )

        for group in groups:
            rows = _matched_rows(ws, header_row, columns, group)
            after_rows = _collapse_repeated_after_rows([row for row in group.rows if _is_valid_after(row)])
            target_codes = {
                row.after_code for row in after_rows if _looks_material_code(row.after_code)
            }
            change_before_rows = [
                row
                for row in group.rows
                if _is_valid_before(row)
                or (group.locator_type == "位号" and _looks_material_code(row.before_code))
            ]
            # In replacement ECNs the target may also be repeated in the
            # before-side list. It is not an old material and must not appear
            # in the deletion/validation rows.
            if group.locator_type == "位号" and target_codes:
                change_before_rows = [
                    row for row in change_before_rows if row.before_code not in target_codes
                ]
            before_rows = (
                _matched_before_rows(ws, rows, columns, change_before_rows) if group.locator_type == "位号" else change_before_rows
            )
            existing_after_rows = _matched_after_rows(ws, header_row, columns, after_rows)
            alternative_for_unchanged_primary = _alternative_after_unchanged_primary(
                [row for row in group.rows if _is_valid_before(row)],
                after_rows,
            )
            if alternative_for_unchanged_primary is not None and len(rows) == 1:
                primary_locator = _cell_text(ws, rows[0], columns["\u4f4d\u53f7"])
                after = alternative_for_unchanged_primary
                after_rows = [ChangeRow(
                    before_code=after.before_code,
                    before_model=after.before_model,
                    before_description=after.before_description,
                    before_quantity=after.before_quantity,
                    before_locator=after.before_locator,
                    after_code=after.after_code,
                    after_model=after.after_model,
                    after_description=after.after_description,
                    after_quantity=after.after_quantity,
                    after_locator=primary_locator,
                )]
                existing_after_rows = _matched_after_rows(ws, header_row, columns, after_rows)
                rows = []
                before_rows = []
            is_alternative_addition = (
                alternative_for_unchanged_primary is not None
                or _is_alternative_addition_change(before_rows, after_rows)
            )
            if is_alternative_addition and len(rows) == 1:
                primary_locator = _cell_text(ws, rows[0], columns["位号"])
                after = after_rows[0]
                after_rows = [ChangeRow(
                    before_code=after.before_code,
                    before_model=after.before_model,
                    before_description=after.before_description,
                    before_quantity=after.before_quantity,
                    before_locator=after.before_locator,
                    after_code=after.after_code,
                    after_model=after.after_model,
                    after_description=after.after_description,
                    after_quantity="替代料",
                    after_locator=primary_locator,
                )]
                rows = []
                before_rows = []
            # Some ECNs repeat the target material in the before-side list. If
            # the BOM already contains only that target at the affected
            # locator, this is an already-applied change and must be a no-op.
            if group.locator_type == "位号" and after_rows and rows:
                target_codes = {
                    row.after_code for row in after_rows if _looks_material_code(row.after_code)
                }
                row_codes = {
                    _cell_text(ws, row_index, columns["物料编码"])
                    for row_index in rows
                }
                target_locators = set(_split_locators(group.locator_value))
                if (
                    target_codes
                    and len(rows) == 1
                    and len(
                        _existing_numeric_material_code_rows(
                            ws,
                            header_row=header_row,
                            columns=columns,
                            material_code=next(iter(target_codes)),
                        )
                    ) == 1
                    and row_codes.issubset(target_codes)
                    and existing_after_rows
                    and all(
                        target_locators.issubset(
                            set(_split_locators(ws.cell(row=row_index, column=columns["位号"]).value))
                        )
                        for row_index in existing_after_rows
                    )
                ):
                    continue
            is_add_only = is_alternative_addition or (not change_before_rows and bool(after_rows))
            duplicate_target_rows = (
                _existing_numeric_material_code_rows(
                    ws,
                    header_row=header_row,
                    columns=columns,
                    material_code=after_rows[0].after_code,
                )
                if is_add_only and len(after_rows) == 1
                else []
            )
            if len(duplicate_target_rows) > 1:
                allowed_segments = _allowed_segments_for_change(
                    ws,
                    rows=[duplicate_target_rows[0]],
                    columns=columns,
                    new_materials=[],
                )
                if allowed_segments is not None and segment not in allowed_segments:
                    continue
                merged = _consolidate_duplicate_material_rows(
                    ws,
                    row_indexes=duplicate_target_rows,
                    columns=columns,
                )
                if merged is None:
                    _append_abnormal(
                        report_rows,
                        sheet_name=ws.title,
                        group=group,
                        reason="原BOM同料号重复行的用量不是数字，未自动合并",
                    )
                    continue
                _, merged_quantity, merged_locator = merged
                _append_report_for_change(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    before_rows=[],
                    after_rows=after_rows,
                    after_locator=merged_locator,
                    after_quantity=merged_quantity,
                )
                handled_group_keys.add(group.key)
                sheet_modified = True
                continue
            if (
                is_add_only
                and len(existing_after_rows) == len(after_rows)
                and all(_locator_compare_text(row.after_locator) != "/" for row in after_rows)
            ):
                continue
            existing_target_row = (
                _find_existing_addition_target_row(
                    ws,
                    header_row=header_row,
                    columns=columns,
                    after_rows=after_rows,
                    excluded_rows=deletion_rows,
                )
                if len(after_rows) == 1
                and after_rows[0].after_quantity != UNCHANGED_QUANTITY_MARKER
                and after_rows[0].after_locator != UNCHANGED_LOCATOR_MARKER
                else None
            )
            is_add_only_without_locator = is_add_only and not rows
            if not rows and not is_add_only:
                continue
            expected_before_rows = len(before_rows)
            if group.locator_type == "位号" and target_codes:
                expected_before_rows += sum(
                    1
                    for row_index in rows
                    if _cell_text(ws, row_index, columns["物料编码"]) in target_codes
                )
            if not is_add_only and len(rows) != expected_before_rows:
                _append_abnormal(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    reason=f"BOM命中{len(rows)}行，变更前有效物料{len(before_rows)}行，数量不一致",
                )
                continue
            if not is_add_only and not _is_contiguous(rows):
                _append_abnormal(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    reason="BOM命中行不连续，未自动修改",
                )
                continue
            is_partial_locator_change = not is_add_only and (
                (
                    len(rows) == 1
                    and _can_split_locator_row(
                        ws,
                        row_index=rows[0],
                        columns=columns,
                        group=group,
                        before_rows=before_rows,
                    )
                )
                or _can_split_locator_material_group(
                    ws,
                    rows=rows,
                    header_row=header_row,
                    columns=columns,
                    group=group,
                    before_rows=before_rows,
                )
            )
            actual = [_row_tuple(ws, row_index, columns) for row_index in rows]
            if not is_add_only and group.locator_type != "位号" and not _rows_match_before_rows(actual, before_rows):
                _append_abnormal(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    reason="BOM当前内容与变更前描述不一致",
                )
                continue
            if existing_target_row is not None and (is_add_only or existing_target_row not in rows):
                allowed_segments = _allowed_segments_for_change(
                    ws,
                    rows=[existing_target_row],
                    columns=columns,
                    new_materials=[],
                )
                if allowed_segments is not None and segment not in allowed_segments:
                    continue
                current_locators = _split_locators(
                    ws.cell(row=existing_target_row, column=columns["位号"]).value
                )
                added_locators = _split_locators(after_rows[0].after_locator)
                current_locator_set = set(current_locators)
                added_locator_set = set(added_locators)
                if added_locator_set and added_locator_set.issubset(current_locator_set):
                    if is_add_only:
                        continue
                    added_locators = []
                missing_locators = [locator for locator in added_locators if locator not in current_locator_set]
                added_locators = missing_locators or added_locators
                quantity_to_add: Any = (
                    len(added_locators)
                    if group.locator_type == "位号"
                    and added_locators
                    and current_locator_set.intersection(added_locator_set)
                    else after_rows[0].after_quantity
                )
                merged_quantity = (
                    ws.cell(row=existing_target_row, column=columns["用量"]).value
                    if group.locator_type == "位号" and not added_locators and not is_add_only
                    else _add_numeric_quantity(
                        ws.cell(row=existing_target_row, column=columns["用量"]).value,
                        quantity_to_add,
                    )
                )
                if merged_quantity is None:
                    _append_abnormal(
                        report_rows,
                        sheet_name=ws.title,
                        group=group,
                        reason="原BOM同料号用量或变更单新增用量不是数字，未自动合并",
                    )
                    continue
                report_before_quantities: list[Any] = []
                report_before_locators: list[Any] = []
                if not is_add_only:
                    target_locator_set = set(_split_locators(group.locator_value))
                    rows_to_delete: set[int] = set()
                    partial_removal = False
                    for row_index in rows:
                        group_rows = _material_group_rows(
                            ws,
                            row_index=row_index,
                            header_row=header_row,
                            columns=columns,
                        )
                        primary_locators = _split_locators(
                            ws.cell(row=group_rows[0], column=columns["位号"]).value
                        )
                        if target_locator_set and any(locator not in target_locator_set for locator in primary_locators):
                            partial_removal = True
                            rows_to_delete.update(
                                _remove_locators_from_existing_material_group(
                                    ws,
                                    row_index=row_index,
                                    header_row=header_row,
                                    columns=columns,
                                    target_locators=target_locator_set,
                                    quantity_to_remove=before_rows[0].before_quantity if before_rows else 0,
                                )
                            )
                        else:
                            rows_to_delete.update(group_rows)
                    if not partial_removal:
                        report_before_quantities = [
                            ws.cell(row=row_index, column=columns["用量"]).value for row_index in rows
                        ][: len(before_rows)]
                        report_before_locators = [
                            ws.cell(row=row_index, column=columns["位号"]).value for row_index in rows
                        ][: len(before_rows)]
                    for row_index in sorted(rows_to_delete, reverse=True):
                        ws.delete_rows(row_index, 1)
                        if row_index < existing_target_row:
                            existing_target_row -= 1
                    operations = _shift_operations_after_row_deletions(operations, rows_to_delete)
                if added_locators or is_add_only:
                    ws.cell(row=existing_target_row, column=columns["用量"]).value = merged_quantity
                    ws.cell(row=existing_target_row, column=columns["用量"]).number_format = "General"
                    ws.cell(row=existing_target_row, column=columns["位号"]).value = _join_locators(
                        [*current_locators, *added_locators]
                    )
                    for column in range(1, ws.max_column + 1):
                        ws.cell(row=existing_target_row, column=column).fill = copy(CHANGE_FILL)
                    _synchronize_alternative_locators(ws, row_index=existing_target_row, columns=columns)
                _append_report_for_change(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    before_rows=before_rows,
                    after_rows=after_rows,
                    before_locators=report_before_locators,
                    before_quantities=report_before_quantities,
                    after_locator=ws.cell(row=existing_target_row, column=columns["位号"]).value,
                    after_quantity=merged_quantity,
                )
                handled_group_keys.add(group.key)
                sheet_modified = True
                continue
            missing = [row.after_code for row in after_rows if row.after_code not in materials]
            if missing:
                _append_abnormal(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    reason=f"物料主数据未找到：{'、'.join(missing)}",
                )
                continue
            duplicated = [
                row.after_code
                for row in after_rows
                if row.after_code in materials and len(materials[row.after_code]) > 1
            ]
            if duplicated:
                _append_abnormal(
                    report_rows,
                    sheet_name=ws.title,
                    group=group,
                    reason=f"物料主数据匹配多行：{'、'.join(duplicated)}",
                )
                continue
            new_materials = [materials[row.after_code][0] for row in after_rows]
            allowed_segments = _allowed_segments_for_change(
                ws,
                rows=rows,
                columns=columns,
                new_materials=new_materials,
            )
            if allowed_segments is not None and segment not in allowed_segments:
                if (
                    is_add_only
                    or group.key in handled_group_keys
                    or _has_allowed_segment_match(wb, group, allowed_segments)
                ):
                    continue
            first_row = (
                _material_group_rows(
                    ws,
                    row_index=next(iter(_matched_rows(ws, header_row, columns, group))),
                    header_row=header_row,
                    columns=columns,
                )[-1] + 1
                if is_alternative_addition
                else _find_add_only_insert_row(
                    ws,
                    header_row=header_row,
                    columns=columns,
                    after_rows=after_rows,
                    new_materials=new_materials,
                )
                if is_add_only_without_locator
                else max(rows) + 1 if is_add_only else min(rows)
            )
            operations.append(
                PendingOperation(
                    first_row=first_row,
                    rows=[] if is_add_only else rows,
                    group=group,
                    before_rows=before_rows,
                    after_rows=after_rows,
                    new_materials=new_materials,
                    is_partial_locator_change=is_partial_locator_change,
                    preserve_first_row=is_alternative_addition,
                )
            )

        operations = _merge_same_target_operations(ws, header_row, columns, operations)
        # Apply removals/replacements against the original BOM rows before any
        # add-only insertion can shift those row indexes.
        for operation in sorted(
            operations,
            key=lambda item: (not item.rows, -item.first_row),
        ):
            if operation.is_partial_locator_change:
                _apply_partial_locator_change(
                    ws,
                    row_index=operation.first_row,
                    header_row=header_row,
                    columns=columns,
                    group=operation.group,
                    before_rows=operation.before_rows,
                    after_rows=operation.after_rows,
                    new_materials=operation.new_materials,
                )
                _append_report_for_change(
                    report_rows,
                    sheet_name=ws.title,
                    group=operation.group,
                    before_rows=operation.before_rows,
                    after_rows=operation.after_rows,
                )
                handled_group_keys.add(operation.group.key)
                continue
            if not operation.rows and not operation.preserve_first_row:
                operation = PendingOperation(
                    first_row=_find_add_only_insert_row(
                        ws,
                        header_row=header_row,
                        columns=columns,
                        after_rows=operation.after_rows,
                        new_materials=operation.new_materials,
                    ),
                    rows=operation.rows,
                    group=operation.group,
                    before_rows=operation.before_rows,
                    after_rows=operation.after_rows,
                    new_materials=operation.new_materials,
                    is_partial_locator_change=operation.is_partial_locator_change,
                )
            # Added rows must inherit the preceding material row. The insertion
            # point itself can be blank or belong to the next material section.
            style_source_row = operation.first_row if operation.rows else max(header_row + 1, operation.first_row - 1)
            captured_style = _capture_row_style(ws, style_source_row)
            if operation.rows:
                row_quantities = [ws.cell(row=row_index, column=columns["用量"]).value for row_index in operation.rows]
                row_locators = [ws.cell(row=row_index, column=columns["位号"]).value for row_index in operation.rows]
                report_before_quantities = row_quantities[: len(operation.before_rows)]
                report_before_locators = row_locators[: len(operation.before_rows)]
                preserved_quantities = list(row_quantities)
                preserved_locators = list(row_locators)
                if len(operation.after_rows) == 1 and len(operation.rows) > 1:
                    primary_index = next(
                        (
                            index
                            for index, quantity in enumerate(preserved_quantities)
                            if _quantity_compare_text(quantity) != "替代料"
                        ),
                        0,
                    )
                    preserved_quantities = [preserved_quantities[primary_index]]
                    preserved_locators = [preserved_locators[primary_index]]
                rows_to_delete = set(operation.rows)
                for row_index in operation.rows:
                    rows_to_delete.update(
                        _material_group_rows(
                            ws,
                            row_index=row_index,
                            header_row=header_row,
                            columns=columns,
                        )
                    )
                for row_index in sorted(rows_to_delete, reverse=True):
                    ws.delete_rows(row_index, 1)
                operations = _shift_operations_after_row_deletions(operations, rows_to_delete)
            else:
                preserved_quantities = []
                preserved_locators = []
                report_before_quantities = []
                report_before_locators = []
            if operation.new_materials:
                ws.insert_rows(operation.first_row, len(operation.new_materials))
                _set_inserted_sequence_values(
                    ws,
                    first_row=operation.first_row,
                    inserted_count=len(operation.new_materials),
                    sequence_column=columns.get("序号"),
                )
            for offset, (change_row, material) in enumerate(zip(operation.after_rows, operation.new_materials, strict=True)):
                target_row = operation.first_row + offset
                _apply_row_style(ws, target_row, captured_style)
                _set_bom_values(
                    ws,
                    target_row,
                    columns,
                    material,
                    change_row,
                    preserved_quantity=preserved_quantities[offset] if offset < len(preserved_quantities) else None,
                    preserved_locator=preserved_locators[offset] if offset < len(preserved_locators) else None,
                )
            _append_report_for_change(
                report_rows,
                sheet_name=ws.title,
                group=operation.group,
                before_rows=operation.before_rows,
                after_rows=operation.after_rows,
                before_locators=report_before_locators,
                before_quantities=report_before_quantities,
                after_locator=preserved_locators[0] if preserved_locators else None,
                # A replacement can add several materials with different
                # ECN quantities. Only a one-row replacement inherits the
                # removed row's quantity in the report.
                after_quantity=(
                    preserved_quantities[0]
                    if len(operation.after_rows) == 1 and preserved_quantities
                    else None
                ),
            )
            handled_group_keys.add(operation.group.key)
        for after_code in {
            row.after_code
            for operation in operations
            for row in operation.after_rows
            if _looks_material_code(row.after_code)
            and row.after_quantity != UNCHANGED_QUANTITY_MARKER
            and row.after_locator != UNCHANGED_LOCATOR_MARKER
        }:
            duplicate_rows = _existing_numeric_material_code_rows(
                ws,
                header_row=header_row,
                columns=columns,
                material_code=after_code,
            )
            if len(duplicate_rows) > 1:
                _consolidate_duplicate_material_rows(ws, row_indexes=duplicate_rows, columns=columns)
                sheet_modified = True
        if _synchronize_all_alternative_locators(ws, header_row=header_row, columns=columns):
            sheet_modified = True
        if operations or sheet_modified:
            _renumber_sequence_for_numeric_quantity_rows(ws, header_row=header_row, columns=columns)

    if processed_sheet_count == 0:
        raise ValueError("BOM未找到包含必要列的sheet")

    modified_count = sum(1 for row in report_rows if row.result == "成功")
    if change_record_people is not None and modified_count > 0:
        _append_change_record_sheet(
            wb,
            bom_path=bom_path,
            groups=groups,
            report_rows=report_rows,
            record_date=change_record_date or date.today(),
            people=change_record_people,
        )
    for issue in validate_processed_bom_worksheets(
        wb=wb,
        groups=groups,
        successful_report_rows=report_rows,
    ):
        report_rows.append(
            ReportRow(
                sheet_name=issue.sheet_name,
                locator_type="位号",
                locator_value=issue.locator_value,
                action="异常",
                before_code=issue.before_code,
                after_code=issue.after_code,
                reason=issue.reason,
            )
        )
    abnormal_count = sum(1 for row in report_rows if row.action == "异常")
    checklist_rows = _checklist_rows_from_report_rows(report_rows, processed_sheet_titles)
    if include_report_sheet:
        _write_report_sheet(wb, report_rows)
    elif REPORT_SHEET_NAME in wb.sheetnames:
        del wb[REPORT_SHEET_NAME]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return ProcessResult(
        output_path=output_path,
        modified_count=modified_count,
        abnormal_count=abnormal_count,
        report_count=len(report_rows),
        processed_sheet_count=processed_sheet_count,
        checklist_rows=checklist_rows,
    )

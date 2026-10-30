import re
import shutil
import subprocess
import time
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path
from threading import BoundedSemaphore
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, col, select

from app.core.config import settings
from app.models import (
    BomElectronicEditTask,
    BomElectronicEditTaskStatus,
    get_datetime_utc,
)
from app.services import bom_electronic_edit_processor
from app.services.storage_cleanup import robust_rmtree

ALLOWED_SUFFIXES = {".xls", ".xlsx", ".xlsm"}
OPENPYXL_SUFFIXES = {".xlsx", ".xlsm"}
ARCHIVE_SUFFIXES = {".zip", ".rar", ".7z"}
RESULT_NAME = "BOM电子档修改结果.xlsx"
RESULT_ZIP_NAME = "BOM电子档修改结果.zip"
CHECKLIST_SUMMARY_NAME = "BOM电子档修改清单汇总.xlsx"
INTERRUPTED_TASK_MESSAGE = "服务重启导致任务中断，请重新提交任务"
INTERRUPTED_STATUSES = {
    BomElectronicEditTaskStatus.preparing,
    BomElectronicEditTaskStatus.running,
    BomElectronicEditTaskStatus.uploading_report,
}
PRODUCT_MODEL_PATTERN = re.compile(r"(P[0-9A-Za-z]+-\d+[0-9A-Za-z]*-\d+(?:\.\d+)*)")
FILE_ACCESS_RETRY_COUNT = 10
FILE_ACCESS_RETRY_SECONDS = 2
EXTRACTION_SETTLE_SECONDS = 2
BOM_PROCESSING_WORKERS = 2
_BOM_PROCESSING_SLOTS = BoundedSemaphore(BOM_PROCESSING_WORKERS)


def task_dir(task_id: uuid.UUID) -> Path:
    return Path(settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR) / str(task_id)


def input_dir(task_id: uuid.UUID) -> Path:
    return task_dir(task_id) / "input"


def extracted_dir(task_id: uuid.UUID) -> Path:
    return input_dir(task_id) / "package"


def reports_dir(task_id: uuid.UUID) -> Path:
    return task_dir(task_id) / "reports"


def work_dir(task_id: uuid.UUID) -> Path:
    return task_dir(task_id) / "work"


def mark_interrupted_tasks_failed(session: Session) -> int:
    tasks = session.exec(
        select(BomElectronicEditTask).where(col(BomElectronicEditTask.status).in_(INTERRUPTED_STATUSES))
    ).all()
    now = get_datetime_utc()
    changed_count = 0
    for task in tasks:
        recovered_result = _recover_completed_result(task.id, task.input_files)
        if recovered_result is not None:
            task.result_files = recovered_result["result_files"]  # type: ignore[assignment]
            task.modified_count = int(recovered_result.get("modified_count", 0))
            task.abnormal_count = int(recovered_result.get("abnormal_count", 0))
            task.report_count = int(recovered_result.get("report_count", 0))
            task.processed_sheet_count = int(recovered_result.get("processed_sheet_count", 0))
            task.status = BomElectronicEditTaskStatus.success
            task.error_message = None
            task.finished_at = now
            task.updated_at = now
            session.add(task)
            changed_count += 1
            continue
        task.status = BomElectronicEditTaskStatus.failed
        task.error_message = INTERRUPTED_TASK_MESSAGE
        task.finished_at = now
        task.updated_at = now
        session.add(task)
        changed_count += 1
    if tasks:
        session.commit()
    return changed_count


def _recover_completed_result(
    task_id: uuid.UUID,
    input_files: list[dict[str, Any]],
) -> dict[str, object] | None:
    input_bom_count = sum(1 for file in input_files if file.get("role") == "bom")
    if input_bom_count <= 0:
        return None
    report_dir = reports_dir(task_id)
    if not report_dir.is_dir():
        return None
    report_workbooks = sorted(
        path
        for path in report_dir.glob("*.xls*")
        if path.name != CHECKLIST_SUMMARY_NAME
    )
    if len(report_workbooks) < input_bom_count:
        return None

    zip_path = result_zip_path(task_id)
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for workbook in report_workbooks:
            archive.write(workbook, workbook.name)
        summary_path = checklist_summary_path(task_id)
        if summary_path.is_file():
            archive.write(summary_path, CHECKLIST_SUMMARY_NAME)

    return {
        "result_files": [
            {
                "file_name": RESULT_ZIP_NAME,
                "file_path": str(zip_path),
                "content_type": "application/zip",
                "file_size": zip_path.stat().st_size,
                "role": "modified_bom_bundle",
            }
        ],
        "modified_count": 0,
        "abnormal_count": 0,
        "report_count": len(report_workbooks),
        "processed_sheet_count": 0,
    }


def remove_task_files(task_id: uuid.UUID) -> None:
    storage_dir = Path(settings.BOM_ELECTRONIC_EDIT_STORAGE_DIR).resolve()
    target_dir = (storage_dir / str(task_id)).resolve()
    if target_dir.name != str(task_id) or target_dir.parent != storage_dir:
        raise RuntimeError("Invalid BOM electronic edit task storage path")
    if target_dir.exists():
        if not target_dir.is_dir():
            raise RuntimeError("BOM electronic edit task storage path is not a directory")
        robust_rmtree(target_dir)


def _safe_filename(filename: str | None, fallback: str) -> str:
    name = Path(filename or fallback).name
    for char in '<>:"/\\|?*\x00':
        name = name.replace(char, "_")
    return name.strip() or fallback


async def save_excel_upload(
    task_id: uuid.UUID,
    *,
    upload: UploadFile,
    role: str,
    fallback: str,
    filename_override: str | None = None,
) -> dict[str, str | int | None]:
    content = await upload.read()
    if not content:
        raise HTTPException(status_code=400, detail=f"{fallback}不能为空")
    filename = _safe_filename(filename_override or upload.filename, fallback)
    if Path(filename).suffix.lower() not in ALLOWED_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"{fallback}必须是 .xls、.xlsx 或 .xlsm 文件")
    directory = input_dir(task_id)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / filename
    path.write_bytes(content)
    return {
        "file_name": filename,
        "file_path": str(path),
        "content_type": upload.content_type,
        "file_size": len(content),
        "role": role,
    }


async def save_bom_uploads(
    task_id: uuid.UUID,
    *,
    uploads: list[UploadFile],
) -> list[dict[str, str | int | None]]:
    if not uploads:
        raise HTTPException(status_code=400, detail="BOM表不能为空")
    saved_files: list[dict[str, str | int | None]] = []
    used_names: set[str] = set()
    for index, upload in enumerate(uploads, start=1):
        filename = _safe_filename(upload.filename, f"BOM-{index}.xlsx")
        stem = Path(filename).stem
        suffix = Path(filename).suffix
        unique_name = filename
        counter = 2
        while unique_name.lower() in used_names:
            unique_name = f"{stem}-{counter}{suffix}"
            counter += 1
        used_names.add(unique_name.lower())
        saved_files.append(
            await save_excel_upload(
                task_id,
                upload=upload,
                role="bom",
                fallback=f"BOM-{index}.xlsx",
                filename_override=unique_name,
            )
        )
    return saved_files


async def save_project_archive_upload(
    task_id: uuid.UUID,
    *,
    upload: UploadFile,
) -> list[dict[str, str | int | None]]:
    content = await upload.read()
    if not content:
        raise HTTPException(status_code=400, detail="项目压缩包不能为空")
    directory = input_dir(task_id)
    directory.mkdir(parents=True, exist_ok=True)
    filename = _safe_filename(upload.filename, "BOM电子档修改资料.zip")
    archive_path = directory / filename
    archive_path.write_bytes(content)
    if archive_path.suffix.lower() not in ARCHIVE_SUFFIXES:
        raise HTTPException(status_code=400, detail="项目压缩包必须是 .zip、.rar 或 .7z 文件")
    destination = extracted_dir(task_id)
    if archive_path.suffix.lower() == ".zip":
        _extract_zip(archive_path, destination)
    elif archive_path.suffix.lower() == ".rar":
        _extract_with_rar(archive_path, destination)
    else:
        _extract_with_7zip(archive_path, destination)
    classified = _classify_extracted_excels(destination)
    return [
        {
            "file_name": filename,
            "file_path": str(archive_path),
            "content_type": upload.content_type,
            "file_size": len(content),
            "role": "project_archive",
        },
        *classified,
    ]


async def save_task_uploads(
    task_id: uuid.UUID,
    *,
    change_notice_upload: UploadFile,
    bom_uploads: list[UploadFile],
    material_master_upload: UploadFile,
) -> list[dict[str, str | int | None]]:
    saved_files = [
        await save_excel_upload(
            task_id,
            upload=change_notice_upload,
            role="change_notice",
            fallback="工程变更通知单.xlsx",
        ),
        await save_excel_upload(
            task_id,
            upload=material_master_upload,
            role="material_master",
            fallback="物料主数据查询.xlsx",
        ),
    ]
    saved_files[1:1] = await save_bom_uploads(task_id, uploads=bom_uploads)
    return saved_files


def _classify_extracted_excels(root: Path) -> list[dict[str, str | int | None]]:
    excel_files = sorted(
        path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in ALLOWED_SUFFIXES and not path.name.startswith("~$")
    )
    if not excel_files:
        raise HTTPException(status_code=400, detail="压缩包内未找到 Excel 文件")

    change_files: list[Path] = []
    master_files: list[Path] = []
    for path in excel_files:
        normalized = "".join(path.relative_to(root).parts)
        if any(token in normalized for token in ("变更通知", "工程变更", "变更单")):
            change_files.append(path)
        elif any(token in normalized for token in ("物料主数据", "主数据")):
            master_files.append(path)
    bom_files = [
        path
        for path in excel_files
        if any(part == "变更BOM" for part in path.relative_to(root).parts[:-1])
    ]

    if len(change_files) != 1:
        names = "、".join(path.name for path in change_files) or "0个"
        raise HTTPException(status_code=400, detail=f"压缩包内工程变更通知单必须正好1个，当前识别到：{names}")
    if len(master_files) != 1:
        names = "、".join(path.name for path in master_files) or "0个"
        raise HTTPException(status_code=400, detail=f"压缩包内物料主数据查询必须正好1个，当前识别到：{names}")
    if not bom_files:
        raise HTTPException(status_code=400, detail="压缩包内未在“变更BOM”文件夹找到待修改BOM表")

    files = [
        _file_metadata(change_files[0], "change_notice"),
        _file_metadata(master_files[0], "material_master"),
    ]
    files.extend(_file_metadata(path, "bom") for path in bom_files)
    return files


def _file_metadata(path: Path, role: str) -> dict[str, str | int | None]:
    return {
        "file_name": path.name,
        "file_path": str(path),
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "file_size": path.stat().st_size,
        "role": role,
    }


def _extract_with_7zip(path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    command = _resolve_7z_command()
    _extract_with_archive_tool(path, destination, [command, "x", str(path), f"-o{destination}", "-y"])


def _extract_with_rar(path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    command = _resolve_rar_command()
    _extract_with_archive_tool(path, destination, [command, "x", "-y", str(path), str(destination) + "\\"])


def _extract_with_archive_tool(path: Path, destination: Path, command: list[str]) -> None:
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=300,
            check=False,
        )
    except FileNotFoundError as exc:
        raise HTTPException(status_code=400, detail=f"服务器未安装可用解压工具，无法解压 {path.suffix.lower()} 文件") from exc
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(status_code=400, detail="压缩包解压超时，请检查文件大小") from exc
    if completed.returncode != 0:
        detail = "\n".join([completed.stderr or "", completed.stdout or ""]).strip().splitlines()
        raise HTTPException(
            status_code=400,
            detail=f"压缩包解压失败，请检查文件是否损坏或带密码：{detail[-1] if detail else '未知错误'}",
        )
    _validate_extracted_paths(destination)


def _resolve_7z_command() -> str:
    candidates = ["7z", "7z.exe", r"C:\Program Files\7-Zip\7z.exe", r"C:\Program Files (x86)\7-Zip\7z.exe"]
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    raise HTTPException(status_code=400, detail="服务器未安装 7-Zip，无法解压 7z 文件")


def _resolve_rar_command() -> str:
    candidates = [
        r"D:\SoftWare\WinRAR\UnRAR.exe",
        r"D:\SoftWare\WinRAR\Rar.exe",
        r"D:\SoftWare\WinRAR\WinRAR.exe",
        "WinRAR",
        "WinRAR.exe",
        "RAR",
        "RAR.exe",
        "UnRAR",
        "UnRAR.exe",
        r"C:\Program Files\WinRAR\WinRAR.exe",
        r"C:\Program Files\WinRAR\Rar.exe",
        r"C:\Program Files\WinRAR\UnRAR.exe",
        r"C:\Program Files (x86)\WinRAR\WinRAR.exe",
        r"C:\Program Files (x86)\WinRAR\Rar.exe",
        r"C:\Program Files (x86)\WinRAR\UnRAR.exe",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    raise HTTPException(status_code=400, detail="服务器未安装 WinRAR/RAR/UnRAR，无法解压 rar 文件")


def _validate_extracted_paths(destination: Path) -> None:
    root = destination.resolve()
    for extracted in destination.rglob("*"):
        if not extracted.resolve().is_relative_to(root):
            raise HTTPException(status_code=400, detail="压缩包包含非法路径")


def _extract_zip(path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    root = destination.resolve()
    with zipfile.ZipFile(path) as archive:
        for item in archive.infolist():
            if item.is_dir():
                continue
            target = (destination / item.filename).resolve()
            if not target.is_relative_to(root):
                raise HTTPException(status_code=400, detail="压缩包包含非法路径")
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(archive.read(item))


def _file_by_role(
    input_files: list[dict[str, object]],
    role: str,
) -> Path:
    for file in input_files:
        if file.get("role") == role:
            path = Path(str(file.get("file_path") or ""))
            if path.is_file():
                return path
    raise RuntimeError(f"缺少输入文件：{role}")


def _files_by_role(
    input_files: list[dict[str, object]],
    role: str,
) -> list[Path]:
    files: list[Path] = []
    for file in input_files:
        if file.get("role") == role:
            path = Path(str(file.get("file_path") or ""))
            if path.is_file():
                files.append(path)
    if not files:
        raise RuntimeError(f"缺少输入文件：{role}")
    return files


def _result_workbook_path(task_id: uuid.UUID, bom_path: Path) -> Path:
    """Build the output workbook name without keeping dates.

    - A trailing three-digit sequence is incremented (`-003` -> `-004`,
      letter tags like `A001` -> `A002`, trailing text is preserved).
    - Otherwise a trailing eight-digit date is replaced with `-001`.
    - Names without a sequence or date get `-001` appended.
    """
    safe_stem = _safe_filename(bom_path.stem, "BOM")
    match = re.search(r"(?P<prefix>.*-)(?P<tag>[A-Za-z]?)(?P<sequence>\d{3})(?!\d)(?P<trailing>[^-]*)$", safe_stem)
    output_suffix = ".xlsx" if bom_path.suffix.lower() == ".xls" else bom_path.suffix
    if match:
        sequence = match.group("sequence")
        next_sequence = str(int(sequence) + 1).zfill(len(sequence))
        return reports_dir(task_id) / f"{match.group('prefix')}{match.group('tag')}{next_sequence}{match.group('trailing')}{output_suffix}"
    date_match = re.search(r"-(?P<date>\d{8})$", safe_stem)
    if date_match:
        stem_without_date = safe_stem[: -(len(date_match.group("date")) + 1)]
        return reports_dir(task_id) / f"{stem_without_date}-001{output_suffix}"
    return reports_dir(task_id) / f"{safe_stem}-001{output_suffix}"


def _unchanged_result_workbook_path(task_id: uuid.UUID, bom_path: Path) -> Path:
    safe_stem = _safe_filename(bom_path.stem, "BOM")
    output_suffix = ".xlsx" if bom_path.suffix.lower() == ".xls" else bom_path.suffix
    return reports_dir(task_id) / f"{safe_stem}{output_suffix}"


def _access_error_workbook_path(task_id: uuid.UUID, bom_path: Path) -> Path:
    safe_stem = _safe_filename(bom_path.stem, "BOM")
    output_suffix = ".xlsx" if bom_path.suffix.lower() == ".xls" else bom_path.suffix
    return reports_dir(task_id) / f"异常_{safe_stem}{output_suffix}"


def _convert_xls_to_xlsx_with_excel(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    source_literal = "'" + str(source.resolve()).replace("'", "''") + "'"
    target_literal = "'" + str(target.resolve()).replace("'", "''") + "'"
    script = f"""
$ErrorActionPreference = 'Stop'
$source = {source_literal}
$target = {target_literal}
$excel = $null
$workbook = $null
try {{
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $workbook = $excel.Workbooks.Open($source)
    $workbook.SaveAs($target, 51)
}} finally {{
    if ($workbook -ne $null) {{ $workbook.Close($false) | Out-Null }}
    if ($excel -ne $null) {{
        $excel.Quit()
        [System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) | Out-Null
    }}
}}
"""
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
        check=False,
    )
    if completed.returncode != 0 or not target.is_file():
        detail = "\n".join([completed.stderr or "", completed.stdout or ""]).strip().splitlines()
        raise RuntimeError(f"转换 .xls 文件失败：{source.name}；{detail[-1] if detail else '未生成转换文件'}")


def _prepare_bom_for_processing(task_id: uuid.UUID, bom_path: Path) -> Path:
    suffix = bom_path.suffix.lower()
    if suffix in OPENPYXL_SUFFIXES:
        return bom_path
    if suffix != ".xls":
        raise RuntimeError(f"不支持的BOM文件格式：{bom_path.name}")
    target = work_dir(task_id) / "converted" / f"{bom_path.stem}.xlsx"
    _convert_xls_to_xlsx_with_excel(bom_path, target)
    return target


def _copy_bom_to_temporary_workfile(task_id: uuid.UUID, bom_path: Path) -> Path:
    processing_dir = work_dir(task_id) / "processing"
    processing_dir.mkdir(parents=True, exist_ok=True)
    temporary_path = processing_dir / bom_path.name
    shutil.copyfile(bom_path, temporary_path)
    return temporary_path


def _write_access_error_workbook(path: Path, *, bom_name: str, error: PermissionError) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "异常"
    ws.append(["BOM文件", "处理结果", "异常原因"])
    ws.append([bom_name, "未处理", f"文件无法访问，已重试{FILE_ACCESS_RETRY_COUNT}次：{error}"])
    wb.save(path)


def _write_bom_processing_error_workbook(path: Path, *, bom_name: str, error: Exception) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "异常"
    ws.append(["BOM文件", "处理结果", "异常原因"])
    ws.append([bom_name, "未处理", f"处理失败：{error}"])
    wb.save(path)


def result_zip_path(task_id: uuid.UUID) -> Path:
    return reports_dir(task_id) / RESULT_ZIP_NAME


def checklist_summary_path(task_id: uuid.UUID) -> Path:
    return reports_dir(task_id) / CHECKLIST_SUMMARY_NAME


def _write_checklist_summary(
    path: Path,
    results: list[tuple[str, bom_electronic_edit_processor.ProcessResult]],
) -> None:
    wb = Workbook()
    checklist = wb.active
    checklist.title = "清单汇总"
    checklist_headers = [
        "BOM文件",
        "Sheet",
        "Sheet类型",
        "位号/定位值",
        "动作",
        "变更前物料编码",
        "变更前型号",
        "变更前用量",
        "变更前位号",
        "变更后物料编码",
        "变更后型号",
        "变更后用量",
        "变更后位号",
        "状态",
        "原因",
    ]
    checklist.append(checklist_headers)

    summary = wb.create_sheet("BOM汇总")
    summary_headers = ["BOM文件", "状态", "总表完成数", "分层完成数", "未更改数", "异常数", "判定结论", "备注"]
    summary.append(summary_headers)

    for bom_name, result in results:
        display_name = _bom_display_name(bom_name, result)
        rows = _rows_with_workbook_sheet_types(result)
        for row in rows:
            checklist.append(
                [
                    display_name,
                    display_name if row.sheet_type == "总表" else row.sheet_name,
                    row.sheet_type,
                    row.locator_value,
                    row.action,
                    row.before_code,
                    row.before_model,
                    row.before_quantity,
                    row.before_locator,
                    row.after_code,
                    row.after_model,
                    row.after_quantity,
                    row.after_locator,
                    row.status,
                    row.reason,
                ]
            )

        total_done = sum(1 for row in rows if row.status == "已完成" and row.sheet_type == "总表")
        layer_done = sum(1 for row in rows if row.status == "已完成" and row.sheet_type != "总表")
        not_changed = sum(1 for row in rows if row.status == "未更改")
        abnormal = sum(1 for row in rows if row.status == "异常")
        if abnormal:
            status = "存在异常"
            conclusion = "不一致"
        elif not_changed:
            status = "存在未更改"
            conclusion = "不一致"
        elif rows:
            status = "全部完成"
            conclusion = "一致"
        else:
            status = "无变更"
            conclusion = "一致"
        reasons = []
        seen_reasons: set[str] = set()
        for row in rows:
            if row.reason and row.reason not in seen_reasons:
                seen_reasons.add(row.reason)
                reasons.append(row.reason)
        summary.append([display_name, status, total_done, layer_done, not_changed, abnormal, conclusion, "；".join(reasons)])

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for ws in (checklist, summary):
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column_index, column_cells in enumerate(ws.columns, start=1):
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 36)
    _merge_repeated_cells(checklist, [1, 2, 3])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _sheet_type_label(sheet_name: str, sheet_titles: list[str] | None = None) -> str:
    segment = bom_electronic_edit_processor._sheet_segment(sheet_name, sheet_titles)  # noqa: SLF001
    labels = {
        "total": "总表",
        "smt": "SMT段",
        "dip": "DIP段",
        "product": "成品段",
    }
    return labels.get(segment, "其他")


def _bom_display_name(
    fallback_name: str,
    result: bom_electronic_edit_processor.ProcessResult,
) -> str:
    try:
        wb = load_workbook(result.output_path, read_only=True)
    except Exception:
        return fallback_name
    try:
        sheet_titles = _effective_bom_sheet_titles(wb)
        for title in sheet_titles:
            if _sheet_type_label(title, sheet_titles) == "总表":
                model_name = _product_model_from_sheet_title(title)
                if model_name:
                    return model_name
        for title in sheet_titles:
            model_name = _product_model_from_sheet_title(title)
            if model_name:
                return model_name
    finally:
        wb.close()
    return fallback_name


def _product_model_from_sheet_title(title: str) -> str | None:
    match = PRODUCT_MODEL_PATTERN.search(str(title).strip())
    return match.group(1) if match else None


def _rows_with_workbook_sheet_types(
    result: bom_electronic_edit_processor.ProcessResult,
) -> list[bom_electronic_edit_processor.ChecklistRow]:
    if not result.checklist_rows:
        return []
    try:
        wb = load_workbook(result.output_path, read_only=True)
    except Exception:
        return list(result.checklist_rows)
    try:
        sheet_titles = _effective_bom_sheet_titles(wb)
    finally:
        wb.close()
    rows: list[bom_electronic_edit_processor.ChecklistRow] = []
    for row in result.checklist_rows:
        rows.append(
            bom_electronic_edit_processor.ChecklistRow(
                sheet_name=row.sheet_name,
                sheet_type=_sheet_type_label(row.sheet_name, sheet_titles),
                locator_value=row.locator_value,
                action=row.action,
                before_code=row.before_code,
                before_model=row.before_model,
                before_quantity=row.before_quantity,
                before_locator=row.before_locator,
                after_code=row.after_code,
                after_model=row.after_model,
                after_quantity=row.after_quantity,
                after_locator=row.after_locator,
                after_action=row.after_action,
                status=row.status,
                reason=row.reason,
            )
        )
    return rows


def _effective_bom_sheet_titles(wb: Any) -> list[str]:
    titles: list[str] = []
    for ws in wb.worksheets:
        if ws.title in {
            bom_electronic_edit_processor.REPORT_SHEET_NAME,
            bom_electronic_edit_processor.CHANGE_RECORD_SHEET_NAME,
        }:
            continue
        try:
            if bom_electronic_edit_processor._find_header_row(ws, bom_electronic_edit_processor.BOM_HEADERS) is not None:  # noqa: SLF001
                titles.append(ws.title)
        except Exception:
            continue
    return titles


def _merge_repeated_cells(ws: Any, columns: list[int]) -> None:
    values = {
        (row, column): ws.cell(row=row, column=column).value
        for row in range(2, ws.max_row + 1)
        for column in range(1, max(columns) + 1)
    }
    for column in columns:
        start_row = 2
        while start_row <= ws.max_row:
            value = values[(start_row, column)]
            end_row = start_row
            while (
                end_row + 1 <= ws.max_row
                and values[(end_row + 1, column)] == value
                and all(values[(end_row + 1, parent)] == values[(start_row, parent)] for parent in range(1, column))
            ):
                end_row += 1
            if value not in (None, "") and end_row > start_row:
                ws.merge_cells(start_row=start_row, start_column=column, end_row=end_row, end_column=column)
                ws.cell(row=start_row, column=column).alignment = Alignment(vertical="center")
            start_row = end_row + 1


def _status_from_report_result(action: str, result: str) -> str:
    if action == "异常":
        return "异常"
    if result == "成功":
        return "已完成"
    if result == "未修改":
        return "未更改"
    return result or "未更改"


def _checklist_rows_from_workbook_report(workbook_path: Path) -> list[bom_electronic_edit_processor.ChecklistRow]:
    try:
        wb = load_workbook(workbook_path, data_only=True, read_only=True)
    except Exception:
        return []
    try:
        if bom_electronic_edit_processor.REPORT_SHEET_NAME not in wb.sheetnames:
            return []
        ws = wb[bom_electronic_edit_processor.REPORT_SHEET_NAME]
        headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        columns = {header: index for index, header in enumerate(headers)}

        def value(row: tuple[object, ...], header: str) -> str:
            index = columns.get(header)
            if index is None or index >= len(row):
                return ""
            cell_value = getattr(row[index], "value", None)
            return "" if cell_value is None else str(cell_value).strip()

        rows: list[bom_electronic_edit_processor.ChecklistRow] = []
        for excel_row in ws.iter_rows(min_row=2):
            row_values = tuple(excel_row)
            sheet_name = value(row_values, "Sheet")
            action = value(row_values, "动作")
            result = value(row_values, "结果")
            if not any(getattr(cell, "value", None) for cell in row_values):
                continue
            rows.append(
                bom_electronic_edit_processor.ChecklistRow(
                    sheet_name=sheet_name,
                    sheet_type=_sheet_type_label(sheet_name),
                    locator_value=value(row_values, "定位值"),
                    action=action,
                    before_code=value(row_values, "变更前物料编码"),
                    before_model=value(row_values, "变更前规格型号"),
                    before_quantity=value(row_values, "变更前用量"),
                    before_locator=value(row_values, "变更前位号"),
                    after_code=value(row_values, "变更后物料编码"),
                    after_model=value(row_values, "变更后规格型号"),
                    after_quantity=value(row_values, "变更后用量"),
                    after_locator=value(row_values, "变更后位号"),
                    after_action="添加" if action == "添加" else "",
                    status=_status_from_report_result(action, result),
                    reason=value(row_values, "原因"),
                )
            )
        return rows
    finally:
        wb.close()


def _checklist_summary_results_from_report_dir(
    report_dir: Path,
) -> list[tuple[str, bom_electronic_edit_processor.ProcessResult]]:
    results: list[tuple[str, bom_electronic_edit_processor.ProcessResult]] = []
    for workbook_path in sorted(report_dir.glob("*.xls*")):
        if workbook_path.name == CHECKLIST_SUMMARY_NAME:
            continue
        rows = _checklist_rows_from_workbook_report(workbook_path)
        results.append(
            (
                workbook_path.name,
                bom_electronic_edit_processor.ProcessResult(
                    output_path=workbook_path,
                    modified_count=sum(1 for row in rows if row.status == "已完成"),
                    abnormal_count=sum(1 for row in rows if row.status == "异常"),
                    report_count=len(rows),
                    processed_sheet_count=0,
                    checklist_rows=rows,
                ),
            )
        )
    return results


def ensure_checklist_summary_in_result_zip(zip_path: Path) -> bool:
    if not zip_path.is_file():
        return False
    summary_path = zip_path.parent / CHECKLIST_SUMMARY_NAME
    _write_checklist_summary(summary_path, _checklist_summary_results_from_report_dir(zip_path.parent))

    tmp_zip = zip_path.with_name(f"{zip_path.stem}.tmp{zip_path.suffix}")
    with zipfile.ZipFile(zip_path) as source, zipfile.ZipFile(tmp_zip, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename == CHECKLIST_SUMMARY_NAME:
                continue
            target.writestr(item, source.read(item.filename))
        target.write(summary_path, CHECKLIST_SUMMARY_NAME)
    tmp_zip.replace(zip_path)
    return True


def _change_record_people_from_task_name(task_name: str | None) -> bom_electronic_edit_processor.ChangeRecordPeople:
    name = task_name or ""
    if "廖晶华" in name:
        return bom_electronic_edit_processor.ChangeRecordPeople(
            preparer="廖晶华",
            reviewer="邓如锦",
            approver="许志杰",
        )
    if "邓如锦" in name:
        return bom_electronic_edit_processor.ChangeRecordPeople(
            preparer="邓如锦",
            reviewer="廖晶华",
            approver="许志杰",
        )
    if "刘德媚" in name:
        return bom_electronic_edit_processor.ChangeRecordPeople(
            preparer="刘德媚",
            reviewer="/",
            approver="/",
        )
    raise ValueError("任务名称未包含廖晶华、邓如锦或刘德媚，无法填写变更记录人员")


def _process_one_bom(
    *,
    task_id: uuid.UUID,
    bom: Path,
    change_notice: Path,
    material_master: Path,
    change_groups: list[bom_electronic_edit_processor.ChangeGroup],
    materials: dict[str, list[bom_electronic_edit_processor.MaterialInfo]],
    change_record_people: bom_electronic_edit_processor.ChangeRecordPeople,
    change_record_date: date,
) -> bom_electronic_edit_processor.ProcessResult:
    with _BOM_PROCESSING_SLOTS:
        return _process_one_bom_with_slot(
            task_id=task_id,
            bom=bom,
            change_notice=change_notice,
            material_master=material_master,
            change_groups=change_groups,
            materials=materials,
            change_record_people=change_record_people,
            change_record_date=change_record_date,
        )


def _process_one_bom_with_slot(
    *,
    task_id: uuid.UUID,
    bom: Path,
    change_notice: Path,
    material_master: Path,
    change_groups: list[bom_electronic_edit_processor.ChangeGroup],
    materials: dict[str, list[bom_electronic_edit_processor.MaterialInfo]],
    change_record_people: bom_electronic_edit_processor.ChangeRecordPeople,
    change_record_date: date,
) -> bom_electronic_edit_processor.ProcessResult:
    output = _result_workbook_path(task_id, bom)
    for attempt in range(FILE_ACCESS_RETRY_COUNT):
        temporary_bom: Path | None = None
        try:
            temporary_bom = _copy_bom_to_temporary_workfile(task_id, bom)
            prepared_bom = _prepare_bom_for_processing(task_id, temporary_bom)
            result = bom_electronic_edit_processor.process_bom_workbook(
                change_notice_path=change_notice,
                bom_path=prepared_bom,
                material_master_path=material_master,
                output_path=output,
                change_groups=change_groups,
                materials=materials,
                change_record_people=change_record_people,
                change_record_date=change_record_date,
                include_report_sheet=False,
            )
            break
        except PermissionError as error:
            if attempt + 1 < FILE_ACCESS_RETRY_COUNT:
                time.sleep(FILE_ACCESS_RETRY_SECONDS * (attempt + 1))
                continue
            access_error_output = _access_error_workbook_path(task_id, bom)
            _write_access_error_workbook(access_error_output, bom_name=bom.name, error=error)
            result = bom_electronic_edit_processor.ProcessResult(
                output_path=access_error_output,
                modified_count=0,
                abnormal_count=1,
                report_count=1,
                processed_sheet_count=0,
                checklist_rows=[
                    bom_electronic_edit_processor.ChecklistRow(
                        sheet_name="",
                        sheet_type="",
                        locator_value="",
                        action="异常",
                        status="异常",
                        reason=f"文件无法访问，已重试{FILE_ACCESS_RETRY_COUNT}次：{error}",
                    )
                ],
            )
            break
        except Exception as error:
            error_output = _access_error_workbook_path(task_id, bom)
            _write_bom_processing_error_workbook(error_output, bom_name=bom.name, error=error)
            result = bom_electronic_edit_processor.ProcessResult(
                output_path=error_output,
                modified_count=0,
                abnormal_count=1,
                report_count=1,
                processed_sheet_count=0,
                checklist_rows=[
                    bom_electronic_edit_processor.ChecklistRow(
                        sheet_name="",
                        sheet_type="",
                        locator_value="",
                        action="异常",
                        status="异常",
                        reason=f"处理失败：{error}",
                    )
                ],
            )
            break
        finally:
            if temporary_bom is not None:
                temporary_bom.unlink(missing_ok=True)
    if result.modified_count == 0 and result.abnormal_count == 0:
        unchanged_output = _unchanged_result_workbook_path(task_id, bom)
        if result.output_path != unchanged_output:
            unchanged_output.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(result.output_path, unchanged_output)
        return bom_electronic_edit_processor.ProcessResult(
            output_path=unchanged_output,
            modified_count=result.modified_count,
            abnormal_count=result.abnormal_count,
            report_count=result.report_count,
            processed_sheet_count=result.processed_sheet_count,
            checklist_rows=result.checklist_rows,
        )
    return result


def run_task(
    task_id: uuid.UUID,
    input_files: list[dict[str, object]],
    task_name: str | None = None,
) -> dict[str, object]:
    change_notice = _file_by_role(input_files, "change_notice")
    boms = _files_by_role(input_files, "bom")
    material_master = _file_by_role(input_files, "material_master")
    change_record_people = _change_record_people_from_task_name(task_name)
    change_record_date = date.today()
    reports_dir(task_id).mkdir(parents=True, exist_ok=True)
    change_groups = bom_electronic_edit_processor.parse_change_notice(change_notice)
    materials = bom_electronic_edit_processor.load_material_master(material_master)
    time.sleep(EXTRACTION_SETTLE_SECONDS)
    with ThreadPoolExecutor(max_workers=BOM_PROCESSING_WORKERS) as executor:
        results = list(
            executor.map(
                lambda bom: _process_one_bom(
                    task_id=task_id,
                    bom=bom,
                    change_notice=change_notice,
                    material_master=material_master,
                    change_groups=change_groups,
                    materials=materials,
                    change_record_people=change_record_people,
                    change_record_date=change_record_date,
                ),
                boms,
            )
        )
    summary_path = checklist_summary_path(task_id)
    _write_checklist_summary(summary_path, [(result.output_path.name, result) for result in results])
    zip_path = result_zip_path(task_id)
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for result in results:
            archive.write(result.output_path, result.output_path.name)
        archive.write(summary_path, CHECKLIST_SUMMARY_NAME)
    return {
        "result_files": [
            {
                "file_name": RESULT_ZIP_NAME,
                "file_path": str(zip_path),
                "content_type": "application/zip",
                "file_size": zip_path.stat().st_size,
                "role": "modified_bom_bundle",
            }
        ],
        "modified_count": sum(result.modified_count for result in results),
        "abnormal_count": sum(result.abnormal_count for result in results),
        "report_count": sum(result.report_count for result in results),
        "processed_sheet_count": sum(result.processed_sheet_count for result in results),
    }
